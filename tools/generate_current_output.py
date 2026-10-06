"""Genera la vista navegable ``output/current`` a partir de runs inmutables.

La vista contiene exclusivamente README, un hard link al scorecard oficial y
junctions de Windows hacia ``output/assessments``. No copia ni reescribe
artefactos históricos y rechaza reemplazar una carpeta ``current`` que no haya
sido creada por este generador.
"""

from __future__ import annotations

import argparse
import json
import os
import re
import stat
import subprocess
import sys
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable

import yaml
from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]
OUTPUT_ROOT = PROJECT_ROOT / "output"
ASSESSMENTS_ROOT = OUTPUT_ROOT / "assessments"
VIEW_ROOT = OUTPUT_ROOT / "current"
STAGING_ROOT = OUTPUT_ROOT / ".current-staging"
DEFAULT_CATALOG = PROJECT_ROOT / "config" / "current-output-catalog.yaml"
GENERATOR_ID = "tools/generate_current_output.py"
VIEW_MANIFEST_NAME = "current-output-manifest.json"
EXPECTED_WORKSHEET_NAMES = ["Variable assessment", "Coverage", "Score rationale"]


class ViewGenerationError(RuntimeError):
    """Error de contrato que debe detener una generación segura."""


@dataclass(frozen=True)
class SourceArtifact:
    """Un directorio de variable dentro de un run validado."""

    run_id: str
    generated_at: datetime
    folder: Path
    scope: str


@dataclass(frozen=True)
class AssessmentRun:
    """Metadatos mínimos leídos de un manifiesto validado."""

    run_id: str
    generated_at: datetime
    root: Path
    variables: dict[str, Path]


@dataclass(frozen=True)
class VariableView:
    """Contrato de una carpeta navegable de variable."""

    canonical_name: str
    output_relative: Path
    base_artifact: SourceArtifact
    review_artifacts: tuple[SourceArtifact, ...]
    is_scorecard_variable: bool

    @property
    def history(self) -> tuple[SourceArtifact, ...]:
        if self.is_scorecard_variable:
            artifacts = (self.base_artifact, *self.review_artifacts)
        else:
            artifacts = self.review_artifacts
        return tuple(sorted(artifacts, key=lambda artifact: (artifact.generated_at, artifact.run_id)))

    @property
    def latest_review(self) -> SourceArtifact | None:
        if self.is_scorecard_variable:
            candidates = self.review_artifacts
        else:
            candidates = tuple(
                artifact for artifact in self.review_artifacts if artifact != self.base_artifact
            )
        if not candidates:
            return None
        return max(candidates, key=lambda artifact: (artifact.generated_at, artifact.run_id))


def read_yaml(path: Path) -> dict[str, Any]:
    try:
        value = yaml.safe_load(path.read_text(encoding="utf-8"))
    except FileNotFoundError as exc:
        raise ViewGenerationError(f"No existe el catálogo requerido: {path}") from exc
    except yaml.YAMLError as exc:
        raise ViewGenerationError(f"El catálogo YAML no es válido: {path}: {exc}") from exc
    if not isinstance(value, dict):
        raise ViewGenerationError(f"El catálogo debe ser un objeto YAML: {path}")
    return value


def read_json(path: Path, label: str) -> dict[str, Any]:
    try:
        value = json.loads(path.read_text(encoding="utf-8"))
    except FileNotFoundError as exc:
        raise ViewGenerationError(f"No existe {label}: {path}") from exc
    except json.JSONDecodeError as exc:
        raise ViewGenerationError(f"{label} no contiene JSON válido: {path}: {exc}") from exc
    if not isinstance(value, dict):
        raise ViewGenerationError(f"{label} debe ser un objeto JSON: {path}")
    return value


def assert_inside(parent: Path, candidate: Path, label: str) -> Path:
    """Resuelve una ruta y verifica que permanezca debajo de ``parent``."""

    resolved_parent = parent.resolve(strict=True)
    resolved_candidate = candidate.resolve(strict=True)
    try:
        resolved_candidate.relative_to(resolved_parent)
    except ValueError as exc:
        raise ViewGenerationError(
            f"{label} debe resolver dentro de {resolved_parent}; se recibió {resolved_candidate}"
        ) from exc
    return resolved_candidate


def relative_to_project(path: Path) -> str:
    return path.resolve(strict=True).relative_to(PROJECT_ROOT.resolve(strict=True)).as_posix()


def parse_timestamp(value: Any, run_id: str) -> datetime:
    if not isinstance(value, str):
        raise ViewGenerationError(f"El manifiesto de {run_id} no declara generated_at.")
    try:
        result = datetime.fromisoformat(value)
    except ValueError as exc:
        raise ViewGenerationError(f"generated_at no es ISO-8601 en {run_id}: {value!r}") from exc
    if result.tzinfo is None:
        raise ViewGenerationError(f"generated_at debe incluir zona horaria en {run_id}: {value!r}")
    return result


def safe_run_id(value: Any) -> str:
    if not isinstance(value, str) or not re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9._-]*", value):
        raise ViewGenerationError(f"run_id no permitido en el catálogo: {value!r}")
    return value


def load_validated_run(run_id: str, required_status: str) -> AssessmentRun:
    run_id = safe_run_id(run_id)
    run_root = ASSESSMENTS_ROOT / run_id
    if not run_root.is_dir():
        raise ViewGenerationError(f"No existe el run configurado: {run_root}")
    run_root = assert_inside(ASSESSMENTS_ROOT, run_root, f"El run {run_id}")
    manifest = read_json(run_root / "manifest.json", f"el manifiesto de {run_id}")
    if manifest.get("run_id") != run_id:
        raise ViewGenerationError(
            f"El manifiesto en {run_root} declara run_id={manifest.get('run_id')!r}, no {run_id!r}."
        )
    if manifest.get("status") != required_status:
        raise ViewGenerationError(
            f"El run {run_id} debe tener status={required_status!r}; se encontró {manifest.get('status')!r}."
        )
    generated_at = parse_timestamp(manifest.get("generated_at"), run_id)

    index_path = run_root / "vars" / "index.json"
    try:
        index = json.loads(index_path.read_text(encoding="utf-8"))
    except FileNotFoundError as exc:
        raise ViewGenerationError(f"No existe el índice de variables de {run_id}: {index_path}") from exc
    except json.JSONDecodeError as exc:
        raise ViewGenerationError(f"El índice de variables de {run_id} no es JSON válido: {exc}") from exc
    if not isinstance(index, list):
        raise ViewGenerationError(f"El índice de variables de {run_id} debe ser una lista JSON.")

    variables: dict[str, Path] = {}
    for item in index:
        if not isinstance(item, dict):
            raise ViewGenerationError(f"El índice de {run_id} contiene una entrada no válida.")
        name = item.get("variable")
        folder = item.get("folder")
        if not isinstance(name, str) or not name.strip() or not isinstance(folder, str) or not folder:
            raise ViewGenerationError(f"El índice de {run_id} tiene una variable o carpeta inválida.")
        artifact_folder = assert_inside(PROJECT_ROOT, PROJECT_ROOT / Path(folder), f"La variable {name!r}")
        assert_inside(ASSESSMENTS_ROOT, artifact_folder, f"La variable {name!r}")
        if not artifact_folder.is_dir():
            raise ViewGenerationError(f"La carpeta de {name!r} no existe: {artifact_folder}")
        if name in variables:
            raise ViewGenerationError(f"El run {run_id} declara dos veces la variable {name!r}.")
        variables[name] = artifact_folder

    return AssessmentRun(run_id=run_id, generated_at=generated_at, root=run_root, variables=variables)


def require_string(mapping: dict[str, Any], key: str, context: str) -> str:
    value = mapping.get(key)
    if not isinstance(value, str) or not value.strip():
        raise ViewGenerationError(f"{context} debe declarar {key!r} como texto no vacío.")
    return value


def qvoice_entries(catalog: dict[str, Any], key: str, expected_count: int) -> dict[str, str]:
    qvoice = catalog.get("qvoice")
    if not isinstance(qvoice, dict):
        raise ViewGenerationError("El catálogo debe declarar el bloque qvoice.")
    entries = qvoice.get(key)
    if not isinstance(entries, list) or len(entries) != expected_count:
        raise ViewGenerationError(f"qvoice.{key} debe contener exactamente {expected_count} entradas.")
    result: dict[str, str] = {}
    for entry in entries:
        if not isinstance(entry, dict):
            raise ViewGenerationError(f"qvoice.{key} contiene una entrada no válida.")
        name = require_string(entry, "canonical_name", f"qvoice.{key}")
        stem = require_string(entry, "folder_stem", f"qvoice.{key}")
        if not re.fullmatch(r"qvoice-[a-z0-9-]+", stem):
            raise ViewGenerationError(f"El folder_stem de {name!r} debe usar el prefijo qvoice-: {stem!r}")
        if name in result or stem in result.values():
            raise ViewGenerationError(f"qvoice.{key} contiene nombres o carpetas duplicados.")
        result[name] = stem
    return result


def load_review_artifacts(
    catalog: dict[str, Any], required_status: str
) -> dict[str, list[SourceArtifact]]:
    configured_reviews = catalog.get("review_runs")
    if not isinstance(configured_reviews, list) or not configured_reviews:
        raise ViewGenerationError("El catálogo debe declarar al menos un review_run.")

    review_artifacts: dict[str, list[SourceArtifact]] = {}
    seen_run_ids: set[str] = set()
    for configured_review in configured_reviews:
        if not isinstance(configured_review, dict):
            raise ViewGenerationError("review_runs contiene una entrada no válida.")
        run_id = safe_run_id(configured_review.get("run_id"))
        if run_id in seen_run_ids:
            raise ViewGenerationError(f"El catálogo repite el review run {run_id!r}.")
        seen_run_ids.add(run_id)
        scope = require_string(configured_review, "scope", f"review_runs[{run_id}]")
        variable_names = configured_review.get("variables")
        if not isinstance(variable_names, list) or not variable_names:
            raise ViewGenerationError(f"review_runs[{run_id}] debe declarar sus variables.")
        run = load_validated_run(run_id, required_status)
        seen_variables: set[str] = set()
        for variable_name in variable_names:
            if not isinstance(variable_name, str) or not variable_name:
                raise ViewGenerationError(f"review_runs[{run_id}] contiene un nombre de variable inválido.")
            if variable_name in seen_variables:
                raise ViewGenerationError(
                    f"review_runs[{run_id}] declara más de una vez {variable_name!r}."
                )
            seen_variables.add(variable_name)
            try:
                folder = run.variables[variable_name]
            except KeyError as exc:
                raise ViewGenerationError(
                    f"El run {run_id} no contiene la variable configurada {variable_name!r}."
                ) from exc
            review_artifacts.setdefault(variable_name, []).append(
                SourceArtifact(run_id=run_id, generated_at=run.generated_at, folder=folder, scope=scope)
            )

    for artifacts in review_artifacts.values():
        artifacts.sort(key=lambda artifact: (artifact.generated_at, artifact.run_id))
    return review_artifacts


def build_view_model(catalog: dict[str, Any]) -> tuple[AssessmentRun, list[VariableView]]:
    if catalog.get("schema_version") != 1:
        raise ViewGenerationError("Se esperaba schema_version: 1 en el catálogo.")
    master = catalog.get("master_scorecard")
    policy = catalog.get("selection_policy")
    if not isinstance(master, dict) or not isinstance(policy, dict):
        raise ViewGenerationError("El catálogo debe declarar master_scorecard y selection_policy.")
    required_status = require_string(policy, "required_manifest_status", "selection_policy")
    if policy.get("visible_review") != "latest_validated_by_manifest_timestamp":
        raise ViewGenerationError("La política visible_review no coincide con el contrato esperado.")
    if policy.get("visible_review_is_not_score_replacement") is not True:
        raise ViewGenerationError("La política debe prohibir que la revisión visible sustituya el score.")
    if policy.get("score_recalculation") != "prohibited":
        raise ViewGenerationError("La vista no puede habilitar recálculo de scores.")

    baseline = load_validated_run(safe_run_id(master.get("run_id")), required_status)
    expected_variable_count = master.get("expected_variable_count")
    if not isinstance(expected_variable_count, int) or expected_variable_count != 43:
        raise ViewGenerationError("master_scorecard.expected_variable_count debe ser 43.")
    if len(baseline.variables) != expected_variable_count:
        raise ViewGenerationError(
            f"El scorecard maestro debe contener {expected_variable_count} variables; contiene {len(baseline.variables)}."
        )
    workbook_name = require_string(master, "workbook", "master_scorecard")
    workbook = assert_inside(ASSESSMENTS_ROOT, baseline.root / workbook_name, "El scorecard oficial")
    if not workbook.is_file():
        raise ViewGenerationError(f"No existe el workbook oficial: {workbook}")

    qvoice_base = qvoice_entries(catalog, "base_variables", expected_count=7)
    qvoice_extensions = qvoice_entries(catalog, "extensions", expected_count=12)
    missing_qvoice_base = set(qvoice_base).difference(baseline.variables)
    if missing_qvoice_base:
        raise ViewGenerationError(
            "Las variables QVoice base no están en el scorecard maestro: "
            + ", ".join(sorted(missing_qvoice_base))
        )
    overlap = set(qvoice_base).intersection(qvoice_extensions)
    if overlap:
        raise ViewGenerationError("Una extensión QVoice no puede ser parte del scorecard: " + ", ".join(sorted(overlap)))

    review_artifacts = load_review_artifacts(catalog, required_status)
    missing_extensions = set(qvoice_extensions).difference(review_artifacts)
    if missing_extensions:
        raise ViewGenerationError(
            "Las extensiones QVoice no tienen un run validado configurado: "
            + ", ".join(sorted(missing_extensions))
        )

    ordered_baseline = sorted(
        baseline.variables.items(),
        key=lambda item: int(item[1].name.split("-", 1)[0]),
    )
    records: list[VariableView] = []
    baseline_scope = "Evaluación base incluida en el scorecard oficial de 43 variables."
    for number, (name, folder) in enumerate(ordered_baseline, start=1):
        try:
            expected_number = int(folder.name.split("-", 1)[0])
        except ValueError as exc:
            raise ViewGenerationError(f"La carpeta base de {name!r} no inicia con su número: {folder.name!r}") from exc
        if expected_number != number:
            raise ViewGenerationError(
                f"El orden del scorecard no es consecutivo en {name!r}: se esperaba {number}, se encontró {expected_number}."
            )
        stem = qvoice_base.get(name)
        output_name = f"{number:02d}-{stem}" if stem else folder.name
        base_artifact = SourceArtifact(
            run_id=baseline.run_id,
            generated_at=baseline.generated_at,
            folder=folder,
            scope=baseline_scope,
        )
        records.append(
            VariableView(
                canonical_name=name,
                output_relative=Path("portfolio-43") / "variables" / output_name,
                base_artifact=base_artifact,
                review_artifacts=tuple(review_artifacts.get(name, [])),
                is_scorecard_variable=True,
            )
        )

    for number, (name, stem) in enumerate(qvoice_extensions.items(), start=1):
        artifacts = tuple(review_artifacts[name])
        base_artifact = artifacts[0]
        records.append(
            VariableView(
                canonical_name=name,
                output_relative=Path("qvoice-extensions") / f"{number:02d}-{stem}",
                base_artifact=base_artifact,
                review_artifacts=artifacts,
                is_scorecard_variable=False,
            )
        )

    if len([record for record in records if record.is_scorecard_variable]) != 43:
        raise ViewGenerationError("La vista debe contener exactamente las 43 variables del scorecard.")
    if len([record for record in records if not record.is_scorecard_variable]) != 12:
        raise ViewGenerationError("La vista debe contener exactamente 12 extensiones QVoice.")
    return baseline, records


def history_link_name(artifact: SourceArtifact) -> str:
    timestamp = artifact.generated_at.strftime("%Y-%m-%d-%H%M%S%z")
    return f"{timestamp}--{artifact.run_id}"


def directory_link_map(records: Iterable[VariableView]) -> dict[str, str]:
    links: dict[str, str] = {}
    for record in records:
        links[(record.output_relative / "base-assessment").as_posix()] = relative_to_project(
            record.base_artifact.folder
        )
        for artifact in record.history:
            history_path = record.output_relative / "history" / history_link_name(artifact)
            links[history_path.as_posix()] = relative_to_project(artifact.folder)
        if record.latest_review is not None:
            links[(record.output_relative / "latest-validated-review").as_posix()] = relative_to_project(
                record.latest_review.folder
            )
    if len(links) != len(set(links)):
        raise ViewGenerationError("La configuración produciría enlaces de directorio duplicados.")
    return dict(sorted(links.items()))


def generated_metadata(baseline: AssessmentRun, records: list[VariableView]) -> dict[str, Any]:
    master_workbook = baseline.root / "assessment.xlsx"
    return {
        "schema_version": 1,
        "generated_by": GENERATOR_ID,
        "master_scorecard": {
            "run_id": baseline.run_id,
            "source_workbook": relative_to_project(master_workbook),
            "view_workbook": "00-scorecard-43.xlsx",
            "official": True,
        },
        "portfolio_variable_count": 43,
        "qvoice_base_variable_count": 7,
        "qvoice_extension_count": 12,
        "directory_links": directory_link_map(records),
    }


def _is_reparse_point(path: Path) -> bool:
    try:
        attributes = os.lstat(path).st_file_attributes
    except (AttributeError, FileNotFoundError):
        return False
    return bool(attributes & stat.FILE_ATTRIBUTE_REPARSE_POINT)


def is_link_like(path: Path) -> bool:
    return path.is_symlink() or _is_reparse_point(path)


def remove_generated_path(path: Path) -> None:
    """Elimina sólo un árbol generado, sin atravesar reparse points/junctions."""

    if not path.exists() and not path.is_symlink():
        return
    if is_link_like(path):
        if path.is_dir():
            path.rmdir()
        else:
            path.unlink()
        return
    if path.is_dir():
        for child in path.iterdir():
            remove_generated_path(child)
        path.rmdir()
        return
    path.unlink()


def assert_removable_generated_view(path: Path) -> None:
    if not path.exists():
        return
    metadata_path = path / VIEW_MANIFEST_NAME
    metadata = read_json(metadata_path, "el manifiesto de la vista actual")
    if metadata.get("generated_by") != GENERATOR_ID:
        raise ViewGenerationError(
            f"Se rehúsa a reemplazar {path}: no fue generado por {GENERATOR_ID}."
        )


def create_directory_link(link: Path, target: Path) -> None:
    if link.exists() or link.is_symlink():
        raise ViewGenerationError(f"Ya existe el enlace de directorio: {link}")
    target = assert_inside(ASSESSMENTS_ROOT, target, f"El destino de {link.name}")
    if os.name == "nt":
        result = subprocess.run(
            ["cmd.exe", "/d", "/c", "mklink", "/J", str(link), str(target)],
            check=False,
            capture_output=True,
            text=True,
        )
        if result.returncode != 0:
            detail = (result.stderr or result.stdout).strip()
            raise ViewGenerationError(f"No se pudo crear la junction {link}: {detail}")
    else:
        os.symlink(target, link, target_is_directory=True)


def create_scorecard_hardlink(link: Path, target: Path) -> None:
    if link.exists() or link.is_symlink():
        raise ViewGenerationError(f"Ya existe el enlace al scorecard: {link}")
    target = assert_inside(ASSESSMENTS_ROOT, target, "El scorecard oficial")
    if not target.is_file():
        raise ViewGenerationError(f"No existe el scorecard oficial: {target}")
    try:
        os.link(target, link)
    except OSError as exc:
        raise ViewGenerationError(
            f"No se pudo crear el hard link del scorecard. No se permite crear copias: {exc}"
        ) from exc


def write_text(path: Path, contents: str) -> None:
    path.write_text(contents.rstrip() + "\n", encoding="utf-8", newline="\n")


def render_history_row(artifact: SourceArtifact) -> str:
    return (
        f"| {artifact.generated_at.isoformat()} | `{artifact.run_id}` | {artifact.scope} | "
        f"[abrir](history/{history_link_name(artifact)}/) |"
    )


def render_variable_readme(record: VariableView) -> str:
    history_rows = "\n".join(render_history_row(artifact) for artifact in record.history)
    lines = [f"# {record.canonical_name}", "", "## Estado de la vista", ""]
    if record.is_scorecard_variable:
        lines.extend(
            [
                "- Forma parte del scorecard oficial de 43 variables.",
                "- Análisis base: [base-assessment](base-assessment/).",
                f"- Run base: `{record.base_artifact.run_id}` ({record.base_artifact.generated_at.isoformat()}).",
            ]
        )
    else:
        lines.extend(
            [
                "- Extensión QVoice: no forma parte del scorecard oficial de 43 variables.",
                "- Análisis base validado: [base-assessment](base-assessment/).",
                f"- Run base: `{record.base_artifact.run_id}` ({record.base_artifact.generated_at.isoformat()}).",
            ]
        )
    lines.extend(["", "## Revisión validada más reciente", ""])
    if record.latest_review is None:
        lines.append("No existe una revisión validada posterior al análisis base visible en esta vista.")
    else:
        latest = record.latest_review
        lines.extend(
            [
                "- Enlace: [latest-validated-review](latest-validated-review/).",
                f"- Run: `{latest.run_id}`.",
                f"- Fecha de generación: {latest.generated_at.isoformat()}.",
                f"- Alcance: {latest.scope}",
                "- Esta revisión es evidencia más reciente con alcance específico; no reemplaza automáticamente el score oficial ni recalcula el portfolio.",
            ]
        )
    lines.extend(
        [
            "",
            "## Historial validado",
            "",
            "| Fecha de generación | Run | Alcance | Artefactos |",
            "| --- | --- | --- | --- |",
            history_rows,
        ]
    )
    return "\n".join(lines)


def render_current_readme() -> str:
    return """# Vista actual de resultados

Esta carpeta es una vista generada: sus enlaces apuntan a artefactos dentro de
`output/assessments/`, que permanecen como runs históricos inmutables.

- [00-scorecard-43.xlsx](00-scorecard-43.xlsx) es un hard link al workbook oficial del run maestro de 43 variables.
- `portfolio-43/variables/` contiene las 43 variables en el orden del scorecard. Las siete de QVoice llevan el prefijo `qvoice-`.
- `qvoice-extensions/` contiene doce variables adicionales de QVoice; no se incorporan ni consolidan con el scorecard oficial.

Una carpeta `latest-validated-review` muestra la revisión validada más reciente
por fecha de manifiesto. Es evidencia navegable con alcance explícito, no una
sustitución automática del score ni un portfolio recalculado.

El scorecard es un hard link, no una copia. Ábralo en modo solo lectura: una
edición directa modificaría también el archivo histórico de origen.

Para reconstruir esta vista, ejecute `tools/generate_current_output.py`. El
generador sólo puede reemplazar una vista marcada como generada y no modifica
los runs de `output/assessments/`.
"""


def render_portfolio_readme() -> str:
    return """# Portfolio oficial de 43 variables

Las carpetas de `variables/` siguen el orden del scorecard oficial. Cada una
separa su análisis base, cualquier revisión validada posterior y el historial
cronológico de runs. La presencia de una revisión parcial no cambia por sí sola
el score oficial ni produce un scorecard consolidado.

Las siete variables QVoice que ya pertenecen al portfolio se identifican por
su prefijo `qvoice-` y conservan el mismo nombre canónico dentro de su README.
"""


def render_extensions_readme() -> str:
    return """# Extensiones QVoice

Estas doce variables fueron evaluadas en runs QVoice validados, pero no forman
parte del scorecard oficial de 43 variables. Sus carpetas mantienen separados
el análisis base, una revisión posterior cuando existe y el historial completo.

Incorporarlas a un portfolio consolidado requeriría una revalidación v4 formal;
esta vista no combina ni recalcula scores.
"""


def write_json(path: Path, value: dict[str, Any]) -> None:
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2, sort_keys=True) + "\n", encoding="utf-8")


def build_staging_view(baseline: AssessmentRun, records: list[VariableView]) -> None:
    if STAGING_ROOT.exists():
        assert_removable_generated_view(STAGING_ROOT)
        remove_generated_path(STAGING_ROOT)
    STAGING_ROOT.mkdir(parents=False, exist_ok=False)
    write_json(STAGING_ROOT / VIEW_MANIFEST_NAME, {"generated_by": GENERATOR_ID, "state": "building"})

    create_scorecard_hardlink(STAGING_ROOT / "00-scorecard-43.xlsx", baseline.root / "assessment.xlsx")
    write_text(STAGING_ROOT / "README.md", render_current_readme())

    portfolio_root = STAGING_ROOT / "portfolio-43"
    portfolio_root.mkdir()
    write_text(portfolio_root / "README.md", render_portfolio_readme())
    (portfolio_root / "variables").mkdir()

    extensions_root = STAGING_ROOT / "qvoice-extensions"
    extensions_root.mkdir()
    write_text(extensions_root / "README.md", render_extensions_readme())

    for record in records:
        folder = STAGING_ROOT / record.output_relative
        folder.mkdir()
        write_text(folder / "README.md", render_variable_readme(record))
        create_directory_link(folder / "base-assessment", record.base_artifact.folder)
        history_root = folder / "history"
        history_root.mkdir()
        for artifact in record.history:
            create_directory_link(history_root / history_link_name(artifact), artifact.folder)
        if record.latest_review is not None:
            create_directory_link(folder / "latest-validated-review", record.latest_review.folder)

    write_json(STAGING_ROOT / VIEW_MANIFEST_NAME, generated_metadata(baseline, records))


def install_staging_view() -> None:
    if VIEW_ROOT.exists():
        assert_removable_generated_view(VIEW_ROOT)
        remove_generated_path(VIEW_ROOT)
    try:
        os.replace(STAGING_ROOT, VIEW_ROOT)
    except OSError as exc:
        raise ViewGenerationError(f"No se pudo instalar la vista generada en {VIEW_ROOT}: {exc}") from exc


def collect_directory_links(root: Path) -> set[str]:
    links: set[str] = set()
    for child in root.iterdir():
        if child.is_dir() and is_link_like(child):
            links.add(child.relative_to(root).as_posix())
            continue
        if child.is_dir():
            nested = collect_directory_links(child)
            links.update((child.relative_to(root) / item).as_posix() for item in nested)
    return links


def verify_directory_link(link: Path, target: Path) -> None:
    if not link.is_dir() or not is_link_like(link):
        raise ViewGenerationError(f"No existe una junction/enlace de directorio válida: {link}")
    target = assert_inside(ASSESSMENTS_ROOT, target, f"El destino esperado de {link}")
    try:
        same_target = os.path.samefile(link, target)
    except OSError as exc:
        raise ViewGenerationError(f"No se pudo resolver el enlace {link}: {exc}") from exc
    if not same_target:
        raise ViewGenerationError(f"El enlace {link} no resuelve al artefacto esperado {target}.")


def verify_view(baseline: AssessmentRun, records: list[VariableView]) -> None:
    if not VIEW_ROOT.is_dir():
        raise ViewGenerationError(f"No existe la vista generada: {VIEW_ROOT}")
    actual_metadata = read_json(VIEW_ROOT / VIEW_MANIFEST_NAME, "el manifiesto de la vista generada")
    expected_metadata = generated_metadata(baseline, records)
    if actual_metadata != expected_metadata:
        raise ViewGenerationError(
            "El manifiesto de la vista no coincide con el catálogo y los manifiestos fuente esperados."
        )

    scorecard_link = VIEW_ROOT / "00-scorecard-43.xlsx"
    scorecard_source = baseline.root / "assessment.xlsx"
    if not scorecard_link.is_file():
        raise ViewGenerationError("No existe el enlace al scorecard oficial.")
    try:
        if not os.path.samefile(scorecard_link, scorecard_source):
            raise ViewGenerationError("00-scorecard-43.xlsx no es un enlace al scorecard oficial.")
    except OSError as exc:
        raise ViewGenerationError(f"No se pudo verificar el enlace al scorecard oficial: {exc}") from exc
    workbook = load_workbook(scorecard_link, read_only=True, data_only=False)
    try:
        if workbook.sheetnames != EXPECTED_WORKSHEET_NAMES:
            raise ViewGenerationError(
                "El scorecard enlazado debe conservar exactamente las hojas "
                + ", ".join(repr(name) for name in EXPECTED_WORKSHEET_NAMES)
                + f"; se encontraron {workbook.sheetnames!r}."
            )
    finally:
        workbook.close()

    expected_links = expected_metadata["directory_links"]
    actual_links = collect_directory_links(VIEW_ROOT)
    if actual_links != set(expected_links):
        missing = sorted(set(expected_links).difference(actual_links))
        stale = sorted(actual_links.difference(expected_links))
        details = []
        if missing:
            details.append("faltantes: " + ", ".join(missing))
        if stale:
            details.append("obsoletos: " + ", ".join(stale))
        raise ViewGenerationError("Los enlaces de directorio no coinciden (" + "; ".join(details) + ").")
    for relative_link, relative_target in expected_links.items():
        verify_directory_link(VIEW_ROOT / relative_link, PROJECT_ROOT / Path(relative_target))

    portfolio_variables = [record for record in records if record.is_scorecard_variable]
    extension_variables = [record for record in records if not record.is_scorecard_variable]
    actual_portfolio_folders = {path.name for path in (VIEW_ROOT / "portfolio-43" / "variables").iterdir() if path.is_dir()}
    expected_portfolio_folders = {record.output_relative.name for record in portfolio_variables}
    if actual_portfolio_folders != expected_portfolio_folders:
        raise ViewGenerationError("Las 43 carpetas principales no coinciden con el scorecard maestro.")
    actual_extension_folders = {path.name for path in (VIEW_ROOT / "qvoice-extensions").iterdir() if path.is_dir()}
    expected_extension_folders = {record.output_relative.name for record in extension_variables}
    if actual_extension_folders != expected_extension_folders:
        raise ViewGenerationError("Las 12 extensiones QVoice no coinciden con el catálogo.")

    qvoice_portfolio = [record for record in portfolio_variables if "-qvoice-" in record.output_relative.name]
    if len(qvoice_portfolio) != 7:
        raise ViewGenerationError("La vista no identifica exactamente las siete variables QVoice del scorecard.")
    for record in records:
        if not (VIEW_ROOT / record.output_relative / "README.md").is_file():
            raise ViewGenerationError(f"Falta el README de la variable {record.canonical_name!r}.")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Genera o verifica la vista output/current de Attrition.")
    parser.add_argument(
        "--catalog",
        type=Path,
        default=DEFAULT_CATALOG,
        help="Ruta al catálogo YAML (por defecto: config/current-output-catalog.yaml).",
    )
    parser.add_argument(
        "--verify",
        action="store_true",
        help="Verifica una vista existente sin regenerarla.",
    )
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    catalog_path = args.catalog.resolve()
    try:
        catalog = read_yaml(catalog_path)
        baseline, records = build_view_model(catalog)
        if args.verify:
            verify_view(baseline, records)
            print("Verificación correcta: 43 variables, 7 QVoice integradas, 12 extensiones y scorecard enlazado.")
            return 0
        build_staging_view(baseline, records)
        install_staging_view()
        verify_view(baseline, records)
        print("Vista generada y validada: output/current (43 variables + 12 extensiones QVoice).")
        return 0
    except ViewGenerationError as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())

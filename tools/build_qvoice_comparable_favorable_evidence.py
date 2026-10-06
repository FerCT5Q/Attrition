"""Freeze aggregate-only evidence for the comparable QVoice favorable trend.

The raw workbooks are processed in memory.  Participant keys are used only to
check one row per participant-wave and are never written to the JSON artifact.
The artifact contains aggregate counts, percentages, and item-header metadata
only; it never contains identifiers, response rows, or individual scores.
"""

from __future__ import annotations

import argparse
import json
from collections import Counter
from datetime import date
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from qvoice_excel_assessment import (
    build_columns,
    coerce_number,
    header_base,
    load_survey_structure,
)


ROOT = Path(__file__).resolve().parents[1]
QVOICE_ROOT = ROOT / "data" / "QVoice"
OUTPUT = ROOT / "output" / "distribution" / "json" / "qvoice-comparable-favorable-evidence-20260909.json"
MINIMUM_COHORT = 30
ITEM_COUNT = 33

# The source files use different technical decorations around the same flag
# headers (for example, `_favorable` versus `._Favorable`).  `header_base` is
# the existing deterministic QVoice header contract: it removes only those
# technical decorations, normalizes case/whitespace, and preserves the item
# wording.  It does not apply synonyms, token similarity, semantic mappings,
# or internal punctuation substitutions (for example, well-being != wellbeing).
STRICT_HEADER_RULE = (
    "Coincidencia exacta de header_base tras retirar sólo el prefijo técnico de "
    "pregunta, el sufijo técnico binario y separadores terminales; conserva las "
    "palabras y la puntuación interna. No usa similitud, sinónimos ni mapeo semántico."
)

WAVES = (
    {
        "year": 2024,
        "file": "Jan 2024 Qvoice raw data.xlsx",
        "sheet": "Raw data",
        "key_column": "Participant's Unique Identifier",
    },
    {
        "year": 2025,
        "file": "Jan 2025_Qvoice_raw data.xlsx",
        "sheet": "Raw data",
        "key_column": "EmpID",
    },
    {
        "year": 2026,
        "file": "Jan 2026 Qvoice_raw data.xlsx",
        "sheet": "Raw data",
        "key_column": "Emp ID",
    },
)

# This is deliberately a fixed whitelist rather than an inferred similarity
# match.  Any questionnaire wording change changes the strict intersection and
# stops the frozen build for review.
EXPECTED_COMMON_BASES = (
    "at qualcomm, i can challenge the traditional ways of doing things",
    "i am able to balance my work and personal life to my satisfaction",
    "i am comfortable voicing my ideas and opinions, even if they are different from others",
    "i am treated with respect at work",
    "i can freely raise ethical concerns without negative consequences",
    "i can see a clear link between my work and my business function's objectives",
    "i feel comfortable being myself at work",
    "i feel encouraged to come up with new and better ways of doing things",
    "i feel energized when i am working",
    "i have a sense of belonging at qualcomm",
    "i have confidence in qualcomm's senior leadership team to make the right decisions for this company",
    "i have confidence in the future of qualcomm",
    "i have enough energy after work to enjoy other activities (e.g. socializing, hobbies, exercise)",
    "i rarely think about looking for a new job with another company",
    "i receive meaningful recognition for my contributions",
    "i see myself working at qualcomm in a few years",
    "i understand how my rewards are determined (merit increase, bonus, stock)",
    "i understand what it takes to advance in my career at qualcomm",
    "i would recommend qualcomm as a great place to work",
    "my job makes good use of my skills and abilities",
    "my manager gives me timely feedback throughout the year that helps me improve my performance",
    "my manager helps remove barriers to getting my work done",
    "my manager supports my efforts to balance my work and personal life",
    "my team looks for ways to change processes to improve productivity",
    "my work gives me a sense of personal accomplishment",
    "qualcomm does a good job of communicating the reasons behind important changes",
    "qualcomm is making the changes necessary to compete effectively",
    "qualcomm motivates me to go above and beyond in my work",
    "the amount of stress in my job is manageable",
    "the people i work with cooperate to get the job done",
    "the work processes at qualcomm allow employees to be as productive as possible",
    "there is effective collaboration between teams at qualcomm",
    "there is open and honest communication at qualcomm",
)

EXPECTED_WAVES = {
    2024: {"participants": 36_806, "complete": 27_563, "favorable_percent": 77.3663},
    2025: {"participants": 37_837, "complete": 30_342, "favorable_percent": 78.2771},
    2026: {"participants": 42_203, "complete": 33_067, "favorable_percent": 79.3780},
}
EXPECTED_BINARY_ITEM_COUNTS = {2024: 52, 2025: 52, 2026: 53}

BAND_LABELS = ("<50", "50–69", "70–79", "80–89", "90–99", "100")


def _headers_and_binary_bases(config: dict[str, object], survey_structure: dict[str, list[dict[str, str]]]) -> tuple[list[str], dict[str, int], set[str]]:
    """Return a unique deterministic binary-header base mapping for one wave."""
    path = QVOICE_ROOT / str(config["file"])
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook[str(config["sheet"])]
        headers = [str(value or "").replace("\n", " ").strip() for value in next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True))]
        columns = build_columns(headers, survey_structure)
    finally:
        workbook.close()

    binary_indices = list(columns["binary_items"])
    if len(binary_indices) != EXPECTED_BINARY_ITEM_COUNTS[int(config["year"])]:
        raise ValueError(f"El conteo binario de {config['year']} cambió: {len(binary_indices)}.")

    by_base: dict[str, int] = {}
    for index in binary_indices:
        base = header_base(headers[index])
        if base in by_base:
            raise ValueError(f"El header canónico binario se repite en {config['year']}: {base!r}")
        by_base[base] = index
    return headers, by_base, set(by_base)


def _cell_status(value: Any) -> tuple[int | None, str]:
    """Return a binary value or a non-favorable missing/invalid status.

    The documented 255 sentinel is missing, never a zero/unfavorable response.
    """
    if value is None or (isinstance(value, str) and not value.strip()):
        return None, "missing"
    number = coerce_number(value)
    if number is None:
        return None, "invalid"
    if number == 255:
        return None, "sentinel_255"
    if number in (0, 1):
        return int(number), "valid"
    return None, "invalid"


def _band_for_favorable_count(favorable_count: int) -> str:
    if favorable_count < 17:
        return "<50"
    if favorable_count < 24:
        return "50–69"
    if favorable_count < 27:
        return "70–79"
    if favorable_count < 30:
        return "80–89"
    if favorable_count < ITEM_COUNT:
        return "90–99"
    if favorable_count == ITEM_COUNT:
        return "100"
    raise ValueError(f"Conteo favorable fuera de rango: {favorable_count}")


def _scan_wave(config: dict[str, object], headers: list[str], column_by_base: dict[str, int], common_bases: tuple[str, ...]) -> dict[str, object]:
    """Read one source workbook and return privacy-safe aggregate evidence."""
    path = QVOICE_ROOT / str(config["file"])
    key_column = str(config["key_column"])
    try:
        key_index = headers.index(key_column)
    except ValueError as exc:
        raise ValueError(f"No se encontró la llave declarada para {config['year']}.") from exc
    response_index = headers.index("Response Type") if "Response Type" in headers else None
    selected_indices = [column_by_base[base] for base in common_bases]

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook[str(config["sheet"])]
        seen_signatures: dict[str, tuple[tuple[int | None, str], ...]] = {}
        row_count = 0
        rows_without_key = 0
        duplicate_rows = 0
        duplicate_conflicts = 0
        response_types: Counter[str] = Counter()
        missing_cells = 0
        sentinel_255_cells = 0
        invalid_cells = 0
        complete = 0
        favorable_total = 0
        bands: Counter[str] = Counter()
        incomplete_rows = 0

        for row in worksheet.iter_rows(min_row=2, values_only=True):
            row_count += 1
            if response_index is not None:
                response_types[str(row[response_index] if row[response_index] is not None else "NULL")] += 1
            raw_key = row[key_index]
            key = str(raw_key).strip() if raw_key is not None else ""
            if not key:
                rows_without_key += 1
                continue

            signature = tuple(_cell_status(row[index]) for index in selected_indices)
            previous = seen_signatures.get(key)
            if previous is not None:
                duplicate_rows += 1
                if previous != signature:
                    duplicate_conflicts += 1
                continue
            seen_signatures[key] = signature

            statuses = [status for _, status in signature]
            missing_cells += sum(status == "missing" for status in statuses)
            sentinel_255_cells += sum(status == "sentinel_255" for status in statuses)
            invalid_cells += sum(status == "invalid" for status in statuses)
            values = [value for value, status in signature if status == "valid"]
            if len(values) != ITEM_COUNT or len(statuses) != ITEM_COUNT:
                incomplete_rows += 1
                continue

            favorable_count = sum(values)
            complete += 1
            favorable_total += favorable_count
            bands[_band_for_favorable_count(favorable_count)] += 1
    finally:
        workbook.close()

    participants = len(seen_signatures)
    if participants + duplicate_rows + rows_without_key != row_count:
        raise ValueError(f"La reconciliación de filas no cuadra para {config['year']}.")
    if complete + incomplete_rows != participants:
        raise ValueError(f"La reconciliación de completitud no cuadra para {config['year']}.")
    if sum(bands.values()) != complete:
        raise ValueError(f"Las bandas no suman los participantes completos para {config['year']}.")

    participant_scores_percent = (favorable_total / (complete * ITEM_COUNT) * 100) if complete else None
    band_rows = [
        {
            "band": label,
            "participant_count": int(bands[label]),
            "percent_of_complete": round(bands[label] / complete * 100, 4) if complete else None,
        }
        for label in BAND_LABELS
    ]
    return {
        "year": int(config["year"]),
        "source_file": str(config["file"]),
        "sheet": str(config["sheet"]),
        "participant_rows": participants,
        "source_rows_scanned": row_count,
        "rows_without_participant_key": rows_without_key,
        "duplicate_participant_rows": duplicate_rows,
        "duplicate_conflicting_rows": duplicate_conflicts,
        "response_type_summary": dict(sorted(response_types.items())),
        "complete_comparable_participants": complete,
        "complete_comparable_percent": round(complete / participants * 100, 4) if participants else None,
        "incomplete_comparable_participants": incomplete_rows,
        "mean_favorable_percent": round(participant_scores_percent, 4) if participant_scores_percent is not None else None,
        "valid_binary_cells": complete * ITEM_COUNT,
        "missing_cells": missing_cells,
        "sentinel_255_cells_treated_as_missing": sentinel_255_cells,
        "out_of_domain_non_sentinel_cells": invalid_cells,
        "score_bands": band_rows,
        "band_count_total": sum(item["participant_count"] for item in band_rows),
    }


def _validate_against_frozen_contract(waves: list[dict[str, object]], comparable_bases: tuple[str, ...]) -> None:
    if comparable_bases != EXPECTED_COMMON_BASES:
        raise ValueError("Los 33 headers comparables no corresponden al contrato estricto congelado.")
    if len(comparable_bases) != ITEM_COUNT:
        raise ValueError("El contrato requiere exactamente 33 items comparables.")
    for wave in waves:
        year = int(wave["year"])
        expected = EXPECTED_WAVES[year]
        if int(wave["participant_rows"]) != expected["participants"]:
            raise ValueError(f"La población de participantes de {year} cambió.")
        if int(wave["complete_comparable_participants"]) != expected["complete"]:
            raise ValueError(f"El conteo 33/33 de {year} cambió.")
        if abs(float(wave["mean_favorable_percent"]) - expected["favorable_percent"]) > 0.00005:
            raise ValueError(f"El promedio favorable de {year} cambió.")
        if int(wave["duplicate_participant_rows"]) != 0 or int(wave["duplicate_conflicting_rows"]) != 0:
            raise ValueError(f"Se detectaron duplicados o conflictos en {year}.")
        if int(wave["out_of_domain_non_sentinel_cells"]) != 0:
            raise ValueError(f"Se detectaron valores binarios inválidos en {year}.")
        if int(wave["band_count_total"]) != int(wave["complete_comparable_participants"]):
            raise ValueError(f"Las bandas no reconcilian en {year}.")
        if int(wave["complete_comparable_participants"]) < MINIMUM_COHORT:
            raise ValueError(f"La cohorte de {year} no alcanza n >= {MINIMUM_COHORT}.")


def build_evidence(output_path: Path = OUTPUT, *, overwrite: bool = False) -> dict[str, object]:
    """Create one immutable, aggregate-only evidence file for the slide."""
    if output_path.exists() and not overwrite:
        raise FileExistsError(f"La evidencia congelada ya existe: {output_path}")

    survey_structure = load_survey_structure(QVOICE_ROOT / "survey_structure.xlsx")
    metadata: list[tuple[dict[str, object], list[str], dict[str, int], set[str]]] = []
    for config in WAVES:
        headers, by_base, bases = _headers_and_binary_bases(config, survey_structure)
        metadata.append((config, headers, by_base, bases))

    common_bases = tuple(sorted(set.intersection(*(bases for _, _, _, bases in metadata))))
    if common_bases != EXPECTED_COMMON_BASES:
        raise ValueError("La intersección estricta de headers cambió; no se aplicará un mapeo fuzzy.")

    waves = [
        _scan_wave(config, headers, by_base, common_bases)
        for config, headers, by_base, _ in metadata
    ]
    _validate_against_frozen_contract(waves, common_bases)

    result = {
        "artifact": "internal_analysis",
        "evidence_type": "aggregate-only comparable QVoice favorable average",
        "generated_on": date.today().isoformat(),
        "outcome": "pass",
        "minimum_shared_cohort": MINIMUM_COHORT,
        "source_rule": "QVoice raw Excel workbooks only; Redshift is not used because it contains only the 2026 wave.",
        "scope": "Respondentes con 33/33 ítems comparables válidos; no representa a toda la plantilla.",
        "grain": "participant-wave derived aggregate",
        "identifier_or_row_data_persisted": False,
        "formula": {
            "item_coding": {"0": "desfavorable", "1": "favorable", "255": "faltante"},
            "individual_score": "100 × favorable_items / 33",
            "annual_result": "Promedio aritmético de los puntajes individuales incluidos",
            "inclusion_rule": "Exactamente 33 respuestas válidas 0/1 en los 33 ítems comunes.",
            "no_official_completion_flag": "Response Type=EX en 2026 no se interpreta como completitud; 2024–2025 no aportan un flag usable.",
        },
        "strict_item_contract": {
            "binary_items_by_wave": {str(year): count for year, count in EXPECTED_BINARY_ITEM_COUNTS.items()},
            "strict_common_item_count": ITEM_COUNT,
            "header_matching_rule": STRICT_HEADER_RULE,
            "common_header_bases": list(common_bases),
        },
        "waves": waves,
        "acceptance_rule": {
            "expected_years": [2024, 2025, 2026],
            "strict_common_item_count": ITEM_COUNT,
            "participant_and_complete_counts": "Deben coincidir con el contrato congelado por wave.",
            "mean_favorable_percent_tolerance": 0.00005,
            "zero_duplicate_participant_rows": True,
            "zero_duplicate_conflicting_rows": True,
            "zero_out_of_domain_non_sentinel_cells": True,
            "band_counts_equal_complete_participants": True,
            "minimum_published_cohort": MINIMUM_COHORT,
            "no_identifier_or_row_persistence": True,
        },
    }
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    return result


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Genera evidencia agregada comparable de QVoice.")
    parser.add_argument("--output", type=Path, default=OUTPUT)
    parser.add_argument("--overwrite", action="store_true")
    args = parser.parse_args()
    evidence = build_evidence(args.output, overwrite=args.overwrite)
    print(json.dumps({"output": str(args.output), "waves": evidence["waves"], "outcome": evidence["outcome"]}, ensure_ascii=False))

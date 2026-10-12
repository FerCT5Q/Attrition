"""Suplemento de documentación desde código congelado y artefactos agregados.

No consulta fuentes individuales, reentrena ni cambia scores. El workbook es un
diccionario de modelado, separado de los workbooks de Variable assessment.
"""
from __future__ import annotations

import argparse
from collections import Counter, defaultdict
import csv
from datetime import datetime, timezone
import json
import math
import os
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import engineering_b_incremental as inc
import engineering_model_c as base
import b_playbook_variable_catalog as definitions
import b_playbook_variable_decisions as decisions_source
import b_playbook_documentation_report as report_source

ROOT=base.ROOT
POPULATION='Empleados regulares activos de Engineering, excluyendo RF360, en cada foto mensual.'
CODING=('Todas las entradas son categorías con indicadores 0/1. La referencia es la categoría más frecuente '
        'por filas de entrenamiento y se omite. Cada categoría restante tiene su propio indicador. '
        'Una categoría no aprendida deja todos los indicadores de ese término en cero. '
        'Las bandas no se usan como números ordinales ni se estandarizan como variables continuas.')
POOLING=('Sólo incorporaciones RH: conserva categorías con al menos 300 personas distintas en todo el '
         'entrenamiento; el resto, incluidas categorías nuevas o desconocidos infrecuentes, se transforma '
         'en (other). B y QVoice no usan este pooling. Una persona puede contar en más de una categoría '
         'si cambió durante el entrenamiento; no es un mínimo por foto.')
GUARDS=('Ajuste válido y elegibilidad; Brier no empeora; PR-AUC ≥99% del control. En validación de seis '
        'meses, además: lift al 10% ≥99% del control, MAE de los mismos grupos ≤control+0.1 pp y '
        'captura ≥control−0.5 pp. Entre las opciones que cumplen se minimiza log-loss y se aplica '
        'parsimonia dentro de una tolerancia relativa exacta de 0.0005. Se escoge la opción con menos '
        'incorporaciones dentro de esa ventana; después log-loss e identificador.')
SEARCH=('B conserva ocho predictores y dos interacciones. Se prueban todas las incorporaciones '
        'individuales; pares y tríos sólo entre las seis primeras de la reserva interna. RH ordena '
        'las individuales ajustadas y elegibles; QVoice ordena las ajustadas, y filtra disponibilidad '
        'y condiciones de validación al elegir el resultado. No se crean interacciones nuevas entre candidatas. '
        'Las tres penalizaciones de B y QVoice son 0.1, 1 y 10; RH usa la elegida para B en cada corte. '
        'Para cada conjunto QVoice se elige primero la penalización con menor log-loss interno; '
        'la comparación de disponibilidad y las condiciones se aplican a ese ajuste representativo del conjunto.')


def require(condition, message):
    if not condition:
        raise RuntimeError(message)


def read(path):
    return json.loads(Path(path).read_text(encoding='utf-8-sig'))


def write_json(path, value):
    Path(path).write_text(json.dumps(value,ensure_ascii=False,indent=2,allow_nan=False)+'\n',encoding='utf-8')


def plain(value):
    if value is None:
        return ''
    if isinstance(value,bool):
        return 'Sí' if value else 'No'
    if isinstance(value,(list,tuple)):
        return '; '.join(plain(v) for v in value)
    if isinstance(value,dict):
        return json.dumps(value,ensure_ascii=False,allow_nan=False)
    return value


def fmt(value):
    if value is None:
        return '—'
    if isinstance(value,bool):
        return plain(value)
    if isinstance(value,(int,float)):
        if 0<abs(value)<0.05:
            return ('−' if value<0 else '')+'<0.1'
        rounded=round(value,1)
        return str(int(rounded)) if float(rounded).is_integer() else f'{rounded:.1f}'
    return str(plain(value)).replace('|','\\|').replace('\n',' ')


def link(path, folder, label=None):
    absolute=Path(path) if Path(path).is_absolute() else ROOT/path
    require(absolute.is_file(),'Referencia ausente: '+str(absolute))
    relative=Path(os.path.relpath(absolute,folder)).as_posix()
    return f'[{label or absolute.name}](<{relative}>)'


def markdown_table(headers, rows):
    return ['| '+' | '.join(headers)+' |','|'+'|'.join('---' for _ in headers)+'|']+[
        '| '+' | '.join(fmt(v) for v in row)+' |' for row in rows]


def csv_file(path, rows):
    require(bool(rows),'No se puede exportar una tabla sin filas.')
    columns=list(rows[0])
    with Path(path).open('w',encoding='utf-8-sig',newline='') as handle:
        writer=csv.DictWriter(handle,fieldnames=columns)
        writer.writeheader()
        for row in rows:
            require(set(row)==set(columns),'Columnas inconsistentes en '+str(path))
            writer.writerow({k:plain(v) for k,v in row.items()})


def evidence_mapping(config):
    mapping=defaultdict(list)
    def add(names,features):
        for feature in features:
            mapping[feature]+=names
    contracts=read(ROOT/config['input_contracts'])['contracts']
    for name,contract in contracts.items():
        add([name],contract['features'])
    for feature in ['region','comp_ratio_band','grade_family','acquisition_recorded','transfer_12m']:
        add([feature],[feature])
    add(['transfer_history'],['transfer_12m'])
    for name,features in {
        'history-manager_recency':['manager_changes_3m','manager_changes_6m'],
        'history-team_recency':['team_exits_3m','team_exits_6m','team_exit_recency_12m'],
        'history-manager_span':['manager_team_size_band'],
        'history-team_intensity':['team_exit_intensity_12m'],
        'v11-city':['work_city'], 'v11-labor-market':['labor_market_detail'],
        'v11-arrangement-v2':['work_arrangement_v11'], 'v11-acquisition':['acquisition_detail'],
        'v11-grade-entry':['current_grade_entry_path','time_since_grade_entry_band','pay_track_detail'],
        'v11-hierarchy-v3':['hierarchy_depth_band'],
        'v11-hierarchy-synthetic-v3':['hierarchy_depth_band'],
        'v11-manager6-v2':['manager_observed_expected6_band','manager_team_exits6_v11'],
        'v11-manager18-v3':['manager_above_average18'],
        'v11-manager-feed':['manager_observed_expected6_band','manager_team_exits6_v11','manager_above_average18'],
        'v11-manager-gate-synthetic':['manager_observed_expected6_band','manager_team_exits6_v11','manager_above_average18'],
        'v11-timezone-v3':['manager_colocation_v11','manager_timezone_relation'],
        'v11-timezone-dst':['manager_timezone_relation'],
        'model-v11-timezone-transform':['manager_colocation_v11','manager_timezone_relation'],
        'v11-commute-inputs':['commute_distance_onsite_v11'],
        'v11-commute-mapping':['commute_distance_onsite_v11'],
        'model-v11-commute-transform':['commute_distance_onsite_v11'],
        'v11-stored-role':['time_in_role_stored_v11'],
    }.items():
        add([name],features)
    add(['team-pay-independent','team-pay-distinct'],['within_team_pay_position','manager_pay_spread_band'])
    add(['qvoice-regional'],['qv_mean33','qv_availability','qv_region_relative'])
    comments=[f for f in config['qvoice']['extra_features'] if f.startswith('qv_comment_') or
              f in {'qv_improvement_comment_presence','qv_remote_hybrid_request','qv_onsite_policy_concern'}]
    add(['v11-qvoice-comments','v11-qvoice-comment-keys','model-qvoice-comments'],comments)
    return mapping


def prepare_catalog(config,out,catalog,scope,decisions,encodings):
    hr=set(inc.hr_candidates(config))
    q=set(inc.q_candidates(config))
    support=read(out/'input-evidence/hr-feature-support.json')['rows']+read(out/'input-evidence/qvoice-feature-support.json')['rows']
    evidence=evidence_mapping(config)
    for feature,row in catalog.items():
        family='RH' if feature in hr else 'QVoice' if feature in q else 'B'
        row['family']=family
        row['population']=POPULATION
        row['grain']='Una representación por persona-foto; sólo agregados en estos entregables.'
        row['model_encoding']=CODING
        row['pooling_rule']=POOLING if family=='RH' else 'No se aplica agrupación de categorías raras a este término.'
        group_rules={
            'tenure_playbook_band':'Define grupos con las mismas seis bandas, traducidas a español; desconocido → sin dato.',
            'team_exits_12m':'Para agrupar: 0 → sin salidas; (unknown) → sin dato; 1/2+ → con salidas. El modelo conserva 0/1/2+ por separado.',
            'level_grouped':'Para agrupar toma el prefijo anterior a « - »: 1/2/3 → alto; 5/6 → medio; 7/8 → entrada; otros → otro. El modelo conserva sus etiquetas detalladas.',
        }
        row['group_role']=group_rules.get(feature,
            'Puede modificar la predicción, pero no añade una dimensión de grupo en esta ronda. Los grupos siguen definidos por antigüedad, salidas del equipo y nivel.')
        if feature in group_rules:
            row['sources']+=['tools/engineering_playbook_group_selection.py','tools/engineering_c_group_selection.py']
        row['assessment_score_changed']=False
        row['decisions']=[r for r in decisions if r['feature']==feature]
        history=[]
        for anchor in sorted({r['anchor'] for r in row['decisions']}):
            cuts=[r for r in row['decisions'] if r['anchor']==anchor and r['period']!='final']
            if cuts:
                history.append(f"{anchor}: seleccionada en {sum(r['selected'] for r in cuts)} de {len(cuts)} cortes")
        row['selection_history_summary']='; '.join(history)
        row['encodings']=[r for r in encodings if r['feature']==feature]
        row['annual_support']=[r for r in support if r['feature']==feature]
        views=[v for v in scope['views'] if feature in v['representaciones']]
        if feature.startswith('qv_M'):
            views+=[v for v in scope['views'] if 'qv_mean33' in v['representaciones'] and v not in views]
        if feature in config['interactions']:
            views=[v for v in scope['views'] if set(config['interactions'][feature]) & set(v['representaciones'])]
        row['playbook_views']=[{'view_id':v['source_stage_id'],'label':v['vista'],'slide':v['slide']} for v in views]
        first=[r['first_observed_snapshot'] for r in row['annual_support'] if r.get('first_observed_snapshot')]
        row['first_observed_snapshot']=min(first) if first else None
        final=[r for r in row['decisions'] if r['period']=='final']
        require(bool(final),'Falta decisión final de '+feature)
        row['final_selected_anchors']=[r['anchor'] for r in final if r['selected']]
        row['final_summary']=' '.join(f"Sobre {r['anchor']}: {r['decision']}. {r['reason']}" for r in final)
        row['history_flow']=(
            f"Primera foto con valor observable en esta carga: {row['first_observed_snapshot']}. "
            'El soporte se comprueba por año y la elegibilidad se decide sólo con el entrenamiento de cada corte. '
            if first else
            'Término obligatorio heredado de B; se conserva su definición y se reajustan sus coeficientes con el pasado de cada corte. '
            'Esta ronda no guarda una nueva tabla de cobertura individual para los términos B. ')
        row['history_flow']+=row['known_at']
        row['evidence_files']=[(out/'input-evidence'/f'{name}.json').relative_to(ROOT).as_posix() for name in evidence[feature]]
        if family!='B':
            name='hr-feature-support' if family=='RH' else 'qvoice-feature-support'
            row['evidence_files'].append((out/'input-evidence'/f'{name}.json').relative_to(ROOT).as_posix())
        for source in row['sources']+row['evidence_files']:
            require((ROOT/source).is_file(),'No existe la fuente '+source)
        if family=='B':
            row['action']='Conservar el término de B; no se realizó una prueba para eliminarlo.'
        elif row['final_selected_anchors']:
            row['action']='Conservar como candidata experimental y confirmar en datos posteriores; no sustituye automáticamente el modelo vigente.'
        elif any(not r.get('individual_eligible',True) for r in final):
            row['action']='Resolver la restricción indicada en las decisiones antes de interpretar su ausencia como falta de señal.'
        else:
            row['action']='Mantener el registro de la prueba; no concluir que carece de señal en toda representación o combinación posible.'
    return catalog


def dictionary_rows(catalog):
    return [{
        'Variable':r['feature'],'Nombre':r['label'],'Familia':r['family'],'Tipo de entrada':r['input_type'],
        'Campos fuente':r['source_fields'],'Valores de origen':r['source_values'],
        'Valores antes de agrupar categorías raras':r['values_before_pooling'],'Transformación exacta':r['transformation'],
        'Faltantes y defaults':r['missing_and_default'],'Fecha de conocimiento':r['known_at'],
        'Codificación del modelo':r['model_encoding'],'Categorías poco frecuentes':r['pooling_rule'],
        'Consideraciones':r['considerations'],'Primera foto observable':r['first_observed_snapshot'],
        'Uso en grupos':r['group_role'],
        'Selección para periodos externos':r['selection_history_summary'],
        'Selección final':r['final_summary'],'Acción':r['action'],
        'Vistas V11':[v['label'] for v in r['playbook_views']],
        'Fuentes':r['sources'],'Evidencia':r['evidence_files'],
        'Ficha':'fichas/'+r['feature']+'.md',
    } for r in catalog.values()]


def decision_rows(decisions):
    return [{
        'Variable':r['feature'],'Nombre':r['label'],'Familia':r['family'],'Base':r['anchor'],
        'Periodo':r['period'],'Fase':r['phase'],'Corte':r['cutoff'],'Modo interno':r['selection_mode'],
        'Decisión':r['decision'],'Motivo de decisión':r['reason'],
        'Resultado individual':r.get('individual_reason'),'Ajuste individual':r.get('individual_fit_status'),
        'L2 del ajuste individual':r.get('individual_l2'),
        'Elegible individual':r.get('individual_eligible'),'Cumple condiciones individual':r.get('individual_passes_guards'),
        'Condiciones incumplidas':r.get('individual_failed_guards'),
        'Fotos observables entrenamiento':r.get('training_support',{}).get('observable_training_snapshots'),
        'Categorías de entrenamiento':r.get('training_support',{}).get('training_categories'),
        'Lista de seis':r.get('shortlisted'),'Combinaciones probadas':r.get('combinations_tested'),
        'Combinaciones que cumplen':r.get('combinations_passing_guards'),
        'Mejora log-loss interna (%)':r.get('individual_log_loss_gain_relative_percent'),
        'Mejora Brier interna (%)':r.get('individual_brier_gain_relative_percent'),
        'Cambio captura interno (pp)':r.get('individual_capture_gain_pp'),
        'Cambio captura externo (pp)':r.get('external_capture_gain_pp'),
        'Mejora Brier externa (%)':r.get('external_brier_gain_relative_percent'),
        'Historia externa suficiente':r.get('external_eligible_history'),
        'Modelo ganador':r.get('winner_id'),'Archivo de selección':r['selection_file'],
        'Índice de prueba (desde 0)':r.get('selection_trial_index_zero_based'),
    } for r in decisions]


def encoding_rows(encodings):
    return [{
        'Modelo':r['model_name'],'Componente':r['component'],'Corte':r['training_cutoff'],
        'Variable':r['feature'],'Nombre':r['label'],'Referencia omitida':r['reference'],
        'Niveles aprendidos':r['learned_levels'],'Indicadores':r['indicator_columns'],
        'Número de indicadores':r['indicator_count'],'Aplica pooling':r['rare_pool_applied'],
        'Categorías conservadas por pooling':r['rare_pool_keep'],
        'Valores nuevos al predecir':r['unknown_at_prediction'],'Archivo':r['model_file'],
    } for r in encodings]


def workbook(path,tables):
    book=Workbook()
    book.remove(book.active)
    for name,rows in tables.items():
        sheet=book.create_sheet(name)
        headers=list(rows[0])
        sheet.append(headers)
        for row in rows:
            sheet.append([plain(row[h]) for h in headers])
        sheet.freeze_panes='C2' if name in {'Diccionario','Decisiones'} else 'B2'
        sheet.auto_filter.ref=sheet.dimensions
        sheet.sheet_view.zoomScale=80
        for cell in sheet[1]:
            cell.fill=PatternFill('solid',fgColor='16354A')
            cell.font=Font(color='FFFFFF',bold=True)
            cell.alignment=Alignment(wrap_text=True,vertical='top')
        sheet.row_dimensions[1].height=32
        for i,header in enumerate(headers,1):
            width=22 if any(word in header for word in ['(%)','(pp)','Número','Fotos','Corte','Periodo','Fase']) else 42
            if any(word in header for word in ['Motivo','Transformación','Consideraciones','Selección','Regla','Descripción']):
                width=70
            sheet.column_dimensions[get_column_letter(i)].width=width
        for cells in sheet.iter_rows(min_row=2):
            sheet.row_dimensions[cells[0].row].height=90 if name in {'Diccionario','Exclusiones'} else 62
            for cell in cells:
                cell.alignment=Alignment(wrap_text=True,vertical='top')
                if isinstance(cell.value,str):
                    # Texto de metadatos, nunca fórmula de Excel.
                    cell.data_type='s'
                    if headers[cell.column-1]=='Ficha':
                        cell.hyperlink=cell.value
                        cell.font=Font(color='126879',underline='single')
                elif isinstance(cell.value,(int,float)):
                    value=float(cell.value)
                    if 0<abs(value)<0.05:
                        cell.number_format='"'+('−' if value<0 else '')+'<0.1"'
                    else:
                        cell.number_format='0' if round(value,1).is_integer() else '0.0'
                if cells[0].row%2==0:
                    cell.fill=PatternFill('solid',fgColor='EEF4F7')
        sheet.print_options.horizontalCentered=False
        sheet.sheet_properties.pageSetUpPr.fitToPage=True
        sheet.page_setup.orientation='landscape'
    book.save(path)
    check=load_workbook(path,read_only=True,data_only=False)
    try:
        require(check.sheetnames==list(tables),'Cambió el índice del workbook.')
        for name,rows in tables.items():
            require(check[name].max_row==len(rows)+1,'Faltan filas en '+name)
            headers=list(rows[0])
            require([c.value for c in next(check[name].iter_rows())]==headers,'Cabeceras distintas en '+name)
            for expected,cells in zip(rows,check[name].iter_rows(min_row=2)):
                require(all(c.data_type!='f' for c in cells),'Se generó una fórmula no prevista.')
                for header,cell in zip(headers,cells):
                    value=plain(expected[header])
                    if isinstance(value,(int,float)):
                        equal=isinstance(cell.value,(int,float)) and math.isclose(cell.value,value,rel_tol=1e-12,abs_tol=1e-12)
                    else:
                        equal=(cell.value or '')==value
                    require(equal,'El Excel cambió o truncó un valor en '+name+' / '+header)
    finally:
        check.close()


def render_feature(folder,row,out):
    feature=row['feature']
    lines=[f"# {row['label']}",'',row['history_flow'],'',
        f"Modo `presentation` · variable `{feature}` · familia {row['family']}.",
        POPULATION+' Grain: persona-foto. Corte de conocimiento: 2026-09-07.','',
        '[Índice del diccionario](../diccionario.md) · [Decisiones completas](../decisiones.csv) · [Codificaciones de cada modelo](../codificaciones.csv)','',
        '## Tratamiento efectivo','']
    lines+=markdown_table(['Elemento','Qué se hizo'],[
        ['Tipo',row['input_type']],['Campos de origen',row['source_fields']],
        ['Valores de origen',row['source_values']],['Valores antes de agrupar categorías raras',row['values_before_pooling']],
        ['Transformación',row['transformation']],['Faltantes y valores por defecto',row['missing_and_default']],
        ['Codificación para el modelo',row['model_encoding']],['Categorías raras',row['pooling_rule']],
        ['Consideraciones',row['considerations']],
        ['Uso en la definición de grupos',row['group_role']],
    ])
    if row.get('wave_mapping'):
        lines+=['','### Encabezados efectivamente utilizados','']
        lines+=markdown_table(['Wave','Año de uso','Encabezado binario de origen'],[
            [m['wave'],m['model_year'],m['header']] for m in row['wave_mapping']])
        if row.get('question_aliases'):
            lines+=['','Aliases del diccionario: '+', '.join(row['question_aliases'])+'. Son la misma pregunta, no incorporaciones adicionales.']
    lines+=['','## Comprobación y decisión','',
        '**Pregunta:** '+('¿Se preservó el término obligatorio y su codificación?' if row['family']=='B' else
                           '¿Qué aporta esta representación al añadirla a la base, con datos conocidos antes de cada corte?'),
        '', '**Motivo:** separar una transformación de datos de su elegibilidad y de su mejora predictiva; dejar verificable por qué se conservó o quedó fuera.',
        '', '**Método:** inspección del SQL/código congelado, soporte agregado de la carga, registro de selección interna y codificaciones guardadas. '
        'La documentación no reconstruye filas ni cambia decisiones a partir de resultados externos.',
        '', '**Regla de aceptación:** '+('Los diez términos de B deben aparecer en cada artefacto de modelo; no participan en una eliminación de variables.'
                                      if row['family']=='B' else GUARDS)]
    if row['family']=='RH':
        lines+=['','La elegibilidad requiere al menos seis fotos observables dentro del entrenamiento de selección y más de una categoría después de agrupar las poco frecuentes. '
                   'Ese entrenamiento termina antes de la reserva interna; puede ser más corto que el reajuste utilizado para evaluar el periodo externo.']
    elif row['family']=='QVoice':
        lines+=['','El contenido debe además reducir log-loss frente a un ajuste que sólo conserva los mismos estados de disponibilidad. '
                   'La reserva mensual del primer corte QVoice usa Brier y PR-AUC; no aplica las tres condiciones de seis meses. '
                   'La regla de seis fotos RH no se aplica automáticamente a QVoice.']
    if row['annual_support']:
        lines+=['','### Métricas observadas de soporte','',
                'Son conteos de la representación ya transformada. «Observable» sigue los estados explícitos de ausencia del runner; '
                'no prueba por sí solo la definición oficial ni la completitud de eventos. Personas repetidas entre fotos no son personas distintas.','']
        lines+=markdown_table(['Año','Persona-fotos','Observables','Cobertura (%)','Fotos observables','Categorías observables'],[
            [r['year'],r['person_snapshots'],r['observable_person_snapshots'],
             100*r['observable_person_snapshots']/r['person_snapshots'],r['observable_snapshots'],r['observed_categories']]
            for r in row['annual_support']])
    lines+=['','### Selección en cada corte','']
    lines+=markdown_table(['Periodo','Base','Fotos internas observables','Cumple condiciones individual','Lista de seis','Decisión'],[
        [r['period'],r['anchor'],r.get('training_support',{}).get('observable_training_snapshots'),
         r.get('individual_passes_guards'),r.get('shortlisted'),r['decision']]
        for r in row['decisions']])
    lines+=['']
    for r in row['decisions']:
        lines.append(f"- **{r['period']} / {r['anchor']}:** {r['reason']}")
        if r.get('individual_reason') and r['individual_reason'] not in r['reason']:
            lines[-1]+=' Prueba individual: '+r['individual_reason']
    external=[r for r in row['decisions'] if r.get('external_fit_status')=='ok' and r.get('external_comparable')]
    if external:
        lines+=['','### Resultado externo de añadirla por separado','',
            'Cambios respecto a la misma base y población de cada periodo; valores positivos indican mejora de pérdidas o mayor captura. '
            'Estas cifras describen el resultado externo y no se emplean para seleccionar.','']
        lines+=markdown_table(['Periodo','Base','Mejora log-loss (%)','Mejora Brier (%)','Cambio captura (pp)'],[
            [r['period'],r['anchor'],r['external_log_loss_gain_relative_percent'],
             r['external_brier_gain_relative_percent'],r['external_capture_gain_pp']] for r in external])
    skipped=[r for r in row['decisions'] if r.get('external_fit_status')=='ok' and not r.get('external_comparable')]
    if skipped:
        lines+=['','Se omiten de la tabla externa los ajustes con historia insuficiente: '+', '.join(r['period'] for r in skipped)+
                '. Sus intentos y soporte siguen en decisiones.json; un ajuste numérico no equivale a una prueba temporal admisible.']
    lines+=['','**Frecuencia en selecciones anteriores:** '+row['selection_history_summary']+'.',
            '', '**Conclusión de selección final:** '+row['final_summary'],
            '', '**Impacto en score de assessment:** ninguno. Los scores previos no se reasignan con estas pruebas.',
            '', '**Acción:** '+row['action']]
    final_codes=[r for r in row['encodings'] if r['model_name'].startswith('final-')]
    lines+=['','## Codificación de los modelos finales','']
    if final_codes:
        lines+=markdown_table(['Modelo','Referencia omitida','Número de indicadores','Niveles aprendidos'],[
            [r['model_name'],r['reference'],r['indicator_count'],r['learned_levels']] for r in final_codes])
    else:
        lines+=['No aparece en los modelos finales seleccionados. Sus valores candidatos están descritos arriba; '
                'no se inventa una referencia final para una variable que no quedó en esos artefactos.']
    lines+=['','## Fuentes trazables','']
    for source in row['sources']:
        lines.append('- '+link(source,folder))
    for source in row['evidence_files']:
        record=read(ROOT/source)
        when=record.get('executed_at',record.get('generated_at','registro de carga de esta ejecución'))
        text='- '+link(source,folder)+f' · {when}'
        if record.get('outcome'):
            text+=' · resultado del control: '+record['outcome']
        if record.get('query'):
            text+=' · SQL: '+link(record['query'],folder)
        lines.append(text)
    if not row['evidence_files']:
        lines+=['','La referencia de este término es el SQL y los artefactos del modelo B. '
                'No se presenta esta ficha como una nueva verificación semántica o de lineage.']
    lines+=['','La fecha de la consulta no reemplaza la fecha de conocimiento de la variable. '
            'Un control técnico pass sólo respalda la afirmación acotada de su contrato.','']
    (folder/(feature+'.md')).write_text('\n'.join(lines),encoding='utf-8')


def render_exclusion(folder,row,out):
    lines=[f"# {row['label']}",'',row['history'],'',
        'Modo `presentation`. '+POPULATION+' Corte de conocimiento: 2026-09-07.','',
        '[Índice del diccionario](../diccionario.md)','']
    lines+=markdown_table(['Elemento','Registro'],[
        ['Fuente',row['source']],['Campos',row['source_fields']],['Valores de origen',row['source_values']],
        ['Representación considerada',row['proposed_type']],['Transformación aplicada',row['transformation']],
        ['Decisión',row['decision']],['Motivo',row['reason']],['Acción',row['action']],
    ])
    lines+=['','**Pregunta:** ¿se puede construir una representación autorizada, distinguible y conocida en cada foto?',
        '', '**Motivo:** evitar llamar «cero» a falta de fuente o presentar una exclusión previa como una prueba de bajo desempeño.',
        '', '**Método y regla:** se conserva el alcance exacto del audit enlazado y de la autorización. '
        'Un timestamp técnico o una estructura parcial no acredita disponibilidad histórica de la variable.',
        '', '**Métricas y conclusión:** '+row['history']+' '+row['reason'],
        '', '**Resultado predictivo:** no estimado para esta representación; no se afirma que empeore o que no tenga señal.',
        '', '**Impacto en score:** ninguno.','', '**Acción:** '+row['action'],'', '## Evidencia','']
    for name in row['evidence']:
        path=out/'input-evidence'/name
        record=read(path)
        lines.append('- '+link(path,folder)+' · '+str(record.get('executed_at',record.get('generated_at','fecha en manifest')))+
                     ' · '+str(record.get('outcome','inspección')))
        if record.get('query'):
            lines.append('- '+link(record['query'],folder))
    for path in row['documentation']:
        lines.append('- '+link(path,folder))
    (folder/('exclusion-'+row['id']+'.md')).write_text('\n'.join(lines)+'\n',encoding='utf-8')


def render_index(folder,catalog,excluded,results,config,scope):
    final=results['final']
    lines=['# Qué se hizo con cada variable del playbook V11','',
        'Diccionario de modelado y decisiones de la ejecución `'+config['run_id']+'`. Modo `presentation`; sólo agregados y metadatos, sin filas ni predicciones individuales.','',
        'Se documentan **102 representaciones**: los **10 términos obligatorios de B**, **41 incorporaciones de RH** y **51 de QVoice**. '
        'El archivo llamado «43 variables» contiene **42 vistas analíticas**, además de portada/navegación; varias vistas comparten variables y otras contienen múltiples preguntas.','',
        'Cada ficha contiene valores de origen, bandas/categorías, transformación exacta, tratamiento de faltantes, fecha de conocimiento, '
        'pruebas, decisiones por periodo y motivo. Las referencias y niveles concretos se extraen de los modelos guardados.','',
        'Las fichas describen la implementación ejecutada. Cuando el alcance previo requiere una precisión, '
        'se añade una nota del tratamiento efectivo en la correspondencia del playbook.','',
        '[Abrir diccionario en Excel](diccionario.xlsx) · [Tabla CSV](diccionario.csv) · [Decisiones CSV](decisiones.csv) · '
        '[Codificaciones CSV](codificaciones.csv) · [Informe de resultados](informe.md)','',
        '## Selección final experimental','',
        '**RH añadido a B:** '+(', '.join(catalog[f]['label'] for f in final['hr_additions']) or 'ninguna incorporación')+'.','']
    for anchor,choice in final['qvoice'].items():
        selected=', '.join(catalog[f]['label'] for f in choice['features']) or 'ninguna variable adicional'
        lines.append(f'- **QVoice sobre {anchor}:** {selected}. Tipo de ajuste guardado: `{choice["kind"]}`.')
    lines+=['',f"Corte final: {final['cutoff']}. La evaluación temporal corresponde a selecciones hechas en cortes anteriores; "
        'no se evalúan retrospectivamente estos coeficientes finales como si fueran externos.','',
        '## Cómo leer las decisiones','',
        '«Obligatoria en B» significa que se conservó por instrucción. «No seleccionada» tiene un motivo por corte: '
        'historia insuficiente, ajuste fallido, condiciones incumplidas o una alternativa con mejor resultado bajo la regla de parsimonia. '
        'Una variable puede ser seleccionada en combinación aunque no cumpla las condiciones por sí sola. '
        'Las exclusiones anteriores a la selección aparecen separadas, con su causa y lo que falta resolver.','',
        CODING,'','Por ejemplo, una binaria de origen con valores sí/no y un estado desconocido puede producir tres categorías y dos indicadores. '
        'El número concreto depende de los niveles presentes y de la referencia que aparece en el archivo de codificaciones.',
        '',POOLING,'',SEARCH,'',GUARDS,'',
        'QVoice añade un ajuste a la predicción de B o B+RH, conservando los términos de esa base. Se compara con la base fija, '
        'un ajuste sólo de intercepto y un control de disponibilidad emparejado. El contenido debe superar en log-loss a su control '
        'de disponibilidad. Los estados de ausencia no se rellenan como respuestas desfavorables.','',
        'Sobre B+RH, la búsqueda QVoice está condicionada a las incorporaciones RH elegidas para ese corte externo. '
        'Los coeficientes de la base se vuelven a ajustar antes de la reserva QVoice; no es una búsqueda conjunta de todas las combinaciones RH×QVoice.','',
        'Las decisiones internas reservan tres fotos previas con seguimiento conocido. El primer corte QVoice usa una reserva mensual '
        'por historia corta; después se usa seguimiento de seis meses. Las variables RH requieren seis fotos observables de entrenamiento '
        'y más de una categoría después de agrupar las poco frecuentes.','',
        'Los grupos conservan antigüedad × salidas del equipo × nivel para comparar el aporte de variables. '
        'Se recalcula su cantidad por foto; esta ronda no optimiza nuevas dimensiones de agrupación. '
        'El informe adjunto contiene el diagnóstico agregado de captura sin la restricción de estos grupos.','',
        'La captura principal usa capacidad del 30% de la población en cada foto. Se ordenan grupos elegibles de al menos 300 personas por su riesgo medio; '
        'el último grupo puede contribuir parcialmente y los empates se reparten proporcionalmente. Es un cálculo agregado, sin lista de personas priorizadas.','',
        '## Fichas de las variables','']
    for family,title in [('B','Base B obligatoria'),('RH','Incorporaciones RH'),('QVoice','Incorporaciones QVoice')]:
        lines+=['### '+title,'']
        rows=[]
        for row in catalog.values():
            if row['family']!=family:
                continue
            selected=('Obligatoria' if family=='B' else
                      ', '.join('sobre '+a for a in row['final_selected_anchors']) or 'No seleccionada al corte final')
            rows.append([f"[{row['label']}](fichas/{row['feature']}.md)",row['feature'],row['input_type'],row['selection_history_summary'],selected])
        lines+=markdown_table(['Ficha','Identificador','Tipo de entrada','Selección para periodos externos','Selección final'],rows)+['']
    lines+=['## Conceptos excluidos o sin nueva columna','']
    lines+=markdown_table(['Ficha','Decisión','Motivo'],[
        [f"[{r['label']}](fichas/exclusion-{r['id']}.md)",r['decision'],r['reason']] for r in excluded])
    lines+=['','## Archivos y trazabilidad','',
        '- `diccionario.xlsx`: guía, transformaciones, decisiones, valores aprendidos, correspondencia V11, exclusiones y soporte anual.',
        '- `diccionario.csv` y `fichas/`: consulta por variable.',
        '- `decisiones.json` y `decisiones.csv`: decisiones por variable, periodo y base; el JSON conserva métricas y límites exactos.',
        '- `codificaciones.json` y `codificaciones.csv`: referencias y categorías realmente aprendidas por cada modelo, incluidos los externos.',
        '- `vistas-playbook.csv`: correspondencia de las 42 vistas del V11.',
        '- `manifest.json`: hashes propios y del manifest de entrenamiento. Este suplemento no altera los artefactos congelados.',
        '', 'Las fichas documentan el código ejecutado; no elevan scores ni afirman lineage oficial donde sigue pendiente. '
        'Las fórmulas y JSON/CSV conservan la precisión de la evidencia; las tablas de resultados visibles usan como máximo un decimal.','']
    (folder/'diccionario.md').write_text('\n'.join(lines),encoding='utf-8')


def support_rows(catalog):
    return [{
        'Variable':row['feature'],'Nombre':row['label'],'Familia':row['family'],'Año':s['year'],
        'Persona-fotos':s['person_snapshots'],'Persona-fotos observables':s['observable_person_snapshots'],
        'Cobertura (%)':100*s['observable_person_snapshots']/s['person_snapshots'],
        'Fotos observables':s['observable_snapshots'],'Categorías observables':s['observed_categories'],
        'Primera foto observable del año':s['first_observed_snapshot'],
    } for row in catalog.values() for s in row['annual_support']]


def view_rows(scope,catalog):
    return [{
        'Posición en V11':view['slide'],'ID en HTML':view['source_stage_id'],
        'Vista':view['vista'],'Título original':view['source_title'],
        'Estado de preparación':view['estado'],'Representaciones':view['representaciones'],
        'Seleccionadas al corte final':[f for f in view['representaciones'] if catalog[f]['final_selected_anchors']],
        'Alcance previsto (registro congelado)':view['explicacion'],
        'Nota del tratamiento efectivo':(
            'La fórmula de manager_left_6m exige evento conocido y ventana de seis meses, pero no añade '
            'un filtro de service_date para reiniciarlo por recontratación. La mención a episodio vigente '
            'del alcance previo no debe interpretarse como un filtro implementado.'
            if 'manager_left_6m' in view['representaciones'] else
            'Ver las fichas de cada representación para la fórmula ejecutada, ausencias y límites.'),
    } for view in scope['views']]


def exclusion_rows(excluded):
    return [{
        'Concepto':r['label'],'Fuente':r['source'],'Campos':r['source_fields'],
        'Valores de origen':r['source_values'],'Representación considerada':r['proposed_type'],
        'Qué se hizo':r['transformation'],'Historia':r['history'],
        'Decisión':r['decision'],'Motivo':r['reason'],'Acción':r['action'],
        'Evidencia':r['evidence'],'Ficha':'fichas/exclusion-'+r['id']+'.md',
    } for r in excluded]


def guide_rows(config,results,catalog,generated):
    values=[
        ('Propósito','Documentar qué se hizo con cada variable del playbook V11, incluyendo las excluidas y los motivos.'),
        ('Clasificación','presentation; sólo agregados, metadatos de variables y categorías de los modelos.'),
        ('Ejecución',config['run_id']),('Generado',generated),('Corte de conocimiento',config['knowledge_boundary']),
        ('Población',POPULATION),('Grain','Una persona-foto mensual; ninguna fila de persona se incluye en este archivo.'),
        ('Alcance','102 representaciones: 10 términos de B, 41 RH y 51 QVoice; 42 vistas analíticas del playbook llamado 43 variables.'),
        ('Base B','Ocho predictores y dos interacciones obligatorios; no se eliminan en esta ronda.'),
        ('Diccionario','Tipo, origen, valores, transformación, faltantes, conocimiento, consideraciones y decisión final; una fila por representación.'),
        ('Decisiones','Una fila por variable/periodo/base. Las individuales y las combinaciones pueden tener resultados diferentes.'),
        ('Codificaciones','Referencia omitida y valores aprendidos en cada modelo guardado; se incluye pooling real, no uno inventado.'),
        ('Playbook V11','Relaciona cada vista con sus representaciones y límites. Varias vistas pueden compartir una variable.'),
        ('Exclusiones','Siete conceptos/representaciones sin nueva columna; no tienen un efecto predictivo estimado en esta ronda.'),
        ('Soporte anual','Cobertura y categorías observables después de preparar la variable; no reemplaza su definición o validación semántica.'),
        ('Codificación',CODING),('Pooling RH',POOLING),('Búsqueda',SEARCH),('Reglas de selección',GUARDS),
        ('Control QVoice','El contenido debe reducir log-loss frente a su disponibilidad emparejada; se evalúan bases fijas y ajuste sólo de intercepto.'),
        ('Condición B+RH','QVoice se selecciona condicionado a las incorporaciones RH de ese corte externo; los coeficientes de la base se reajustan antes de la reserva QVoice. No es una búsqueda conjunta exhaustiva RH×QVoice.'),
        ('Tiempo QVoice','Wave 2024→2025; 2025→2026; 2026→2027. No se afirma que esa sea la fecha original de publicación.'),
        ('Positivo en resultados','Mejora de Brier/log-loss (%) positiva indica menor pérdida; cambio de captura positivo indica más salidas captadas.'),
        ('Comparación externa','La columna historia externa suficiente debe ser Sí para interpretar una prueba RH como temporalmente admisible; las métricas no deciden la selección.'),
        ('Precisión','El Excel muestra un decimal en resultados; CSV/JSON conservan la precisión técnica. Las fórmulas no se redondean.'),
        ('Scores','Este archivo es un diccionario de modelado, no un workbook de assessment. No se reasignó ningún score.'),
        ('Grupos','Se conservan dimensiones de B y se recalcula elegibilidad por foto. No se fuerza 22; no se declara un óptimo global de agrupaciones.'),
        ('Informe','informe.md'),('Fichas legibles','diccionario.md'),('Trazabilidad','manifest.json y el manifest original de entrenamiento.'),
        ('RH final',', '.join(catalog[f]['label'] for f in results['final']['hr_additions']) or 'Sin incorporaciones'),
    ]
    for anchor,spec in results['final']['qvoice'].items():
        values.append(('QVoice final sobre '+anchor,', '.join(catalog[f]['label'] for f in spec['features']) or 'Sin variables adicionales'))
    return [{'Elemento':key,'Descripción':value} for key,value in values]


def check_catalog(config):
    metadata=read(ROOT/config['source_qvoice_metadata'])
    catalog=definitions.catalog(config,metadata)
    expected=config['base_features']+inc.hr_candidates(config)+inc.q_candidates(config)
    require(len(expected)==102 and len(set(expected))==102,'Conteo de representaciones no esperado.')
    require(set(catalog)==set(expected),'El diccionario no cubre todas las representaciones.')
    require('ios_12m' not in catalog,'IOS no debe aparecer como candidata.')
    required=['feature','label','input_type','source_fields','source_values','values_before_pooling',
              'transformation','missing_and_default','known_at','considerations','sources']
    for key,row in catalog.items():
        require(all(row.get(field) for field in required),'Ficha incompleta: '+key)
        for path in row['sources']:
            require((ROOT/path).is_file(),'Fuente ausente: '+path)
    scope=read(ROOT/config['playbook_scope'])
    require(len(scope['views'])==42 and len({v['source_stage_id'] for v in scope['views']})==42,
            'Correspondencia del V11 incompleta.')
    for view in scope['views']:
        require(set(view['representaciones'])<=set(catalog),'Representación no documentada en V11.')
    require(base.sha(ROOT/scope['source'])==scope['source_sha256']==config['source_playbook_sha256'],
            'El playbook no coincide con el alcance congelado.')
    return {feature:catalog[feature] for feature in expected},scope


def build(config_path):
    active=read(config_path)
    out=(ROOT/active['output']).resolve()
    require(out.is_relative_to(ROOT.resolve()),'Salida fuera del workspace.')
    config=read(out/'frozen-config.json')
    require(config==active,'La configuración actual difiere de la congelada.')
    original_manifest=read(out/'manifest.json')
    require(original_manifest['status']=='completed_experimental','Esperar a que termine la ejecución para documentar decisiones finales.')
    destination=out/'documentacion'
    require(not destination.exists(),'El suplemento ya existe; conservarlo y generar otra versión si cambia la documentación.')
    for relative,digest in original_manifest['source_hashes'].items():
        require(base.sha(ROOT/relative)==digest,'Cambió una fuente congelada: '+relative)
    for relative,digest in original_manifest['files'].items():
        require(base.sha(out/relative)==digest,'Cambió un resultado congelado: '+relative)
    catalog,scope=check_catalog(config)
    results=read(out/'results.json')
    decisions,encodings=decisions_source.build(out,config,catalog,results)
    final=results['final']
    selected_rh={r['feature'] for r in decisions if r['period']=='final' and r['family']=='RH' and r['selected']}
    require(selected_rh==set(final['hr_additions']),'La selección RH final difiere entre resultados y decisiones.')
    expected_models={
        'final-B':set(config['base_features']),
        'final-B-RH':set(config['base_features'])|selected_rh,
    }
    for anchor,choice in final['qvoice'].items():
        selected_q={r['feature'] for r in decisions if r['period']=='final' and r['family']=='QVoice'
                    and r['anchor']==anchor and r['selected']}
        require(selected_q==set(choice['features']),'La selección QVoice final difiere para '+anchor)
        expected_models['final-'+anchor+'-Q']=set(config['base_features'])|selected_q|(selected_rh if anchor=='B_RH' else set())
    for name,expected in expected_models.items():
        actual={r['feature'] for r in encodings if r['model_name']==name}
        require(actual==expected,'La codificación del modelo final difiere de su selección: '+name)
    catalog=prepare_catalog(config,out,catalog,scope,decisions,encodings)
    excluded=definitions.exclusions()
    require(len(excluded)==7,'Falta un concepto excluido.')
    inputs={out/'manifest.json',out/'results.json',out/'frozen-config.json',out/'source-hashes.json',
            out/'candidate-summary.json',out/'informe.md',
            ROOT/config['playbook_scope'],ROOT/config['source_qvoice_metadata'],ROOT/config['source_playbook']}
    inputs.update((out/'selection').glob('*.json'))
    inputs.update((out/'models').glob('*.json'))
    for row in catalog.values():
        inputs.update(ROOT/path for path in row['sources']+row['evidence_files'])
    for row in excluded:
        inputs.update(out/'input-evidence'/name for name in row['evidence'])
        inputs.update(ROOT/path for path in row['documentation'])
    source_files=[Path(__file__),ROOT/'tools/b_playbook_variable_catalog.py',ROOT/'tools/b_playbook_variable_decisions.py',
                  ROOT/'tools/b_playbook_documentation_report.py']
    inputs.update(source_files)
    hashes={p.relative_to(ROOT).as_posix():base.sha(p) for p in sorted(inputs)}
    generated=datetime.now(timezone.utc).isoformat()
    staging=out/('documentacion-en-preparacion-'+datetime.now(timezone.utc).strftime('%Y%m%dT%H%M%S'))
    require(staging.resolve().is_relative_to(out) and not staging.exists(),'Directorio de preparación no válido.')
    (staging/'fichas').mkdir(parents=True)
    dictionaries=dictionary_rows(catalog)
    decision_table=decision_rows(decisions)
    codes=encoding_rows(encodings)
    views=view_rows(scope,catalog)
    exclusions=exclusion_rows(excluded)
    support=support_rows(catalog)
    require(all(r['Persona-fotos']>=30 and (r['Persona-fotos observables']==0 or r['Persona-fotos observables']>=30)
                for r in support),'El soporte anual contiene una cohorte pequeña que requiere supresión.')
    tables={'Guía':guide_rows(config,results,catalog,generated),'Diccionario':dictionaries,
            'Decisiones':decision_table,'Codificaciones':codes,'Playbook V11':views,
            'Exclusiones':exclusions,'Soporte anual':support}
    for name,rows in [('diccionario',dictionaries),('decisiones',decision_table),('codificaciones',codes),
                      ('vistas-playbook',views),('exclusiones',exclusions),('soporte-anual',support)]:
        csv_file(staging/(name+'.csv'),rows)
    write_json(staging/'diccionario.json',{'reporting_mode':'presentation','run_id':config['run_id'],
        'generated_at':generated,'variables':list(catalog.values()),'assessment_score_changed':False,
        'employee_level_rows_persisted':False})
    write_json(staging/'decisiones.json',{'reporting_mode':'presentation','rows':decisions})
    write_json(staging/'codificaciones.json',{'reporting_mode':'presentation','rows':encodings})
    write_json(staging/'exclusiones.json',{'reporting_mode':'presentation','rows':excluded})
    for row in catalog.values():
        render_feature(staging/'fichas',row,out)
    for row in excluded:
        render_exclusion(staging/'fichas',row,out)
    render_index(staging,catalog,excluded,results,config,scope)
    visual_checks=report_source.render(staging,out,config,results,read(out/'candidate-summary.json'))
    workbook(staging/'diccionario.xlsx',tables)
    require(len(list((staging/'fichas').glob('*.md')))==109,'Número de fichas incorrecto.')
    for name,rows in [('diccionario',dictionaries),('decisiones',decision_table),('codificaciones',codes),
                      ('vistas-playbook',views),('exclusiones',exclusions),('soporte-anual',support)]:
        with (staging/(name+'.csv')).open(encoding='utf-8-sig',newline='') as handle:
            restored=list(csv.DictReader(handle))
        require(len(restored)==len(rows),'El CSV no reproduce las filas: '+name)
    require(all(base.sha(ROOT/p)==digest for p,digest in hashes.items()),'Una fuente cambió durante la documentación.')
    verification={'outcome':'pass','dictionary_representations':len(catalog),'playbook_views':len(views),
        'excluded_concepts_or_representations':len(excluded),'decision_rows':len(decisions),
        'decisions_by_family':dict(Counter(r['family'] for r in decisions)),
        'encoding_rows':len(encodings),'feature_cards':109,
        'saved_winners_reproduced_from_trials':True,'guard_reasons_reconciled':True,
        'csv_and_workbook_row_counts_verified':True,'workbook_contains_formulas':False,
        'workbook_values_match_source_tables':True,'final_results_decisions_and_encodings_agree':True,
        'positive_annual_support_cohorts_at_least_30':True,
        'visual_checks':visual_checks,
        'all_input_hashes_unchanged':True,'training_artifacts_modified':False,
        'assessment_score_changed':False,'employee_level_rows_persisted':False}
    write_json(staging/'verification.json',verification)
    manifest={'schema_version':1,'status':'completed_documentation','reporting_mode':'presentation',
        'run_id':config['run_id'],'generated_at':generated,'training_manifest':
        (out/'manifest.json').relative_to(ROOT).as_posix(),'training_manifest_sha256':base.sha(out/'manifest.json'),
        'method':'Documentación del código ejecutado y decisiones agregadas; no es una nueva evaluación de scores.',
        'source_hashes':hashes,'files':{p.relative_to(staging).as_posix():base.sha(p)
                                     for p in sorted(staging.rglob('*')) if p.is_file()},
        'assessment_score_changed':False,'employee_level_rows_persisted':False,'individual_scores_published':False}
    write_json(staging/'manifest.json',manifest)
    # Ambos paths se resuelven dentro de la salida autorizada; no se reemplaza una versión previa.
    require(staging.resolve().parent==destination.resolve().parent and not destination.exists(),'Destino no válido para publicar el suplemento local.')
    staging.rename(destination)
    return {'outcome':'pass','documentation':destination.relative_to(ROOT).as_posix(),**verification}


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--config',type=Path,default=inc.DEFAULT)
    parser.add_argument('--check-catalog',action='store_true')
    args=parser.parse_args()
    if args.check_catalog:
        catalog,scope=check_catalog(read(args.config))
        print(json.dumps({'outcome':'pass','representations':len(catalog),'playbook_views':len(scope['views']),
                          'excluded_concepts':len(definitions.exclusions()),'training_modified':False},ensure_ascii=False))
    else:
        print(json.dumps(build(args.config),ensure_ascii=False))


if __name__=='__main__':
    main()

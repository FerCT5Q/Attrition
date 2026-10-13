"""Publica tablas, gráficos y fichas desde agregados congelados; sin consultas."""
from __future__ import annotations

import json
import os
from pathlib import Path
from datetime import datetime, timezone
import hashlib
import re

import numpy as np
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.ticker import FuncFormatter
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'output/analysis/engineering/20261012/julia-rh-review-v1'
LABELS={'CQ':'CQ','FQ':'FQ','FQ_BU_L3':'FQ + BU L3'}
COLORS={'CQ':'#777777','FQ':'#1670ae','FQ_BU_L3':'#a24d16'}


def read(p):return json.loads(p.read_text(encoding='utf-8-sig'))
def write(p,x):p.write_text(json.dumps(x,ensure_ascii=False,indent=2,allow_nan=False)+'\n',encoding='utf-8')
def fmt(v):
    if v is None:return 'No publicable / no acreditado'
    if isinstance(v,(bool,np.bool_)):return 'Sí' if v else 'No'
    v=float(v)
    if v==int(v):return str(int(v))
    if 0<abs(v)<.05:return '<0.1' if v>0 else '>−0.1'
    return f'{v:.1f}'.rstrip('0').rstrip('.')
def ci(v,scale=1):return '['+', '.join('inferior a 0' if -.05<x*scale<0 else fmt(x*scale) for x in v)+']'
def table(headers,rows):return '\n'.join(['| '+' | '.join(headers)+' |','|'+'|'.join(['---']*len(headers))+'|']+['| '+' | '.join(str(v).replace('|',' / ').replace('\n','<br>') for v in row)+' |' for row in rows])
def xl(v):return json.dumps(v,ensure_ascii=False) if isinstance(v,(dict,list)) else v
def save_md(name,lines):
    path=OUT/name;path.parent.mkdir(parents=True,exist_ok=True);path.write_text('\n'.join(lines)+'\n',encoding='utf-8')


def style_sheet(ws):
    ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
    for cell in ws[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='245C80')
    for col in ws.columns:
        letter=col[0].column_letter;ws.column_dimensions[letter].width=14 if letter=='A' else 42
        for cell in col:
            cell.alignment=Alignment(vertical='top',wrap_text=True)
            if isinstance(cell.value,float):cell.number_format='0.#'


def documents(result):
    d=read(OUT/'documentacion/variables-44.json')
    dictionary={r['feature']:r for r in read(ROOT/'output/analysis/engineering/20261011/b-playbook-v11-v1/documentacion/diccionario.json')['variables']}
    extra={'qv_Q1':'qv_M19','qv_Q2':'qv_M25','qv_Q3':'qv_M28','qv_Q13':'qv_M12','qv_Q17':'qv_M27'}
    d['rows'][12]['representaciones']+=list(extra)
    d['rows'][12]['limitacion']+=' Se documentan aquí también las preguntas individuales que componen la vista, no sólo el promedio. El promedio no entra en CQ/FQ históricos.'
    for f,alias in extra.items():
        t={k:dictionary[alias].get(k) for k in ['feature','source_fields','input_type','source_values','transformation','missing_and_default','known_at','sources','considerations']}
        t.update(feature=f,alias_v11=alias)
        d['rows'][12]['tratamientos_existentes'].append(t)
    d['rows'][9]['representaciones']=[f for f in dictionary if f.startswith('qv_comment_') or f=='qv_improvement_comment_presence']
    d['rows'][9]['tratamientos_existentes']=[dictionary[f] for f in d['rows'][9]['representaciones']]
    d['rows'][43]['tratamientos_existentes']=[{'feature':'bu_l3','source_fields':['post_hire_monthly_worker_dtls.cost_centerid','d_department.nk_dept_id','d_department.ude_bu_reporting_l3'],
        'input_type':'Categoría nominal; indicador 0/1 por categoría, referencia del entrenamiento',
        'source_values':'Etiqueta BU L3 del mapping histórico válido a la foto',
        'transformation':'TRIM; join por cost center y vigencia [edw_effective_dt, edw_end_dt). BU de menos de 300 personas distintas de entrenamiento → (other); categorías y referencias exactas en Codificaciones de modelos.',
        'missing_and_default':'Sin match → (unknown). No se sustituye por el mapping futuro. Categoría no aprendida → indicadores cero para ese término.',
        'known_at':'Mapping vigente por foto; controles as-of y multi-match ejecutados',
        'considerations':'Los grupos de reporte usan BU identificable, no el pool de predictores. No es una variable causal.'}]
    d['rows'][0]['tratamientos_existentes'].append({'feature':'ios_12m','source_fields':['workers_ios_transfers'],
        'input_type':'Categoría nominal','source_values':'Antes de la fuente, none, lateral u otro, reorg',
        'transformation':'Antes de 22-oct-2023 → (antes de la fuente); sin movimiento → none; algún motivo distinto de Reorg Department and/or Manager Change → lateral u otro; resto → reorg. Ventana retrospectiva de 12 meses.',
        'missing_and_default':'Sin movimiento enlazado → none; no acredita ausencia real ni publicación histórica',
        'known_at':'No acreditada; fuente actual sin versión histórica suficiente',
        'considerations':'Conservado sólo para reproducir la composición de C/CQ; no confundir con transferencia genérica.'})
    for r in d['rows']:
        if not r['tratamientos_existentes']:
            r['tratamientos_existentes']=[{'feature':'No construida como predictor','source_fields':[],
                'input_type':'Sin entrada al modelo; filtro en Engineering', 'source_values':r['preparacion_v11'],
                'transformation':'No se transforma ni imputa para esta comparación','missing_and_default':'No aplica',
                'known_at':'Ver motivo de exclusión o historia insuficiente en la ficha', 'considerations':r['resultado_historico_B_V11_no_CQ']}]
        for p in result['periods']:
            for n,m in p['models'].items():
                roles=r['uso_por_modelo_corte'][p['period']['id']+'/'+n]
                used=set(m['base_features']+m['selected']['features'])
                roles['predictor']=[v for v in r['representaciones'] if v in used]
                roles['excluidas_de_esta_comparacion']=[v for v in r['representaciones'] if v not in used]
        r['detalle']=[str((ROOT/'planes/modelo_c'/p).resolve().relative_to(ROOT)).replace('\\','/') for p in r['detalle']]
    d.update(stage='documentacion_final_verificada',source_initial='variables-44.json',assessment_score_changed=False)
    write(OUT/'documentacion/variables-44-final.json',d)
    wb=Workbook();ws=wb.active;ws.title='44 variables y acciones'
    headers=['Número','Variable','Representaciones','Potencialmente accionable','Útil para priorizar','Inadecuada para explicación individual','Requiere cautela',
             'Preparación V11 (contexto)','Resultado histórico B V11, no CQ','Acción candidata','Responsable por función','Contraste RH','Indicador','Plazo','Limitación','Detalle']
    ws.append(headers)
    roles=wb.create_sheet('Uso por modelo y corte');roles.append(['Número','Variable','Modelo y corte','Predictores','Definición de grupos','Diagnóstico','Excluidas','Subdivisión BU experimental','Motivo'])
    treatment=wb.create_sheet('Tratamientos existentes');treatment.append(['Número','Representación','Fuente','Valores','Tipo de entrada','Transformación','Faltantes y defecto','Conocimiento','Cautela'])
    encoding=wb.create_sheet('Codificaciones de modelos');encoding.append(['Periodo','Modelo','Componente','Representación','Referencia','Categorías aprendidas','Pool BU'])
    executed=wb.create_sheet('Tratamiento ejecutado');executed.append(['Periodo','Modelo','Componente','Representación','Transformación y faltantes ejecutados','Agrupación de categorías','Categoría desconocida al predecir'])
    for r in d['rows']:
        cl=r['clasificacion'];ws.append([r['numero'],r['variable'],', '.join(r['representaciones']),
            *['Sí' if cl[k] else 'No' for k in ['potencialmente_accionable','util_para_priorizar','inadecuada_para_explicacion_individual','requiere_cautela']],
            r['preparacion_v11'],r['resultado_historico_B_V11_no_CQ'],r['accion_candidata'],r['responsable_funcion'],r['evidencia_RH'],r['indicador'],r['plazo_sugerido'],r['limitacion'],'; '.join(r['detalle'])])
        for k,v in r['uso_por_modelo_corte'].items():
            roles.append([r['numero'],r['variable'],k,', '.join(v['predictor']),', '.join(v['definicion_grupos']),', '.join(v['diagnostico']),', '.join(v['excluidas_de_esta_comparacion']),'Sí' if v['subdivision_experimental'] else 'No',v['motivo']])
        for t in r['tratamientos_existentes']:
            treatment.append([xl(v) for v in [r['numero'],t['feature'],', '.join(t.get('source_fields') or []),t.get('source_values'),t.get('input_type'),t.get('transformation'),t.get('missing_and_default'),t.get('known_at'),t.get('considerations')]])
    sql=(ROOT/'queries/profiles/engineering_bu_l3_person_month_20261012.sql').read_text(encoding='utf-8')
    select=sql.split('\nSELECT\n  f.snapshot_date,',1)[1].split('\nFROM features f',1)[0]
    expressions=[];depth=0;quoted=False;start=0
    for i,char in enumerate(select):
        if char=="'":quoted=not quoted
        if not quoted:
            if char=='(':depth+=1
            elif char==')':depth-=1
            elif char==',' and depth==0:expressions.append(select[start:i].strip());start=i+1
    expressions.append(select[start:].strip())
    mapping={}
    for expression in expressions:
        match=re.search(r'\bAS\s+(\w+)\s*$',expression,re.I)
        if match:mapping[match.group(1)]=expression
    executed_rows=[]
    for r in read(OUT/'documentacion/codificaciones-modelos.json')['rows']:
        encoding.append([r['period'],r['model'],r['component'],r['feature'],r['reference'],'; '.join(r['levels']),'; '.join(r['pool_keep'] or [])])
        f=r['feature']
        transform=mapping.get(f,'')
        if r['component']=='QVoice':transform='Indicador favorable 1 → v1, 0 → v0. missing, no_linked_row, different_employment_episode, no_available_wave y question_absent se conservan. Wave Y → fotos Y+1, mismo episodio laboral. Sin imputar ausencia como 0 favorable.'
        elif f=='calendar_month':transform='Mes calendario de snapshot_date como categoría; avanza durante los seis hazards del R6.'
        elif f=='level_x_tenure':transform='Indicador de combinación level_grouped × tenure_playbook_band; categorías observadas, junto a los efectos principales.'
        assert transform, 'Falta definición ejecutada: '+f
        pooling='BU de menos de 300 personas de entrenamiento → (other); conservar pool aprendido' if f=='bu_l3' else 'Sin pool adicional en CQ/FQ de esta revisión; niveles observados del encoder'
        executed_rows.append({'period':r['period'],'model':r['model'],'component':r['component'],'feature':f,'expression_or_rule':transform,'pooling':pooling})
        executed.append([r['period'],r['model'],r['component'],f,transform,pooling,'Indicadores cero para el término; permanecen intercepto y otros términos'])
    write(OUT/'documentacion/tratamientos-ejecutados.json',{'source_query':'queries/profiles/engineering_bu_l3_person_month_20261012.sql','rows':executed_rows})
    for sheet in wb:style_sheet(sheet)
    target=OUT/'documentacion/variables-44.xlsx';wb.save(target)
    check=load_workbook(target,read_only=True);assert check.worksheets[0].max_row==45 and check.worksheets[1].max_row==265;check.close()
    save_md('documentacion/variables-44-final.md',['# Tabla final de las 44 filas','',
        'Modo `presentation`. Las 43 filas V11 más BU L3 no equivalen a 44 predictores. [Excel completo](variables-44.xlsx), [tratamientos y usos JSON](variables-44-final.json), [codificaciones exactas](codificaciones-modelos.json). La [primera entrega](variables-44.md) se conserva sin reescribir.','',
        table(['#','Variable y representaciones','Uso como predictor en cada periodo','Acción candidata / responsable','Detalle'],[
          [str(r['numero']),r['variable']+'; '+', '.join(r['representaciones']),'<br>'.join(k+': '+(', '.join(v['predictor']) or 'sin predictor') for k,v in r['uso_por_modelo_corte'].items()),r['accion_candidata']+' / '+r['responsable_funcion'],', '.join('[Ficha '+str(i+1)+']('+Path(os.path.relpath(ROOT/p,OUT/'documentacion')).as_posix()+')' for i,p in enumerate(r['detalle']))] for r in d['rows']]),'',
        'Las transformaciones, faltantes, codificaciones y decisiones por corte están en el Excel y JSON. Calendario e interacción nivel × antigüedad son términos auxiliares, documentados en Codificaciones de modelos. Las preguntas Q1/Q2/Q3/Q13/Q17 se incluyen como componentes de la vista de preguntas comparables, sin crear filas nuevas. No se modifican scores.'])


def charts(result):
    dest=OUT/'charts';dest.mkdir(exist_ok=True)
    plt.rcParams.update({'font.size':11,'axes.spines.top':False,'axes.spines.right':False,'axes.spines.left':False,'axes.titlelocation':'left'})
    fig,ax=plt.subplots(figsize=(9,4.8))
    names=list(LABELS);ys=np.arange(3)
    for i,name in enumerate(names):
        delta=0 if name=='CQ' else result['paired_differences'][name+'_minus_CQ']['30']['log_loss']['difference']*1000
        bounds=[0,0] if name=='CQ' else [v*1000 for v in result['paired_differences'][name+'_minus_CQ']['30']['log_loss']['ci95']]
        ax.errorbar(delta,i,xerr=[[max(0,delta-bounds[0])],[max(0,bounds[1]-delta)]],fmt='o',color=COLORS[name],capsize=5,markersize=8)
        ax.annotate(fmt(delta)+'  '+ci(bounds),(delta,i),xytext=(8,13),textcoords='offset points',fontsize=10)
    ax.axvline(0,color='#777777',linewidth=1);ax.set_yticks(ys,[LABELS[n] for n in names]);ax.invert_yaxis();ax.set_ylim(2.6,-.65)
    ax.xaxis.set_major_formatter(FuncFormatter(lambda x,_:fmt(x)))
    ax.set_xlabel('Diferencia de log-loss × 1000 frente a CQ; menor es mejor')
    ax.set_title('La precisión se compara con intervalos pareados',fontweight='bold',pad=24)
    fig.text(.02,.025,'Engineering · dos periodos con igual peso · IC95: 500 remuestreos por persona\nFuente: results.json y SQL guardado · revisión de octubre de 2026; retrospectiva',fontsize=9)
    fig.tight_layout(rect=[0,.14,1,1]);fig.savefig(dest/'diferencia-log-loss.png',dpi=160);plt.close(fig)
    fig,ax=plt.subplots(figsize=(9,4.8))
    for name in names:
        vals=[result['summary'][name][cap]['capture_percent'] for cap in ['10','20','30']]
        ax.plot([10,20,30],vals,'o-',label=LABELS[name],color=COLORS[name])
        ax.annotate(LABELS[name]+' '+fmt(vals[-1])+'%',(30,vals[-1]),xytext=(8,(-16 if name=='CQ' else 6 if name=='FQ' else 28)),textcoords='offset points',arrowprops={'arrowstyle':'-','color':COLORS[name],'lw':.7})
    ax.set_xlim(8,39);ax.set_ylim(0,max(result['summary'][n]['30']['capture_percent'] for n in names)*1.2)
    ax.set_xticks([10,20,30],['10%','20%','30%']);ax.yaxis.set_major_formatter(FuncFormatter(lambda x,_:fmt(x)+'%'))
    ax.set_xlabel('Capacidad exacta de atención por grupos');ax.set_ylabel('Captura de exposiciones con salida')
    ax.set_title('La captura cambia con la capacidad de atención',fontweight='bold',pad=20)
    fig.text(.02,.025,'Mismos grupos: antigüedad × salidas del equipo × nivel; padres ≥300\nPonderación uniforme del grupo límite; no son salidas únicas de una cohorte conjunta',fontsize=9)
    fig.tight_layout(rect=[0,.14,1,1]);fig.savefig(dest/'capacidad-captura.png',dpi=160);plt.close(fig)


def report(result):
    summary=result['summary'];diff=result['paired_differences'];periods=result['periods']
    pairs=[]
    for name in LABELS:
        m=summary[name]['30'];pairs.append([LABELS[name],fmt(m['log_loss']*1000),fmt(m['brier']*1000),fmt(m['capture_percent'])+'%',fmt(m['group_snapshot_mae_pp'])+' pp'])
    best=min(LABELS,key=lambda n:summary[n]['30']['log_loss'])
    primary_pass=[r for r in result['subdivision'] if r['minimum']==100 and r['passes_both_periods']]
    lines=['# Revisión de Julia: estimación, grupos y preparación para RH','',
           f"**{LABELS[best]} obtiene el menor log-loss medio de esta comparación.** La recomendación predictiva se separa de la adopción: CQ sigue elegido. "+('Ninguna subdivisión principal de 100 personas cumple todas las guardas en ambos periodos.' if not primary_pass else 'Los escenarios de subdivisión que cumplen todas las guardas se identifican abajo.'),'',
           'Modo `presentation`. Resultado retrospectivo en Engineering, Employee, Regular, activa y sin RF360. Fotos de agosto–diciembre de 2025 y enero–febrero de 2026; salida voluntaria a seis meses. Cada periodo pesa la mitad. CQ conserva su condición de modelo elegido; esta ejecución evalúa candidatos y escenarios.', '',
           '**Comparabilidad del ajuste:** CQ histórico conserva su base RH de 1-ago-2025 en ambos periodos. Las bases FQ/FQ + BU se reajustan a 1-ago-2025 y 1-ene-2026. Se comparan estimadores completos; la diferencia de enero frente a CQ reúne ciudad y actualización de la base, por lo que no aísla el aporte de ciudad. FQ frente a FQ + BU sí comparte origen RH en ambos periodos. Los tres comparten selección QVoice, población, fechas de evaluación y grupos.', '',
           table(['Modelo','Log-loss ×1000 ↓','Brier ×1000 ↓','Captura al 30% ↑','Error de grupo-mes ↓'],pairs),'',
           '![Diferencias de estimación](charts/diferencia-log-loss.png)','']
    for name,base in [('FQ','CQ'),('FQ_BU_L3','CQ'),('FQ_BU_L3','FQ')]:
        d=diff[name+'_minus_'+base]['30']
        lines.append(f"{LABELS[name]} frente a {LABELS[base]}: cambio de log-loss ×1000 {fmt(d['log_loss']['difference']*1000)}, IC95 {ci(d['log_loss']['ci95'],1000)}; cambio de captura {fmt(d['capture_percent']['difference'])} pp, IC95 {ci(d['capture_percent']['ci95'])}. El empate o una diferencia pequeña de captura no invalida por sí solo una mejora de estimación.")
        lines.append('')
    lines+=['## Capacidad y esfuerzo','', '![Capacidad y captura](charts/capacidad-captura.png)','',
            'La capacidad ponderada es exacta en cada foto. Para conversaciones de RH se incluyen también los escenarios de grupos completos: excluir o incluir todo el bloque empatado que cruza el límite. Las personas efectivas ponderadas son capacidad equivalente; no una lista de personas seleccionadas.', '',
            table(['Modelo','Capacidad','Captura','Error grupo-mes','Personas/foto exactas','Grupos/foto','Rango de personas con grupos completos'],[
              [LABELS[name],cap+'%',fmt(summary[name][cap]['capture_percent'])+'%',fmt(summary[name][cap]['group_snapshot_mae_pp'])+' pp',
               fmt(np.mean([np.mean([r['effective_people_selected'] for r in p['models'][name]['capacities'][cap]['monthly']]) for p in periods])),
               fmt(np.mean([np.mean([r['priority_groups'] for r in p['models'][name]['capacities'][cap]['monthly']]) for p in periods])),
               fmt(min(r['whole_groups_excluding_people'] for p in periods for r in p['models'][name]['capacities'][cap]['monthly']))+'–'+fmt(max(r['whole_groups_including_people'] for p in periods for r in p['models'][name]['capacities'][cap]['monthly']))]
              for name in LABELS for cap in ['10','20','30']]),'',
            '## Subdivisión selectiva por BU L3','',
            'Se conservan los padres no priorizados y se subdividen sólo los padres inicialmente priorizados. Se mantienen al menos dos hijos BU identificables con soporte, reuniendo categorías restantes en un remanente; si éste queda pequeño se conserva el padre. La prioridad se vuelve a ordenar sobre la partición completa para mantener capacidad.', '',
            table(['Modelo','Capacidad','Mínimo hijo','Mediana de personas/grupo priorizado','Captura','Error','Pasa en ambos periodos'],[
              [LABELS[r['model']],str(r['capacity'])+'%',str(r['minimum']),fmt(r['median_people']),fmt(r['capture_percent'])+'%',fmt(r['error_pp'])+' pp',fmt(r['passes_both_periods'])] for r in result['subdivision']]),'',
            'Las guardas exigen mediana menor, capacidad conservada, pérdida de captura ≤2% relativo, aumento de error ≤0.1 pp y ningún aumento del verdadero REST. Las filas de 100 personas son el escenario principal; 150 y 200 son sensibilidades, sin escoger retrospectivamente un umbral para declarar éxito.', '',
            '## QVoice: ausencia, información parcial y cobertura','']
    for p in periods:
        lines += ['### '+p['period']['id'],'']
        for name in LABELS:
            rows=p['models'][name]['diagnostics']['qvoice']
            lines += [f'**{LABELS[name]}**. Cobertura = personas de la foto con todos los campos requeridos válidos / todas las personas de la foto. Se informa además disponibilidad por campo en las fichas; no equivale a participación entre invitados.','',
                      table(['Estado','Personas distintas en el periodo','Observaciones persona-mes','Riesgo medio','Observado','Error absoluto'],[
                         [r['cell'],str(r.get('distinct_people','Suprimido')),str(r.get('person_month_observations','Suprimido')),fmt(r.get('predicted_risk_percent')),fmt(r.get('observed_percent')),fmt(r.get('absolute_error_pp'))] for r in rows]),'']
            for r in p['models'][name]['qvoice_vs_RH_same_cohorts']:
                if r.get('suppressed'):continue
                lines += [f"En información **{r['state']}**, el cambio de riesgo QVoice menos RH es {fmt(r['mean_Q_minus_RH_risk_pp'])} pp; cambio de log-loss ×1000 {fmt(r['log_loss_delta']*1000)}, IC95 {ci(r['log_loss_delta_ci95'],1000)}. Coincidencia exacta con RH: {fmt(r['equals_RH'])}.",'']
    lines += ['Los estados `missing`, `no_linked_row`, `different_employment_episode`, `no_available_wave` y `question_absent` se conservan según la fuente. Un valor nunca aprendido deja en cero los indicadores de ese término, manteniendo intercepto y otros términos. La ausencia no demuestra descontento ni implica volver automáticamente al modelo RH.', '',
              '## Calibración y estabilidad','',table(['Periodo','Modelo','Sesgo global','IC95 sesgo','Etapa R6','Recomendar corrección'],[
                  [p['period']['id'],LABELS[n],fmt(p['models'][n]['calibration']['bias_pp'])+' pp',ci(p['models'][n]['calibration']['bias_ci95_pp']),
                   {'identity_no_trigger':'Identidad: no activa umbral','identity_insufficient_history':'Identidad: menos de 3 orígenes','scalar_tested':'Escalar contrastado'}[p['models'][n]['calibration']['status']],fmt(p['models'][n]['calibration']['recommended'])]
                  for p in periods for n in LABELS]),'',
              'Las cinco bandas usan límites comunes fijados con CQ en la última foto anterior a cada periodo. Los límites, diagnósticos por mes/banda/grupo y supresiones están guardados. Un escalar global no garantiza corregir brechas locales.', '',
              table(['Modelo','Transición','Personas comunes','Cambio absoluto riesgo','Correlación ranking','Solapamiento prioridad ponderada','Cambio de encuesta/modelo'],[
                [LABELS[n],r['from']+' → '+r['to'],str(r['common_people']),fmt(r['mean_absolute_risk_change_pp'])+' pp',fmt(r['rank_spearman']),fmt(r['weighted_priority_jaccard']*100)+'%',fmt(r['wave_transition'])]
                for n in LABELS for r in result['monthly_stability'][n]['30']]),'',
              'El anexo separa altas/bajas, cambio sobre personas comunes, cambio del promedio total y residual de composición. El salto entre diciembre de 2025 y enero de 2026 coincide con una nueva encuesta y un nuevo ajuste; no se atribuye todo el movimiento a cambios de opinión.', '',
              '## Denominadores, privacidad y límites','',
              table(['Periodo','Personas distintas','Observaciones persona-mes','Exposiciones con evento','Personas con evento','Salidas únicas'],[
                  [p['period']['id'],str(p['people']),str(p['observations']),str(p['event_exposures']),str(p['distinct_people_with_event']),str(p['unique_exit_events'])] for p in periods]),'',
              'Las salidas únicas no coinciden exactamente con personas con evento: una persona puede tener más de una salida tras una recontratación. El conteo independiente se guardó en `input-evidence/unique-events.json`. Las fichas que sólo permiten identificar personas con evento no las rotulan como salidas únicas.', '',
              'Los intervalos de tasas observadas, sesgos y diferencias de métricas usan 500 remuestreos de personas compartidos entre modelos, manteniendo juntos todos sus meses y el peso de cada periodo. Las prioridades quedan fijas. El intervalo aproximado del riesgo medio usa el gradiente del R6 y la covarianza agrupada por persona de RH y QVoice, incluyendo su dependencia. Ninguno es un intervalo de predicción individual; no incluyen toda la incertidumbre de elegir variables/grupos ni toda la dependencia entre equipos.', '',
              'Los grupos observados requieren 30 personas y 30 personas con evento. Si hay celdas pequeñas se suprime el bloque observado completo para evitar reconstruirlas por resta; esto incluye eventos, tasas, errores e intervalos. El riesgo medio y cobertura sólo aparecen donde el soporte de población permite publicarlos. Las limitaciones de soporte no equivalen a buen o mal desempeño.', '',
              'IOS en C carece de acreditación histórica de publicación; QVoice usa una política anual, pero su fecha real de publicación no está acreditada. No se afirma que toda la información estuviera operacionalmente disponible en el pasado. La fecha RH, el año de encuesta y la disponibilidad acreditada se mantienen separados.', '',
              '## Preparación para RH','',
              'La comparación permite valorar estimadores y agrupaciones retrospectivas. La preparación para RH exige, además, responsables funcionales, contraste con evidencia organizacional, fechas acreditadas y seguimiento posterior. Las acciones de las fichas son hipótesis: revisar soporte, vacantes, desarrollo y coordinación, sin prometer que cambiar una señal reducirá salidas.', '',
              'Las diferencias descriptivas por región, nivel, antigüedad y QVoice no se etiquetan automáticamente como sesgo injustificado. Sus agregados, representación en prioridad y errores publicables se consultan en [diagnósticos por segmento](segmentos.md).', '',
              '## Archivos y reproducción','',
              '- [Tabla de 44 variables en Excel](documentacion/variables-44.xlsx), con clasificaciones, acciones, tratamientos y categorías extraídas de los modelos.',
              '- [Primera matriz documental](documentacion/matriz-comentarios.md) y [cierre de comentarios](cierre-comentarios.md).',
              '- [Fichas históricas de grupos](fichas/indice.md).',
              '- [Resultados completos](results.json), [verificación](verification.json) y [anexo técnico](anexo-tecnico.md).',
              '- Ejecutable: `tools/engineering_julia_rh_review.py`; configuración: `config/julia_rh_review_20261012.json`; render: `tools/render_julia_rh_review.py`.', '',
              'La ejecución conserva modelos y agregados; ninguna fila, llave ni predicción individual se persiste. Los modelos CQ previos permanecen intactos.']
    save_md('informe.md',lines)


def extended_diagnostics(result,supplement):
    lines=['# Diagnósticos completos y capacidades','',
       'Modo `presentation`. Las medias dan el mismo peso a ambos periodos. Las capturas cuentan exposiciones con evento a seis meses. Capacidad exacta por ponderación uniforme en el grupo límite; los escenarios completos incluyen o excluyen todo el bloque empatado.','',
       '## Separar estimador de priorización','',
       'Mismos grupos, mismas personas y misma prioridad de CQ al 30%; sólo cambia el estimador.','',
       table(['Periodo','Modelo','Error grupo-mes en prioridad fija CQ'],[[p['period']['id'],LABELS[n],fmt(p['models'][n]['on_common_CQ_priority']['group_snapshot_mae_pp'])+' pp'] for p in result['periods'] for n in LABELS]),'',
       '## Grupos completos en el límite','',
       table(['Periodo','Modelo','Capacidad nominal','Escenario','Cobertura real','Captura','Error grupo-mes'],[
         [p['period']['id'],LABELS[n],cap+'%',label,fmt(m.get('coverage_percent'))+'%',fmt(m.get('capture_percent'))+'%',fmt(m.get('group_snapshot_mae_pp'))+' pp']
         for p in result['periods'] for n in LABELS for cap in ['10','20','30'] for key,label in [('low','Excluir frontera'),('high','Incluir frontera')] for m in [p['models'][n]['capacities'][cap][key]]]),'',
       '## QVoice por capacidad sobre las mismas cohortes','',
       'Completa = todos los campos que requiere el modelo son válidos; parcial = algunos; ausente = ninguno. Cobertura usa observaciones persona-mes como denominador en la ventana. La selección se hace globalmente: la fracción priorizada dentro de cada estado puede diferir de 10/20/30%. RH usa su propia prioridad por los mismos grupos, sobre esas mismas cohortes.','',
       table(['Periodo','Modelo','Capacidad','Estado','Cobertura QVoice','Captura Q','Captura RH','Δ captura IC95','Error Q','Error RH'],[
         [r['period'],LABELS[r['model']],str(r['capacity'])+'%',r['state'],fmt(r['coverage_percent'])+'%',fmt(r.get('metrics',{}).get('capture_percent')),fmt(r.get('RH_same_cohort',{}).get('capture_percent')),
          ci(r['Q_minus_RH']['capture_percent']['ci95']) if r.get('Q_minus_RH',{}).get('capture_percent') else 'Suprimido',fmt(r.get('metrics',{}).get('group_snapshot_mae_pp')),fmt(r.get('RH_same_cohort',{}).get('group_snapshot_mae_pp'))] for r in supplement['qvoice_by_capacity']]),'',
       '## Cobertura QVoice en segmentos','',
       table(['Periodo','Modelo','Dimensión','Segmento','Denominador persona-mes','Completa / parcial / ausente'],[
         [r['period'],LABELS[r['model']],r['dimension'],r['segment'],str(r['denominator_person_months']),' / '.join(fmt(v)+'%' for v in r['complete_partial_absent_percent']) if r['complete_partial_absent_percent'] is not None else 'Suprimido'] for r in supplement['segments_qvoice_coverage']]),'',
       '## Calibración R6 condicionada','',
       table(['Periodo','Modelo','Multiplicador','Sesgo antes','Sesgo después','Log-loss antes/después ×1000','Brier antes/después ×1000','Recomendar'],[
         [p['period']['id'],LABELS[n],fmt(c['multiplier']),fmt(c['bias_pp']),fmt(c.get('posterior_bias_pp')),fmt(p['models'][n]['capacities']['30']['metrics']['log_loss']*1000)+' / '+fmt(c['posterior_log_loss']*1000),fmt(p['models'][n]['capacities']['30']['metrics']['brier']*1000)+' / '+fmt(c['posterior_brier']*1000),fmt(c['recommended'])]
         for p in result['periods'] for n in LABELS for c in [p['models'][n]['calibration']] if c['status']=='scalar_tested']),'',
       'Se prueban tres orígenes maduros como máximo por modelo y periodo. La regla exige menor sesgo absoluto sin empeorar ninguna de las dos pérdidas en el periodo posterior. Los resultados principales y las fichas conservan los riesgos sin corrección. Los intervalos están condicionados a la estructura elegida; la activación por el test es retrospectiva.','',
       '## Estabilidad y composición','',
       table(['Modelo','Escenario capacidad/mínimo','Transición','Altas','Bajas','Entradas prioridad equivalentes','Salidas prioridad equivalentes','Cambio personas comunes pp','Residual composición pp','Cambio definición grupo %'],[
         [LABELS[n],s,r['from']+' → '+r['to'],str(r['new_people']),str(r['departed_people']),fmt(r['priority_entry_effective_people']),fmt(r['priority_exit_effective_people']),fmt(r['common_mean_change_pp']),fmt(r['composition_residual_pp']),fmt(r['group_definition_change_percent'])]
         for n,scenarios in result['monthly_stability'].items() for s,rows in scenarios.items() for r in rows]),'',
       'El residual de composición es la diferencia entre el cambio del promedio total y el cambio medio en personas comunes. No es una descomposición causal. Una persona puede cambiar de grupo aunque la regla de agrupación siga igual.','',
       '## Tamaño, cobertura organizacional y REST','',
       table(['Periodo','Modelo','Capacidad/mínimo','Foto','Mediana grupo','Máximo grupo','BU en prioridad','Prioridad en hijos BU %','REST','No priorizados inicialmente','Padres conservados','Remanente','Hijos BU'],[
         [p['period']['id'],LABELS[n],s,r['snapshot'],fmt(r['priority_median_group_people']),str(r['priority_maximum_group_people']),str(r['bu_categories_in_priority']),fmt(r['identifiable_bu_priority_percent']),*[str(r['categories'][k]) for k in ['REST','no_priorizado','padre_conservado','remanente','hijo_BU']]]
         for p in result['periods'] for n in LABELS for s,sc in p['models'][n]['subdivision'].items() for r in sc['monthly']]),'',
       'Los tipos describen la partición inicial: un padre inicialmente no priorizado puede entrar al reordenar la partición final. Ninguna persona aparece simultáneamente en padre e hijo. REST conserva exactamente la población cuyo padre no llega a 300.']
    save_md('diagnosticos-completos.md',lines)


def cards_and_segments(result):
    index=['# Fichas históricas de grupos','', 'Modo `presentation`. Cada fila es una cohorte histórica por foto; se conserva su definición. Los riesgos corresponden a salida voluntaria a seis meses. No constituyen pronósticos actuales ni listas de personas.','']
    segment=['# Diagnósticos de concentración por segmento','', 'Modo `presentation`. Mismas personas, fechas y capacidad del 30%. Pesos y riesgos en porcentaje; error en puntos porcentuales. Una diferencia descriptiva no demuestra sesgo injustificado. Las supresiones se conservan para impedir reconstrucciones.','']
    for p in result['periods']:
        for name,m in p['models'].items():
            file=p['period']['id']+'-'+name+'.md'
            index.append(f'- [{p["period"]["id"]} · {LABELS[name]}]({file})')
            lines=['# '+LABELS[name]+' · '+p['period']['id'],'','Fichas históricas. Modo `presentation`. Los enteros se muestran sin decimales y los otros valores con máximo un decimal.','',
                   'Los nombres técnicos de cortes no describen un perfil individual: «entrada» reúne niveles 7–8 (Senior & Lead y Assoc & Int/Eng); «medio», 5–6; «alto», 1–3. «Con salidas» indica salidas anteriores bajo el manager actual en 12 meses.','']
            lines+=['## Personas distintas durante la ventana','',
              'Una persona puede cambiar de grupo durante la ventana; los conteos por grupo no son sumables. Las fichas siguientes describen cada foto por separado.','',
              table(['Grupo','Personas distintas en ventana','Persona-mes','Número de fotos'],[[r['cell'],str(r['distinct_people']),str(r['person_month_observations']),str(r['snapshot_count'])] for r in m['diagnostics']['group_window'] if not r.get('suppressed')]),'']
            for i,r in enumerate(m['group_cards']):
                if r.get('suppressed'):continue
                lines+=['## '+r['cell'],'',
                  f"Personas en la foto: **{r['distinct_people']}**. Observaciones persona-mes: {r['person_month_observations']}. Riesgo medio R6: **{fmt(r['predicted_risk_percent'])}%**, IC95 aproximado por coeficientes {ci(r['estimated_mean_risk_ci95_delta_percent'])}%.",'',
                  ('Tasa observada: '+fmt(r['observed_percent'])+'%; IC95 por remuestreo de personas '+ci(r['observed_rate_ci95_cluster_percent'])+'%; personas con evento: '+str(r['distinct_people_with_event'])+'; exposiciones con evento: '+str(r['event_exposures'])+'; salidas únicas persona-fecha: '+str(r.get('unique_exit_events') if r.get('unique_exit_events') is not None else 'suprimidas')+'.' if r['observed_publishable'] else r['observed_suppression_reason']),'',
                  (f"QVoice completa/parcial/ausente: {fmt(r['qvoice_complete_percent'])}% / {fmt(r['qvoice_partial_percent'])}% / {fmt(r['qvoice_absent_percent'])}%." if r['qvoice_complete_percent'] is not None else 'Desglose de cobertura QVoice suprimido por soporte pequeño.')+f" Denominador: {r['qvoice_coverage_denominator']} personas de la foto. Wave: {r['survey_wave']}. Fecha de encuesta: desconocida; disponibilidad real: no acreditada. Foto RH: {r['rh_last_snapshot']}.",'',
                  'Prioridad equivalente en la foto: '+fmt(r['selected_effective_person_months']/r['person_month_observations']*100)+'% del grupo. Los grupos con menos de 300 personas pertenecen al verdadero REST y no son elegibles.','',
                  'Tema a validar: '+r['topic_to_validate']+'. Responsable propuesto: '+r['owner_function']+'. Acción candidata: '+r['candidate_action']+'.','',
                  'Indicador: '+r['followup_indicator']+'. Plazo sugerido: '+r['suggested_deadline']+'.','',
                  'La asociación no demuestra causalidad ni describe a cada persona. [Estabilidad mensual](../results.json) y [fuente y método](../anexo-tecnico.md).','']
            save_md('fichas/'+file,lines)
            for dimension in ['region','level','tenure','qvoice','month','risk_band']:
                segment+=['## '+p['period']['id']+' · '+LABELS[name]+' · '+dimension,'',
                    table(['Segmento','Personas distintas','Persona-mes','Peso en población','Peso en prioridad','Riesgo','Observado','Error'],[
                      [r['cell'],str(r.get('distinct_people','Suprimido')),str(r.get('person_month_observations','Suprimido')),fmt(r.get('population_share_percent')),fmt(r.get('priority_share_percent')),fmt(r.get('predicted_risk_percent')),fmt(r.get('observed_percent')),fmt(r.get('absolute_error_pp'))] for r in m['diagnostics'][dimension]]),'']
    save_md('fichas/indice.md',index);save_md('segmentos.md',segment)


def technical(result):
    repro=result['reproduction']
    save_md('anexo-tecnico.md',['# Anexo reproducible','',
        'Modo `presentation`. SQL principal: `queries/profiles/engineering_bu_l3_person_month_20261012.sql`; metadata y join BU en `input-evidence/bu-l3-*.json`; eventos únicos en `input-evidence/unique-events.json`. Consulta read-only con la autorización de persona-mes sólo en RAM de 1-oct-2026.', '',
        'El manifiesto conserva fechas, configuración, hashes y tiempos. La carpeta `source/` congela ejecutables y especificaciones. Los modelos históricos CQ se restauran; FQ y FQ + BU se ajustan conservando variables y L2. No se ejecuta otra búsqueda de variables.', '',
        table(['Periodo','Métrica','Archivado','Reproducido','Diferencia','Tolerancia 1e-8'],[
          [r['period'],r['metric'],fmt(r['archived']),fmt(r['current']),fmt(r['difference']),fmt(r['within_1e_8'])] for r in repro]),'',
        'El `results.json` conserva precisión numérica para auditar estas diferencias; el texto usa como máximo un decimal. La covarianza se calcula mediante ecuaciones apiladas: bloque RH, bloque QVoice y derivada cruzada del offset. La matriz meat suma scores centrados por persona; bread incluye curvatura observada y penalización. Se usa pseudoinversa con tolerancia declarada y se guarda rango/eigenvalor mínimo.', '',
        'El gradiente del riesgo conserva el avance de calendario de los seis hazards y el recorte de cada predictor lineal. La validación sintética compara derivadas con diferencias finitas y comprueba la invariancia del sandwich al duplicar filas y penalización dentro de la misma persona.', '',
        'Bootstrap: multinomial exacto de personas, semilla 20261012, 500 réplicas compartidas. Las métricas se recalculan por réplica y después se promedian los dos periodos con peso 1/2. Intervalos percentiles 2.5%/97.5%, condicionados a modelos y prioridades fijas. No se guarda la matriz de pesos ni filas.', '',
        'La subdivisión considera un grupo base válido con al menos 300 personas por foto. Los mínimos hijos 100/150/200 se aplican a las categorías de BU reales; `(unknown)` y `(other)` se agrupan en remanente. Exige al menos dos hijos identificables; un remanente positivo por debajo del mínimo hace conservar el padre. REST se determina sólo por invalidez del padre base, y por eso su población se conserva exactamente.', '',
        'Reproducción de cálculo en una carpeta nueva: copie la configuración con otro `run_id` y `output`, prepare allí la entrega documental, ejecute el cálculo y después el render sobre esa salida. `--resume-models` recupera modelos guardados tras una interrupción, pero exige volver a cargar datos en RAM porque no existen cachés individuales.', '',
        'Pruebas: `.venv/Scripts/python.exe -B -m unittest discover -s tests -p test_julia_review_statistics.py -v`. Render: `.venv/Scripts/python.exe -B tools/render_julia_rh_review.py`. No hay cambios en reglas o contratos de los skills; esta ejecución no es un assessment y no asigna scores.', '',
        'La lectura principal ajustó cuatro bases RH y cuatro correcciones QVoice, además de nueve bases y nueve correcciones auxiliares para el escalar (tres orígenes de enero por modelo). CQ principal se restauró. El suplemento hizo una lectura adicional enfocada de las fechas de evaluación y QVoice, sin ajustes, porque los primeros agregados no conservaban captura/error por estado y observado publicable en fichas. Esto se aparta del objetivo de una sola carga compartida, conservando el máximo de ajustes principales. SQL y seis reconciliaciones exactas están en `cohort-diagnostics-supplement.json`.', '',
        'Salidas únicas por ficha: `queries/audits/julia_review_unique_group_events_20261012.sql`, agregado persona-fecha con los mismos cortes; se reconcilian personas y exposiciones con evento contra RAM antes de publicar. Una salida puede aparecer en varias fotos de seis meses: no se suman las fichas para obtener salidas únicas de toda la ventana. Los márgenes de conteos de ventana se suprimen cuando permitirían reconstruir una foto de menos de 30 personas.', '',
        'Pendientes que limitan uso por RH: acreditar publicación histórica de IOS y fecha real de publicación QVoice; confirmar responsables funcionales y capacidad con RH; validar prospectivamente una cohorte nueva. La supresión es por particiones publicadas; no equivale a una garantía formal de anonimización frente a otras bases externas.'])
    states=[
        ('Resultado accionable','Fichas por grupo con función responsable y seguimiento','Los temas son hipótesis; contrastar con RH en 30 días.'),
        ('Ausencia QVoice','Estados separados, riesgo y diferencias de pérdidas frente a RH en mismas cohortes','Soporte pequeño suprimido; cobertura no es tasa de participación.'),
        ('44 variables','Excel y JSON con usos por periodo, acciones y categorías de modelos','B V11 es contexto; no se atribuye su selección a CQ/FQ.'),
        ('Personas y salidas','Conteo independiente separa exposiciones, personas y eventos, también por ficha publicable','No sumar fotos de horizontes superpuestos; márgenes restringidos suprimidos.'),
        ('Capacidad','10/20/30 exacto y grupos completos de frontera','El esfuerzo depende también de responsables y número de grupos.'),
        ('Calibración','Sesgo e IC por periodo; bandas previas y etapa R6 condicionada','No promete corrección de todas las brechas locales.'),
        ('Estabilidad','Ranking, riesgo, prioridad y composición mensual','La transición de enero cambia encuesta y modelo.'),
        ('Concentración','Región, nivel, antigüedad y QVoice con soporte','Una diferencia no se califica automáticamente como sesgo injustificado.'),
        ('Subdivisión BU','Partición corregida y guardas por modelo/capacidad/mínimo','El verdadero REST no mezcla población fuera de prioridad.'),
        ('Uso organizacional','Agregados, supresión complementaria y fichas','Pendiente acreditación temporal y validación prospectiva antes de uso operativo.')]
    save_md('cierre-comentarios.md',['# Cierre de comentarios','',table(['Comentario','Entrega realizada','Límite y siguiente comprobación'],states)])
    save_md('resumen-julia-rh.md',['# Resumen para Julia y RH','',
        '**FQ + BU L3 estima mejor en promedio; CQ continúa elegido.** FQ reduce log-loss ×1000 de 137.7 a 136.4; añadir BU L3 a FQ lo reduce a 136. Los intervalos pareados del cambio excluyen cero. Al atender 30%, los tres capturan alrededor de 43.1% de las exposiciones con salida. El empate de captura no elimina la mejora de estimación. En enero CQ conserva su base RH de agosto y FQ la reajusta: la comparación frente a CQ no aísla el aporte de ciudad.','',
        '**Conservar la agrupación principal.** La subdivisión de 100 personas no cumple las guardas en ambos periodos. Con FQ + BU L3 y capacidad 10%, las sensibilidades de 150 y 200 sí cumplen; son escenarios limitados para validar, sin adoptar retrospectivamente un mínimo como ganador general. El verdadero REST se conserva.','',
        '**No aplicar el escalar R6 probado.** La corrección aprendida en tres orígenes anteriores empeora el sesgo absoluto y log-loss de enero–febrero de 2026 para los tres modelos. Las brechas locales siguen visibles.','',
        '**La ausencia QVoice requiere atención específica.** No activa automáticamente el modelo RH. Se separan completa/parcial/ausente y se comparan riesgo, pérdidas y captura frente a RH sobre las mismas cohortes. El cambio de diciembre a enero coincide con nueva encuesta y nuevo modelo.','',
        '**Cómo empezar con RH.** Revisar las fichas históricas de los grupos priorizados: continuidad después de salidas, soporte de incorporación o acceso a desarrollo, según el grupo. RH y la función operativa propuesta contrastan la hipótesis en 30 días y siguen los acuerdos a 90 días. No se atribuye causalidad ni se interpretan personas individualmente.','',
        'Antes de un uso operativo, acreditar disponibilidad histórica de IOS y fecha real de publicación QVoice, confirmar capacidad y responsables, y evaluar una cohorte posterior al ajuste. Los resultados son retrospectivos y condicionales a las variables y reglas elegidas.','',
        '- [Informe y comparación](informe.md).\n- [44 variables en Excel](documentacion/variables-44.xlsx) y [tabla final con enlaces](documentacion/variables-44-final.md).\n- [Matriz inicial de comentarios](documentacion/matriz-comentarios.md) y [cierre](cierre-comentarios.md).\n- [Fichas históricas](fichas/indice.md).\n- [Diagnósticos completos](diagnosticos-completos.md), [segmentos](segmentos.md) y [anexo reproducible](anexo-tecnico.md).', '',
        'Validación: CQ histórico reproducido en ambos periodos; capacidad, población y particiones reconciliadas; 500 remuestreos por persona y siete pruebas numéricas/sintéticas. Los intervalos observados y del riesgo medio tienen etiquetas y métodos distintos. Sólo modelos y agregados persistidos.'])


def main():
    result=read(OUT/'results.json')
    supplement=read(OUT/'cohort-diagnostics-supplement.json')
    for p in result['periods']:
        for n,m in p['models'].items():
            updates={r['cell']:r for r in supplement['cards_observed'][p['period']['id']+'/'+n]}
            for r in m['group_cards']:
                if r['cell'] in updates:r.update(updates[r['cell']])
                if r.get('suppressed'):continue
                if 'con salidas' in r['cell']:
                    r.update(topic_to_validate='Continuidad y carga de trabajo tras salidas del equipo',owner_function='RH y liderazgo operativo',candidate_action='Contrastar carga, vacantes y continuidad; acordar refuerzos o redistribución si se confirma la necesidad',followup_indicator='Vacantes y acuerdos de capacidad resueltos; cobertura y error en la siguiente cohorte madura')
                elif '<1 año' in r['cell'] or '1-2 años' in r['cell']:
                    r.update(topic_to_validate='Integración y claridad de expectativas en los primeros años',owner_function='RH, onboarding y liderazgo operativo',candidate_action='Revisar soporte de incorporación y claridad del rol mediante escucha agregada',followup_indicator='Cobertura del acompañamiento y acuerdos de incorporación; revisión de la siguiente cohorte madura')
                else:
                    r.update(topic_to_validate='Desarrollo y continuidad en el rol',owner_function='RH y Talent Development',candidate_action='Contrastar acceso a desarrollo y oportunidades internas mediante escucha agregada',followup_indicator='Acciones de desarrollo acordadas y su seguimiento; revisión de la siguiente cohorte madura')
    documents(result);charts(result);report(result);extended_diagnostics(result,supplement);cards_and_segments(result);technical(result)
    source=OUT/'source/tools'
    source.mkdir(parents=True,exist_ok=True)
    import shutil
    shutil.copy2(Path(__file__),source/Path(__file__).name)
    write(OUT/'report-review.json',{'reporting_mode':'presentation','language':'es','graph_methods':'Puntos e IC para diferencias pequeñas; líneas para capacidad. No barras con origen recortado.',
          'classification':'Comparación retrospectiva; no recomendación causal','tables_44_rows':True,
          'observed_and_estimated_intervals_separate':True,'pending_visual_review':True})
    print('Informe, Excel, fichas, segmentos y gráficos generados desde agregados.')


if __name__=='__main__':
    import argparse
    parser=argparse.ArgumentParser();parser.add_argument('--output',type=Path,default=OUT)
    OUT=parser.parse_args().output.resolve();main()

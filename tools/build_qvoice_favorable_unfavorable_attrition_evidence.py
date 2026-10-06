"""Freeze aggregate QVoice Favorable / No favorable attrition evidence.

The raw QVoice workbooks are processed only in memory. Participant identifiers
are used only for deduplication and in ephemeral aggregate Redshift joins; they
are never printed or written to the evidence artifact. The output contains
only aggregate cohort counts, response shares, exits, and rates.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import re
import subprocess
import sys
from collections import defaultdict
from datetime import date
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from build_qvoice_comment_evidence import _iter_selected_rows, _norm_header
from qvoice_excel_assessment import coerce_number, header_base, is_binary_header


ROOT = Path(__file__).resolve().parents[1]
QVOICE_ROOT = ROOT / "data" / "QVoice"
HELPER = ROOT / ".codex" / "skills" / "attrition-redshift-db" / "scripts" / "run_readonly_query.py"
PYTHON = ROOT / ".venv" / "Scripts" / "python.exe"
WINDOW_AUDIT = ROOT / "queries" / "audits" / "qvoice_favorable_unfavorable_attrition_window_audit_20260921.sql"
DEFAULT_OUTPUT = ROOT / "output" / "distribution" / "json" / "qvoice-favorable-unfavorable-question-attrition-evidence-20260923.json"
MINIMUM_COHORT = 30
RUNTIME_BATCH_SIZE = 12_000


# No fuzzy matching is used. `header_base` only removes documented technical
# suffixes/prefixes and normalizes case/spacing; every semantic phrase below
# must resolve exactly once in every requested workbook.
INTENT_COMPONENT_BASES = {
    "Q4": "i see myself working at qualcomm in a few years",
    "Q12": "i rarely think about looking for a new job with another company",
}
CAREER_BASES = {
    2024: "overall, i feel that my career goals can be met at qualcomm",
    2025: "i feel that my career goals can be met at qualcomm",
    2026: "i feel that my career goals can be met at qualcomm",
}
COMPONENT_DISPLAY = {
    "intent_to_stay": (
        {
            "key": "Q4",
            "short_label": "Q4 · See myself at Qualcomm in a few years",
            "question": "I see myself working at Qualcomm in a few years",
        },
        {
            "key": "Q12",
            "short_label": "Q12 · Rarely think about looking for another job",
            "question": "I rarely think about looking for a new job with another company",
        },
    ),
    "career_goals": (
        {
            "key": "Career goals item",
            "short_label": "Career goals can be met at Qualcomm",
            "question": "I feel that my career goals can be met at Qualcomm",
        },
    ),
}
WAVES = (
    {
        "year": 2024,
        "wave": "Jan-2024",
        "file": "Jan 2024 Qvoice raw data.xlsx",
        "sheet": "Raw data",
        "key_column": "Participant's Unique Identifier",
    },
    {
        "year": 2025,
        "wave": "Jan-2025",
        "file": "Jan 2025_Qvoice_raw data.xlsx",
        "sheet": "Raw data",
        "key_column": "EmpID",
    },
    {
        "year": 2026,
        "wave": "Jan-2026",
        "file": "Jan 2026 Qvoice_raw data.xlsx",
        "sheet": "Raw data",
        "key_column": "Emp ID",
    },
)


def require(condition: bool, message: str) -> None:
    if not condition:
        raise ValueError(message)


def sha256_path(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def parse_employee_id(value: object) -> int | None:
    text = re.sub(r"\s+", "", str(value or ""))
    if not text:
        return None
    if re.fullmatch(r"\d+", text):
        return int(text)
    if re.fullmatch(r"\d+\.0+", text):
        return int(float(text))
    return None


def binary_value(value: object) -> tuple[int | None, str]:
    """Return valid binary value; sentinel 255 is missing, not no-favorable."""
    if value is None or (isinstance(value, str) and not value.strip()):
        return None, "missing"
    number = coerce_number(value)
    if number is None or not math.isfinite(number):
        return None, "invalid"
    if number == 255:
        return None, "sentinel_255"
    if number in (0, 1):
        return int(number), "valid"
    return None, "invalid"


def json_query(*, sql: str | None = None, sql_file: Path | None = None, max_rows: int = 100) -> dict[str, Any]:
    """Run a guarded presentation-mode aggregate query without emitting runtime SQL."""
    require((sql is None) != (sql_file is None), "Specify exactly one SQL source.")
    command = [
        str(PYTHON if PYTHON.exists() else sys.executable),
        str(HELPER),
        "--max-rows",
        str(max_rows),
        "--reporting-mode",
        "presentation",
    ]
    if sql_file is not None:
        command.extend(["--sql-file", str(sql_file)])
    result = subprocess.run(
        command,
        cwd=ROOT,
        input=sql,
        text=True,
        capture_output=True,
        encoding="utf-8",
        check=False,
    )
    if result.returncode != 0:
        # Do not surface helper diagnostics because runtime SQL contains IDs.
        raise RuntimeError("The guarded aggregate Redshift query did not complete.")
    payload = json.loads(result.stdout)
    if payload.get("truncated"):
        raise ValueError("The aggregate query was truncated.")
    return payload


def resolve_binary_columns(config: dict[str, object]) -> tuple[list[str], dict[str, int]]:
    """Resolve the three display inputs by exact normalized header base."""
    path = QVOICE_ROOT / str(config["file"])
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook[str(config["sheet"])]
        headers = [str(value or "").replace("\n", " ").strip() for value in next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True))]
    finally:
        workbook.close()

    key_column = str(config["key_column"])
    require(key_column in headers, f"{config['wave']} is missing its declared participant key.")
    expected = {
        "intent_q4": INTENT_COMPONENT_BASES["Q4"],
        "intent_q12": INTENT_COMPONENT_BASES["Q12"],
        "career": CAREER_BASES[int(config["year"])],
    }
    resolved: dict[str, int] = {"participant_key": headers.index(key_column)}
    for semantic_key, base in expected.items():
        matches = [
            index
            for index, header in enumerate(headers)
            if is_binary_header(header) and header_base(header) == base
        ]
        require(len(matches) == 1, f"{config['wave']} must resolve exactly one binary header for {semantic_key}.")
        resolved[semantic_key] = matches[0]
    require(len({resolved["intent_q4"], resolved["intent_q12"], resolved["career"]}) == 3, f"{config['wave']} resolved overlapping item columns.")
    return headers, resolved


def scan_wave(config: dict[str, object]) -> tuple[dict[str, set[int]], dict[str, Any]]:
    """Create mutually exclusive raw response groups without persisting any ID."""
    headers, columns = resolve_binary_columns(config)
    seen: dict[int, tuple[tuple[int | None, str], tuple[int | None, str], tuple[int | None, str]]] = {}
    source_rows = missing_key_rows = invalid_key_rows = duplicate_rows = duplicate_conflicts = 0
    path = QVOICE_ROOT / str(config["file"])
    selected_headers = {_norm_header(headers[columns[key]]) for key in columns}
    normalized_columns = {key: _norm_header(headers[index]) for key, index in columns.items()}
    # The data pass uses the streaming XML reader and holds only the four
    # declared columns. This avoids materializing the entire employee-level
    # workbook while preserving the exact header contract above.
    for header_map, row in _iter_selected_rows(path, str(config["sheet"]), selected_headers):
        source_rows += 1
        key_value = row.get(header_map[normalized_columns["participant_key"]])
        employee_id = parse_employee_id(key_value)
        if employee_id is None:
            if str(key_value or "").strip():
                invalid_key_rows += 1
            else:
                missing_key_rows += 1
            continue
        signature = (
            binary_value(row.get(header_map[normalized_columns["intent_q4"]])),
            binary_value(row.get(header_map[normalized_columns["intent_q12"]])),
            binary_value(row.get(header_map[normalized_columns["career"]])),
        )
        previous = seen.get(employee_id)
        if previous is not None:
            duplicate_rows += 1
            if previous != signature:
                duplicate_conflicts += 1
            continue
        seen[employee_id] = signature

    require(seen, f"{config['wave']} has no usable participant keys.")
    require(duplicate_rows == 0 and duplicate_conflicts == 0, f"{config['wave']} has duplicate participant rows or conflicts.")
    require(len(seen) + missing_key_rows + invalid_key_rows == source_rows, f"{config['wave']} row reconciliation failed.")

    groups = {
        "intent_to_stay:Favorable": set(),
        "intent_to_stay:No favorable": set(),
        "intent_to_stay:Q4:Favorable": set(),
        "intent_to_stay:Q4:No favorable": set(),
        "intent_to_stay:Q12:Favorable": set(),
        "intent_to_stay:Q12:No favorable": set(),
        "career_goals:Favorable": set(),
        "career_goals:No favorable": set(),
        "career_goals:Career goals item:Favorable": set(),
        "career_goals:Career goals item:No favorable": set(),
    }
    quality: dict[str, int] = defaultdict(int)
    for employee_id, (q4, q12, career) in seen.items():
        quality[f"intent_q4_status_{q4[1]}"] += 1
        quality[f"intent_q12_status_{q12[1]}"] += 1
        quality[f"career_status_{career[1]}"] += 1
        if q4[1] == "valid":
            groups[f"intent_to_stay:Q4:{'Favorable' if q4[0] == 1 else 'No favorable'}"].add(employee_id)
        if q12[1] == "valid":
            groups[f"intent_to_stay:Q12:{'Favorable' if q12[0] == 1 else 'No favorable'}"].add(employee_id)
        if q4[1] == "valid" and q12[1] == "valid":
            quality["intent_complete_valid"] += 1
            intent_label = "Favorable" if q4[0] == 1 and q12[0] == 1 else "No favorable"
            groups[f"intent_to_stay:{intent_label}"].add(employee_id)
        elif q4[1] == "valid" or q12[1] == "valid":
            quality["intent_partial_valid_omitted"] += 1
        else:
            quality["intent_without_valid_component_omitted"] += 1
        if career[1] == "valid":
            career_label = "Favorable" if career[0] == 1 else "No favorable"
            groups[f"career_goals:{career_label}"].add(employee_id)
            groups[f"career_goals:Career goals item:{career_label}"].add(employee_id)
            quality["career_valid"] += 1
        else:
            quality["career_not_valid_omitted"] += 1

    for metric in ("intent_to_stay", "career_goals"):
        favorable = groups[f"{metric}:Favorable"]
        no_favorable = groups[f"{metric}:No favorable"]
        require(not favorable.intersection(no_favorable), f"{config['wave']} {metric} groups overlap.")
        expected = quality["intent_complete_valid"] if metric == "intent_to_stay" else quality["career_valid"]
        require(len(favorable) + len(no_favorable) == expected, f"{config['wave']} {metric} groups do not reconcile.")
        require(expected >= MINIMUM_COHORT, f"{config['wave']} {metric} raw cohort is below n={MINIMUM_COHORT}.")

    component_expected = {
        "intent_to_stay:Q4": quality["intent_q4_status_valid"],
        "intent_to_stay:Q12": quality["intent_q12_status_valid"],
        "career_goals:Career goals item": quality["career_status_valid"],
    }
    for component_group, expected in component_expected.items():
        favorable = groups[f"{component_group}:Favorable"]
        no_favorable = groups[f"{component_group}:No favorable"]
        require(not favorable.intersection(no_favorable), f"{config['wave']} {component_group} groups overlap.")
        require(len(favorable) + len(no_favorable) == expected, f"{config['wave']} {component_group} groups do not reconcile.")
        require(expected >= MINIMUM_COHORT, f"{config['wave']} {component_group} raw cohort is below n={MINIMUM_COHORT}.")

    quality.update(
        source_rows=source_rows,
        distinct_usable_participant_count=len(seen),
        missing_key_rows=missing_key_rows,
        invalid_key_rows=invalid_key_rows,
        duplicate_participant_rows=duplicate_rows,
        duplicate_conflicting_rows=duplicate_conflicts,
    )
    return groups, dict(sorted(quality.items()))


def load_windows() -> dict[int, dict[str, Any]]:
    payload = json_query(sql_file=WINDOW_AUDIT, max_rows=10)
    rows = payload.get("rows")
    require(isinstance(rows, list) and len(rows) == 3, "The QVoice favorable/no-favorable window audit must return three years.")
    windows = {int(row["survey_year"]): dict(row) for row in rows}
    require(set(windows) == {2024, 2025, 2026}, "The QVoice favorable/no-favorable window audit returned unexpected years.")
    for year, window in windows.items():
        baseline = int(window["baseline_distinct_employee_count"])
        require(baseline >= MINIMUM_COHORT, f"{year} workforce baseline is below n={MINIMUM_COHORT}.")
        require(int(window["duplicate_baseline_employee_count"]) == 0, f"{year} workforce baseline has duplicate employee rows.")
        if year in (2024, 2025):
            require(window["outcome_type"] == "annual_12m" and int(window["followup_complete"]) == 1, f"{year} lacks complete 12-month follow-up.")
        else:
            require(window["outcome_type"] == "ytd", "Only 2026 may use the YTD outcome contract.")
            require(str(window["baseline_snapshot"]) == "2025-12-28", "Unexpected 2026 pre-wave snapshot.")
            require(str(window["outcome_start"]) == "2026-01-07", "Unexpected 2026 YTD start.")
            require(str(window["outcome_end"])[:10] == "2026-09-27", "Unexpected 2026 YTD cutoff.")
    return windows


def runtime_group_sql(records: list[dict[str, int]], window: dict[str, Any]) -> str:
    """Build ephemeral aggregate SQL. The identifier-bearing statement is never saved."""
    source_json = json.dumps(records, separators=(",", ":")).replace("'", "''")
    return f"""WITH source AS (
  SELECT elem.employee_id::BIGINT AS employee_id, elem.group_code::INTEGER AS group_code
  FROM (SELECT JSON_PARSE('{source_json}') AS payload) AS src,
       src.payload AS elem
), baseline AS (
  SELECT employee_id
  FROM hrda.post_hire_monthly_worker_dtls
  WHERE CAST(fis_mo_end_dt AS DATE) = DATE '{str(window['baseline_snapshot'])[:10]}'
    AND worker_type = 'Employee'
    AND employee_type = 'Regular'
    AND active_status = 1
    AND employee_id IS NOT NULL
  GROUP BY 1
), voluntary_events AS (
  SELECT employee_id
  FROM hrda.post_hire_monthly_worker_dtls
  WHERE worker_type = 'Employee'
    AND employee_type = 'Regular'
    AND employee_id IS NOT NULL
    AND UPPER(NULLIF(TRIM(CAST(term_type AS VARCHAR(256))), '')) = 'VOLUNTARY'
    AND CAST(term_date AS DATE) > DATE '{str(window['outcome_start'])[:10]}'
    AND CAST(term_date AS DATE) <= DATE '{str(window['outcome_end'])[:10]}'
  GROUP BY 1
)
SELECT
  s.group_code,
  COUNT(*) AS source_employee_count,
  SUM(CASE WHEN b.employee_id IS NOT NULL THEN 1 ELSE 0 END) AS matched_employee_count,
  SUM(CASE WHEN b.employee_id IS NOT NULL AND e.employee_id IS NOT NULL THEN 1 ELSE 0 END) AS voluntary_exit_count
FROM source AS s
LEFT JOIN baseline AS b ON b.employee_id = s.employee_id
LEFT JOIN voluntary_events AS e ON e.employee_id = b.employee_id
GROUP BY 1
ORDER BY 1"""


def aggregate_groups(groups: dict[str, set[int]], window: dict[str, Any], year: int) -> dict[str, dict[str, int | float]]:
    labels = list(groups)
    records = [
        {"employee_id": employee_id, "group_code": labels.index(label)}
        for label in labels
        for employee_id in sorted(groups[label])
    ]
    totals: dict[int, dict[str, int]] = {index: defaultdict(int) for index in range(len(labels))}
    for start in range(0, len(records), RUNTIME_BATCH_SIZE):
        payload = json_query(sql=runtime_group_sql(records[start : start + RUNTIME_BATCH_SIZE], window), max_rows=10)
        rows = payload.get("rows")
        require(isinstance(rows, list), f"{year} group join did not return aggregate rows.")
        for row in rows:
            group_code = int(row["group_code"])
            for column in ("source_employee_count", "matched_employee_count", "voluntary_exit_count"):
                totals[group_code][column] += int(row.get(column) or 0)
        print(f"  {year}: linked response cohorts {start + 1:,}-{min(start + RUNTIME_BATCH_SIZE, len(records)):,}", flush=True)

    result: dict[str, dict[str, int | float]] = {}
    for code, label in enumerate(labels):
        source = int(totals[code]["source_employee_count"])
        matched = int(totals[code]["matched_employee_count"])
        exits = int(totals[code]["voluntary_exit_count"])
        require(source == len(groups[label]), f"{year} source total does not reconcile for {label}.")
        require(MINIMUM_COHORT <= matched <= source and 0 <= exits <= matched, f"{year} linked cohort is not publishable for {label}.")
        result[label] = {
            "source_employee_count": source,
            "matched_employee_count": matched,
            "unmatched_employee_count": source - matched,
            "voluntary_exit_count": exits,
            "voluntary_attrition_rate_percent": round(100 * exits / matched, 4),
        }
    return result


def wave_metric_rows(year: int, groups: dict[str, set[int]], window: dict[str, Any]) -> dict[str, dict[str, Any]]:
    joined = aggregate_groups(groups, window, year)
    result: dict[str, dict[str, Any]] = {}
    for metric in ("intent_to_stay", "career_goals"):
        favorable = joined[f"{metric}:Favorable"]
        no_favorable = joined[f"{metric}:No favorable"]
        total = int(favorable["matched_employee_count"]) + int(no_favorable["matched_employee_count"])
        require(total >= MINIMUM_COHORT, f"{year} has no publishable linked {metric} cohort.")
        group_rows = {}
        for label, row in (("Favorable", favorable), ("No favorable", no_favorable)):
            share = round(100 * int(row["matched_employee_count"]) / total, 4)
            group_rows[label] = {**row, "linked_response_share_percent": share}
        require(abs(sum(float(row["linked_response_share_percent"]) for row in group_rows.values()) - 100) < 0.0002, f"{year} {metric} shares do not sum to 100%.")
        components = {}
        for component in COMPONENT_DISPLAY[metric]:
            component_key = str(component["key"])
            component_favorable = joined[f"{metric}:{component_key}:Favorable"]
            component_no_favorable = joined[f"{metric}:{component_key}:No favorable"]
            component_total = int(component_favorable["matched_employee_count"]) + int(component_no_favorable["matched_employee_count"])
            require(component_total >= MINIMUM_COHORT, f"{year} has no publishable linked {metric} component {component_key}.")
            component_rows = {}
            for label, row in (("Favorable", component_favorable), ("No favorable", component_no_favorable)):
                share = round(100 * int(row["matched_employee_count"]) / component_total, 4)
                component_rows[label] = {**row, "linked_response_share_percent": share}
            require(
                abs(sum(float(row["linked_response_share_percent"]) for row in component_rows.values()) - 100) < 0.0002,
                f"{year} {metric} component {component_key} shares do not sum to 100%.",
            )
            components[component_key] = {
                **component,
                "groups": component_rows,
                "linked_response_cohort_count": component_total,
            }
        result[metric] = {
            "year": year,
            "wave": f"Jan-{year}",
            "outcome_type": str(window["outcome_type"]),
            "baseline_snapshot": str(window["baseline_snapshot"]),
            "outcome_start": str(window["outcome_start"]),
            "outcome_end": str(window["outcome_end"])[:10],
            "linked_response_cohort_count": total,
            "groups": group_rows,
            "components": components,
        }
    return result


def build_evidence(output_path: Path = DEFAULT_OUTPUT, *, overwrite: bool = False) -> dict[str, Any]:
    if output_path.exists() and not overwrite:
        raise FileExistsError(f"The frozen evidence already exists: {output_path}")
    require(WINDOW_AUDIT.exists(), "The saved QVoice window audit is missing.")
    windows = load_windows()
    variables: dict[str, list[dict[str, Any]]] = {"intent_to_stay": [], "career_goals": []}
    source_quality: dict[str, dict[str, Any]] = {}

    for config in WAVES:
        year = int(config["year"])
        print(f"Scanning {config['wave']} response groups…", flush=True)
        groups, quality = scan_wave(config)
        print(f"Linking {config['wave']} response cohorts to the pre-wave workforce…", flush=True)
        metrics = wave_metric_rows(year, groups, windows[year])
        for metric, row in metrics.items():
            variables[metric].append(row)
        source_quality[str(year)] = quality

    for metric, rows in variables.items():
        require([int(row["year"]) for row in rows] == [2024, 2025, 2026], f"{metric} is missing a requested year.")
        for row in rows:
            groups = row["groups"]
            require(set(groups) == {"Favorable", "No favorable"}, f"{metric} group names are incomplete.")
            for group in groups.values():
                matched = int(group["matched_employee_count"])
                exits = int(group["voluntary_exit_count"])
                rate = float(group["voluntary_attrition_rate_percent"])
                require(matched >= MINIMUM_COHORT and 0 <= exits <= matched, f"{metric} contains an unpublishable group.")
                require(abs(rate - 100 * exits / matched) < 0.00011, f"{metric} rate does not reconcile.")
            components = row.get("components")
            require(isinstance(components, dict) and set(components) == {str(item["key"]) for item in COMPONENT_DISPLAY[metric]}, f"{metric} components are incomplete.")
            for component in COMPONENT_DISPLAY[metric]:
                component_row = components[str(component["key"])]
                component_groups = component_row.get("groups")
                require(set(component_groups) == {"Favorable", "No favorable"}, f"{metric} component groups are incomplete.")
                component_share = 0.0
                for group in component_groups.values():
                    matched = int(group["matched_employee_count"])
                    exits = int(group["voluntary_exit_count"])
                    rate = float(group["voluntary_attrition_rate_percent"])
                    share = float(group["linked_response_share_percent"])
                    require(matched >= MINIMUM_COHORT and 0 <= exits <= matched, f"{metric} component contains an unpublishable group.")
                    require(abs(rate - 100 * exits / matched) < 0.00011, f"{metric} component rate does not reconcile.")
                    component_share += share
                require(abs(component_share - 100) < 0.0002, f"{metric} component shares do not sum to 100%.")

    evidence = {
        "artifact": "presentation",
        "evidence_type": "aggregate-only QVoice favorable/no-favorable response distribution with time-safe voluntary attrition",
        "generated_on": date.today().isoformat(),
        "outcome": "pass",
        "minimum_shared_cohort": MINIMUM_COHORT,
        "scope": "QVoice raw-workbook respondents in Jan-2024, Jan-2025, and Jan-2026 who have the displayed valid response and match the active Regular pre-wave workforce cohort.",
        "mapping_contract": {
            "method": "Deterministic exact header_base matching after removal of technical binary suffixes/prefixes; no fuzzy matching, synonym substitution, or row persistence.",
            "intent_to_stay": {
                "components": INTENT_COMPONENT_BASES,
                "display_components": COMPONENT_DISPLAY["intent_to_stay"],
                "display_rule": "Favorable requires both valid Q4 and Q12 binary components equal to 1. No favorable requires both components valid and at least one equal to 0. Partial component responses are omitted from this two-bar display.",
                "status": "Operational display mapping; the broader business definition and partial-response treatment remain documented as pending, not score-2 claims.",
            },
            "career_goals": {
                "header_base_by_year": {str(year): value for year, value in CAREER_BASES.items()},
                "display_components": COMPONENT_DISPLAY["career_goals"],
                "display_rule": "Favorable is stored binary value 1; No favorable is stored binary value 0. Invalid, blank, and 255 sentinel values are omitted.",
                "status": "The 2024 header starts with 'Overall'; 2025-2026 differ. The display is each-wave response mix and does not claim normalized cross-wave wording equivalence.",
            },
        },
        "outcome_contract": {
            "2024_2025": "Annual voluntary attrition: matched active Regular employees at the pre-wave monthly snapshot; strict VOLUNTARY exits in (survey-month start, survey-month start + 12 months].",
            "2026": "Voluntary attrition YTD for the same pre-wave cohort in (2026-01-07, 2026-09-27]; not annualized.",
            "population": "worker_type='Employee'; employee_type='Regular'; active_status=1 at each pre-wave baseline snapshot.",
            "grain": "Distinct QVoice respondent in a response group after a time-safe employee-date join.",
        },
        "variables": variables,
        "window_audit": [windows[year] for year in (2024, 2025, 2026)],
        "validation": {
            "outcome": "pass",
            "header_mapping_exactly_once_per_wave": True,
            "all_source_participant_keys_unique": True,
            "all_display_groups_mutually_exclusive": True,
            "linked_response_shares_sum_to_100_percent": True,
            "component_groups_mutually_exclusive": True,
            "component_response_shares_sum_to_100_percent": True,
            "all_component_groups_n_at_least": MINIMUM_COHORT,
            "all_published_groups_n_at_least": MINIMUM_COHORT,
            "historical_waves_have_complete_12_month_followup": True,
            "jan_2026_is_ytd_not_annualized": True,
            "runtime_join_sql_persisted": False,
            "reason_runtime_join_sql_not_persisted": "It contains participant identifiers only in process memory and returns aggregate counts.",
            "source_quality_by_year": source_quality,
        },
        "privacy": {
            "employee_identifiers_persisted": False,
            "employee_rows_persisted": False,
            "individual_qvoice_values_persisted": False,
            "runtime_join_sql_persisted": False,
        },
        "provenance": {
            "window_audit": str(WINDOW_AUDIT.relative_to(ROOT)).replace("\\", "/"),
            "window_audit_sha256": sha256_path(WINDOW_AUDIT),
            "raw_workbooks": [str(config["file"]) for config in WAVES],
        },
    }
    encoded = json.dumps(evidence, ensure_ascii=False)
    for prohibited in ('"employee_id"', '"emp_id"', '"participant_key"'):
        require(prohibited not in encoded, f"Evidence must not persist {prohibited}.")
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(json.dumps(evidence, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return evidence


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Genera evidencia agregada QVoice Favorable / No favorable.")
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--overwrite", action="store_true")
    args = parser.parse_args()
    evidence = build_evidence(args.output, overwrite=args.overwrite)
    summary = {
        "output": str(args.output),
        "variables": list(evidence["variables"]),
        "years": [2024, 2025, 2026],
        "outcome": evidence["outcome"],
    }
    print(json.dumps(summary, ensure_ascii=False))

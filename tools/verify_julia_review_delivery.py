"""Puerta de publicación de los agregados de esta ejecución aún no cerrada."""
from pathlib import Path
from datetime import datetime,timezone
import hashlib
import importlib.util
import json
import re
import sys
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'output/analysis/engineering/20261012/julia-rh-review-v1'
DECK=ROOT/'output/presentations/20261012-2350_revision-julia-rh/presentacion-eng.html'
def read(p):return json.loads(p.read_text(encoding='utf-8-sig'))
def write(p,v):p.write_text(json.dumps(v,ensure_ascii=False,indent=2,allow_nan=False)+'\n',encoding='utf-8')
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()

def guard():
    result=read(OUT/'results.json');changes=[]
    for period in result['periods']:
        for name,model in period['models'].items():
            dates={r['snapshot'] for r in model['capacities']['30']['monthly']}
            present={}
            for card in model['group_cards']:
                if card.get('suppressed'):continue
                date,parent=card['cell'].split(' | ',1)
                present.setdefault(parent,set()).add(date)
            for j,row in enumerate(model['diagnostics']['group_window']):
                if row.get('suppressed'):continue
                # Un total de ventana menos las fotos públicas no debe revelar
                # el headcount de una foto restringida. Exigir todas las fotos.
                if present.get(row['cell'],set())!=dates:
                    changes.append({'period':period['period']['id'],'model':name,'reason':'Margen de ventana con foto no publicable; evita reconstrucción por resta'})
                    model['diagnostics']['group_window'][j]={'cell':'celda_suprimida','suppressed':True,'reason':'Margen de ventana suprimido porque no todas sus fotos son publicables'}
            write(OUT/'periods'/f'{period["period"]["id"]}-{name}.json',model)
    if changes:
        before=sha(OUT/'results.json');write(OUT/'results.json',result)
        write(OUT/'publication-adjustments.json',{'reporting_mode':'presentation','executed_at':datetime.now(timezone.utc).isoformat(),
             'before_sha256':before,'after_sha256':sha(OUT/'results.json'),'adjustments':changes,
             'metrics_changed':False,'note':'Revisión de publicación antes del cierre; sólo suprime márgenes de conteos. No cambia modelos ni métricas.'})
    print('Margenes suprimidos:',len(changes))

def verify():
    r=read(OUT/'results.json');sup=read(OUT/'cohort-diagnostics-supplement.json');checks=[]
    def check(name,truth):
        checks.append({'check':name,'pass':bool(truth)})
        if not truth:raise AssertionError(name)
    check('Reproducción CQ archivado',all(x['within_1e_8'] for x in r['reproduction']))
    check('Cuatro ajustes RH y cuatro Q principales',r['primary_fits']=={'base':4,'q':4})
    check('Seis reproducciones del suplemento; cero ajustes',len(sup['checks'])==6 and max(x['log_loss_reproduction_difference'] for x in sup['checks'])<1e-10 and sup['additional_model_fits']==0)
    check('Ninguna subdivisión principal pasa ambos periodos',not any(x['passes_both_periods'] for x in r['subdivision'] if x['minimum']==100))
    passed={(x['model'],x['capacity'],x['minimum']) for x in r['subdivision'] if x['passes_both_periods']}
    check('Sensibilidades BU restringidas al 10%',passed=={('FQ_BU_L3',10,150),('FQ_BU_L3',10,200)})
    for p in r['periods']:
        for n,m in p['models'].items():
            check(p['period']['id']+'/'+n+' familia excluida','job_family_detail' not in m['base_features'])
            check(p['period']['id']+'/'+n+' escalar sin recomendar',not m['calibration']['recommended'])
            for cap,c in m['capacities'].items():
                check(p['period']['id']+'/'+n+'/'+cap+' capacidad',abs(c['metrics']['coverage_percent']-int(cap))<1e-8)
            for rows in m['diagnostics'].values():
                for row in rows:
                    if row.get('suppressed'):continue
                    check('Cohorte publicada >=30',row['distinct_people']>=30)
                    if row['observed_publishable']:check('Tasa observada con >=30 personas con evento',row['distinct_people_with_event']>=30)
    check('Fichas observadas con >=30 eventos y salidas reconciliadas',all(x['distinct_people_with_event']>=30 and x['unique_exit_events']>=x['event_exposures'] for rows in sup['cards_observed'].values() for x in rows))
    wb=load_workbook(OUT/'documentacion/variables-44.xlsx',read_only=True)
    check('44 filas y 264 usos',wb.worksheets[0].max_row==45 and wb.worksheets[1].max_row==265)
    check('Todas las codificaciones tienen regla ejecutada',wb['Tratamiento ejecutado'].max_row==wb['Codificaciones de modelos'].max_row)
    wb.close()
    docs=read(OUT/'documentacion/variables-44-final.json')['rows']
    check('Todas las filas documentan tratamiento',all(x['tratamientos_existentes'] for x in docs))
    covered={f for x in docs for f in x['representaciones']}|{'calendar_month','level_x_tenure'}
    check('Todos los términos del modelo tienen correspondencia',all(x['feature'] in covered for x in read(OUT/'documentacion/codificaciones-modelos.json')['rows']))
    spec=importlib.util.spec_from_file_location('deck_verify',ROOT/'.agents/skills/build-html-presentation/scripts/verify_deck.py')
    module=importlib.util.module_from_spec(spec);spec.loader.exec_module(module)
    errors,warnings=module.check(DECK.read_text(encoding='utf-8'),'presentation')
    allowed=['El elemento <html> debe declarar lang="es".']
    check('Deck: estructura, privacidad, números y plantilla',not [e for e in errors if e not in allowed])
    check('Deck en inglés según AGENTS',bool(re.search(r'<html[^>]+lang="en"',DECK.read_text(encoding='utf-8'))))
    broken=[]
    for doc in OUT.rglob('*.md'):
        if 'source' in doc.parts or 'input-evidence' in doc.parts:continue
        for link in re.findall(r'\]\(([^)]+)\)',doc.read_text(encoding='utf-8')):
            if '://' in link or link.startswith('#'):continue
            target=(doc.parent/link.split('#')[0]).resolve()
            if not target.exists():broken.append({'file':str(doc.relative_to(OUT)),'link':link})
    check('Enlaces de la entrega existen',not broken)
    write(OUT/'verification-delivery.json',{'status':'pass','checks':checks,'broken_links':broken,
        'static_deck_errors':errors,'static_deck_warnings':warnings,'accepted_exception':'AGENTS.md exige conservar texto proyectado de decks de modelos en inglés. Único error del validador genérico: lang=es.',
        'visual_review_record':'output/preview/julia-rh-review/verification-browser.json',
        'employee_level_rows_persisted':False,'reporting_mode':'presentation'})
    print('Verificaciones de entrega:',len(checks),'correctas')

if __name__=='__main__':
    import argparse
    parser=argparse.ArgumentParser();parser.add_argument('--guard',action='store_true');parser.add_argument('--output',type=Path,default=OUT);parser.add_argument('--deck',type=Path,default=DECK)
    args=parser.parse_args();OUT=args.output.resolve();DECK=args.deck.resolve()
    sys.stdout.reconfigure(encoding='utf-8')
    if args.guard:guard()
    else:verify()

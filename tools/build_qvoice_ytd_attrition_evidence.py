"""Build aggregate-only Jan-2026 QVoice voluntary-attrition YTD evidence.

Employee identifiers and comment text are processed only in memory.  Runtime
SQL joins contain identifiers only in the process input sent to the guarded
read-only helper; that SQL is neither printed nor persisted.  The frozen JSON
contains only cohort-level aggregates for the existing QVoice slide groups.
"""

from __future__ import annotations

import hashlib
import json
import math
import re
import subprocess
import sys
from collections import Counter, defaultdict
from datetime import date
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook

from build_qvoice_comment_evidence import (
    QVOICE_ROOT,
    TAXONOMY as COMMENT_TAXONOMY,
    WAVES as COMMENT_WAVES,
    _classify_theme,
    _find_column as find_comment_column,
    _iter_selected_rows,
    _norm_header,
    _norm_text,
)
from build_qvoice_comparable_favorable_evidence import (
    EXPECTED_COMMON_BASES,
    EXPECTED_WAVES,
    ITEM_COUNT,
    WAVES as COMPARABLE_WAVES,
    _band_for_favorable_count,
    _cell_status,
    _headers_and_binary_bases,
)
from build_qvoice_heartbeat_evidence import read_structure, scan_source
from build_qvoice_nonrespondent_proxy_evidence import WAVES as PARTICIPATION_WAVES, scan_raw_wave
from build_qvoice_normalized_region_evidence import Z_BANDS, z_band_label
from build_qvoice_retention_concerns_evidence import (
    WAVES as RETENTION_WAVES,
    _find_column as find_retention_column,
    classify_themes,
)
from qvoice_excel_assessment import load_survey_structure


ROOT = Path(__file__).resolve().parents[1]
HELPER = ROOT / ".codex" / "skills" / "attrition-redshift-db" / "scripts" / "run_readonly_query.py"
PYTHON = ROOT / ".venv" / "Scripts" / "python.exe"
WINDOW_AUDIT = ROOT / "queries" / "audits" / "qvoice_2026_ytd_voluntary_attrition_window_audit_20260917.sql"
STRUCTURED_QUERY = ROOT / "queries" / "audits" / "qvoice_2026_structured_ytd_attrition_20260917.sql"
RETENTION_EVIDENCE = ROOT / "output" / "distribution" / "json" / "qvoice-retention-concerns-evidence-20260910.json"
OUTPUT = ROOT / "output" / "distribution" / "json" / "qvoice-2026-ytd-attrition-evidence-20260917.json"

SURVEY_LAUNCH = date(2026, 1, 7)
BASELINE_SNAPSHOT = date(2025, 12, 28)
MINIMUM_COHORT = 30


def require(condition: bool, message: str) -> None:
    if not condition:
        raise ValueError(message)


def sha256_path(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def parse_employee_id(value: object) -> int | None:
    value_text = _norm_text(value)
    if not value_text:
        return None
    if re.fullmatch(r"\d+", value_text):
        return int(value_text)
    if re.fullmatch(r"\d+\.0+", value_text):
        return int(float(value_text))
    return None


def json_query(*, sql: str | None = None, sql_file: Path | None = None, max_rows: int = 100) -> dict[str, Any]:
    require((sql is None) != (sql_file is None), "Specify exactly one SQL source.")
    command = [
        str(PYTHON if PYTHON.exists() else sys.executable),
        str(HELPER),
        "--max-rows",
        str(max_rows),
        "--reporting-mode",
        "internal_analysis",
    ]
    if sql_file is not None:
        command.extend(["--sql-file", str(sql_file)])
    result = subprocess.run(
        command,
        cwd=ROOT,
        input=sql,
        text=True,
        capture_output=True,
        encoding="utf-8",
        check=False,
    )
    if result.returncode != 0:
        raise RuntimeError(result.stderr[-1600:] or result.stdout[-1600:])
    payload = json.loads(result.stdout)
    if payload.get("truncated"):
        raise ValueError("La consulta agregada se truncó.")
    return payload


def one_row_query(path: Path) -> dict[str, Any]:
    payload = json_query(sql_file=path, max_rows=20)
    rows = payload.get("rows")
    require(isinstance(rows, list) and len(rows) == 1, f"{path.name} debe devolver una sola fila agregada.")
    return dict(rows[0])


def jan_2026_config(configs: Iterable[dict[str, object]], date_key: str = "wave") -> dict[str, object]:
    for config in configs:
        if str(config.get(date_key)) == "Jan-2026" or int(config.get("year", 0) or 0) == 2026:
            return dict(config)
    raise ValueError("No se encontró la configuración QVoice Jan-2026.")


def headers_for(path: Path) -> list[str]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook["Raw data"]
        return [_norm_header(value) for value in next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True))]
    finally:
        workbook.close()


def scan_comment_groups() -> tuple[dict[str, set[int]], dict[str, set[int]], dict[str, int]]:
    """Return Jan-2026 comment-theme memberships, retaining no raw comments."""
    comment_config = jan_2026_config(COMMENT_WAVES)
    retention_config = jan_2026_config(RETENTION_WAVES)
    path = QVOICE_ROOT / str(comment_config["file"])
    headers = headers_for(path)
    key_column = find_comment_column(headers, [str(comment_config["key_column"])])
    require(key_column is not None, "No se encontró la llave de Jan-2026 para comentarios.")
    prompt_columns = {
        role: find_comment_column(headers, list(candidates))
        for role, candidates in dict(comment_config["prompt_candidates"]).items()
    }
    require(prompt_columns.get("doing_well") and prompt_columns.get("can_improve"), "Faltan prompts de comentario requeridos.")
    retention_improve = find_retention_column(headers, tuple(retention_config["can_improve_candidates"]))
    retention_processes = find_retention_column(headers, tuple(retention_config["processes_to_improve_candidates"]))
    require(retention_improve is not None and retention_processes is not None, "Faltan prompts de mejora Jan-2026.")
    selected = {key_column, retention_improve, retention_processes, *(column for column in prompt_columns.values() if column)}

    comment_groups = {theme: set() for theme in COMMENT_TAXONOMY}
    retention_groups: dict[str, set[int]] = defaultdict(set)
    raw_rows = invalid_key_rows = duplicate_key_rows = 0
    seen_keys: set[int] = set()
    for header_map, row in _iter_selected_rows(path, "Raw data", selected):
        raw_rows += 1
        employee_id = parse_employee_id(row.get(header_map[key_column]))
        if employee_id is None:
            if _norm_text(row.get(header_map[key_column])):
                invalid_key_rows += 1
            continue
        if employee_id in seen_keys:
            duplicate_key_rows += 1
        seen_keys.add(employee_id)
        all_text = " ".join(
            _norm_text(row.get(header_map[column]))
            for column in prompt_columns.values()
            if column and _norm_text(row.get(header_map[column]))
        )
        improvement_text = " ".join(
            _norm_text(row.get(header_map[column]))
            for column in (retention_improve, retention_processes)
            if _norm_text(row.get(header_map[column]))
        )
        for theme in COMMENT_TAXONOMY:
            if _classify_theme(all_text, theme):
                comment_groups[theme].add(employee_id)
        for theme in classify_themes(improvement_text):
            retention_groups[theme].add(employee_id)

    require(invalid_key_rows == 0, "La fuente de comentarios Jan-2026 contiene llaves no numéricas.")
    return comment_groups, dict(retention_groups), {
        "raw_rows": raw_rows,
        "distinct_usable_employee_keys": len(seen_keys),
        "duplicate_key_rows": duplicate_key_rows,
        "invalid_key_rows": invalid_key_rows,
    }


def scan_comparable_groups() -> tuple[dict[str, set[int]], dict[str, set[int]], dict[str, int]]:
    """Rebuild the 33-item and within-region bins from the Jan-2026 raw wave."""
    config = jan_2026_config(COMPARABLE_WAVES)
    structure = load_survey_structure(QVOICE_ROOT / "survey_structure.xlsx")
    headers, column_by_base, bases = _headers_and_binary_bases(config, structure)
    require(set(EXPECTED_COMMON_BASES) <= bases, "Los 33 ítems comparables no están disponibles en Jan-2026.")
    key_index = headers.index(str(config["key_column"]))
    region_index = headers.index("Region")
    selected_indices = [column_by_base[base] for base in EXPECTED_COMMON_BASES]
    key_column = _norm_header(headers[key_index])
    region_column = _norm_header(headers[region_index])
    item_columns = [_norm_header(headers[index]) for index in selected_indices]
    selected_columns = {key_column, region_column, *item_columns}
    path = QVOICE_ROOT / str(config["file"])
    signatures: dict[int, tuple[str, tuple[tuple[int | None, str], ...]]] = {}
    duplicate_rows = duplicate_conflicts = missing_key_rows = invalid_key_rows = 0

    source_rows = 0
    for header_map, row in _iter_selected_rows(path, str(config["sheet"]), selected_columns):
            source_rows += 1
            employee_id = parse_employee_id(row.get(header_map[key_column]))
            if employee_id is None:
                if _norm_text(row.get(header_map[key_column])):
                    invalid_key_rows += 1
                else:
                    missing_key_rows += 1
                continue
            region = _norm_header(row.get(header_map[region_column]))
            statuses = tuple(
                _cell_status(row.get(header_map[column]))
                for column in item_columns
            )
            signature = (region, statuses)
            prior = signatures.get(employee_id)
            if prior is not None:
                duplicate_rows += 1
                if prior != signature:
                    duplicate_conflicts += 1
                continue
            signatures[employee_id] = signature

    require(duplicate_rows == 0 and duplicate_conflicts == 0, "Hay duplicados o conflictos en QVoice comparable Jan-2026.")
    require(missing_key_rows == 0 and invalid_key_rows == 0, "Faltan llaves usables en QVoice comparable Jan-2026.")
    score_groups: dict[str, set[int]] = defaultdict(set)
    scores_by_region: dict[str, list[tuple[int, int]]] = defaultdict(list)
    complete = 0
    for employee_id, (region, statuses) in signatures.items():
        values = [value for value, status in statuses if status == "valid"]
        if len(values) != ITEM_COUNT:
            continue
        require(region != "", "Un respondente comparable Jan-2026 no tiene Region.")
        score = sum(values)
        score_groups[_band_for_favorable_count(score)].add(employee_id)
        scores_by_region[region].append((employee_id, score))
        complete += 1
    expected = EXPECTED_WAVES[2026]
    require(len(signatures) == expected["participants"] and complete == expected["complete"], "La población comparable Jan-2026 no reconcilia al contrato congelado.")
    require(len(scores_by_region) == 5 and all(len(rows) >= MINIMUM_COHORT for rows in scores_by_region.values()), "Los cohortes Region × wave no son publicables.")

    z_groups: dict[str, set[int]] = defaultdict(set)
    for rows in scores_by_region.values():
        mean = math.fsum(score for _, score in rows) / len(rows)
        variance = math.fsum((score - mean) ** 2 for _, score in rows) / len(rows)
        standard_deviation = math.sqrt(variance)
        require(standard_deviation > 0, "Un cohorte Region × wave tiene desviación estándar cero.")
        for employee_id, score in rows:
            z_groups[z_band_label((score - mean) / standard_deviation)].add(employee_id)
    require(sum(len(rows) for rows in z_groups.values()) == complete, "Las bandas z no reconcilian a la cohorte comparable.")
    return dict(score_groups), dict(z_groups), {
        "source_rows": source_rows,
        "participant_rows": len(signatures),
        "complete_comparable_participants": complete,
        "region_wave_cohort_count": len(scores_by_region),
        "duplicate_participant_rows": duplicate_rows,
    }


def scan_heartbeat_group() -> tuple[set[int], dict[str, int]]:
    """Reuse the governed 53-scale heartbeat eligibility rule in memory."""
    scanned = scan_source(read_structure())
    personal_rows = scanned.pop("personal_rows")
    scanned.pop("question_headers", None)
    employee_ids = {int(row["employee_id"]) for row in personal_rows if bool(row["eligible"])}
    require(employee_ids, "No hay empleados elegibles para el heartbeat Jan-2026.")
    return employee_ids, {key: int(value) for key, value in scanned.items() if isinstance(value, int)}


def build_runtime_group_sql(records: list[dict[str, int]], latest_snapshot: str) -> str:
    source_json = json.dumps(records, separators=(",", ":")).replace("'", "''")
    return f"""WITH source AS (
  SELECT
    elem.employee_id::BIGINT AS employee_id,
    elem.group_code::INTEGER AS group_code
  FROM (SELECT JSON_PARSE('{source_json}') AS payload) AS src,
       src.payload AS elem
), baseline AS (
  SELECT employee_id
  FROM hrda.post_hire_monthly_worker_dtls
  WHERE CAST(fis_mo_end_dt AS DATE) = DATE '{BASELINE_SNAPSHOT.isoformat()}'
    AND worker_type = 'Employee'
    AND employee_type = 'Regular'
    AND active_status = 1
    AND employee_id IS NOT NULL
  GROUP BY 1
), voluntary_events AS (
  SELECT employee_id
  FROM hrda.post_hire_monthly_worker_dtls
  WHERE worker_type = 'Employee'
    AND employee_type = 'Regular'
    AND employee_id IS NOT NULL
    AND UPPER(NULLIF(TRIM(CAST(term_type AS VARCHAR(256))), '')) = 'VOLUNTARY'
    AND CAST(term_date AS DATE) > DATE '{SURVEY_LAUNCH.isoformat()}'
    AND CAST(term_date AS DATE) <= DATE '{latest_snapshot}'
  GROUP BY 1
)
SELECT
  s.group_code,
  COUNT(*) AS source_employee_count,
  SUM(CASE WHEN b.employee_id IS NOT NULL THEN 1 ELSE 0 END) AS matched_employee_count,
  SUM(CASE WHEN b.employee_id IS NOT NULL AND e.employee_id IS NOT NULL THEN 1 ELSE 0 END) AS voluntary_exit_count
FROM source AS s
LEFT JOIN baseline AS b
  ON b.employee_id = s.employee_id
LEFT JOIN voluntary_events AS e
  ON e.employee_id = b.employee_id
GROUP BY 1
ORDER BY 1"""


def aggregate_runtime_groups(groups: dict[str, dict[str, set[int]]], latest_snapshot: str) -> dict[str, dict[str, dict[str, int | float]]]:
    ordered: list[tuple[str, str]] = []
    records: list[dict[str, int]] = []
    for metric, category_map in groups.items():
        for label, employee_ids in category_map.items():
            require(len(employee_ids) >= MINIMUM_COHORT, f"La cohorte {metric}:{label} es menor a n=30.")
            code = len(ordered)
            ordered.append((metric, label))
            records.extend({"employee_id": employee_id, "group_code": code} for employee_id in sorted(employee_ids))
    totals: dict[int, dict[str, int]] = {code: defaultdict(int) for code in range(len(ordered))}
    batch_size = 16_000
    for start in range(0, len(records), batch_size):
        payload = json_query(sql=build_runtime_group_sql(records[start : start + batch_size], latest_snapshot), max_rows=100)
        rows = payload.get("rows")
        require(isinstance(rows, list), "El join runtime no devolvió filas agregadas.")
        for row in rows:
            code = int(row["group_code"])
            for column in ("source_employee_count", "matched_employee_count", "voluntary_exit_count"):
                totals[code][column] += int(row.get(column) or 0)
        print(f"  runtime attrition groups {start + 1:,}-{min(start + batch_size, len(records)):,}", flush=True)
    result: dict[str, dict[str, dict[str, int | float]]] = defaultdict(dict)
    for code, (metric, label) in enumerate(ordered):
        metric_totals = totals[code]
        source = int(metric_totals["source_employee_count"])
        matched = int(metric_totals["matched_employee_count"])
        exits = int(metric_totals["voluntary_exit_count"])
        require(source == len(groups[metric][label]), f"El source runtime no reconcilia para {metric}:{label}.")
        require(MINIMUM_COHORT <= matched <= source and 0 <= exits <= matched, f"El join runtime es inválido para {metric}:{label}.")
        result[metric][label] = {
            "source_employee_count": source,
            "matched_employee_count": matched,
            "unmatched_employee_count": source - matched,
            "voluntary_exit_count": exits,
            "ytd_voluntary_attrition_rate_percent": round(100.0 * exits / matched, 4),
        }
    return {metric: dict(categories) for metric, categories in result.items()}


def structured_metrics() -> tuple[dict[str, dict[str, dict[str, int | float]]], dict[str, str]]:
    payload = json_query(sql_file=STRUCTURED_QUERY, max_rows=100)
    rows = payload.get("rows")
    require(isinstance(rows, list), "La consulta estructurada no devolvió filas.")
    metrics: dict[str, dict[str, dict[str, int | float]]] = defaultdict(dict)
    quality: dict[str, str] = {}
    for row in rows:
        if row["result_type"] == "quality_metric":
            quality[str(row["metric"])] = str(row["metric_value"])
            continue
        metric = str(row["metric"])
        response_value = int(row["response_value"])
        count = int(row["eligible_employee_count"])
        exits = int(row["voluntary_exit_count"])
        require(response_value in (0, 1) and count >= MINIMUM_COHORT and 0 <= exits <= count, "La categoría QVoice estructurada no es publicable.")
        metrics[metric][str(response_value)] = {
            "matched_employee_count": count,
            "voluntary_exit_count": exits,
            "ytd_voluntary_attrition_rate_percent": float(row["ytd_voluntary_attrition_rate_percent"]),
        }
    require(quality.get("qvoice_duplicate_employee_count") == "0", "La fuente QVoice estructurada contiene duplicados.")
    require(quality.get("baseline_duplicate_employee_count") == "0", "La cohorte HRDA estructurada contiene duplicados.")
    require(set(metrics) == {"intent_see_myself", "intent_rarely_look", "career_goals"}, "Faltan items estructurados para el deck.")
    require(all(set(values) == {"0", "1"} for values in metrics.values()), "Los items estructurados deben conservar 0 y 1.")
    return {metric: dict(values) for metric, values in metrics.items()}, quality


def static_window() -> dict[str, Any]:
    row = one_row_query(WINDOW_AUDIT)
    latest = str(row["latest_snapshot_date"])
    require(str(row["survey_launch_date"]) == SURVEY_LAUNCH.isoformat(), "La fecha de QVoice no coincide con la definición Jan-2026.")
    require(str(row["baseline_snapshot"]) == BASELINE_SNAPSHOT.isoformat(), "El snapshot pre-wave no coincide con la definición.")
    require(latest > SURVEY_LAUNCH.isoformat(), "El corte YTD debe ser posterior a la wave.")
    baseline = int(row["baseline_distinct_employee_count"])
    exits = int(row["ytd_voluntary_exit_count"])
    require(baseline >= MINIMUM_COHORT and int(row["duplicate_baseline_employee_count"]) == 0 and 0 <= exits <= baseline, "La cohorte YTD no pasó el control de unicidad.")
    return row


def main() -> None:
    require(WINDOW_AUDIT.exists() and STRUCTURED_QUERY.exists() and RETENTION_EVIDENCE.exists(), "Falta una entrada necesaria de QVoice YTD.")
    window = static_window()
    latest_snapshot = str(window["latest_snapshot_date"])
    comment_groups, retention_groups, comment_quality = scan_comment_groups()
    score_groups, z_groups, comparable_quality = scan_comparable_groups()
    heartbeat_group, heartbeat_quality = scan_heartbeat_group()
    participation_config = jan_2026_config(PARTICIPATION_WAVES)
    participation = scan_raw_wave(participation_config)
    flags = participation.pop("flags")
    require(isinstance(flags, dict), "No se pudieron obtener flags de participación Jan-2026.")
    retention_display = json.loads(RETENTION_EVIDENCE.read_text(encoding="utf-8"))["jan_2026_display"]["ranked_themes"]
    require(all(theme in retention_groups for theme in retention_display), "Las themes visibles no se reconstruyeron desde el workbook Jan-2026.")

    groups = {
        "slide_20_comment_theme": {theme: comment_groups[theme] for theme in COMMENT_TAXONOMY},
        "slide_21_retention_theme": {theme: retention_groups[theme] for theme in retention_display},
        "slide_22_participation": {
            "Responded": {employee_id for employee_id, answered in flags.items() if answered},
            "QVoice row / no answer": {employee_id for employee_id, answered in flags.items() if not answered},
        },
        "slide_25_score_band": score_groups,
        "slide_26_heartbeat": {"Heartbeat eligible": heartbeat_group},
        "slide_27_z_score_band": z_groups,
    }
    runtime = aggregate_runtime_groups(groups, latest_snapshot)
    participation_rates = runtime["slide_22_participation"]
    baseline = int(window["baseline_distinct_employee_count"])
    baseline_exits = int(window["ytd_voluntary_exit_count"])
    source_matched = sum(int(row["matched_employee_count"]) for row in participation_rates.values())
    source_exits = sum(int(row["voluntary_exit_count"]) for row in participation_rates.values())
    no_qvoice_count = baseline - source_matched
    no_qvoice_exits = baseline_exits - source_exits
    require(no_qvoice_count >= MINIMUM_COHORT and 0 <= no_qvoice_exits <= no_qvoice_count, "No QVoice row no reconcilia a la cohorte pre-wave.")
    participation_rates["No QVoice row"] = {
        "source_employee_count": no_qvoice_count,
        "matched_employee_count": no_qvoice_count,
        "unmatched_employee_count": 0,
        "voluntary_exit_count": no_qvoice_exits,
        "ytd_voluntary_attrition_rate_percent": round(100.0 * no_qvoice_exits / no_qvoice_count, 4),
    }
    structured, structured_quality = structured_metrics()

    evidence = {
        "artifact": "internal_analysis",
        "evidence_type": "aggregate-only QVoice Jan-2026 voluntary attrition YTD",
        "generated_on": date.today().isoformat(),
        "outcome": "pass",
        "minimum_shared_cohort": MINIMUM_COHORT,
        "outcome_contract": {
            "baseline_snapshot": str(window["baseline_snapshot"]),
            "survey_launch_date": str(window["survey_launch_date"]),
            "outcome_window": f"({window['survey_launch_date']}, {window['latest_snapshot_date']}]",
            "label": "Voluntary attrition YTD 2026; not annualized",
            "target": "At least one term_type exactly VOLUNTARY after the survey launch and on or before the latest snapshot.",
            "population": "worker_type='Employee'; employee_type='Regular'; active_status=1 at the pre-wave snapshot.",
            "grain": "Distinct employee in each slide-specific QVoice group.",
        },
        "window_audit": window,
        "slides": {
            "20": runtime["slide_20_comment_theme"],
            "21": runtime["slide_21_retention_theme"],
            "22": participation_rates,
            "23": {"intent_see_myself": structured["intent_see_myself"], "intent_rarely_look": structured["intent_rarely_look"]},
            "24": structured["career_goals"],
            "25": runtime["slide_25_score_band"],
            "26": runtime["slide_26_heartbeat"],
            "27": runtime["slide_27_z_score_band"],
        },
        "validation": {
            "outcome": "pass",
            "comment_source": comment_quality,
            "comparable_source": comparable_quality,
            "heartbeat_source": heartbeat_quality,
            "participation_source": {key: value for key, value in participation.items() if isinstance(value, int)},
            "structured_source": structured_quality,
            "retention_display_order": retention_display,
            "runtime_join_sql_persisted": False,
            "reason_runtime_join_sql_not_persisted": "It contains in-memory employee identifiers and only returns aggregate counts.",
            "all_displayed_groups_have_n_at_least": MINIMUM_COHORT,
        },
        "privacy": {
            "employee_identifiers_persisted": False,
            "employee_rows_persisted": False,
            "raw_comment_text_persisted": False,
            "individual_qvoice_values_persisted": False,
            "runtime_join_sql_persisted": False,
        },
        "provenance": {
            "window_audit": str(WINDOW_AUDIT.relative_to(ROOT)).replace("\\", "/"),
            "window_audit_sha256": sha256_path(WINDOW_AUDIT),
            "structured_query": str(STRUCTURED_QUERY.relative_to(ROOT)).replace("\\", "/"),
            "structured_query_sha256": sha256_path(STRUCTURED_QUERY),
            "raw_workbooks": "data/QVoice/Jan 2026 Qvoice_raw data.xlsx, read in memory only",
        },
    }
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT.write_text(json.dumps(evidence, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"output": str(OUTPUT), "cutoff": latest_snapshot, "slides": sorted(evidence["slides"]), "outcome": "pass"}, ensure_ascii=False))


if __name__ == "__main__":
    main()

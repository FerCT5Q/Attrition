"""Build an aggregate-only QVoice non-response proxy.

Raw employee IDs are read and joined in memory through the read-only Redshift
helper. The generated SQL containing IDs is never written to disk or printed;
only aggregate counts are persisted.
"""

from __future__ import annotations

import json
import re
import subprocess
import sys
from collections import Counter
from datetime import date
from pathlib import Path

from openpyxl import load_workbook

from build_qvoice_comment_evidence import QVOICE_ROOT, _iter_selected_rows


ROOT = Path(__file__).resolve().parents[1]
RUN_READONLY = ROOT / ".codex" / "skills" / "attrition-redshift-db" / "scripts" / "run_readonly_query.py"
PYTHON = ROOT / ".venv" / "Scripts" / "python.exe"
OUTPUT = ROOT / "output" / "distribution" / f"qvoice-nonrespondent-proxy-evidence-{date.today().isoformat()}.json"

WAVES = [
    {"wave": "Mar-2023", "file": "March 2023 Pulse_raw data.xlsx", "key": "EmployeeID", "snapshot": "2023-02-19", "proposed": "2023-02-15"},
    {"wave": "Jan-2024", "file": "Jan 2024 Qvoice raw data.xlsx", "key": "Participant's Unique Identifier", "snapshot": "2023-12-24", "proposed": "2023-12-15"},
    {"wave": "Jan-2025", "file": "Jan 2025_Qvoice_raw data.xlsx", "key": "EmpID", "snapshot": "2024-12-29", "proposed": "2024-12-15"},
    {"wave": "Jan-2026", "file": "Jan 2026 Qvoice_raw data.xlsx", "key": "Emp ID", "snapshot": "2025-12-28", "proposed": "2025-12-15"},
]


def norm(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "").replace("\n", " ")).strip()


def nonempty(value: object) -> bool:
    return value is not None and bool(str(value).strip())


def is_response_header(header: str) -> bool:
    value = norm(header).casefold()
    if not value or " en" in value or value.endswith("en"):
        return False
    if "strong/moderate factor" in value or "sentiment" in value or "theme" in value:
        return False
    if value.endswith(("_fav", "_unfav", "_favorable", "_unfavorable", "_s")):
        return True
    if "_1-5 scale" in value or value.endswith("scale"):
        return True
    return any(
        phrase in value
        for phrase in (
            "what is one thing qualcomm is doing really well",
            "what is one thing qualcomm can do to improve",
            "if desired, please describe",
            "the work processes at qualcomm allow employees",
        )
    )


def parse_employee_id(value: object) -> int | None:
    text = norm(value)
    if not text:
        return None
    if re.fullmatch(r"\d+", text):
        return int(text)
    if re.fullmatch(r"\d+\.0+", text):
        return int(float(text))
    return None


def scan_raw_wave(config: dict[str, str]) -> dict[str, object]:
    path = QVOICE_ROOT / config["file"]
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook["Raw data"]
    header_row = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True))
    headers = [norm(value) for value in header_row]
    header_by_name = {header.casefold(): header for header in headers if header}
    key_header = header_by_name.get(config["key"].casefold())
    if key_header is None:
        workbook.close()
        raise ValueError(f"Missing key column {config['key']} in {config['file']}")
    response_headers = {header for header in headers if is_response_header(header)}
    if not response_headers:
        workbook.close()
        raise ValueError(f"No response columns identified in {config['file']}")
    workbook.close()

    flags: dict[int, bool] = {}
    raw_rows = 0
    missing_key_rows = 0
    missing_key_rows_with_any_answer = 0
    invalid_key_rows = 0
    response_columns = len(response_headers)

    for header_map, row in _iter_selected_rows(path, "Raw data", response_headers | {key_header}):
        raw_rows += 1
        employee_id = parse_employee_id(row.get(header_map[key_header]))
        has_response = any(nonempty(row.get(header_map[header])) for header in response_headers)
        if employee_id is None:
            if norm(row.get(header_map[key_header])):
                invalid_key_rows += 1
            else:
                missing_key_rows += 1
                if has_response:
                    missing_key_rows_with_any_answer += 1
            continue
        if employee_id in flags:
            flags[employee_id] = flags[employee_id] or has_response
        else:
            flags[employee_id] = has_response

    duplicate_key_rows = raw_rows - missing_key_rows - invalid_key_rows - len(flags)
    return {
        "raw_rows": raw_rows,
        "raw_distinct_usable_keys": len(flags),
        "raw_duplicate_key_rows": duplicate_key_rows,
        "raw_missing_key_rows": missing_key_rows,
        "raw_missing_key_rows_with_any_answer": missing_key_rows_with_any_answer,
        "raw_invalid_key_rows": invalid_key_rows,
        "raw_keys_with_any_answer": sum(1 for value in flags.values() if value),
        "raw_keys_with_no_answer": sum(1 for value in flags.values() if not value),
        "response_columns_checked": response_columns,
        "flags": flags,
    }


def run_aggregate_query(flags: dict[int, bool], snapshot: str) -> dict[str, object]:
    if not flags:
        raise ValueError("No usable QVoice keys available for join")
    command = [str(PYTHON if PYTHON.exists() else sys.executable), str(RUN_READONLY), "--max-rows", "10", "--reporting-mode", "internal_analysis"]

    def execute(sql: str) -> dict[str, object]:
        result = subprocess.run(command, cwd=ROOT, input=sql, text=True, capture_output=True, check=False)
        if result.returncode != 0:
            raise RuntimeError(f"Aggregate join failed with exit code {result.returncode}: {result.stderr[-1000:]}")
        return json.loads(result.stdout)["rows"][0]

    population = execute(f"""SELECT COUNT(DISTINCT employee_id) AS eligible_snapshot_employees
FROM hrda.post_hire_monthly_worker_dtls
WHERE fis_mo_end_dt = DATE '{snapshot}'
  AND worker_type = 'Employee'
  AND employee_type = 'Regular'
  AND active_status = 1;""")
    totals = {"responded_proxy": 0, "qvoice_row_without_answer": 0, "qvoice_keys_inside_snapshot": 0}
    ordered = sorted(flags.items())
    chunk_size = 4000
    for start in range(0, len(ordered), chunk_size):
        chunk = ordered[start : start + chunk_size]
        all_ids = ",".join(str(employee_id) for employee_id, _ in chunk)
        response_ids = ",".join(str(employee_id) for employee_id, has_response in chunk if has_response) or "NULL"
        no_answer_ids = ",".join(str(employee_id) for employee_id, has_response in chunk if not has_response) or "NULL"
        row = execute(f"""SELECT
    COUNT(DISTINCT CASE WHEN employee_id IN ({response_ids}) THEN employee_id END) AS responded_proxy,
    COUNT(DISTINCT CASE WHEN employee_id IN ({no_answer_ids}) THEN employee_id END) AS qvoice_row_without_answer,
    COUNT(DISTINCT employee_id) AS qvoice_keys_inside_snapshot
FROM hrda.post_hire_monthly_worker_dtls
WHERE fis_mo_end_dt = DATE '{snapshot}'
  AND worker_type = 'Employee'
  AND employee_type = 'Regular'
  AND active_status = 1
  AND employee_id IN ({all_ids});""")
        for key in totals:
            totals[key] += int(row.get(key) or 0)
    eligible = int(population["eligible_snapshot_employees"])
    return {
        "eligible_snapshot_employees": eligible,
        **totals,
        "no_qvoice_row": eligible - totals["responded_proxy"] - totals["qvoice_row_without_answer"],
    }


def main() -> None:
    evidence_waves = []
    for config in WAVES:
        print(f"Scanning {config['wave']}...", flush=True)
        raw = scan_raw_wave(config)
        print(f"  raw keys={raw['raw_distinct_usable_keys']} any_answer={raw['raw_keys_with_any_answer']} no_answer={raw['raw_keys_with_no_answer']}", flush=True)
        aggregate = run_aggregate_query(raw.pop("flags"), config["snapshot"])
        print(f"  aggregate snapshot={aggregate['eligible_snapshot_employees']} responded={aggregate['responded_proxy']} row_without_answer={aggregate['qvoice_row_without_answer']} no_row={aggregate['no_qvoice_row']}", flush=True)
        aggregate["raw_qvoice_keys"] = raw["raw_distinct_usable_keys"]
        aggregate["qvoice_keys_outside_snapshot"] = raw["raw_distinct_usable_keys"] - int(aggregate["qvoice_keys_inside_snapshot"])
        evidence_waves.append(
            {
                "wave": config["wave"],
                "source_file": config["file"],
                "proposed_snapshot_assumption": config["proposed"],
                "stored_snapshot_used": config["snapshot"],
                **raw,
                **aggregate,
            }
        )
    evidence = {
        "artifact": "internal_analysis",
        "evidence_type": "aggregate-only QVoice non-response proxy",
        "generated_on": date.today().isoformat(),
        "outcome": "pass",
        "assumptions": [
            "The selected stored snapshot is treated as the prior-month day-15 population proxy for internal review.",
            "Every Employee + Regular + active_status=1 employee in that snapshot is treated as invited/eligible.",
            "A raw QVoice key with at least one original answer field populated is treated as responded for this proxy.",
            "A missing QVoice key is labelled no QVoice row, not a confirmed non-response, until the invitation roster is confirmed.",
        ],
        "privacy": {
            "identifiers_persisted": False,
            "response_values_persisted": False,
            "free_text_persisted": False,
            "runtime_join_sql_persisted": False,
        },
        "waves": evidence_waves,
    }
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT.write_text(json.dumps(evidence, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"output": str(OUTPUT), "waves": [wave["wave"] for wave in evidence_waves], "outcome": "pass"}, ensure_ascii=False))


if __name__ == "__main__":
    main()

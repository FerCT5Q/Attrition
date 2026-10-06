"""Build aggregate-only QVoice improvement-theme evidence for slide 19.

The raw workbooks are scanned in memory.  This module deliberately persists
only aggregate participant-wave metrics: never comment text, participant keys,
or input rows.  It is a descriptive internal-analysis artifact, not an
attrition feature, causal analysis, or individual-risk signal.
"""

from __future__ import annotations

import json
import re
from collections import Counter
from datetime import date
from pathlib import Path
from typing import Iterable, Mapping
from zipfile import ZipFile

from lxml import etree
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
QVOICE_ROOT = ROOT / "data" / "QVoice"
DEFAULT_OUTPUT = (
    ROOT / "output" / "distribution" / "json" / "qvoice-retention-concerns-evidence-20260910.json"
)
MINIMUM_COHORT = 30
TOP_THEME_COUNT = 5
TREND_THEME_COUNT = 3


WAVES: tuple[dict[str, object], ...] = (
    {
        "wave": "Mar-2023",
        "file": "March 2023 Pulse_raw data.xlsx",
        "key_column": "EmployeeID",
        "can_improve_candidates": (
            "What is one thing Qualcomm can do to improve? Please describe something that you would like Qualcomm to do to improve your experience.",
        ),
        "processes_to_improve_candidates": (),
        "processes_expected": False,
    },
    {
        "wave": "Jan-2024",
        "file": "Jan 2024 Qvoice raw data.xlsx",
        "key_column": "Participant's Unique Identifier",
        "can_improve_candidates": (
            "What is one thing Qualcomm can do to improve? Please describe something that you would like Qualcomm to do to improve your experience.",
        ),
        "processes_to_improve_candidates": (),
        "processes_expected": False,
    },
    {
        "wave": "Jan-2025",
        "file": "Jan 2025_Qvoice_raw data.xlsx",
        "key_column": "EmpID",
        "can_improve_candidates": (
            "What is one thing Qualcomm can do to improve? Please describe something that you would like Qualcomm to do to improve your experience.",
        ),
        "processes_to_improve_candidates": (
            "If desired, please describe 1-2 work processes that you would like to see improved or made more efficient.",
        ),
        "processes_expected": True,
    },
    {
        "wave": "Jan-2026",
        "file": "Jan 2026 Qvoice_raw data.xlsx",
        "key_column": "Emp ID",
        "can_improve_candidates": (
            "What is one thing Qualcomm can do to improve? Please describe something that you would like Qualcomm to do to improve your experience.",
        ),
        "processes_to_improve_candidates": (
            "If desired, please describe 1-2 work processes that you would like to see improved or made more efficient.",
        ),
        "processes_expected": True,
    },
)


# Terms deliberately identify broad improvement themes rather than sentiment,
# complaints, or causality.  The labels align conceptually to existing QVoice
# and exit-survey reporting domains, but this builder never reads exit data.
TAXONOMY: dict[str, dict[str, object]] = {
    "Rewards": {
        "alignment": "QVoice rewards / exit-survey rewards concept; no exit-survey data used.",
        "patterns": (
            "compensation", "salary", "salaries", "pay raise", "pay increase", "bonus", "bonuses",
            "benefits", "equity", "stock grant", "rsu", "compensación", "salario", "sueldos",
            "aumento salarial", "bono", "bonos", "beneficios", "acciones",
        ),
    },
    "Career growth and promotion": {
        "alignment": "QVoice career-growth / exit-survey advancement concept; no exit-survey data used.",
        "patterns": (
            "career growth", "career development", "career path", "career opportunity", "career opportunities",
            "growth opportunity", "growth opportunities", "promotion", "promotions", "promote", "advancement",
            "professional development", "development opportunity", "desarrollo profesional", "plan de carrera",
            "oportunidad de carrera", "oportunidades de carrera", "oportunidad de crecimiento",
            "oportunidades de crecimiento", "promoción", "promociones", "ascenso", "ascensos",
        ),
    },
    "Manager, leadership and respect": {
        "alignment": "QVoice leadership / exit-survey manager concept; no exit-survey data used.",
        "patterns": (
            "my manager", "direct manager", "manager support", "manager training", "manager behavior",
            "manager communication", "manager feedback", "leadership", "leadership team", "senior leaders",
            "senior leadership", "management", "respect", "respectful", "micromanagement", "mi manager",
            "mi gerente", "liderazgo", "equipo directivo", "alta dirección", "gerencia", "respeto",
            "trato respetuoso", "microgestión", "microgestion",
        ),
    },
    "Workload, staffing and work-life": {
        "alignment": "QVoice wellbeing / exit-survey workload concept; no exit-survey data used.",
        "patterns": (
            "workload", "too much work", "heavy workload", "work life balance", "work-life balance",
            "worklife balance", "burnout", "overworked", "overwork", "long hours", "overtime",
            "understaffed", "understaffing", "staffing", "headcount", "carga de trabajo",
            "demasiado trabajo", "equilibrio vida trabajo", "balance vida trabajo", "agotamiento",
            "horas extra", "jornadas largas", "falta de personal", "dotación", "dotacion",
        ),
    },
    "Work arrangement/RTO/commute": {
        "alignment": "QVoice flexibility / work-arrangement concept; no exit-survey data used.",
        "patterns": (
            "rto", "return to office", "return to the office", "back to office", "back in the office",
            "in office", "in-office", "office attendance", "office days", "days in office", "remote work",
            "work remotely", "working remotely", "work from home", "working from home", "wfh", "hybrid work",
            "hybrid-working", "hybrid schedule", "commute", "commuting", "parking", "onsite", "on-site",
            "presencial", "volver a la oficina", "regreso a la oficina", "ir a la oficina", "trabajo remoto",
            "trabajo desde casa", "teletrabajo", "trabajo híbrido", "trabajo hibrido", "modelo híbrido",
            "modelo hibrido", "traslado al trabajo", "estacionamiento",
        ),
    },
    "Tools, processes and bureaucracy": {
        "alignment": "QVoice processes / tools concept; no exit-survey data used.",
        "patterns": (
            "tools", "tooling", "systems", "software", "process", "processes", "workflow", "workflows",
            "bureaucracy", "red tape", "approval process", "approvals", "automation", "inefficient process",
            "herramientas", "sistemas", "proceso", "procesos", "flujo de trabajo", "burocracia",
            "trámites", "tramites", "aprobaciones", "automatización", "automatizacion",
        ),
    },
    "Organizational confidence/job security": {
        "alignment": "QVoice confidence / exit-survey organizational-security concept; no exit-survey data used.",
        "patterns": (
            "job security", "job stability", "layoff", "layoffs", "reorganization", "reorganisation",
            "restructuring", "company direction", "business outlook", "company stability", "organizational stability",
            "future of the company", "uncertainty", "job loss", "seguridad laboral", "estabilidad laboral",
            "despido", "despidos", "reorganización", "reorganizacion", "reestructuración", "reestructuracion",
            "dirección de la empresa", "direccion de la empresa", "futuro de la empresa", "incertidumbre",
        ),
    },
}


# Intent is purposefully narrower than dissatisfaction.  The patterns are
# multiword English/Spanish expressions, and a nearby negation suppresses a
# match.  This aggregate contextual flag is never an employee label.
LEAVE_INTENT_PATTERNS: tuple[str, ...] = (
    "leave the company", "leave this company", "leave qualcomm", "quit my job", "quit this job",
    "resign from my job", "resign from qualcomm", "looking for another job", "look for another job",
    "find another job", "find a new job", "search for another job", "searching for another job",
    "buscar otro trabajo", "buscando otro trabajo", "buscar un nuevo trabajo", "buscando un nuevo trabajo",
    "dejar la empresa", "dejar qualcomm", "renunciar a mi trabajo", "renunciar a qualcomm",
    "cambiar de trabajo", "buscando empleo fuera",
)
NEGATION_WORDS = frozenset({
    "no", "not", "never", "without", "dont", "don't", "do not", "wouldn't", "wouldnt",
    "won't", "wont", "nunca", "jamás", "jamas", "sin", "ni",
})
PRIVATE_KEY_TOKENS = ("employee_id", "participant_id", "participant_key", "comment_text", "raw_text", "response_text")


def _norm_header(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "").replace("\n", " ")).strip()


def _norm_text(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip().casefold()


def _is_nonempty(value: object) -> bool:
    return bool(_norm_text(value))


def _has_any(text: str, patterns: Iterable[str]) -> bool:
    return any(re.search(r"(?<!\w)" + re.escape(pattern.casefold()) + r"(?!\w)", text) for pattern in patterns)


def classify_themes(text: object) -> set[str]:
    """Return all non-exclusive taxonomy matches for one in-memory comment."""
    normalized = _norm_text(text)
    if not normalized:
        return set()
    return {theme for theme, rule in TAXONOMY.items() if _has_any(normalized, rule["patterns"])}


def _negated_nearby(text: str, match_start: int) -> bool:
    preceding = text[max(0, match_start - 48):match_start]
    words = re.findall(r"[\w']+", preceding)
    if any(word in NEGATION_WORDS for word in words[-5:]):
        return True
    # Covers multiword negations after punctuation-normalization without
    # allowing an unrelated negation earlier in a long comment to suppress it.
    return bool(re.search(r"\b(?:do not|not planning to|not going to|no quiero|no pienso)\s*$", preceding))


def contains_explicit_intent_to_leave(text: object) -> bool:
    """Conservatively identify non-negated multiword leave-intent wording."""
    normalized = _norm_text(text)
    for phrase in LEAVE_INTENT_PATTERNS:
        match = re.search(r"(?<!\w)" + re.escape(phrase) + r"(?!\w)", normalized)
        if match and not _negated_nearby(normalized, match.start()):
            return True
    return False


def _find_column(headers: list[str], candidates: Iterable[str]) -> str | None:
    normalized = {_norm_header(header): header for header in headers}
    for candidate in candidates:
        candidate_norm = _norm_header(candidate)
        if candidate_norm in normalized:
            return normalized[candidate_norm]
    # The 2026 prompt appends an opt-out note.  Prefix resolution is limited to
    # the documented prompt candidates and runs before response values are read.
    for candidate in candidates:
        candidate_norm = _norm_header(candidate).casefold()
        matches = [header for header in headers if _norm_header(header).casefold().startswith(candidate_norm)]
        if len(matches) == 1:
            return matches[0]
    return None


def _column_number(cell_ref: str) -> int:
    letters = re.match(r"[A-Z]+", cell_ref.upper())
    if not letters:
        raise ValueError(f"Invalid XLSX cell reference: {cell_ref}")
    result = 0
    for char in letters.group(0):
        result = result * 26 + ord(char) - ord("A") + 1
    return result - 1


def _xml_tag(value: str) -> str:
    return value.rsplit("}", 1)[-1]


def _workbook_sheet_path(archive: ZipFile, sheet_name: str) -> str:
    workbook_root = etree.fromstring(archive.read("xl/workbook.xml"))
    sheet_rid = None
    for element in workbook_root.iter():
        if _xml_tag(element.tag) == "sheet" and element.attrib.get("name") == sheet_name:
            sheet_rid = element.attrib.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")
            break
    if not sheet_rid:
        raise ValueError(f"Sheet not found: {sheet_name}")
    rel_root = etree.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
    for element in rel_root.iter():
        if _xml_tag(element.tag) == "Relationship" and element.attrib.get("Id") == sheet_rid:
            target = element.attrib["Target"].lstrip("/")
            return target if target.startswith("xl/") else f"xl/{target}"
    raise ValueError(f"Worksheet relationship not found: {sheet_name}")


def _shared_strings(archive: ZipFile) -> list[str]:
    try:
        stream = archive.open("xl/sharedStrings.xml")
    except KeyError:
        return []
    result: list[str] = []
    for _, element in etree.iterparse(stream, events=("end",), tag="{*}si", recover=True):
        result.append("".join(node.text or "" for node in element.iter() if _xml_tag(node.tag) == "t"))
        element.clear()
    stream.close()
    return result


def _cell_value(cell, shared: list[str]) -> object:
    kind = cell.attrib.get("t")
    if kind == "inlineStr":
        return "".join(node.text or "" for node in cell.iter() if _xml_tag(node.tag) == "t")
    value_node = next((node for node in cell if _xml_tag(node.tag) == "v"), None)
    if value_node is None or value_node.text is None:
        return None
    value = value_node.text
    return shared[int(value)] if kind == "s" else value


def _iter_selected_rows(path: Path, sheet_name: str, required_columns: set[str]):
    """Yield selected in-memory cells without materializing an entire sheet."""
    with ZipFile(path) as archive:
        shared = _shared_strings(archive)
        sheet_path = _workbook_sheet_path(archive, sheet_name)
        stream = archive.open(sheet_path)
        header_map: dict[str, int] | None = None
        selected_indices: set[int] | None = None
        for _, element in etree.iterparse(stream, events=("end",), tag="{*}row", recover=True):
            values: dict[int, object] = {}
            for cell in element:
                if _xml_tag(cell.tag) != "c" or "r" not in cell.attrib:
                    continue
                column = _column_number(cell.attrib["r"])
                if header_map is None or (selected_indices is not None and column in selected_indices):
                    values[column] = _cell_value(cell, shared)
            if header_map is None:
                header_map = {}
                for index, value in values.items():
                    header_map.setdefault(_norm_header(value), index)
                missing = required_columns - set(header_map)
                if missing:
                    raise ValueError(f"Required columns not found in {path.name}: {sorted(missing)}")
                selected_indices = {header_map[column] for column in required_columns}
            else:
                yield header_map, values
            element.clear()
        stream.close()


def publication_metric(count: int, denominator: int) -> dict[str, object]:
    """Apply the shared n>=30 publication gate to one aggregate metric."""
    if denominator <= 0:
        raise ValueError("The improvement-comment denominator must be positive.")
    if count < 0 or count > denominator:
        raise ValueError("A theme count must be within its wave denominator.")
    if count < MINIMUM_COHORT:
        return {
            "participant_count": None,
            "percent_of_improvement_comment_participants": None,
            "publication_status": "suppressed_below_n_30",
        }
    return {
        "participant_count": count,
        "percent_of_improvement_comment_participants": round(100 * count / denominator, 2),
        "publication_status": "published",
    }


def rank_publishable_themes(theme_metrics: Mapping[str, Mapping[str, object]], limit: int = TOP_THEME_COUNT) -> list[str]:
    """Rank published themes by count, resolving ties with fixed taxonomy order."""
    taxonomy_index = {theme: index for index, theme in enumerate(TAXONOMY)}
    eligible = [
        theme for theme, metric in theme_metrics.items()
        if metric.get("publication_status") == "published"
    ]
    eligible.sort(key=lambda theme: (-int(theme_metrics[theme]["participant_count"]), taxonomy_index[theme]))
    return eligible[:limit]


def _aggregate_records(records: Iterable[Mapping[str, object]]) -> dict[str, object]:
    """Aggregate one synthetic or streamed wave without retaining raw records."""
    all_keys: Counter[str] = Counter()
    denominator_keys: set[str] = set()
    theme_keys = {theme: set() for theme in TAXONOMY}
    leave_intent_keys: set[str] = set()
    rows_with_improvement_comment = 0
    comments_without_key = 0

    for record in records:
        key = _norm_text(record.get("participant_key"))
        can_improve = _norm_text(record.get("can_improve"))
        processes = _norm_text(record.get("processes_to_improve"))
        improvement_text = " ".join(value for value in (can_improve, processes) if value)
        if key:
            all_keys[key] += 1
        if not improvement_text:
            continue
        rows_with_improvement_comment += 1
        if not key:
            comments_without_key += 1
            continue
        denominator_keys.add(key)
        for theme in classify_themes(improvement_text):
            theme_keys[theme].add(key)
        if contains_explicit_intent_to_leave(improvement_text):
            leave_intent_keys.add(key)

    duplicate_rows = sum(count - 1 for count in all_keys.values() if count > 1)
    denominator = len(denominator_keys)
    if denominator == 0:
        raise ValueError("No keyed improvement comments were found for a wave.")
    themes = {
        theme: publication_metric(len(theme_keys[theme]), denominator)
        for theme in TAXONOMY
    }
    for theme, metric in themes.items():
        if metric["publication_status"] == "published":
            count = int(metric["participant_count"])
            expected_pct = round(100 * count / denominator, 2)
            if metric["percent_of_improvement_comment_participants"] != expected_pct:
                raise ValueError(f"Percentage reconciliation failed for {theme}.")

    return {
        "participant_rows": sum(all_keys.values()),
        "distinct_participant_keys": len(all_keys),
        "improvement_comment_participants": denominator,
        "rows_with_improvement_comment": rows_with_improvement_comment,
        "improvement_comments_without_key": comments_without_key,
        "duplicate_participant_rows": duplicate_rows,
        "themes": themes,
        "explicit_intent_to_leave": publication_metric(len(leave_intent_keys), denominator),
    }


def _scan_wave(config: Mapping[str, object]) -> dict[str, object]:
    path = QVOICE_ROOT / str(config["file"])
    workbook = load_workbook(path, read_only=True, data_only=True)
    if "Raw data" not in workbook.sheetnames:
        workbook.close()
        raise ValueError(f"Missing Raw data sheet: {path.name}")
    worksheet = workbook["Raw data"]
    headers = [_norm_header(value) for value in next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True))]
    workbook.close()

    key_column = _find_column(headers, (str(config["key_column"]),))
    can_improve_column = _find_column(headers, tuple(config["can_improve_candidates"]))
    processes_column = _find_column(headers, tuple(config["processes_to_improve_candidates"]))
    if not key_column or not can_improve_column:
        raise ValueError(f"Required participant key or can-improve prompt not found in {path.name}.")
    processes_expected = bool(config["processes_expected"])
    if processes_expected != bool(processes_column):
        raise ValueError(f"Processes-prompt availability does not match the documented contract for {config['wave']}.")

    selected_headers = {key_column, can_improve_column}
    if processes_column:
        selected_headers.add(processes_column)

    def records():
        for header_map, row in _iter_selected_rows(path, "Raw data", selected_headers):
            yield {
                "participant_key": row.get(header_map[key_column]),
                "can_improve": row.get(header_map[can_improve_column]),
                "processes_to_improve": row.get(header_map[processes_column]) if processes_column else None,
            }

    aggregate = _aggregate_records(records())
    return {
        "wave": str(config["wave"]),
        "source_file": str(config["file"]),
        "sheet": "Raw data",
        "prompt_availability": {
            "can_improve": True,
            "processes_to_improve": bool(processes_column),
        },
        "selected_prompt_rule": (
            "can improve only"
            if not processes_column
            else "can improve plus processes to improve"
        ),
        **aggregate,
    }


def _contains_private_payload(value: object) -> bool:
    if isinstance(value, dict):
        for key, nested in value.items():
            if str(key).casefold() in PRIVATE_KEY_TOKENS:
                return True
            if _contains_private_payload(nested):
                return True
    elif isinstance(value, list):
        return any(_contains_private_payload(item) for item in value)
    return False


def validate_evidence(evidence: Mapping[str, object]) -> dict[str, object]:
    """Validate privacy, contracts, reconciliation, ranking, and suppression."""
    if evidence.get("artifact") != "internal_analysis" or evidence.get("outcome") != "pass":
        raise ValueError("Retention-concerns evidence must be a passing internal_analysis artifact.")
    if evidence.get("comment_text_persisted") or evidence.get("identifiers_persisted") or evidence.get("employee_rows_persisted"):
        raise ValueError("Retention-concerns evidence contains prohibited persisted content.")
    if _contains_private_payload(evidence):
        raise ValueError("Retention-concerns evidence contains a prohibited private payload key.")

    waves = evidence.get("waves")
    if not isinstance(waves, list) or [wave.get("wave") for wave in waves] != [config["wave"] for config in WAVES]:
        raise ValueError("Retention-concerns evidence must contain the four expected waves in order.")
    for wave, config in zip(waves, WAVES):
        if wave.get("duplicate_participant_rows") != 0:
            raise ValueError(f"Duplicate participant rows found in {wave.get('wave')}.")
        if wave.get("improvement_comments_without_key") != 0:
            raise ValueError(f"Keyed-comment validation failed in {wave.get('wave')}.")
        availability = wave.get("prompt_availability", {})
        if availability.get("can_improve") is not True or availability.get("processes_to_improve") is not bool(config["processes_expected"]):
            raise ValueError(f"Prompt contract failed in {wave.get('wave')}.")
        denominator = int(wave.get("improvement_comment_participants", 0))
        if denominator < MINIMUM_COHORT:
            raise ValueError(f"The improvement-comment denominator is below n=30 in {wave.get('wave')}.")
        themes = wave.get("themes", {})
        if set(themes) != set(TAXONOMY):
            raise ValueError(f"Taxonomy contract failed in {wave.get('wave')}.")
        for theme, metric in themes.items():
            status = metric.get("publication_status")
            count = metric.get("participant_count")
            percent = metric.get("percent_of_improvement_comment_participants")
            if status == "published":
                if count is None or int(count) < MINIMUM_COHORT or percent is None:
                    raise ValueError(f"Published theme failed n>=30 suppression in {wave.get('wave')}: {theme}.")
                if float(percent) != round(100 * int(count) / denominator, 2):
                    raise ValueError(f"Theme percentage does not reconcile in {wave.get('wave')}: {theme}.")
            elif status == "suppressed_below_n_30":
                if count is not None or percent is not None:
                    raise ValueError(f"Suppressed theme leaked a value in {wave.get('wave')}: {theme}.")
            else:
                raise ValueError(f"Unknown publication status in {wave.get('wave')}: {theme}.")
        intent = wave.get("explicit_intent_to_leave", {})
        if intent.get("publication_status") == "published":
            if intent.get("participant_count") is None or int(intent["participant_count"]) < MINIMUM_COHORT:
                raise ValueError(f"Published leave-intent metric failed n>=30 in {wave.get('wave')}.")
        elif intent.get("publication_status") == "suppressed_below_n_30":
            if intent.get("participant_count") is not None or intent.get("percent_of_improvement_comment_participants") is not None:
                raise ValueError(f"Suppressed leave-intent metric leaked a value in {wave.get('wave')}.")
        else:
            raise ValueError(f"Leave-intent suppression status is invalid in {wave.get('wave')}.")

    latest = waves[-1]
    ranked = rank_publishable_themes(latest["themes"])
    display = evidence.get("jan_2026_display", {})
    if display.get("ranked_themes") != ranked:
        raise ValueError("Jan-2026 theme ranking does not match the frozen metrics.")
    if len(ranked) != TOP_THEME_COUNT:
        raise ValueError("Fewer than five publishable Jan-2026 themes are available.")
    if display.get("trend_themes") != ranked[:TREND_THEME_COUNT]:
        raise ValueError("Trend themes must be the top three publishable Jan-2026 themes.")
    return dict(evidence)


def build_evidence(output_path: Path = DEFAULT_OUTPUT, *, overwrite: bool = False) -> dict[str, object]:
    """Build a new immutable aggregate artifact; overwrite requires an explicit opt-in."""
    if output_path.exists() and not overwrite:
        raise FileExistsError(f"Refusing to overwrite immutable evidence: {output_path}")
    waves = [_scan_wave(config) for config in WAVES]
    latest_themes = waves[-1]["themes"]
    ranked_themes = rank_publishable_themes(latest_themes)
    result: dict[str, object] = {
        "artifact": "internal_analysis",
        "evidence_type": "aggregate-only QVoice improvement themes relevant to retention",
        "generated_on": date.today().isoformat(),
        "outcome": "pass",
        "minimum_cohort": MINIMUM_COHORT,
        "scope": "Descriptive improvement-comment themes only; no attrition join, causal claim, model feature, or individual-risk label.",
        "source_rule": "Original can-improve prompt in every wave; original processes-to-improve prompt additionally in Jan-2025 and Jan-2026 where available.",
        "denominator_rule": "Distinct participants with at least one selected improvement comment in the same wave.",
        "grain": "Distinct participant-wave-theme; themes are non-exclusive.",
        "deduplication_rule": "A participant contributes at most once per wave and taxonomy theme across selected improvement prompts.",
        "comment_text_persisted": False,
        "identifiers_persisted": False,
        "employee_rows_persisted": False,
        "taxonomy": {theme: {"alignment": rule["alignment"]} for theme, rule in TAXONOMY.items()},
        "explicit_intent_to_leave_rule": "Conservative non-negated multiword English/Spanish wording, calculated separately as an aggregate contextual signal and never as an employee label.",
        "waves": waves,
        "jan_2026_display": {
            "ranked_themes": ranked_themes,
            "trend_themes": ranked_themes[:TREND_THEME_COUNT],
            "explicit_intent_to_leave": waves[-1]["explicit_intent_to_leave"],
            "trend_comparability_note": "Jan-2025 and Jan-2026 additionally include the processes-to-improve prompt; the 2023-2026 trend is descriptive and not strictly comparable.",
        },
        "acceptance_rule": {
            "expected_waves": [config["wave"] for config in WAVES],
            "can_improve_available_in_every_wave": True,
            "processes_available_only_in_jan_2025_and_jan_2026": True,
            "zero_duplicate_participant_rows": True,
            "zero_improvement_comments_without_key": True,
            "theme_percentages_reconcile_to_wave_denominator": True,
            "published_values_minimum_n": MINIMUM_COHORT,
            "top_five_jan_2026_publishable_themes": TOP_THEME_COUNT,
            "no_identifiers_employee_rows_or_free_text": True,
        },
    }
    validate_evidence(result)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    return result


if __name__ == "__main__":
    import argparse

    parser = argparse.ArgumentParser(description="Build aggregate-only QVoice retention-concerns evidence.")
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--overwrite", action="store_true", help="Explicitly replace an existing evidence artifact.")
    args = parser.parse_args()
    evidence = build_evidence(args.output, overwrite=args.overwrite)
    print(json.dumps({"output": str(args.output), "waves": [wave["wave"] for wave in evidence["waves"]], "outcome": evidence["outcome"]}, ensure_ascii=False))

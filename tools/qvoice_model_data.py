"""Fuentes y reglas temporales de QVoice; datos individuales sólo en memoria."""

from __future__ import annotations

import hashlib
import json
import math
import re
from collections import Counter
from dataclasses import dataclass, field
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Any
from zipfile import ZipFile

import numpy as np
import pandas as pd
import yaml
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from lxml import etree

from build_qvoice_comment_evidence import _cell_value, _shared_strings, _workbook_sheet_path, _norm_header
from build_qvoice_comparable_favorable_evidence import EXPECTED_COMMON_BASES
from qvoice_excel_assessment import coerce_number, header_base, is_binary_header, is_scale_header

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_CONFIG = ROOT / "config/qvoice_model_experiment_20261006.yaml"
STATUSES = ("valid", "question_absent", "missing", "sentinel_255", "invalid", "no_linked_row", "duplicate")
BANDS = ("<50", "50–69", "70–79", "80–89", "90–99", "100")


def utc_now() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def write_json(path: Path, value: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2, allow_nan=False, default=str) + "\n", encoding="utf-8")


def load_config(path: Path = DEFAULT_CONFIG) -> dict:
    config = yaml.safe_load(path.read_text(encoding="utf-8"))
    if config.get("reporting_mode") != "presentation":
        raise ValueError("Este experimento publica exclusivamente agregados en modo presentation.")
    if config["privacy"].get("employee_level_rows_persisted") is not False:
        raise ValueError("No se permite persistir filas de empleados.")
    if config["control"]["source_key"] != "specs.b":
        raise ValueError("El control acordado es el modelo B.")
    if len(config["candidates"]) != len({c["id"] for c in config["candidates"]}):
        raise ValueError("Hay identificadores de candidatos duplicados.")
    for wave in config["sources"]["waves"]:
        availability_date(wave)
    return config


def availability_date(wave: dict) -> pd.Timestamp | None:
    """Fecha efectiva de uso; falla cerrado sin evidencia o politica aprobada."""
    availability = wave.get("availability") or {}
    status = availability.get("status")
    if status not in ("verified", "policy_assigned"):
        return None
    if not availability.get("date") or not availability.get("evidence_reference"):
        raise ValueError("Una fecha efectiva necesita fecha y referencia documental.")
    reference = ROOT / str(availability["evidence_reference"])
    if not reference.is_file():
        raise ValueError("No existe la evidencia documental o politica temporal de QVoice.")
    if status == "policy_assigned":
        policy = json.loads(reference.read_text(encoding="utf-8"))
        if policy.get("status") != "approved" or availability.get("basis") != "annual_lag_policy":
            raise ValueError("La asignacion temporal por politica no esta aprobada o no declara su base.")
        expected = policy.get("wave_to_model_year", {}).get(str(wave["year"]))
        if expected is None or int(expected) != int(wave["year"]) + 1:
            raise ValueError("La politica anual no asigna la wave al ano siguiente de forma consistente.")
    value = pd.Timestamp(availability["date"]).normalize()
    if value < pd.Timestamp(wave["descriptive_anchor"]):
        raise ValueError("La fecha efectiva no puede ser anterior al ancla de la encuesta.")
    if status == "policy_assigned" and value != pd.Timestamp(year=int(wave["year"]) + 1, month=1, day=1):
        raise ValueError("La politica anual debe iniciar el 1 de enero del ano siguiente.")
    return value


def select_wave(waves: list[dict], snapshot: Any, episode_start: Any, maximum_age_months: int = 12) -> dict | None:
    """La elección es global por wave: nunca rescata una respuesta anterior."""
    when = pd.Timestamp(snapshot)
    eligible = [(availability_date(w), w) for w in waves]
    eligible = [(d, w) for d, w in eligible if d is not None and d <= when]
    if not eligible:
        return None
    starts = [d for d, _ in eligible]
    if len(starts) != len(set(starts)):
        raise ValueError("Dos waves tienen la misma fecha de disponibilidad; se necesita desempate documentado.")
    known, wave = max(eligible, key=lambda pair: pair[0])
    if when >= known + pd.DateOffset(months=maximum_age_months):
        return None
    # Sin fecha de episodio no se supone continuidad laboral.
    if pd.isna(episode_start) or pd.Timestamp(episode_start) > pd.Timestamp(wave["descriptive_anchor"]):
        return None
    return wave


def parse_key(value: object) -> int | None:
    if value is None:
        return None
    text = str(value).strip()
    if re.fullmatch(r"[0-9]+(?:\.0+)?", text):
        return int(text.split(".")[0])
    return None


def cell_value(value: object, form: str = "binary", present: bool = True) -> tuple[int | None, str]:
    if not present:
        return None, "question_absent"
    if value is None or (isinstance(value, str) and not value.strip()):
        return None, "missing"
    number = coerce_number(value)
    if number == 255:
        return None, "sentinel_255"
    domain = (0, 1) if form == "binary" else (1, 2, 3, 4, 5)
    if number is None or not math.isfinite(number) or number not in domain:
        return None, "invalid"
    return int(number), "valid"


def favorable_band(values: list[int | None], *, required_count: int) -> str | None:
    if len(values) != required_count or any(v not in (0, 1) for v in values):
        return None
    total = sum(values)
    # Comparación entera exacta: conserva las seis bandas del playbook.
    for upper, label in zip((50, 70, 80, 90, 100), BANDS):
        if 100 * total < upper * required_count:
            return label
    return "100"


def relative_sign(item: int | None, scale_values: list[int | None], *, minimum_valid: int = 38, expected_items: int = 53) -> int | None:
    if len(scale_values) != expected_items or item not in (1, 2, 3, 4, 5):
        return None
    valid = [v for v in scale_values if v in (1, 2, 3, 4, 5)]
    if len(valid) < minimum_valid:
        return None
    difference = item * len(valid) - sum(valid)
    return (difference > 0) - (difference < 0)


def dictionary_items(path: Path) -> dict[str, dict]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = workbook.worksheets[0].iter_rows(values_only=True)
        headers = [str(v or "").strip().lower() for v in next(rows)]
        positions = {name: headers.index(name) for name in ("question_id", "question_text", "kpis")}
        items = {}
        for row in rows:
            key = str(row[positions["question_id"]] or "").strip()
            question = str(row[positions["question_text"]] or "").strip()
            if not key or not question:
                continue
            if key in items:
                raise ValueError("El diccionario contiene una pregunta duplicada.")
            items[key] = {"question_text": question, "base": header_base(question), "kpi": str(row[positions["kpis"]] or "")}
        return items
    finally:
        workbook.close()


def inspect_sources(config: dict) -> dict:
    """Preflight: sólo diccionario, encabezados y propiedades de archivo."""
    dictionary_path = ROOT / config["sources"]["dictionary"]
    items = dictionary_items(dictionary_path)
    waves = []
    for wave in config["sources"]["waves"]:
        path = ROOT / wave["file"]
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            sheet = workbook[wave["sheet"]]
            headers = [str(v or "").replace("\n", " ").strip() for v in next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))]
            properties = {"created": str(workbook.properties.created), "modified": str(workbook.properties.modified)}
            sheets = workbook.sheetnames
        finally:
            workbook.close()
        if wave["key_column"] not in headers:
            raise ValueError(f"Falta la columna de enlace declarada en {wave['year']}.")
        binary = {}
        scale = {}
        for form, target, predicate in (("binary", binary, is_binary_header), ("scale", scale, is_scale_header)):
            for index, header in enumerate(headers):
                if predicate(header):
                    target.setdefault(header_base(header), []).append(index)
        mapping = {}
        for key, item in items.items():
            mapping[key] = {
                "base": item["base"],
                "binary_indices": binary.get(item["base"], []),
                "scale_indices": scale.get(item["base"], []),
            }
        career_base = config["source_rules"]["career_wording_2024" if wave["year"] == 2024 else "career_wording_later"]
        mapping["career"] = {"base": career_base, "binary_indices": binary.get(career_base, []), "scale_indices": scale.get(career_base, [])}
        # Cada pregunta del promedio usa su texto exacto, sin equivalencias semánticas.
        for index, base in enumerate(EXPECTED_COMMON_BASES, 1):
            mapping[f"M{index:02d}"] = {"base": base, "binary_indices": binary.get(base, []), "scale_indices": scale.get(base, [])}
        for key, preferred in (config["source_rules"].get("binary_header_overrides", {}).get(str(wave["year"]), {})).items():
            entry = mapping[key]
            if headers.count(preferred) != 1 or headers.index(preferred) not in entry["binary_indices"]:
                raise ValueError(f"El encabezado principal declarado para {key} no se resuelve de forma única.")
            entry["alternative_binary_indices"] = [i for i in entry["binary_indices"] if headers[i] != preferred]
            entry["binary_indices"] = [headers.index(preferred)]
            entry["explicit_header"] = preferred
        waves.append({
            "year": wave["year"], "file": wave["file"], "sha256": sha256_file(path), "sheet": wave["sheet"],
            "sheets": sheets, "file_properties": properties, "key_column": wave["key_column"],
            "headers": headers, "mapping": mapping,
            "date_headers": [h for h in headers if re.search(r"recorded.?date|start.?date|end.?date|publish|availab|launch|survey.?date|close.?date", h, re.I)],
            "binary_base_count": len(binary), "scale_base_count": len(scale),
            "availability": wave["availability"],
        })
    return {
        "generated_at": utc_now(), "reporting_mode": "presentation", "metadata_only": True,
        "dictionary": items, "dictionary_sha256": sha256_file(dictionary_path), "waves": waves,
        "limitation": "Las propiedades created/modified, Recorded Date y el lanzamiento no acreditan la publicación de resultados. No se usaron como fecha predictiva.",
    }


@dataclass(repr=False)
class WaveRows:
    year: int
    rows: dict[int, dict[str, tuple[int | None, str]]] = field(repr=False)
    duplicate_keys: set[int] = field(repr=False)
    source_quality: dict
    mapping: dict


def scan_wave(wave: dict, metadata: dict, requested_keys: set[str]) -> WaveRows:
    """Extraer columnas estructuradas una vez; ninguna fila sale de este proceso."""
    headers = metadata["headers"]
    mapping = metadata["mapping"]
    selected = {}
    for key in sorted(requested_keys):
        entry = mapping.get(key, {})
        indices = entry.get("binary_indices", [])
        # No elegir silenciosamente entre dos encabezados del mismo concepto.
        if len(indices) > 1:
            raise ValueError(f"Hay encabezados binarios ambiguos para {key} en {wave['year']}.")
        if indices:
            selected[key] = (_norm_header(headers[indices[0]]), "binary")
        for alternate_index, index in enumerate(entry.get("alternative_binary_indices", [])):
            selected[f"alternate_{key}_{alternate_index}"] = (_norm_header(headers[index]), "binary")
    # Heartbeat usa todos los ítems del diccionario, no las preguntas de salida.
    for key, entry in mapping.items():
        if key.startswith("Q") and entry.get("scale_indices"):
            if len(entry["scale_indices"]) != 1:
                raise ValueError(f"Hay encabezados ordinales ambiguos para {key} en {wave['year']}.")
            selected[f"scale_{key}"] = (_norm_header(headers[entry["scale_indices"][0]]), "scale")
    key_header = _norm_header(wave["key_column"])
    response_header = _norm_header("Response Type") if "Response Type" in headers else None
    required = {key_header} | {h for h, _ in selected.values()}
    if response_header:
        required.add(response_header)
    rows = {}
    duplicates = set()
    quality = Counter({"source_rows": 0, "missing_key_rows": 0, "invalid_key_rows": 0, "duplicate_rows": 0, "conflicting_rows": 0, "excluded_response_type_rows": 0})
    for header_map, row in iter_structured_rows(ROOT / wave["file"], wave["sheet"], headers, required):
        quality["source_rows"] += 1
        if response_header:
            response_type = str(row.get(header_map[response_header]) or "").strip().upper()
            if response_type and response_type not in ("EX",):
                quality["excluded_response_type_rows"] += 1
                continue
        raw_key = row.get(header_map[key_header])
        key = parse_key(raw_key)
        if key is None:
            quality["missing_key_rows" if raw_key is None or not str(raw_key).strip() else "invalid_key_rows"] += 1
            continue
        parsed = {(h, form): cell_value(row.get(header_map[h]), form) for h, form in set(selected.values())}
        values = {k: parsed[(h, form)] for k, (h, form) in selected.items()}
        for k in requested_keys - set(selected):
            values[k] = (None, "question_absent")
        if key in rows:
            quality["duplicate_rows"] += 1
            quality["conflicting_rows"] += int(rows[key] != values)
            duplicates.add(key)
        else:
            rows[key] = values
    for key in duplicates:
        rows.pop(key, None)
    quality["distinct_duplicate_keys"] = len(duplicates)
    quality["unique_usable_keys"] = len(rows)
    return WaveRows(wave["year"], rows, duplicates, dict(quality), mapping)


def iter_structured_rows(path: Path, sheet_name: str, headers: list[str], required: set[str]):
    """Stream de celdas seleccionadas; evita convertir referencias de columnas ajenas."""
    header_map = {}
    for index, header in enumerate(headers):
        header_map.setdefault(_norm_header(header), index)
    if required - set(header_map):
        raise ValueError("No están todos los encabezados estructurados declarados.")
    selected = {get_column_letter(header_map[key] + 1): header_map[key] for key in required}
    with ZipFile(path) as archive:
        shared = _shared_strings(archive)
        with archive.open(_workbook_sheet_path(archive, sheet_name)) as stream:
            for _, element in etree.iterparse(stream, events=("end",), tag="{*}row"):
                if int(element.get("r", "0")) == 1:
                    element.clear()
                    continue
                values = {}
                for cell in element:
                    reference = cell.get("r", "")
                    index = selected.get(reference.rstrip("0123456789"))
                    if index is not None:
                        values[index] = _cell_value(cell, shared)
                yield header_map, values
                element.clear()
                while element.getprevious() is not None:
                    del element.getparent()[0]


def components_for(candidate: dict) -> list[str]:
    return [f"M{i:02d}" for i in range(1, 34)] if candidate["components"] == "strict_common_33" else list(candidate["components"])


def primary_components(candidate: dict) -> list[str]:
    return list(candidate.get("primary_components") or components_for(candidate))


def component_result(values: dict, key: str, *, heartbeat: bool = False) -> tuple[int | None, str]:
    if not heartbeat:
        return values.get(key, (None, "question_absent"))
    scale_values = [values.get(f"scale_Q{i}", (None, "question_absent"))[0] for i in range(1, 54)]
    present = sum(f"scale_Q{i}" in values for i in range(1, 54))
    item, status = values.get(f"scale_{key}", (None, "question_absent"))
    if present != 53:
        return None, "question_absent"
    result = relative_sign(item, scale_values)
    return (result, "valid") if result is not None else (None, status if status != "valid" else "missing")


def require_no_individual_columns(columns: list[str], allowed: list[str]) -> None:
    if columns != allowed:
        # Se muestran únicamente nombres de columnas, nunca sus valores.
        raise ValueError(f"La consulta no respeta el contrato de columnas: {columns}")

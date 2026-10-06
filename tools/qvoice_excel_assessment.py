"""Build aggregate evidence for QVoice Excel forms.

The raw workbooks contain employee-level survey data. This adapter reads the
workbooks locally, keeps identifiers only in memory for wave deduplication, and
writes aggregate rows suitable for the v4 assessment runner. No source rows or
identifiers are persisted.
"""

from __future__ import annotations

import argparse
import json
import math
import re
import unicodedata
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path
from statistics import mean, pstdev
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel


PROJECT_ROOT = Path(__file__).resolve().parents[1]
DATA_ROOT = PROJECT_ROOT / "data" / "QVoice"
DEFAULT_STRUCTURE = DATA_ROOT / "survey_structure.xlsx"


WORKBOOKS = [
    {
        "file": "Longitudinal Pulse File - May 2019 to June 2022.xlsx",
        "sheet": "Raw Data",
        "date_mode": "row",
        "date_names": ["snapshot date"],
        "wave_date": None,
    },
    {
        "file": "March 2023 Pulse_raw data.xlsx",
        "sheet": "Raw data",
        "date_mode": "fixed",
        "date_names": ["recorded date", "start date"],
        "wave_date": "2023-03-01",
    },
    {
        "file": "Jan 2024 Qvoice raw data.xlsx",
        "sheet": "Raw data",
        "date_mode": "fixed",
        "date_names": ["recorded date"],
        "wave_date": "2024-01-01",
    },
    {
        "file": "Jan 2025_Qvoice_raw data.xlsx",
        "sheet": "Raw data",
        "date_mode": "fixed",
        "date_names": ["recorded date"],
        "wave_date": "2025-01-01",
    },
    {
        "file": "Jan 2026 Qvoice_raw data.xlsx",
        "sheet": "Raw data",
        "date_mode": "fixed",
        "date_names": ["recorded date"],
        "wave_date": "2026-01-01",
    },
]


VARIABLES = {
    "Engagement items": {"form": "binary", "kind": "kpi", "kpi": "Engagement"},
    "Engagement items-scale": {"form": "scale", "kind": "kpi", "kpi": "Engagement"},
    "Intent to Stay items": {"form": "binary", "kind": "kpi", "kpi": "Intent to Stay"},
    "Intent to Stay items-scale": {"form": "scale", "kind": "kpi", "kpi": "Intent to Stay"},
    "Workplace Connection items": {"form": "binary", "kind": "kpi", "kpi": "Workplace Connection"},
    "Workplace Connection items-scale": {"form": "scale", "kind": "kpi", "kpi": "Workplace Connection"},
    "Wellbeing items": {"form": "binary", "kind": "kpi", "kpi": "Wellbeing"},
    "Wellbeing items-scale": {"form": "scale", "kind": "kpi", "kpi": "Wellbeing"},
    "Confidence items": {"form": "binary", "kind": "kpi", "kpi": "Confidence"},
    "Confidence items-scale": {"form": "scale", "kind": "kpi", "kpi": "Confidence"},
    "Career goals item": {"form": "binary", "kind": "career"},
    "Career goals item-scale": {"form": "scale", "kind": "career"},
    "Average of all Qvoice items": {"form": "binary", "kind": "average"},
    "Average of all Qvoice items-scale": {"form": "scale", "kind": "average"},
    "Standardized / normalized within region": {"form": "binary", "kind": "standardized"},
    "Standardized / normalized within region-scale": {"form": "scale", "kind": "standardized"},
}

QUESTION_PATTERNS = {
    "intent_1": "see myself working at qualcomm in a few years",
    "intent_2": "rarely think about looking for a new job",
    "career": "career goals can be met at qualcomm",
}

EXIT_ITEM_MARKERS = (
    "experience working at qualcomm meet your expectations",
    "right opportunity came along in the future",
    "consider rejoining qualcomm",
)


def ascii_text(value: Any) -> str:
    text = str(value or "")
    text = unicodedata.normalize("NFKD", text)
    text = "".join(char for char in text if not unicodedata.combining(char))
    return text


def norm(value: Any) -> str:
    text = ascii_text(value).lower().replace("\n", " ")
    return re.sub(r"\s+", " ", text).strip()


def load_survey_structure(path: Path) -> dict[str, list[dict[str, str]]]:
    """Load the KPI dictionary without persisting survey response data."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook[workbook.sheetnames[0]]
        rows = sheet.iter_rows(values_only=True)
        header_values = next(rows, None)
        if not header_values:
            raise ValueError(f"The survey structure workbook is empty: {path}")
        headers = [norm(value) for value in header_values]
        required = {"question_text", "kpis"}
        missing = required.difference(headers)
        if missing:
            raise ValueError(
                f"The survey structure workbook is missing columns: {sorted(missing)}"
            )
        question_index = headers.index("question_text")
        kpi_index = headers.index("kpis")
        question_id_index = headers.index("question_id") if "question_id" in headers else None
        definitions: dict[str, list[dict[str, str]]] = defaultdict(list)
        for row in rows:
            kpi = str(row[kpi_index] or "").strip()
            question = str(row[question_index] or "").strip()
            if not kpi or not question:
                continue
            definitions[kpi].append(
                {
                    "question_id": str(row[question_id_index] or "").strip()
                    if question_id_index is not None
                    else "",
                    "question_text": question,
                }
            )
        if not definitions:
            raise ValueError(f"No KPI mappings were found in: {path}")
        return dict(definitions)
    finally:
        workbook.close()


def header_base(header: str) -> str:
    text = norm(header)
    for marker in ("_1-5 scale", "_favorable", "_unfavorable", "_fav", "_s"):
        if text.endswith(marker):
            text = text[: -len(marker)]
            break
    text = re.sub(r"^q[0-9.]+[fu]_", "", text)
    text = re.sub(r"^q[0-9.]+_1-5_", "", text)
    text = text.split(" note:", 1)[0]
    return re.sub(r"\s+", " ", text).strip(" ._")


def is_binary_header(header: str) -> bool:
    text = norm(header)
    return text.endswith(("_favorable", "_fav"))


def is_legacy_binary_header(header: str) -> bool:
    return norm(header).endswith("_s")


def is_scale_header(header: str) -> bool:
    return "_1-5 scale" in norm(header)


def is_exit_item(header: str) -> bool:
    text = norm(header)
    return any(marker in text for marker in EXIT_ITEM_MARKERS)


def choose_raw_sheet(workbook: Any, configured: str) -> Any:
    if configured in workbook.sheetnames:
        return workbook[configured]
    for sheet in workbook.worksheets:
        if norm(sheet.title) == "raw data":
            return sheet
    for sheet in workbook.worksheets:
        if "raw" in norm(sheet.title) and "item_codes" not in norm(sheet.title):
            return sheet
    raise ValueError("No raw QVoice sheet was found")


def find_index(headers: list[str], names: list[str]) -> int | None:
    normalized = [norm(value) for value in headers]
    for name in names:
        target = norm(name)
        if target in normalized:
            return normalized.index(target)
    return None


def build_columns(
    headers: list[str], survey_structure: dict[str, list[dict[str, str]]]
) -> dict[str, Any]:
    binary_candidates = [
        index
        for index, header in enumerate(headers)
        if is_binary_header(header) and not is_exit_item(header)
    ]
    scale_candidates = [
        index for index, header in enumerate(headers) if is_scale_header(header)
    ]

    # When both forms exist, pairing the base text removes the four 2026 exit
    # items and avoids counting the duplicated trailing favorable columns.
    scale_bases = {header_base(headers[index]) for index in scale_candidates}
    if scale_bases:
        binary_candidates = [
            index for index in binary_candidates if header_base(headers[index]) in scale_bases
        ]

    def unique_by_base(indices: list[int]) -> list[int]:
        selected: list[int] = []
        seen: set[str] = set()
        for index in indices:
            key = header_base(headers[index])
            if key in seen:
                continue
            seen.add(key)
            selected.append(index)
        return selected

    binary_candidates = unique_by_base(binary_candidates)
    scale_candidates = unique_by_base(scale_candidates)

    def target(pattern: str, indices: list[int], allow_legacy_s: bool = False) -> int | None:
        matches = [index for index in indices if pattern in norm(headers[index])]
        if matches:
            return matches[0]
        if allow_legacy_s:
            legacy = [
                index
                for index, header in enumerate(headers)
                if pattern in norm(header) and norm(header).endswith("_s")
            ]
            if legacy:
                return legacy[0]
        return None

    def question_target(
        question_text: str, form: str, allow_legacy_s: bool = False
    ) -> int | None:
        question_base = header_base(question_text)
        if form == "binary":
            candidate_indices = [
                index
                for index, header in enumerate(headers)
                if (is_binary_header(header) or (allow_legacy_s and is_legacy_binary_header(header)))
                and not is_exit_item(header)
            ]
        else:
            candidate_indices = [
                index for index, header in enumerate(headers) if is_scale_header(header)
            ]
        exact = [
            index
            for index in candidate_indices
            if header_base(headers[index]) == question_base
        ]
        if exact:
            return exact[0]
        contains = [
            index
            for index in candidate_indices
            if question_base in header_base(headers[index])
            or header_base(headers[index]) in question_base
        ]
        return contains[0] if contains else None

    kpi_targets: dict[str, dict[str, list[int | None]]] = {}
    for kpi, questions in survey_structure.items():
        kpi_targets[kpi] = {
            "binary": [
                question_target(question["question_text"], "binary", allow_legacy_s=True)
                for question in questions
            ],
            "scale": [
                question_target(question["question_text"], "scale")
                for question in questions
            ],
        }

    binary_targets = {
        "intent_1": target(QUESTION_PATTERNS["intent_1"], binary_candidates, True),
        "intent_2": target(QUESTION_PATTERNS["intent_2"], binary_candidates, True),
        "career": target(QUESTION_PATTERNS["career"], binary_candidates, True),
    }
    scale_targets = {
        "intent_1": target(QUESTION_PATTERNS["intent_1"], scale_candidates),
        "intent_2": target(QUESTION_PATTERNS["intent_2"], scale_candidates),
        "career": target(QUESTION_PATTERNS["career"], scale_candidates),
    }

    return {
        "binary_items": binary_candidates,
        "scale_items": scale_candidates,
        "binary_targets": binary_targets,
        "scale_targets": scale_targets,
        "binary_kpi_targets": {
            kpi: targets["binary"] for kpi, targets in kpi_targets.items()
        },
        "scale_kpi_targets": {
            kpi: targets["scale"] for kpi, targets in kpi_targets.items()
        },
    }


def coerce_number(value: Any) -> float | None:
    if isinstance(value, bool):
        return float(int(value))
    if isinstance(value, (int, float)):
        number = float(value)
        return number if math.isfinite(number) else None
    if isinstance(value, str):
        text = value.strip()
        if not text:
            return None
        try:
            number = float(text)
        except ValueError:
            return None
        return number if math.isfinite(number) else None
    return None


def domain_value(value: Any, minimum: float, maximum: float) -> tuple[float | None, str]:
    if value is None or (isinstance(value, str) and not value.strip()):
        return None, "missing"
    number = coerce_number(value)
    if number is None or number < minimum or number > maximum or number != round(number):
        return None, "invalid"
    return number, "valid"


def parse_date_value(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)):
        try:
            converted = from_excel(value)
            if isinstance(converted, datetime):
                return converted.date()
            if isinstance(converted, date):
                return converted
        except (TypeError, ValueError, OverflowError):
            return None
    text = str(value or "").strip()
    if not text:
        return None
    text = text[:10]
    for fmt in ("%Y-%m-%d", "%m/%d/%Y", "%m/%d/%y", "%d-%b-%Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def wave_for_row(config: dict[str, Any], row: tuple[Any, ...], date_index: int | None) -> str | None:
    if config["date_mode"] == "fixed":
        return config["wave_date"]
    if date_index is None:
        return None
    parsed = parse_date_value(row[date_index])
    return parsed.isoformat() if parsed else None


def new_wave_state(columns: dict[str, Any], form: str) -> dict[str, Any]:
    item_indices = columns[f"{form}_items"]
    kpi_states = {}
    for kpi, target_indices in columns[f"{form}_kpi_targets"].items():
        kpi_states[kpi] = {
            "expected_component_count": len(target_indices),
            "available_component_count": sum(
                int(index is not None) for index in target_indices
            ),
            "populated": 0,
            "complete": 0,
            "partial": 0,
            "missing": 0,
            "invalid_count": 0,
            "valid_component_cells": 0,
            "distinct": set(),
            "composites": [],
        }
    return {
        "eligible": 0,
        "duplicate_rows": 0,
        "invalid_count": 0,
        "missing_region_count": 0,
        "region_present_count": 0,
        "populated": {"intent": 0, "career": 0, "average": 0},
        "distinct": {"intent": set(), "career": set(), "average": set()},
        "composites": [],
        "item_count": len(item_indices),
        "min_items_required": max(1, math.ceil(0.7 * len(item_indices))),
        "regions": Counter(),
        "kpis": kpi_states,
    }


def accept_response_type(value: Any) -> bool:
    if value is None or (isinstance(value, str) and not value.strip()):
        return True
    return norm(value) in {"ex", "complete", "completed"}


def process_file(
    config: dict[str, Any], survey_structure: dict[str, list[dict[str, str]]]
) -> tuple[dict[str, Any], dict[str, Any]]:
    path = DATA_ROOT / config["file"]
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = choose_raw_sheet(workbook, config["sheet"])
    headers = [str(value or "").replace("\n", " ").strip() for value in next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))]
    columns = build_columns(headers, survey_structure)
    normalized_headers = [norm(value) for value in headers]
    id_index = next(
        (
            index
            for index, header in enumerate(normalized_headers)
            if header in {"emp id", "empid", "employeeid", "employee id", "participant's unique identifier"}
        ),
        None,
    )
    response_index = find_index(headers, ["Response Type"])
    region_index = find_index(headers, ["Region"])
    date_index = find_index(headers, config["date_names"])

    wave_states: dict[str, dict[str, dict[str, Any]]] = defaultdict(dict)
    seen_keys: dict[str, set[tuple[str, str]]] = defaultdict(set)
    row_count = 0
    eligible_row_count = 0
    response_types = Counter()

    selected_indices = set(index for index in columns["binary_items"] + columns["scale_items"])
    selected_indices.update(index for index in columns["binary_targets"].values() if index is not None)
    selected_indices.update(index for index in columns["scale_targets"].values() if index is not None)
    for form in ("binary", "scale"):
        for target_indices in columns[f"{form}_kpi_targets"].values():
            selected_indices.update(index for index in target_indices if index is not None)
    if id_index is not None:
        selected_indices.add(id_index)
    if response_index is not None:
        selected_indices.add(response_index)
    if region_index is not None:
        selected_indices.add(region_index)
    if date_index is not None:
        selected_indices.add(date_index)

    for row in sheet.iter_rows(min_row=2, values_only=True):
        row_count += 1
        response_value = row[response_index] if response_index is not None else None
        response_types[str(response_value if response_value is not None else "NULL")] += 1
        if id_index is None or row[id_index] in (None, "") or not accept_response_type(response_value):
            continue
        wave = wave_for_row(config, row, date_index)
        if wave is None:
            continue
        eligible_row_count += 1
        identifier = str(row[id_index])
        key = (identifier, wave)
        for form in ("binary", "scale"):
            state = wave_states[form].setdefault(wave, new_wave_state(columns, form))
            if key in seen_keys[form]:
                state["duplicate_rows"] += 1
                continue
            seen_keys[form].add(key)
            state["eligible"] += 1

            for kpi, target_indices in columns[f"{form}_kpi_targets"].items():
                kpi_state = state["kpis"][kpi]
                valid_values: list[float] = []
                invalid_cells = 0
                for index in target_indices:
                    if index is None:
                        continue
                    value, status = domain_value(
                        row[index],
                        0 if form == "binary" else 1,
                        1 if form == "binary" else 5,
                    )
                    if status == "invalid":
                        invalid_cells += 1
                    if value is not None:
                        valid_values.append(value)
                kpi_state["invalid_count"] += invalid_cells
                kpi_state["valid_component_cells"] += len(valid_values)
                expected_count = kpi_state["expected_component_count"]
                if valid_values:
                    composite = round(mean(valid_values), 6)
                    kpi_state["populated"] += 1
                    kpi_state["distinct"].add(composite)
                    kpi_state["composites"].append(composite)
                    if len(valid_values) == expected_count:
                        kpi_state["complete"] += 1
                    else:
                        kpi_state["partial"] += 1
                else:
                    kpi_state["missing"] += 1

            targets = columns[f"{form}_targets"]
            for kind in ("career",):
                target_indices = [targets["career"]]
                if any(index is None for index in target_indices):
                    continue
                valid_values: list[float] = []
                statuses: list[str] = []
                for index in target_indices:
                    value, status = domain_value(row[index], 0 if form == "binary" else 1, 1 if form == "binary" else 5)
                    statuses.append(status)
                    if value is not None:
                        valid_values.append(value)
                state["invalid_count"] += sum(status == "invalid" for status in statuses)
                if len(valid_values) == len(target_indices) and all(status == "valid" for status in statuses):
                    state["populated"][kind] += 1
                    state["distinct"][kind].add(tuple(valid_values))

            item_values: list[float] = []
            invalid_cells = 0
            for index in columns[f"{form}_items"]:
                value, status = domain_value(row[index], 0 if form == "binary" else 1, 1 if form == "binary" else 5)
                if status == "invalid":
                    invalid_cells += 1
                if value is not None:
                    item_values.append(value)
            state["invalid_count"] += invalid_cells
            if len(item_values) >= state["min_items_required"] and invalid_cells == 0:
                composite = round(mean(item_values), 6)
                state["populated"]["average"] += 1
                state["distinct"]["average"].add(composite)
                region_value = row[region_index] if region_index is not None else None
                region = str(region_value).strip() if region_value not in (None, "") else ""
                state["composites"].append((composite, region))
                if region:
                    state["region_present_count"] += 1
                    state["regions"][region] += 1
                else:
                    state["missing_region_count"] += 1

    workbook.close()

    inventory = {
        "file": config["file"],
        "sheet": sheet.title,
        "rows_declared": int(sheet.max_row),
        "columns_declared": int(sheet.max_column),
        "rows_scanned": row_count,
        "eligible_rows_scanned": eligible_row_count,
        "response_types": dict(response_types),
        "date_mode": config["date_mode"],
        "date_column": headers[date_index] if date_index is not None else None,
        "id_column": headers[id_index] if id_index is not None else None,
        "region_column": headers[region_index] if region_index is not None else None,
        "binary_item_count": len(columns["binary_items"]),
        "scale_item_count": len(columns["scale_items"]),
        "binary_targets": {
            key: headers[index] if index is not None else None
            for key, index in columns["binary_targets"].items()
        },
        "scale_targets": {
            key: headers[index] if index is not None else None
            for key, index in columns["scale_targets"].items()
        },
        "kpi_components": {
            kpi: {
                "expected_question_count": len(survey_structure[kpi]),
                "binary": [
                    headers[index] if index is not None else None
                    for index in columns["binary_kpi_targets"][kpi]
                ],
                "scale": [
                    headers[index] if index is not None else None
                    for index in columns["scale_kpi_targets"][kpi]
                ],
                "binary_available_count": sum(
                    int(index is not None)
                    for index in columns["binary_kpi_targets"][kpi]
                ),
                "scale_available_count": sum(
                    int(index is not None)
                    for index in columns["scale_kpi_targets"][kpi]
                ),
            }
            for kpi in survey_structure
        },
        "binary_item_headers": [headers[index] for index in columns["binary_items"]],
        "scale_item_headers": [headers[index] for index in columns["scale_items"]],
        "selected_source_columns": [headers[index] for index in sorted(selected_indices)],
    }
    return dict(wave_states), inventory


def empty_audit_row(variable_name: str, form: str, wave: str) -> dict[str, Any]:
    return {
        "variable_name": variable_name,
        "snapshot_date": wave,
        "fiscal_year": int(wave[:4]),
        "eligible_count": 0,
        "populated_count": 0,
        "distinct_count": 0,
        "invalid_count": 0,
        "anomaly_count": 0,
        "unexplained_change_count": 0,
        "matched_count": 0,
        "unmatched_count": 0,
        "duplicate_count": 0,
        "item_count": 0,
        "min_items_required": 0,
        "item_set_change_count": 0,
        "regions_checked": 0,
        "standard_deviation_zero_count": 0,
        "expected_component_count": 0,
        "available_component_count": 0,
        "complete_component_count": 0,
        "partial_component_count": 0,
        "missing_component_count": 0,
        "component_header_coverage_percent": None,
        "component_response_coverage_percent": None,
        "full_kpi_response_percent": None,
        "component_mapping_status": "not_applicable",
        "notes": f"Agregado de Excel para {form}; no se publican filas ni identificadores.",
    }


def build_component_history(
    inventories: list[dict[str, Any]],
    states_by_file: list[dict[str, dict[str, dict[str, Any]]]],
    survey_structure: dict[str, list[dict[str, str]]],
) -> dict[tuple[str, str, str], list[dict[str, Any]]]:
    history: dict[tuple[str, str, str], list[dict[str, Any]]] = {}
    for inventory, states_by_form in zip(inventories, states_by_file):
        for form in ("binary", "scale"):
            for wave in states_by_form.get(form, {}):
                for kpi, questions in survey_structure.items():
                    headers = inventory["kpi_components"][kpi][form]
                    history[(form, wave, kpi)] = [
                        {
                            "question_id": question.get("question_id") or f"item_{index + 1}",
                            "question_text": question["question_text"],
                            "header": headers[index] if index < len(headers) else None,
                            "header_status": "disponible"
                            if index < len(headers) and headers[index]
                            else "no disponible",
                        }
                        for index, question in enumerate(questions)
                    ]
    return history


def component_history_note(
    kpi_name: str,
    form: str,
    wave: str,
    component_history: list[dict[str, Any]],
    complete: int,
    partial: int,
    missing: int,
) -> str:
    form_label = "binaria" if form == "binary" else "1-5"
    components = "; ".join(
        f"{item['question_id']}={item['header_status']}"
        for item in component_history
    ) or "sin componentes definidos"
    return (
        f"Historico de componentes {kpi_name} ({form_label}) en {wave}: {components}. "
        f"Respuestas completas={complete}; parciales={partial}; sin componente valido={missing}."
    )


def finalize_variable_rows(
    variable_name: str,
    form: str,
    kind: str,
    all_states: dict[str, dict[str, dict[str, Any]]],
    kpi_name: str | None = None,
    component_history: dict[tuple[str, str, str], list[dict[str, Any]]] | None = None,
) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    states = all_states.get(form, {})
    rows: list[dict[str, Any]] = []
    for wave in sorted(states):
        state = states[wave]
        if kind == "kpi":
            if not kpi_name or kpi_name not in state["kpis"]:
                raise ValueError(f"Missing KPI mapping for {variable_name}")
            kpi_state = state["kpis"][kpi_name]
            expected = kpi_state["expected_component_count"]
            available = kpi_state["available_component_count"]
            eligible = state["eligible"]
            populated = kpi_state["populated"]
            complete = kpi_state["complete"]
            partial = kpi_state["partial"]
            missing = kpi_state["missing"]
            component_rows = (component_history or {}).get((form, wave, kpi_name), [])
            row = empty_audit_row(variable_name, form, wave)
            row.update(
                {
                    "eligible_count": eligible,
                    "populated_count": populated,
                    "distinct_count": len(kpi_state["distinct"]),
                    "invalid_count": kpi_state["invalid_count"],
                    "matched_count": eligible,
                    "unmatched_count": 0,
                    "duplicate_count": state["duplicate_rows"],
                    "item_count": expected,
                    "min_items_required": 1,
                    "expected_component_count": expected,
                    "available_component_count": available,
                    "complete_component_count": complete,
                    "partial_component_count": partial,
                    "missing_component_count": missing,
                    "component_header_coverage_percent": round(
                        available / expected * 100, 4
                    )
                    if expected
                    else None,
                    "component_response_coverage_percent": round(
                        kpi_state["valid_component_cells"]
                        / (eligible * available)
                        * 100,
                        4,
                    )
                    if eligible and available
                    else None,
                    "full_kpi_response_percent": round(
                        complete / eligible * 100, 4
                    )
                    if eligible
                    else None,
                    "component_mapping_status": (
                        "complete"
                        if available == expected
                        else "partial_header_mapping"
                        if available
                        else "unavailable"
                    ),
                    "component_history": component_rows,
                    "notes": component_history_note(
                        kpi_name,
                        form,
                        wave,
                        component_rows,
                        complete,
                        partial,
                        missing,
                    ),
                }
            )
        elif kind == "standardized":
            region_groups: dict[str, list[float]] = defaultdict(list)
            for score, region in state["composites"]:
                if region:
                    region_groups[region].append(score)
            valid_groups = {
                region: values for region, values in region_groups.items() if len(values) >= 30
            }
            z_scores: list[float] = []
            std_zero_count = 0
            matched = 0
            unmatched = state["eligible"]
            for score, region in state["composites"]:
                values = valid_groups.get(region)
                if not values:
                    continue
                matched += 1
                unmatched -= 1
                deviation = pstdev(values)
                if deviation == 0:
                    std_zero_count += 1
                    z_scores.append(0.0)
                else:
                    z_scores.append(round((score - mean(values)) / deviation, 6))
            row = empty_audit_row(variable_name, form, wave)
            row.update(
                {
                    "eligible_count": state["eligible"],
                    "populated_count": len(z_scores),
                    "distinct_count": len(set(z_scores)),
                    "invalid_count": state["invalid_count"],
                    "matched_count": matched,
                    "unmatched_count": max(0, unmatched),
                    "duplicate_count": state["duplicate_rows"],
                    "item_count": state["item_count"],
                    "min_items_required": state["min_items_required"],
                    "regions_checked": len(valid_groups),
                    "standard_deviation_zero_count": std_zero_count,
                    "notes": (
                        "Z-score poblacional dentro de region y wave; se excluyen cohortes regionales "
                        "menores de 30 y se reporta por agregado."
                    ),
                }
            )
        else:
            populated_kind = "average" if kind == "average" else kind
            populated = state["populated"][populated_kind]
            distinct = len(state["distinct"][populated_kind])
            row = empty_audit_row(variable_name, form, wave)
            row.update(
                {
                    "eligible_count": state["eligible"],
                    "populated_count": populated,
                    "distinct_count": distinct,
                    "invalid_count": state["invalid_count"],
                    "matched_count": state["eligible"],
                    "unmatched_count": 0,
                    "duplicate_count": state["duplicate_rows"],
                    "item_count": state["item_count"],
                    "min_items_required": state["min_items_required"],
                    "regions_checked": len(state["regions"]),
                }
            )
        rows.append(row)

    counts = Counter(row["item_count"] for row in rows)
    if kind == "kpi":
        available_counts = Counter(row["available_component_count"] for row in rows)
        modal_item_count = (
            available_counts.most_common(1)[0][0] if available_counts else 0
        )
    else:
        modal_item_count = counts.most_common(1)[0][0] if counts else 0
    for row in rows:
        observed_count = (
            row["available_component_count"] if kind == "kpi" else row["item_count"]
        )
        changed = int(modal_item_count and observed_count != modal_item_count)
        row["item_set_change_count"] = changed
        row["unexplained_change_count"] = changed
    return rows, {"modal_item_count": modal_item_count, "observations": len(rows)}


def build_results(
    inventories: list[dict[str, Any]],
    states_by_file: list[dict[str, Any]],
    survey_structure: dict[str, list[dict[str, str]]],
    run_date: str,
) -> dict[str, Any]:
    merged: dict[str, dict[str, dict[str, dict[str, Any]]]] = defaultdict(lambda: defaultdict(dict))
    component_history = build_component_history(
        inventories, states_by_file, survey_structure
    )
    for states in states_by_file:
        for form, waves in states.items():
            for wave, state in waves.items():
                target = merged[form][wave]
                if not target:
                    merged[form][wave] = state
                else:
                    # File waves are distinct in the supplied raw inventory. If
                    # a future workbook overlaps, combine only safe aggregates.
                    for key in ("eligible", "duplicate_rows", "invalid_count", "missing_region_count", "region_present_count"):
                        target[key] += state[key]
                    for kind in ("intent", "career", "average"):
                        target["populated"][kind] += state["populated"][kind]
                        target["distinct"][kind].update(state["distinct"][kind])
                    for kpi, source_kpi in state.get("kpis", {}).items():
                        target_kpi = target.setdefault("kpis", {}).setdefault(
                            kpi,
                            {
                                "expected_component_count": source_kpi[
                                    "expected_component_count"
                                ],
                                "available_component_count": source_kpi[
                                    "available_component_count"
                                ],
                                "populated": 0,
                                "complete": 0,
                                "partial": 0,
                                "missing": 0,
                                "invalid_count": 0,
                                "valid_component_cells": 0,
                                "distinct": set(),
                                "composites": [],
                            },
                        )
                        for key in (
                            "populated",
                            "complete",
                            "partial",
                            "missing",
                            "invalid_count",
                            "valid_component_cells",
                        ):
                            target_kpi[key] += source_kpi[key]
                        target_kpi["distinct"].update(source_kpi["distinct"])
                        target_kpi["composites"].extend(source_kpi["composites"])
                    target["composites"].extend(state["composites"])
                    target["regions"].update(state["regions"])

    variables: dict[str, Any] = {}
    for variable_name, metadata in VARIABLES.items():
        form = metadata["form"]
        kind = metadata["kind"]
        row_kind = "average" if kind in {"average", "standardized"} else kind
        rows, summary = finalize_variable_rows(
            variable_name,
            form,
            row_kind,
            merged,
            kpi_name=metadata.get("kpi"),
            component_history=component_history,
        )
        if kind == "kpi":
            kpi_name = metadata["kpi"]
            expected = len(survey_structure[kpi_name])
            available = sum(
                inventory["kpi_components"][kpi_name][f"{form}_available_count"]
                for inventory in inventories
            )
            missing_items = len(inventories) * expected - available
            source_rows = [
                {
                    "variable_name": variable_name,
                    "kpi_name": kpi_name,
                    "expected_component_count": expected,
                    "available_component_count": available,
                    "missing_items": missing_items,
                    "workbooks_checked": len(inventories),
                    "usable_form_workbooks": sum(
                        int(
                            inventory["kpi_components"][kpi_name][
                                f"{form}_available_count"
                            ]
                            > 0
                        )
                        for inventory in inventories
                    ),
                    "observations_checked": len(rows),
                }
            ]
        else:
            source_rows = [
                {
                    "variable_name": variable_name,
                    "missing_items": 0,
                    "workbooks_checked": len(inventories),
                    "usable_form_workbooks": sum(
                        int((inventory[f"{form}_item_count"] or 0) > 0)
                        for inventory in inventories
                    ),
                    "observations_checked": len(rows),
                }
            ]
        variables[variable_name] = {
            "coverage": [
                {
                    "snapshot_date": row["snapshot_date"],
                    "fiscal_year": row["fiscal_year"],
                    "eligible_employees": row["eligible_count"],
                    "observable_employees": row["populated_count"],
                    "event_employees": None,
                    "notes": row["notes"],
                    "expected_component_count": row["expected_component_count"],
                    "available_component_count": row["available_component_count"],
                    "complete_component_count": row["complete_component_count"],
                    "partial_component_count": row["partial_component_count"],
                    "missing_component_count": row["missing_component_count"],
                    "component_header_coverage_percent": row[
                        "component_header_coverage_percent"
                    ],
                    "component_response_coverage_percent": row[
                        "component_response_coverage_percent"
                    ],
                    "full_kpi_response_percent": row["full_kpi_response_percent"],
                    "component_mapping_status": row["component_mapping_status"],
                    "component_history": row.get("component_history", []),
                }
                for row in rows
            ],
            "queries": {
                f"queries/audits/qvoice_excel_source_inventory_{run_date}.sql": source_rows,
                f"queries/audits/qvoice_excel_definition_audit_{run_date}.sql": rows,
                f"queries/audits/qvoice_excel_consistency_audit_{run_date}.sql": rows,
                f"queries/audits/qvoice_excel_join_grain_audit_{run_date}.sql": rows,
            },
            "summary": summary,
        }

    return {
        "schema_version": 4,
        "source": "Excel raw QVoice workbooks; aggregate-only frozen results",
        "generated_at": datetime.now().astimezone().isoformat(timespec="seconds"),
        "minimum_shared_cohort": 30,
        "variables": variables,
    }


def write_mapping_markdown(
    path: Path,
    inventories: list[dict[str, Any]],
    results: dict[str, Any],
    survey_structure: dict[str, list[dict[str, str]]],
) -> None:
    lines = [
        "# Mapeo de QVoice Excel para assessment v4",
        "",
        f"Fecha de generacion: {results['generated_at'][:10]}",
        "",
        "Este inventario se genero leyendo las hojas raw. El extractor usa los identificadores solo en memoria para deduplicar empleado-wave; los artefactos persistidos contienen unicamente conteos y metricas agregadas.",
        "",
        "## Diccionario KPI",
        "",
        "`survey_structure.xlsx` define las preguntas que componen cada KPI. El promedio de un KPI usa las respuestas validas de sus preguntas disponibles; si falta una pregunta en una wave, se permite un promedio parcial solo como evidencia exploratoria y se registra la brecha.",
        "",
    ]
    for kpi, questions in survey_structure.items():
        lines.append(f"### `{kpi}`")
        lines.append("")
        for question in questions:
            question_id = question.get("question_id") or "sin id"
            lines.append(f"- `{question_id}`: {question['question_text']}")
        lines.append("")
    lines.extend(
        [
        "## Formas",
        "",
        "- Las variables sin sufijo usan columnas binarias `_favorable` o `_FAV`; el longitudinal tambien conserva un encabezado con sufijo `_S`. Su semantica de negocio no esta confirmada y los valores observados en ese campo son binarios `0/1`.",
        "- Las variables con sufijo `-scale` usan columnas `_1-5 scale`.",
        "- En 2026 se conservan los 53 items que tienen pareja binaria y escala; los cuatro items adicionales de exit survey quedan fuera.",
        "- El promedio exige al menos 70% de los items de la wave y excluye filas con valores invalidos del composite; es una regla tecnica pendiente de aprobacion.",
        "- La normalizacion calcula un z-score poblacional por region y wave; las cohortes menores de 30 no se publican ni se usan para el z-score.",
        "",
        "## Inventario por workbook",
        "",
        "| Archivo | Hoja | Filas escaneadas | Binarios | Escala 1-5 | Fecha / wave |",
        "|---|---|---:|---:|---:|---|",
        ]
    )
    for inventory in inventories:
        date_label = inventory.get("date_column") or "wave fija"
        lines.append(
            f"| `{inventory['file']}` | `{inventory['sheet']}` | {inventory['rows_scanned']} | "
            f"{inventory['binary_item_count']} | {inventory['scale_item_count']} | {date_label} |"
        )
    lines.extend(["", "## Componentes seleccionados", ""])
    for inventory in inventories:
        lines.append(f"### `{inventory['file']}`")
        lines.append("")
        for form_key in ("binary_targets", "scale_targets"):
            lines.append(f"- `{form_key}`:")
            for key, value in inventory[form_key].items():
                lines.append(f"  - `{key}`: {value or 'no disponible'}")
        lines.append("- `kpi_components`:")
        for kpi, component_info in inventory["kpi_components"].items():
            for form in ("binary", "scale"):
                values = component_info[form]
                lines.append(
                    f"  - `{kpi}` `{form}`: "
                    + ", ".join(value or "no disponible" for value in values)
                )
        lines.append("")
    lines.extend(["## Resultado agregado por variable", "", "| Variable | Waves | Primera wave | Ultima wave |", "|---|---:|---|---|"])
    for name, payload in results["variables"].items():
        rows = payload["coverage"]
        lines.append(
            f"| `{name}` | {len(rows)} | {rows[0]['snapshot_date'] if rows else 'n/a'} | "
            f"{rows[-1]['snapshot_date'] if rows else 'n/a'} |"
        )
    lines.extend(
        [
            "",
            "## Limitaciones que pasan al assessment",
            "",
            "- Excel no aporta en este flujo un roster oficial de elegibilidad ni una reconciliacion completa al snapshot mensual de workforce; Join queda como validacion pendiente.",
            "- Los cambios historicos de wording y de item set se conservan como metricas de consistencia; no se tratan como equivalencia automatica entre waves.",
            "- La fuente es confidencial y el reporte del assessment debe interpretarse bajo el blocker de gobernanza configurado para survey data.",
        ]
    )
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def main() -> None:
    parser = argparse.ArgumentParser(description="Aggregate QVoice Excel evidence without persisting employee rows.")
    parser.add_argument("--structure", default=str(DEFAULT_STRUCTURE))
    parser.add_argument("--run-date", default=datetime.now().strftime("%Y%m%d"))
    parser.add_argument("--output")
    parser.add_argument("--inventory")
    parser.add_argument("--mapping-md")
    args = parser.parse_args()

    structure_path = Path(args.structure).resolve()
    output_path = (
        Path(args.output).resolve()
        if args.output
        else DATA_ROOT / f"qvoice_excel_aggregate_results_{args.run_date}.json"
    )
    inventory_path = (
        Path(args.inventory).resolve()
        if args.inventory
        else DATA_ROOT / f"qvoice_excel_source_inventory_{args.run_date}.json"
    )
    mapping_path = (
        Path(args.mapping_md).resolve()
        if args.mapping_md
        else DATA_ROOT / f"qvoice_excel_mapping_{args.run_date}.md"
    )
    survey_structure = load_survey_structure(structure_path)
    inventories: list[dict[str, Any]] = []
    states_by_file: list[dict[str, Any]] = []
    for config in WORKBOOKS:
        print(f"Procesando {config['file']}...", flush=True)
        states, inventory = process_file(config, survey_structure)
        inventories.append(inventory)
        states_by_file.append(states)
        print(
            f"  filas={inventory['rows_scanned']} binarios={inventory['binary_item_count']} escala={inventory['scale_item_count']}",
            flush=True,
        )

    results = build_results(inventories, states_by_file, survey_structure, args.run_date)
    inventory_path.write_text(
        json.dumps(
            {
                "schema_version": 1,
                "generated_at": results["generated_at"],
                "workbooks": inventories,
            },
            ensure_ascii=True,
            indent=2,
        )
        + "\n",
        encoding="utf-8",
    )
    output_path.write_text(json.dumps(results, ensure_ascii=True, indent=2) + "\n", encoding="utf-8")
    write_mapping_markdown(mapping_path, inventories, results, survey_structure)
    print(f"Agregados: {output_path}")
    print(f"Inventario: {inventory_path}")
    print(f"Mapeo: {mapping_path}")


if __name__ == "__main__":
    main()

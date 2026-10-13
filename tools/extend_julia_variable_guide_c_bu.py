"""Añade roles y codificaciones C+BU al diccionario existente, sin nuevos scores."""
from __future__ import annotations
from copy import copy
import json
from pathlib import Path
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]
OLD=ROOT/'output/analysis/engineering/20261012/julia-rh-review-v1/documentacion'
NEW=ROOT/'output/analysis/engineering/20261013/c-bu-isolation-v1'


def read(p): return json.loads(p.read_text(encoding='utf-8-sig'))


def main():
    data=read(OLD/'variables-44-final.json'); assert len(data['rows'])==44
    book=load_workbook(OLD/'variables-44.xlsx')
    usage=book['Uso por modelo y corte']; encoding=book['Codificaciones de modelos']; treatment=book['Tratamiento ejecutado']
    prior_treatments={}
    for row in list(treatment.values)[1:]:
        period,model,component,feature,*details=row
        if model=='FQ_BU_L3':prior_treatments[(period,component,feature)]=details
    new_codes=[]; added=0
    for period in ['aug_dec_2025','jan_feb_2026']:
        for model in ['C_BU','C_BU_Q']:
            artifact=read(NEW/'models'/f'{period}-{model}.json')
            components=[('RH',artifact['base']['encoder'])]
            if model=='C_BU_Q': components.append(('QVoice',artifact['q_encoder']))
            features={f for _,encoder in components for f in encoder['features']}
            covered=set()
            for item in data['rows']:
                candidates=set(item['representaciones'])
                # La representación cruzada ejecutada pertenece a ambos conceptos.
                if candidates & {'tenure_playbook_band','level_grouped'}:
                    candidates.add('level_x_tenure')
                for entry in item['uso_por_modelo_corte'].values():candidates.update(entry['predictor'])
                used=sorted(candidates&features); covered.update(used)
                grouping=item['uso_por_modelo_corte'][period+'/CQ']['definicion_grupos']
                excluded=sorted(set(item['representaciones'])-set(used)-set(grouping))
                reason=('C + BU con estructura y L2 congelados; no se usan respuestas ni indicadores de ausencia QVoice.'
                        if model=='C_BU' else 'C + BU más selección QVoice histórica de CQ; sin ciudad laboral.')
                role={'predictor':used,'definicion_grupos':grouping,'diagnostico':[],
                      'subdivision_experimental':False,'excluidas_de_esta_comparacion':excluded,
                      'filtro_poblacion':item['uso_por_modelo_corte'][period+'/CQ']['filtro_poblacion'],
                      'motivo':reason,'model_artifact':f'models/{period}-{model}.json'}
                item['uso_por_modelo_corte'][period+'/'+model]=role
                usage.append([item['numero'],item['variable'],period+'/'+model,', '.join(used),', '.join(grouping),
                              '',', '.join(excluded),'No',reason]);added+=1
            assert features-covered <= {'calendar_month'},features-covered
            for component,enc in components:
                for feature in enc['features']:
                    pool=artifact['pool_keep'] if feature=='bu_l3' else None
                    encoding.append([period,model,component,feature,enc['references'][feature],'; '.join(enc['levels'][feature]),
                                     '; '.join(pool) if pool is not None else None])
                    details=prior_treatments[(period,component,feature)]
                    treatment.append([period,model,component,feature,*details])
                    new_codes.append({'period':period,'model':model,'component':component,'feature':feature,
                                      'reference':enc['references'][feature],'levels':enc['levels'][feature],'pool_keep':pool,
                                      'transformation_and_missing':details[0],'unseen_category_behavior':details[2]})
    assert added==44*4
    for sheet in [usage,encoding,treatment]:
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                source=sheet.cell(2,cell.column)
                if cell.row>2:
                    cell._style=copy(source._style)
                    cell.alignment=copy(source.alignment)
        sheet.auto_filter.ref=sheet.dimensions
    out=NEW/'documentacion'; out.mkdir(exist_ok=True)
    data.update(source_inventory=str((OLD/'variables-44-final.json').relative_to(ROOT)),
                additional_models=['C_BU','C_BU_Q'],stage='ampliacion_documental_sin_nuevos_scores',
                limitation='Las acciones candidatas no tienen efecto causal estimado; las nuevas pruebas mantienen grupos comunes.')
    (out/'variables-44-ampliada.json').write_text(json.dumps(data,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    (out/'codificaciones-c-bu.json').write_text(json.dumps(new_codes,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    book.save(out/'variables-44-ampliada.xlsx')
    check=load_workbook(out/'variables-44-ampliada.xlsx',read_only=True)
    assert check['44 variables y acciones'].max_row==45
    assert check['Uso por modelo y corte'].max_row==441
    assert all(not f['feature'].startswith('qv_') for f in new_codes if f['model']=='C_BU')
    (out/'verificacion.json').write_text(json.dumps({'inventory_rows':44,'added_model_period_role_rows':added,
         'model_artifact_encodings_used':True,'C_BU_has_no_QVoice':True,'assessment_score_changed':False,
         'employee_level_rows_persisted':False},indent=2)+'\n',encoding='utf-8')
    print('44 filas conservadas; 176 usos añadidos; codificaciones extraídas de artefactos.')


if __name__=='__main__': main()

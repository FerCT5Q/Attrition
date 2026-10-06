"""Freeze aggregate-only QVoice z-score evidence for the regional-normalization slide.

Participant identifiers, raw responses, raw scores, and stored region labels are
used only while the source workbooks are open.  The persisted artifact contains
only safe, aggregate wave and z-score-band metrics.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import math
from collections import Counter, defaultdict
from datetime import date
from decimal import Decimal
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from build_qvoice_comparable_favorable_evidence import (
    EXPECTED_BINARY_ITEM_COUNTS,
    EXPECTED_COMMON_BASES,
    EXPECTED_WAVES,
    ITEM_COUNT,
    MINIMUM_COHORT,
    STRICT_HEADER_RULE,
    WAVES,
    _cell_status,
    _headers_and_binary_bases,
)
from qvoice_excel_assessment import load_survey_structure


ROOT = Path(__file__).resolve().parents[1]
QVOICE_ROOT = ROOT / "data" / "QVoice"
OUTPUT = ROOT / "output" / "distribution" / "json" / "qvoice-normalized-region-evidence-20260910.json"
EXPECTED_REGION_WAVE_COHORTS = 5
Z_BANDS = (
    ("< −1.0", None, -1.0),
    ("−1.0 to < −0.5", -1.0, -0.5),
    ("−0.5 to < 0.5", -0.5, 0.5),
    ("0.5 to < 1.0", 0.5, 1.0),
    ("≥ 1.0", 1.0, None),
)


def strict_item_contract_hash(common_bases: tuple[str, ...]) -> str:
    """Return a reproducible digest without persisting question wording."""
    payload = "\n".join(common_bases).encode("utf-8")
    return hashlib.sha256(payload).hexdigest()


def z_band_label(value: float) -> str:
    """Assign one standardized score to a deterministic, contiguous display band."""
    if value < -1.0:
        return Z_BANDS[0][0]
    if value < -0.5:
        return Z_BANDS[1][0]
    if value < 0.5:
        return Z_BANDS[2][0]
    if value < 1.0:
        return Z_BANDS[3][0]
    return Z_BANDS[4][0]


def normalized_distribution(scores_by_region: dict[str, list[int]]) -> dict[str, object]:
    """Build a privacy-safe z-score distribution without retaining region labels.

    The population standard deviation is deliberate: every valid stored
    Region × wave reference cohort has mean zero and standard deviation one by
    construction after normalization.
    """
    if not scores_by_region:
        raise ValueError("No complete QVoice Region × wave cohorts were found.")

    cohort_counts = [len(scores) for scores in scores_by_region.values()]
    if any(count < MINIMUM_COHORT for count in cohort_counts):
        raise ValueError("A Region × wave reference cohort is below the minimum publication cohort.")
    if len(scores_by_region) != EXPECTED_REGION_WAVE_COHORTS:
        raise ValueError(
            f"Expected {EXPECTED_REGION_WAVE_COHORTS} stored Region × wave cohorts; found {len(scores_by_region)}."
        )

    z_scores: list[float] = []
    for scores in scores_by_region.values():
        mean = math.fsum(scores) / len(scores)
        variance = math.fsum((score - mean) ** 2 for score in scores) / len(scores)
        standard_deviation = math.sqrt(variance)
        if standard_deviation <= 0:
            raise ValueError("A Region × wave reference cohort has zero population standard deviation.")
        z_scores.extend((score - mean) / standard_deviation for score in scores)

    normalized_count = len(z_scores)
    z_mean = math.fsum(z_scores) / normalized_count
    z_variance = math.fsum((score - z_mean) ** 2 for score in z_scores) / normalized_count
    z_standard_deviation = math.sqrt(z_variance)
    if not math.isclose(z_mean, 0.0, abs_tol=1e-10):
        raise ValueError("The combined normalized population does not have mean zero.")
    if not math.isclose(z_standard_deviation, 1.0, abs_tol=1e-10):
        raise ValueError("The combined normalized population does not have population standard deviation one.")

    bands = Counter(z_band_label(score) for score in z_scores)
    band_counts = [int(bands[label]) for label, _, _ in Z_BANDS]
    # Persist six-decimal percentages that sum to exactly 100.000000 in the
    # frozen artifact.  The final displayed band receives the deterministic
    # residual so rounding cannot create a non-reconciling publication total.
    rounded_percents = [
        round(100 * count / normalized_count, 6)
        for count in band_counts[:-1]
    ]
    rounded_percents.append(round(100 - sum(rounded_percents), 6))
    band_rows = [
        {
            "band": label,
            "participant_count": count,
            "percent_of_normalized": percent,
        }
        for (label, _, _), count, percent in zip(Z_BANDS, band_counts, rounded_percents)
    ]
    if sum(int(row["participant_count"]) for row in band_rows) != normalized_count:
        raise ValueError("Z-score bands do not reconcile to the normalized population.")
    if any(int(row["participant_count"]) < MINIMUM_COHORT for row in band_rows):
        raise ValueError("A displayed z-score band is below the minimum publication cohort.")

    return {
        "region_wave_cohort_count": len(scores_by_region),
        "minimum_region_wave_cohort_count": min(cohort_counts),
        "maximum_region_wave_cohort_count": max(cohort_counts),
        "zero_population_standard_deviation_cohort_count": 0,
        "normalized_participants": normalized_count,
        "weighted_z_score_mean": round(z_mean, 12),
        "weighted_population_z_score_standard_deviation": round(z_standard_deviation, 12),
        "z_score_bands": band_rows,
        "z_score_band_count_total": normalized_count,
        "minimum_z_score_band_count": min(int(row["participant_count"]) for row in band_rows),
    }


def scan_wave(
    config: dict[str, object],
    headers: list[str],
    column_by_base: dict[str, int],
    common_bases: tuple[str, ...],
) -> dict[str, object]:
    """Read one workbook and return aggregate-only normalized evidence."""
    year = int(config["year"])
    try:
        key_index = headers.index(str(config["key_column"]))
        region_index = headers.index("Region")
    except ValueError as exc:
        raise ValueError(f"The required participant key or Region header is missing for {year}.") from exc

    selected_indices = [column_by_base[base] for base in common_bases]
    path = QVOICE_ROOT / str(config["file"])
    seen_signatures: dict[str, tuple[str, tuple[tuple[int | None, str], ...]]] = {}
    scores_by_region: dict[str, list[int]] = defaultdict(list)
    source_rows = rows_without_key = duplicate_rows = duplicate_conflicts = 0
    complete = incomplete = complete_with_missing_region = invalid_cells = 0

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook[str(config["sheet"])]
        for row in worksheet.iter_rows(min_row=2, values_only=True):
            source_rows += 1
            raw_key = row[key_index]
            key = str(raw_key).strip() if raw_key is not None else ""
            if not key:
                rows_without_key += 1
                continue

            region = str(row[region_index]).strip() if row[region_index] is not None else ""
            statuses = tuple(_cell_status(row[index]) for index in selected_indices)
            signature = (region, statuses)
            previous = seen_signatures.get(key)
            if previous is not None:
                duplicate_rows += 1
                if previous != signature:
                    duplicate_conflicts += 1
                continue
            seen_signatures[key] = signature

            invalid_cells += sum(status == "invalid" for _, status in statuses)
            values = [value for value, status in statuses if status == "valid"]
            if len(values) != ITEM_COUNT:
                incomplete += 1
                continue
            complete += 1
            if not region:
                complete_with_missing_region += 1
                continue
            scores_by_region[region].append(sum(values))
    finally:
        workbook.close()

    participant_rows = len(seen_signatures)
    if participant_rows + duplicate_rows + rows_without_key != source_rows:
        raise ValueError(f"Source-row reconciliation failed for {year}.")
    if complete + incomplete != participant_rows:
        raise ValueError(f"Completeness reconciliation failed for {year}.")
    if duplicate_rows or duplicate_conflicts:
        raise ValueError(f"Duplicate participant rows or conflicts were found for {year}.")
    if invalid_cells:
        raise ValueError(f"Out-of-domain non-sentinel binary values were found for {year}.")
    if complete_with_missing_region:
        raise ValueError(f"Complete comparable respondents without Region were found for {year}.")

    expected = EXPECTED_WAVES[year]
    if participant_rows != expected["participants"] or complete != expected["complete"]:
        raise ValueError(f"The frozen 33/33 comparable denominator changed for {year}.")

    normalized = normalized_distribution(scores_by_region)
    if int(normalized["normalized_participants"]) != complete:
        raise ValueError(f"Normalized participant count does not reconcile for {year}.")
    return {
        "year": year,
        "source_file": str(config["file"]),
        "sheet": str(config["sheet"]),
        "participant_rows": participant_rows,
        "source_rows_scanned": source_rows,
        "rows_without_participant_key": rows_without_key,
        "duplicate_participant_rows": duplicate_rows,
        "duplicate_conflicting_rows": duplicate_conflicts,
        "complete_comparable_participants": complete,
        "incomplete_comparable_participants": incomplete,
        "complete_comparable_participants_without_region": complete_with_missing_region,
        "out_of_domain_non_sentinel_cells": invalid_cells,
        **normalized,
    }


def validate_evidence(evidence: dict[str, object]) -> dict[str, object]:
    """Validate the persisted, aggregate-only publication contract."""
    # A closed artifact schema is also a privacy control: participant-level
    # helper data must not be able to slip into the frozen JSON unnoticed.
    expected_top_level_keys = {
        "artifact", "evidence_type", "generated_on", "outcome",
        "minimum_shared_cohort", "source_rule", "scope", "grain", "formula",
        "strict_item_contract", "waves", "privacy_checks", "acceptance_rule",
    }
    expected_formula_keys = {
        "raw_score", "reference_cohort", "standardization", "inclusion_rule",
    }
    expected_contract_keys = {
        "binary_items_by_wave", "strict_common_item_count", "header_matching_rule",
        "common_header_bases_sha256",
    }
    expected_privacy_keys = {
        "identifiers_persisted", "employee_rows_persisted", "response_values_persisted",
        "region_labels_persisted", "all_reference_and_display_cohorts_meet_minimum",
    }
    expected_acceptance_keys = {
        "expected_years", "strict_common_item_count", "region_wave_cohorts_per_wave",
        "minimum_reference_and_display_cohort", "zero_duplicate_participant_rows",
        "zero_duplicate_conflicting_rows", "zero_complete_participants_without_region",
        "zero_out_of_domain_non_sentinel_cells", "zero_zero_standard_deviation_cohorts",
        "normalized_population_reconciles_to_33_of_33_population",
        "weighted_z_mean_zero_and_population_standard_deviation_one",
        "no_identifier_response_or_region_label_persistence",
    }
    expected_wave_keys = {
        "year", "source_file", "sheet", "participant_rows", "source_rows_scanned",
        "rows_without_participant_key", "duplicate_participant_rows",
        "duplicate_conflicting_rows", "complete_comparable_participants",
        "incomplete_comparable_participants",
        "complete_comparable_participants_without_region",
        "out_of_domain_non_sentinel_cells", "region_wave_cohort_count",
        "minimum_region_wave_cohort_count", "maximum_region_wave_cohort_count",
        "zero_population_standard_deviation_cohort_count", "normalized_participants",
        "weighted_z_score_mean", "weighted_population_z_score_standard_deviation",
        "z_score_bands", "z_score_band_count_total", "minimum_z_score_band_count",
    }
    expected_band_keys = {"band", "participant_count", "percent_of_normalized"}
    prohibited_fields = {
        "employee_id",
        "employee_ids",
        "employee_key",
        "employee_keys",
        "emp_id",
        "emp_ids",
        "empid",
        "manager_id",
        "manager_ids",
        "manager_employee_id",
        "manager_employee_ids",
        "participant_id",
        "participant_ids",
        "participant_key",
        "participant_keys",
        "employee_name",
        "employee_names",
        "name",
        "names",
        "region_label",
        "region_labels",
        "individual_score",
        "z_score",
        "response_value",
        "response_values",
    }

    def reject_prohibited_fields(value: object) -> None:
        if isinstance(value, dict):
            for key, nested in value.items():
                if str(key).lower() in prohibited_fields:
                    raise ValueError(f"QVoice normalized evidence persists prohibited field {key!r}.")
                reject_prohibited_fields(nested)
        elif isinstance(value, list):
            for nested in value:
                reject_prohibited_fields(nested)

    if set(evidence) != expected_top_level_keys:
        raise ValueError("QVoice normalized evidence has an unexpected top-level payload.")
    for container, expected_keys, label in (
        (evidence.get("formula"), expected_formula_keys, "formula"),
        (evidence.get("strict_item_contract"), expected_contract_keys, "strict item contract"),
        (evidence.get("privacy_checks"), expected_privacy_keys, "privacy checks"),
        (evidence.get("acceptance_rule"), expected_acceptance_keys, "acceptance rule"),
    ):
        if not isinstance(container, dict) or set(container) != expected_keys:
            raise ValueError(f"QVoice normalized evidence has an unexpected {label} payload.")

    # raw_score is permitted only as the documented formula string.  The closed
    # wave and display-band schemas below prevent any individual score payload.
    reject_prohibited_fields(
        {key: value for key, value in evidence.items() if key != "formula"}
    )
    if evidence.get("artifact") != "internal_analysis" or evidence.get("outcome") != "pass":
        raise ValueError("QVoice normalized evidence must be an internal aggregate passing artifact.")
    if int(evidence.get("minimum_shared_cohort", 0)) != MINIMUM_COHORT:
        raise ValueError("QVoice normalized evidence must retain the minimum publication cohort.")
    privacy = evidence.get("privacy_checks")
    if not isinstance(privacy, dict) or any(
        privacy.get(key) is not False
        for key in ("identifiers_persisted", "employee_rows_persisted", "response_values_persisted", "region_labels_persisted")
    ) or privacy.get("all_reference_and_display_cohorts_meet_minimum") is not True:
        raise ValueError("QVoice normalized evidence contains a prohibited persisted payload.")
    contract = evidence.get("strict_item_contract")
    if not isinstance(contract, dict) or int(contract.get("strict_common_item_count", 0)) != ITEM_COUNT:
        raise ValueError("QVoice normalized evidence does not contain the strict 33-item contract.")
    if contract.get("binary_items_by_wave") != {
        str(year): count for year, count in EXPECTED_BINARY_ITEM_COUNTS.items()
    }:
        raise ValueError("QVoice normalized evidence does not retain the expected binary-item contract.")
    contract_hash = contract.get("common_header_bases_sha256")
    if not isinstance(contract_hash, str) or len(contract_hash) != 64 or any(
        character not in "0123456789abcdef" for character in contract_hash
    ):
        raise ValueError("QVoice normalized evidence has an invalid strict-item digest.")
    waves = evidence.get("waves")
    if not isinstance(waves, list) or [int(wave.get("year", 0)) for wave in waves] != [2024, 2025, 2026]:
        raise ValueError("QVoice normalized evidence must contain the 2024–2026 waves in order.")
    for wave in waves:
        if not isinstance(wave, dict) or set(wave) != expected_wave_keys:
            raise ValueError("QVoice normalized evidence has an unexpected wave-level payload.")
        year = int(wave["year"])
        expected = EXPECTED_WAVES[year]
        if int(wave.get("participant_rows", 0)) != expected["participants"]:
            raise ValueError(f"QVoice normalized participant denominator changed for {year}.")
        if int(wave.get("complete_comparable_participants", 0)) != expected["complete"]:
            raise ValueError(f"QVoice normalized complete denominator changed for {year}.")
        for key in (
            "duplicate_participant_rows",
            "duplicate_conflicting_rows",
            "complete_comparable_participants_without_region",
            "out_of_domain_non_sentinel_cells",
            "zero_population_standard_deviation_cohort_count",
        ):
            if int(wave.get(key, -1)) != 0:
                raise ValueError(f"QVoice normalized evidence failed the {key} gate for {year}.")
        if int(wave.get("region_wave_cohort_count", 0)) != EXPECTED_REGION_WAVE_COHORTS:
            raise ValueError(f"QVoice normalized evidence has an unexpected region-cohort count for {year}.")
        if int(wave.get("minimum_region_wave_cohort_count", 0)) < MINIMUM_COHORT:
            raise ValueError(f"QVoice normalized evidence has a sub-threshold reference cohort for {year}.")
        if int(wave.get("normalized_participants", 0)) != expected["complete"]:
            raise ValueError(f"QVoice normalized evidence does not reconcile normalized participants for {year}.")
        if not math.isclose(float(wave.get("weighted_z_score_mean", float("nan"))), 0.0, abs_tol=1e-9):
            raise ValueError(f"QVoice normalized mean is not zero for {year}.")
        if not math.isclose(float(wave.get("weighted_population_z_score_standard_deviation", float("nan"))), 1.0, abs_tol=1e-9):
            raise ValueError(f"QVoice normalized population standard deviation is not one for {year}.")
        bands = wave.get("z_score_bands")
        if not isinstance(bands, list) or [str(row.get("band")) for row in bands] != [label for label, _, _ in Z_BANDS]:
            raise ValueError(f"QVoice normalized z-score bands are invalid for {year}.")
        if any(not isinstance(row, dict) or set(row) != expected_band_keys for row in bands):
            raise ValueError(f"QVoice normalized evidence has an unexpected display-band payload for {year}.")
        counts = [int(row.get("participant_count", 0)) for row in bands]
        if any(count < MINIMUM_COHORT for count in counts):
            raise ValueError(f"QVoice normalized evidence has a sub-threshold display band for {year}.")
        if sum(counts) != int(wave.get("normalized_participants", 0)):
            raise ValueError(f"QVoice normalized display bands do not reconcile for {year}.")
        if int(wave.get("z_score_band_count_total", -1)) != sum(counts):
            raise ValueError(f"QVoice normalized display-band total is invalid for {year}.")
        if int(wave.get("minimum_z_score_band_count", -1)) != min(counts):
            raise ValueError(f"QVoice normalized minimum display-band count is invalid for {year}.")
        for row, count in zip(bands, counts):
            expected_percent = 100.0 * count / int(wave["normalized_participants"])
            if not math.isclose(
                float(row.get("percent_of_normalized", float("nan"))),
                expected_percent,
                abs_tol=0.000003,
            ):
                raise ValueError(f"QVoice normalized display percentage is invalid for {year}.")
        if sum(Decimal(str(row.get("percent_of_normalized"))) for row in bands) != Decimal("100"):
            raise ValueError(f"QVoice normalized display percentages do not reconcile for {year}.")
    return evidence


def build_evidence(output_path: Path = OUTPUT, *, overwrite: bool = False) -> dict[str, object]:
    """Create one immutable, privacy-safe aggregate evidence artifact."""
    if output_path.exists() and not overwrite:
        raise FileExistsError(f"Frozen evidence already exists: {output_path}")

    survey_structure = load_survey_structure(QVOICE_ROOT / "survey_structure.xlsx")
    metadata: list[tuple[dict[str, object], list[str], dict[str, int], set[str]]] = []
    for config in WAVES:
        headers, by_base, bases = _headers_and_binary_bases(config, survey_structure)
        metadata.append((config, headers, by_base, bases))
    common_bases = tuple(sorted(set.intersection(*(bases for _, _, _, bases in metadata))))
    if common_bases != EXPECTED_COMMON_BASES or len(common_bases) != ITEM_COUNT:
        raise ValueError("The strict comparable 33-item intersection changed; no fuzzy mapping is allowed.")

    waves = [scan_wave(config, headers, by_base, common_bases) for config, headers, by_base, _ in metadata]
    evidence: dict[str, object] = {
        "artifact": "internal_analysis",
        "evidence_type": "aggregate-only QVoice within-region standardized-score distribution",
        "generated_on": date.today().isoformat(),
        "outcome": "pass",
        "minimum_shared_cohort": MINIMUM_COHORT,
        "source_rule": "QVoice raw Excel workbooks only; Redshift is not used because it lacks the 2024–2025 waves.",
        "scope": "Respondents with 33/33 strictly comparable valid binary items and a nonblank stored Region; this is not a workforce denominator or an attrition join.",
        "grain": "participant-wave derived aggregate",
        "formula": {
            "raw_score": "100 × favorable_items / 33",
            "reference_cohort": "stored Region × survey wave, each n ≥ 30",
            "standardization": "z = (raw_score − region-wave mean) / region-wave population standard deviation",
            "inclusion_rule": "Exactly 33 valid 0/1 responses across the strict common-item contract; 255 is missing; Response Type is not a completion filter.",
        },
        "strict_item_contract": {
            "binary_items_by_wave": {str(year): count for year, count in EXPECTED_BINARY_ITEM_COUNTS.items()},
            "strict_common_item_count": ITEM_COUNT,
            "header_matching_rule": STRICT_HEADER_RULE,
            "common_header_bases_sha256": strict_item_contract_hash(common_bases),
        },
        "waves": waves,
        "privacy_checks": {
            "identifiers_persisted": False,
            "employee_rows_persisted": False,
            "response_values_persisted": False,
            "region_labels_persisted": False,
            "all_reference_and_display_cohorts_meet_minimum": True,
        },
        "acceptance_rule": {
            "expected_years": [2024, 2025, 2026],
            "strict_common_item_count": ITEM_COUNT,
            "region_wave_cohorts_per_wave": EXPECTED_REGION_WAVE_COHORTS,
            "minimum_reference_and_display_cohort": MINIMUM_COHORT,
            "zero_duplicate_participant_rows": True,
            "zero_duplicate_conflicting_rows": True,
            "zero_complete_participants_without_region": True,
            "zero_out_of_domain_non_sentinel_cells": True,
            "zero_zero_standard_deviation_cohorts": True,
            "normalized_population_reconciles_to_33_of_33_population": True,
            "weighted_z_mean_zero_and_population_standard_deviation_one": True,
            "no_identifier_response_or_region_label_persistence": True,
        },
    }
    validate_evidence(evidence)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(json.dumps(evidence, ensure_ascii=False, indent=2), encoding="utf-8")
    return evidence


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Generate aggregate-only normalized QVoice evidence.")
    parser.add_argument("--output", type=Path, default=OUTPUT)
    parser.add_argument("--overwrite", action="store_true")
    args = parser.parse_args()
    artifact = build_evidence(args.output, overwrite=args.overwrite)
    print(
        json.dumps(
            {
                "output": str(args.output),
                "waves": [
                    {
                        "year": wave["year"],
                        "normalized_participants": wave["normalized_participants"],
                        "minimum_z_score_band_count": wave["minimum_z_score_band_count"],
                    }
                    for wave in artifact["waves"]
                ],
                "outcome": artifact["outcome"],
            },
            ensure_ascii=False,
        )
    )

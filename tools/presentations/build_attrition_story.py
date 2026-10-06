from __future__ import annotations

import argparse
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_VERTICAL_ANCHOR, PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Pt


COLORS = {
    "navy": "020B3F",
    "blue": "2A2AEA",
    "purple": "843AFF",
    "yellow": "F5F032",
    "line": "D2D7E1",
    "light": "EFF1F5",
    "white": "FFFFFF",
    "text": "252A35",
    "muted": "666D7A",
}
LEGAL = (
    "Confidential - Qualcomm Technologies, Inc. and/or its affiliated companies - "
    "May Contain Trade Secrets"
)


def rgb(name: str) -> RGBColor:
    value = COLORS[name]
    return RGBColor(
        int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16)
    )


def pt(value: float):
    return Pt(value)


def clear_slide(slide) -> None:
    for shape in list(slide.shapes):
        shape._element.getparent().remove(shape._element)


def add_rect(slide, x, y, w, h, fill, line=None):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, pt(x), pt(y), pt(w), pt(h)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    if line:
        shape.line.color.rgb = rgb(line)
        shape.line.width = pt(0.75)
    else:
        shape.line.fill.background()
    return shape


def add_text(
    slide,
    value,
    x,
    y,
    w,
    h,
    size,
    color="text",
    bold=False,
    font="Aptos",
    align=None,
):
    shape = slide.shapes.add_textbox(pt(x), pt(y), pt(w), pt(h))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = 0
    frame.margin_right = 0
    frame.margin_top = 0
    frame.margin_bottom = 0
    frame.vertical_anchor = MSO_VERTICAL_ANCHOR.TOP

    paragraph = frame.paragraphs[0]
    paragraph.space_before = pt(0)
    paragraph.space_after = pt(0)
    paragraph.line_spacing = 1.0
    if align is not None:
        paragraph.alignment = align

    run = paragraph.add_run()
    run.text = value
    run.font.name = font
    run.font.size = pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    return shape


def add_footer(slide, number: int) -> None:
    add_text(slide, LEGAL, 34, 516, 760, 12, 6.5, "muted")
    add_text(
        slide,
        str(number),
        900,
        516,
        26,
        12,
        7,
        "muted",
        align=PP_ALIGN.RIGHT,
    )


def add_base(slide, eyebrow, title, subtitle, number: int) -> None:
    add_rect(slide, 0, 0, 960, 540, "white")
    add_rect(slide, 0, 0, 960, 8, "blue")
    add_text(slide, eyebrow.upper(), 54, 34, 820, 18, 10, "blue", True)
    add_text(
        slide,
        title,
        54,
        58,
        850,
        48,
        31,
        "navy",
        True,
        "Aptos Display",
    )
    add_text(slide, subtitle, 55, 111, 840, 34, 16, "muted")
    add_footer(slide, number)


def add_divider(slide) -> None:
    add_rect(slide, 479, 158, 1, 278, "line")


def add_definition_item(slide, y, title, body, accent, height=68) -> None:
    add_rect(slide, 55, y, 6, height, accent)
    add_text(slide, title, 75, y + 4, 360, 22, 16, "navy", True)
    add_text(slide, body, 75, y + 29, 360, height - 31, 16, "text")


def add_table(
    slide,
    x,
    y,
    headers,
    rows,
    widths,
    row_height=55,
    right_align_last=False,
) -> None:
    header_height = 32
    for column, header in enumerate(headers):
        left = x + sum(widths[:column])
        add_rect(slide, left, y, widths[column], header_height, "blue", "white")
        add_text(
            slide,
            header,
            left + 7,
            y + 8,
            widths[column] - 14,
            18,
            14,
            "white",
            True,
        )

    for row_index, row in enumerate(rows):
        top = y + header_height + row_index * row_height
        fill = "white" if row_index % 2 == 0 else "light"
        for column, value in enumerate(row):
            left = x + sum(widths[:column])
            add_rect(
                slide,
                left,
                top,
                widths[column],
                row_height,
                fill,
                "line",
            )
            alignment = (
                PP_ALIGN.RIGHT
                if right_align_last and column == len(row) - 1
                else None
            )
            add_text(
                slide,
                str(value),
                left + 7,
                top + 10,
                widths[column] - 14,
                row_height - 15,
                14,
                "text",
                align=alignment,
            )


def build_definition_population(slide) -> None:
    clear_slide(slide)
    add_base(
        slide,
        "Definition + population",
        "Attrition definition and study population",
        "Align the outcome and population before model development",
        3,
    )
    add_divider(slide)

    add_text(
        slide,
        "Attrition definition",
        55,
        158,
        390,
        28,
        18,
        "navy",
        True,
        "Aptos Display",
    )
    add_definition_item(
        slide,
        198,
        "In scope",
        "Voluntary termination for regular employees.",
        "blue",
    )
    add_definition_item(
        slide,
        278,
        "Tracked separately",
        "Involuntary exits, internal moves, and contract transitions.",
        "purple",
        82,
    )
    add_definition_item(
        slide,
        372,
        "Why it matters",
        "Cleaner target, better governance, clearer action ownership.",
        "yellow",
        64,
    )

    add_text(
        slide,
        "Study population",
        505,
        158,
        400,
        28,
        18,
        "navy",
        True,
        "Aptos Display",
    )
    add_table(
        slide,
        505,
        198,
        ["Population", "MVP", "Reason"],
        [
            ["Regular employees", "Include", "Stable policy baseline"],
            ["Temporary workers", "Exclude", "Different contract dynamics"],
            ["Interns / contractors", "Exclude", "Different lifecycle and actions"],
        ],
        [145, 76, 184],
        70,
    )
    slide.notes_slide.notes_text_frame.text = (
        "Combined from the original definition and population slides. "
        "Use voluntary exits as the MVP outcome; keep involuntary exits and mobility "
        "separate. Regular employees give the cleanest policy and lifecycle baseline "
        "for launch."
    )


def add_flow_item(slide, y, label, accent) -> None:
    add_rect(slide, 55, y, 390, 35, "light", "line")
    add_rect(slide, 55, y, 6, 35, accent)
    add_text(slide, label, 75, y + 8, 350, 20, 16, "navy", True)


def build_analysis_governance(slide) -> None:
    clear_slide(slide)
    add_base(
        slide,
        "Analysis + governance",
        "Analytical cuts and governed outputs",
        "Prioritize cohorts while keeping outputs actionable and privacy-safe",
        4,
    )
    add_divider(slide)

    add_text(
        slide,
        "Analytical cuts",
        55,
        158,
        390,
        28,
        18,
        "navy",
        True,
        "Aptos Display",
    )
    flow = [
        ("Cohort -> evidence", "blue"),
        ("Evidence -> hypothesis", "purple"),
        ("Hypothesis -> action", "yellow"),
        ("Action -> owner -> metric", "blue"),
    ]
    for index, (label, accent) in enumerate(flow):
        add_flow_item(slide, 198 + index * 43, label, accent)

    add_text(slide, "Potential segments", 55, 378, 390, 22, 16, "navy", True)
    add_text(
        slide,
        "Organization, location, tenure band, job family, manager span, lifecycle indicators.",
        55,
        404,
        390,
        40,
        16,
        "text",
    )

    add_text(
        slide,
        "Governed outputs",
        505,
        158,
        400,
        28,
        18,
        "navy",
        True,
        "Aptos Display",
    )
    add_table(
        slide,
        505,
        198,
        ["Audience", "Permitted", "Not permitted"],
        [
            ["People Analytics", "Employee-level workspace", "Broad sharing"],
            ["HR / leaders", "Aggregated cohort risk", "Individual scores"],
            ["Managers", "Approved action themes", "Ranked employee lists"],
        ],
        [105, 157, 143],
        70,
        right_align_last=True,
    )
    slide.notes_slide.notes_text_frame.text = (
        "Combined from the original analytical cuts and governance slides. Cut evidence "
        "by organization, location, tenure, job family, manager span, and lifecycle "
        "stage. Association is not causation; cuts are used to form testable hypotheses. "
        "Leaders see cohorts and drivers, not employee scores; actions need owners and "
        "metrics."
    )


def delete_slide(prs, index: int) -> None:
    slide_id = prs.slides._sldIdLst[index]
    prs.part.drop_rel(slide_id.rId)
    del prs.slides._sldIdLst[index]


def renumber_footers(prs) -> None:
    for slide_number, slide in enumerate(prs.slides, start=1):
        candidates = []
        for shape in slide.shapes:
            if not getattr(shape, "has_text_frame", False):
                continue
            if shape.top < pt(500) or shape.left < pt(850):
                continue
            candidates.append(shape)
        if len(candidates) != 1:
            raise RuntimeError(
                f"Expected one page-number shape on slide {slide_number}; "
                f"found {len(candidates)}"
            )
        candidates[0].text_frame.text = str(slide_number)
        paragraph = candidates[0].text_frame.paragraphs[0]
        paragraph.alignment = PP_ALIGN.RIGHT
        for run in paragraph.runs:
            run.font.name = "Aptos"
            run.font.size = pt(7)
            run.font.color.rgb = rgb("muted")


def build(source: Path, output: Path) -> None:
    if source.resolve() == output.resolve():
        raise ValueError("Source and output paths must be different")
    if not source.exists():
        raise FileNotFoundError(source)

    prs = Presentation(str(source))
    if len(prs.slides) != 9:
        raise ValueError(f"Expected 9 source slides; found {len(prs.slides)}")

    build_definition_population(prs.slides[2])
    build_analysis_governance(prs.slides[4])

    # Delete original slides 6 and 4 in descending order, preserving references.
    delete_slide(prs, 5)
    delete_slide(prs, 3)
    renumber_footers(prs)

    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output))
    print(output)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    build(args.source, args.output)


if __name__ == "__main__":
    main()

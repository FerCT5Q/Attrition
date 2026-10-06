from __future__ import annotations

import argparse
import sys
from pathlib import Path

LEGAL = "Confidential - Qualcomm Technologies, Inc. and/or its affiliated companies - May Contain Trade Secrets"
C_HEX = {"navy":"020B3F","blue":"2A2AEA","purple":"843AFF","yellow":"F5F032","line":"D2D7E1","light":"EFF1F5","white":"FFFFFF","text":"252A35","muted":"666D7A"}

def add_deps(path: str | None) -> None:
    if not path:
        return
    candidate = Path(path)
    if not candidate.exists():
        raise FileNotFoundError(f"Dependency directory does not exist: {candidate}")
    sys.path.insert(0, str(candidate))

def rgb(name: str):
    from pptx.dml.color import RGBColor
    h = C_HEX[name]
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))

def pt(value: float):
    from pptx.util import Pt
    return Pt(value)

def clear(slide) -> None:
    for shape in list(slide.shapes):
        shape._element.getparent().remove(shape._element)

def rect(slide, x, y, w, h, fill, line=None):
    from pptx.enum.shapes import MSO_SHAPE
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, pt(x), pt(y), pt(w), pt(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    if line:
        shape.line.color.rgb = rgb(line)
    else:
        shape.line.fill.background()
    return shape

def text(slide, body, x, y, w, h, size, color="text", bold=False, font="Aptos", align=None):
    from pptx.enum.text import MSO_VERTICAL_ANCHOR
    shape = slide.shapes.add_textbox(pt(x), pt(y), pt(w), pt(h))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    frame.vertical_anchor = MSO_VERTICAL_ANCHOR.TOP
    paragraph = frame.paragraphs[0]
    paragraph.space_before = pt(0)
    paragraph.space_after = pt(0)
    paragraph.line_spacing = 1.0
    if align:
        paragraph.alignment = align
    run = paragraph.add_run()
    run.text = body
    run.font.name = font
    run.font.size = pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    return shape

def footer(slide, number: int) -> None:
    from pptx.enum.text import PP_ALIGN
    text(slide, LEGAL, 34, 516, 760, 12, 6.5, "muted")
    text(slide, str(number), 900, 516, 26, 12, 7, "muted", align=PP_ALIGN.RIGHT)

def base(slide) -> None:
    clear(slide)
    rect(slide, 0, 0, 960, 540, "white")
    rect(slide, 0, 0, 960, 8, "blue")
    text(slide, "PLAN", 54, 34, 780, 18, 10, "blue", True)
    text(slide, "10-week workplan", 54, 58, 850, 48, 31, "navy", True, "Aptos Display")
    footer(slide, 8)

def add_phase(slide, index, week, title, x, accent):
    from pptx.enum.text import PP_ALIGN
    rect(slide, x, 186, 154, 118, "light", "line")
    rect(slide, x, 186, 154, 6, accent)
    text(slide, str(index), x + 14, 210, 18, 18, 14, accent, True, align=PP_ALIGN.CENTER)
    text(slide, week, x + 42, 210, 78, 18, 14, "muted", True)
    text(slide, title, x + 14, 244, 126, 34, 16, "navy", True, "Aptos Display")

def add_after_mvp(slide, x, title, timing, accent):
    rect(slide, x, 382, 190, 62, "white", "line")
    rect(slide, x, 382, 5, 62, accent)
    text(slide, title, x + 14, 397, 160, 16, 14, "navy", True)
    text(slide, timing, x + 14, 421, 160, 14, 14, "purple", True)

def refine(path: Path, output: Path | None) -> None:
    from pptx import Presentation
    prs = Presentation(path)
    slide = prs.slides[7]
    base(slide)
    phases = [
        (1, "W1", "Alignment", 70, "blue"),
        (2, "W2-3", "Dataset + label QA", 242, "purple"),
        (3, "W4-5", "Baselines + driver scan", 414, "yellow"),
        (4, "W6-9", "Model + validation", 586, "blue"),
        (5, "W10", "Insights + HR review", 758, "purple"),
    ]
    for phase in phases:
        add_phase(slide, *phase)
    text(slide, "AFTER THE MVP", 70, 354, 180, 18, 14, "blue", True)
    after = [
        (70, "Pilot design", "1-2 weeks", "blue"),
        (282, "Action cycle", "6-8 weeks", "purple"),
        (494, "Early readout", "During cycle", "yellow"),
        (706, "Retention impact", "6-month window", "blue"),
    ]
    for item in after:
        add_after_mvp(slide, *item)
    try:
        slide.notes_slide.notes_text_frame.text = (
            "Aligned to the source workplan diagram. W1: measure, population, objective, safeguards, decision rights. "
            "W2-3: monthly snapshots, exclusions, leakage checks, baseline. W4-5: cohort baselines, risk segments, driver scan. "
            "W6-9: interpretable model, challenger, calibration, stability, fairness. W10: priority cohorts, drivers, recommendations, HR review. "
            "After MVP: pilot design, action cycle, early readout, and 6-month retention impact."
        )
    except Exception:
        pass
    prs.save(output or path)

def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("pptx")
    parser.add_argument("--output")
    parser.add_argument("--deps")
    args = parser.parse_args()
    add_deps(args.deps)
    refine(Path(args.pptx), Path(args.output) if args.output else None)

if __name__ == "__main__":
    main()


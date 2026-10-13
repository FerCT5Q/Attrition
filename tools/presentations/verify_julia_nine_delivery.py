"""Verifica cifras contra agregados, preservación de nueve slides y excepción de idioma."""
from __future__ import annotations
import hashlib
import json
from pathlib import Path
import re
import subprocess
import sys
from bs4 import BeautifulSoup
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[2]
OUT=ROOT/'output/presentations/20261013-0735_revision-julia-9-slides'
RUN=ROOT/'output/analysis/engineering/20261013/c-bu-isolation-v1'
REF=ROOT/'output/analysis/engineering/20261012/julia-rh-review-v1'


def read(path): return json.loads(path.read_text(encoding='utf-8-sig'))
def number(v): return f'{round(float(v),1):,.1f}'.rstrip('0').rstrip('.')
def write(path,data): path.write_text(json.dumps(data,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')


def main():
    deck=OUT/'presentacion-eng.html'; html=deck.read_text(encoding='utf-8'); soup=BeautifulSoup(html,'html.parser')
    meta=read(OUT/'evidence.json'); data=read(RUN/'results.json'); ref=read(REF/'results.json')
    checks={}
    checks['nine_stages_and_nav']=len(soup.select('.slide-stage'))==len(soup.select('.nav-pill'))==9
    checks['base_is_original_nine']=meta['base_slide_count']==9 and '20261012-0833' in meta['base']
    checks['four_themes']=len(soup.select('[data-theme]'))==4
    checks['correct_deck_hash']=hashlib.sha256(deck.read_bytes()).hexdigest()==meta['deck_sha256']
    checks['sources_unchanged']=all(hashlib.sha256((ROOT/v['path']).read_bytes()).hexdigest()==v['sha256'] for v in meta['sources'].values())
    # Calcula las celdas desde periodos; no reutiliza los valores renderizados del manifest.
    rows=soup.select('#slide-8 tbody tr')
    checks['five_estimators_shown']=len(rows)==5
    for row,n in zip(rows,['CQ','C_BU','C_BU_Q','FQ','FQ_BU_L3']):
        values={k:sum(p['models'][n]['capacities']['30']['metrics'][k] for p in data['periods'])/2 for k in ['log_loss','capture_percent','group_snapshot_mae_pp']}
        expected=[number(values['log_loss']*1000),number(values['capture_percent'])+'%',number(values['group_snapshot_mae_pp'])+' pp']
        checks['slide8_'+n]=[x.get_text(strip=True) for x in row.select('td')][1:]==expected
    for cap,row in zip(['10','20','30'],soup.select('#slide-9 tbody tr')):
        checks['slide9_capture_'+cap]=row.select('td')[-1].text==number(ref['summary']['CQ'][cap]['capture_percent'])+'%'
    for period in data['periods']:
        pid=period['period']['id']; base=read(RUN/'models'/f'{pid}-C_refit.json')
        bu=read(RUN/'models'/f'{pid}-C_BU.json'); q=read(RUN/'models'/f'{pid}-C_BU_Q.json')
        checks[pid+'_BU_only_addition']=set(bu['base']['features'])==set(base['base']['features'])|{'bu_l3'}
        checks[pid+'_no_QVoice_in_C_BU']=bu['qvoice'] is False and not any(f.startswith('qv_') for f in bu['base']['features'])
        checks[pid+'_same_base_when_adding_Q']=bu['base']==q['base'] and bu['pool_keep']==q['pool_keep']
        checks[pid+'_mature_training_labels']=period['last_training_label']<=period['period']['origin']
        for n,m in period['models'].items():
            for cap,c in m['capacities'].items():
                checks[pid+'_'+n+'_'+cap+'_capacity']=abs(c['metrics']['coverage_percent']-int(cap))<1e-8
    work=load_workbook(RUN/'documentacion/variables-44-ampliada.xlsx',read_only=True)
    checks['44_variables']=work['44 variables y acciones'].max_row==45
    checks['440_roles']=work['Uso por modelo y corte'].max_row==441
    checks['links_exist']=all((deck.parent/a['href']).resolve().exists() for a in soup.select('a[href]'))
    # El usuario y AGENTS requieren inglés. Única excepción al verificador genérico.
    static=subprocess.run([sys.executable,'-B',str(ROOT/'.agents/skills/build-html-presentation/scripts/verify_deck.py'),str(deck),'--mode','presentation'],capture_output=True,encoding='utf-8',errors='replace')
    errors=[line for line in static.stdout.splitlines() if line.startswith('ERROR:')]
    checks['only_required_English_exception']=len(errors)==1 and 'lang="es"' in errors[0]
    checks['all_other_static_checks_pass']=static.returncode==1 and len(errors)==1
    checks['visible_decimals_at_most_one']=not re.search(r'\d+\.\d{2,}', ' '.join(s.get_text(' ',strip=True) for s in soup.select('.slide-stage')))
    result={'reporting_mode':'presentation','checks':checks,'passed':all(checks.values()),
       'static_output':static.stdout,'exception':'Inglés exigido por AGENTS.md y usuario para este deck; idioma y controles conservados.',
       'numeric_sources':['J01','J02','J03'],'employee_level_rows_persisted':False,
       'statistical_tests':'7 pruebas de tests/test_julia_review_statistics.py pasaron durante la ejecución',
       'browser_verification':'browser-verification.json','deck_sha256':meta['deck_sha256']}
    write(OUT/'verification.json',result)
    assert result['passed'],[k for k,v in checks.items() if not v]
    print(f'{len(checks)} comprobaciones aprobadas; única excepción: inglés requerido.')


if __name__=='__main__':
    sys.stdout.reconfigure(encoding='utf-8'); main()

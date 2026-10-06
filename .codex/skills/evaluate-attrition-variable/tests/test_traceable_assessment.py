from __future__ import annotations

import importlib.util
import json
import sys
import tempfile
import unittest
from datetime import date
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook
import yaml


SKILL_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(SKILL_ROOT / "scripts"))
WORKSPACE_TEMP = SKILL_ROOT.parents[2] / "output"


def load_module(name: str, path: Path):
    module_spec = importlib.util.spec_from_file_location(name, path)
    assert module_spec and module_spec.loader
    module = importlib.util.module_from_spec(module_spec)
    module_spec.loader.exec_module(module)
    return module


profiler = load_module("profile_variable", SKILL_ROOT / "scripts" / "profile_variable.py")
renderer = load_module(
    "render_assessment", SKILL_ROOT / "scripts" / "render_assessment.py"
)


runner = load_module(
    "run_assessment", SKILL_ROOT / "scripts" / "run_assessment.py"
)

visual_renderer = load_module(
    "render_variable_assessment_v2",
    SKILL_ROOT.parents[2] / "tools" / "presentations" / "render_variable_assessment_v2.py",
)

def base_variable(name: str = "Test variable") -> dict:
    return {
        "category": "Tests",
        "variable": name,
        "definition": "A test snapshot variable.",
        "source": {
            "type": "Database",
            "table": "hrda.post_hire_monthly_worker_dtls",
            "columns": ["job_family"],
            "status": "confirmed_accessible",
        },
        "level_data": "Employee snapshot",
        "measurement": {
            "type": "snapshot",
            "date_column": "fis_mo_end_dt",
            "fiscal_year_column": "fiscal_year",
            "employee_id_column": "employee_id",
            "observable_columns": ["job_family"],
        },
        "population": {"mode": "default"},
        "numerator_definition": "Employees with a populated test field.",
        "evidence": {
            "definition": "clear_stable",
            "consistency": "clean_standardized",
            "join": "reliable_employee_date",
            "effort": "direct_or_simple_sql",
        },
        "verification": {
            "semantic_type": "direct_field",
            "source": {
                "method": "schema_access_check",
                "evidence": "information_schema and a read-only aggregate query were checked.",
                "result": "The table and required columns are accessible.",
                "as_of_date": "2026-08-14",
                "query_reference": "queries/audits/current_input_quality_audit.sql",
                "acceptance_rule": "The source table and every required column must be accessible.",
                "outcome": "pass",
            },
            "definition": {
                "method": "authoritative_definition_and_data_audit",
                "evidence": "The source-owned field definition was paired with an aggregate domain audit.",
                "result": "Observed values conform to the documented snapshot meaning.",
                "as_of_date": "2026-08-14",
                "query_reference": "queries/audits/current_input_quality_audit.sql",
                "acceptance_rule": "The documented grain and domain must match the aggregate audit.",
                "minimum_match_rate_percent": 99.5,
                "outcome": "pass",
                "metrics": {
                    "eligible_rows": 100,
                    "populated_rows": 100,
                    "distinct_values": 1,
                },
            },
            "consistency": {
                "method": "longitudinal_aggregate_audit",
                "evidence": "Monthly normalized cardinality and transition counts were audited.",
                "result": "No unexplained normalization or temporal anomalies were found.",
                "as_of_date": "2026-08-14",
                "query_reference": "queries/audits/current_input_quality_audit.sql",
                "acceptance_rule": "No unexplained invalid values or longitudinal anomalies are allowed.",
                "outcome": "pass",
                "metrics": {"months_checked": 3, "anomaly_count": 0},
            },
            "join": {
                "method": "aggregate_join_reconciliation",
                "evidence": "Employee and snapshot keys were reconciled at aggregate level.",
                "result": "All eligible test records joined once at the intended grain.",
                "as_of_date": "2026-08-14",
                "query_reference": "queries/audits/current_input_quality_audit.sql",
                "acceptance_rule": "Every eligible row must match once at the employee-date grain.",
                "outcome": "pass",
                "metrics": {
                    "eligible_rows": 100,
                    "matched_rows": 100,
                    "unmatched_rows": 0,
                    "duplicate_rows": 0,
                },
            },
            "effort": {
                "method": "direct_field_inspection",
                "evidence": "The requested feature is the audited source field without transformation.",
                "result": "No derivation beyond projection and approved filtering is required.",
                "as_of_date": "2026-08-14",
                "query_reference": "queries/audits/current_input_quality_audit.sql",
                "acceptance_rule": "The feature must require only projection and approved filters.",
                "outcome": "pass",
            },
        },
        "blocker": "none",
    }


def raw_rows(dates: list[str], observable: int = 100, eligible: int = 100) -> list[dict]:
    return [
        {
            "snapshot_date": value,
            "fiscal_year": value[:4],
            "eligible_employees": eligible,
            "observable_employees": observable,
            "event_employees": None,
        }
        for value in dates
    ]


def audit_row(**overrides) -> dict:
    row = {
        "missing_items": 0,
        "eligible_rows": 100,
        "eligible_count": 100,
        "populated_rows": 100,
        "distinct_values": 2,
        "months_checked": 3,
        "anomaly_count": 0,
        "matched_rows": 100,
        "unmatched_rows": 0,
        "duplicate_rows": 0,
        "comparable_count": 100,
        "matched_within_tolerance_count": 100,
        "mismatch_count": 0,
        "match_rate_percent": 100.0,
        "reconstructed_count": 100,
        "unresolved_count": 0,
        "root_count": 1,
        "cycle_count": 0,
        "convergence_rate_percent": 100.0,
        "root_semantics_pass": 1,
    }
    row.update(overrides)
    return row


def derivation_option(status: str = "proposed") -> dict:
    return {
        "name": "Trailing event derivation",
        "status": status,
        "fields_and_tables": "hrda.post_hire_monthly_worker_dtls transfer fields",
        "formula": "Count dated transfer events in the trailing 12 months.",
        "grain": "employee-month",
        "temporal_window": "Trailing 12 months ending at the snapshot date.",
        "deduplication": "One event per employee and effective date.",
        "nulls_and_edge_cases": "Null event dates are non-events when the flag is observable.",
        "leakage_controls": "Use only dates on or before the snapshot date.",
        "snapshot_join": "Join by employee_id and fis_mo_end_dt snapshot month.",
        "sensitivity": "No employee-level outputs; aggregate validation only.",
        "validations_needed": "Boundary windows, deduplication, leakage, and aggregate reconciliation.",
        "estimated_effort": "Moderate SQL transformation.",
        "decision_reason": "Components exist, but the full derivation has not been reconciled.",
    }


class TraceableScoringTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls.rules = profiler.load_yaml(SKILL_ROOT / "references" / "scoring-rules.yaml")

    def profile(
        self, variable: dict, rows: list[dict], audit: dict | None = None
    ) -> dict:
        offline = {
            "variables": {
                variable["variable"]: {
                    "coverage": rows,
                    "queries": {
                        "queries/audits/current_input_quality_audit.sql": [
                            audit_row() if audit is None else audit
                        ]
                    },
                }
            }
        }
        result = profiler.profile_one(
            variable,
            self.rules,
            Path("assessment-spec.yaml"),
            offline_data=offline,
        )
        assert result is not None
        return result

    def test_coverage_boundaries(self) -> None:
        self.assertEqual(profiler.coverage_score(69.999, self.rules), 0)
        self.assertEqual(profiler.coverage_score(70.0, self.rules), 1)
        self.assertEqual(profiler.coverage_score(94.999, self.rules), 1)
        self.assertEqual(profiler.coverage_score(95.0, self.rules), 2)

    def test_missingness_boundaries(self) -> None:
        dates = ["2026-01-01", "2026-02-01", "2026-03-01"]
        at_five = self.profile(base_variable("Five percent"), raw_rows(dates, 95))
        at_thirty = self.profile(base_variable("Thirty percent"), raw_rows(dates, 70))
        above_thirty = self.profile(base_variable("Above thirty"), raw_rows(dates, 69))
        self.assertEqual(at_five["scores"]["missingness"], 2)
        self.assertEqual(at_thirty["scores"]["missingness"], 1)
        self.assertEqual(above_thirty["scores"]["missingness"], 0)

    def test_history_boundaries_use_completed_calendar_months(self) -> None:
        self.assertEqual(
            profiler.elapsed_months(date(2024, 1, 1), date(2024, 7, 1)), 6
        )
        self.assertEqual(
            profiler.elapsed_months(date(2024, 1, 1), date(2026, 1, 1)), 24
        )
        six = self.profile(
            base_variable("Six months"),
            raw_rows(["2024-01-01", "2024-02-01", "2024-03-01", "2024-07-01"]),
        )
        twenty_four = self.profile(
            base_variable("Twenty four months"),
            raw_rows(["2024-01-01", "2024-02-01", "2024-03-01", "2026-01-01"]),
        )
        self.assertEqual(six["scores"]["history"], 1)
        self.assertEqual(twenty_four["scores"]["history"], 2)

    def test_systematic_missingness_caps_score(self) -> None:
        rows = []
        for index, observable in enumerate([100, 100, 100, 100, 100, 79], start=1):
            rows.extend(raw_rows([f"2026-{index:02d}-01"], observable))
        result = self.profile(base_variable("Systematic"), rows)
        self.assertTrue(result["diagnostics"]["systematic_missingness"])
        self.assertEqual(result["scores"]["missingness"], 1)

    def test_leakage_forces_history_zero(self) -> None:
        variable = base_variable("Leakage")
        variable["quality"] = {"leakage_risk": True}
        result = self.profile(
            variable,
            raw_rows(["2023-01-01", "2023-02-01", "2023-03-01", "2026-01-01"]),
        )
        self.assertEqual(result["scores"]["history"], 0)
        rationale = {
            item["criterion_key"]: item for item in result["score_rationales"]
        }
        self.assertIn("leakage del outcome", rationale["history"]["why_below_2"])

    def test_configured_override_is_explicit(self) -> None:
        variable = base_variable("Override")
        variable["score_overrides"] = {
            "coverage": {
                "score": 1,
                "reason": "The latest feed is under acceptance review.",
                "evidence": "Aggregate counts reconcile, but business sign-off is pending.",
                "recommended_action": "Complete source-owner sign-off and rerun profiling.",
            }
        }
        result = self.profile(
            variable,
            raw_rows(["2023-01-01", "2023-02-01", "2023-03-01", "2026-01-01"]),
        )
        rationale = {
            item["criterion_key"]: item for item in result["score_rationales"]
        }
        self.assertEqual(rationale["coverage"]["scoring_method"], "Override configurado")
        self.assertIn("Override configurado", rationale["coverage"]["status_rule"])

    def test_missing_qualitative_rationale_stops_profiling(self) -> None:
        variable = base_variable("Missing rationale")
        variable["evidence"]["definition"] = "needs_alignment"
        with self.assertRaisesRegex(ValueError, "requires structured rationale"):
            self.profile(
                variable,
                raw_rows(["2023-01-01", "2023-02-01", "2023-03-01"]),
            )

    def test_score_two_requires_structured_verification(self) -> None:
        variable = base_variable("Missing score two evidence")
        del variable["verification"]["definition"]
        with self.assertRaisesRegex(
            ValueError, "score 2 for qualitative criterion definition"
        ):
            self.profile(
                variable,
                raw_rows(["2026-01-01", "2026-02-01", "2026-03-01"]),
            )

    def test_precomputed_metric_requires_independent_recalculation(self) -> None:
        variable = base_variable("Precomputed metric")
        variable["verification"]["semantic_type"] = "precomputed_metric"
        with self.assertRaisesRegex(
            ValueError, "requires independent recalculation"
        ):
            self.profile(
                variable,
                raw_rows(["2026-01-01", "2026-02-01", "2026-03-01"]),
            )

        variable["verification"]["definition"]["method"] = (
            "independent_recalculation"
        )
        variable["verification"]["definition"]["metrics"] = {
            "comparable_count": 100,
            "matched_within_tolerance_count": 100,
            "mismatch_count": 0,
            "match_rate_percent": 100.0,
        }
        result = self.profile(
            variable,
            raw_rows(["2026-01-01", "2026-02-01", "2026-03-01"]),
        )
        self.assertEqual(result["scores"]["definition"], 2)

    def test_score_two_metrics_are_computed_from_frozen_rows(self) -> None:
        variable = base_variable("Computed verification metrics")
        del variable["verification"]["consistency"]["metrics"]
        result = self.profile(
            variable,
            raw_rows(["2026-01-01", "2026-02-01", "2026-03-01"]),
        )
        verification = result["score_verification"]["consistency"]
        self.assertEqual(verification["metrics"]["anomaly_count"], 0)
        self.assertTrue(verification["provenance"]["executed"])

    def test_declared_outcome_cannot_enable_or_disable_score_two(self) -> None:
        variable = base_variable("Declared outcome ignored")
        variable["verification"]["definition"]["outcome"] = "fail"
        passed = self.profile(variable, raw_rows(["2026-01-01"]))
        self.assertEqual(passed["scores"]["definition"], 2)

        failed = self.profile(
            base_variable("Calculated failure"),
            raw_rows(["2026-01-01"]),
            audit=audit_row(populated_rows=0, distinct_values=0),
        )
        self.assertEqual(failed["scores"]["definition"], 0)

    def test_score_two_requires_persisted_query(self) -> None:
        variable = base_variable("Unpersisted validation")
        variable["verification"]["definition"]["query_reference"] = "ad hoc query"
        with self.assertRaisesRegex(ValueError, "existing project-relative SQL file"):
            self.profile(variable, raw_rows(["2026-01-01"]))

    def test_hierarchy_requires_root_cycle_and_convergence_metrics(self) -> None:
        variable = base_variable("Hierarchy")
        variable["verification"]["semantic_type"] = "hierarchy"
        variable["verification"]["definition"]["method"] = "independent_recalculation"
        incomplete = audit_row()
        del incomplete["root_semantics_pass"]
        pending = self.profile(
            variable, raw_rows(["2026-01-01"]), audit=incomplete
        )
        self.assertEqual(pending["scores"]["definition"], 1)
        result = self.profile(
            variable, raw_rows(["2026-01-01"]), audit=audit_row()
        )
        self.assertEqual(result["scores"]["definition"], 2)

    def test_override_cannot_increase_score(self) -> None:
        variable = base_variable("No score inflation")
        variable["quality"] = {"leakage_risk": True}
        variable["score_overrides"] = {
            "history": {
                "score": 2,
                "reason": "Manual attempt to restore history score.",
            }
        }
        with self.assertRaisesRegex(ValueError, "overrides may only reduce"):
            self.profile(
                variable,
                raw_rows(["2023-01-01", "2023-02-01", "2023-03-01", "2026-01-01"]),
            )

    def test_no_direct_field_requires_derivation_or_reason(self) -> None:
        variable = base_variable("Derived without plan")
        variable["measurement"]["type"] = "rolling_derived"
        with self.assertRaisesRegex(ValueError, "provide derivation_options"):
            self.profile(variable, raw_rows(["2026-01-01", "2026-02-01", "2026-03-01"]))

    def test_proposed_derivation_caps_definition_score(self) -> None:
        variable = base_variable("Proposed derivation")
        variable["measurement"]["type"] = "rolling_derived"
        variable["derivation_options"] = [derivation_option("proposed")]
        result = self.profile(
            variable,
            raw_rows(["2024-01-01", "2024-02-01", "2024-03-01", "2026-03-01"]),
        )
        self.assertEqual(result["scores"]["definition"], 1)
        rationale = {item["criterion_key"]: item for item in result["score_rationales"]}
        self.assertEqual(rationale["definition"]["scoring_method"], "Puerta de validación de derivación")
        self.assertNotIn("definition", result["score_verification"])

    def test_executable_validation_caps_score(self) -> None:
        variable = base_variable("Validation cap")
        variable["scoring_validations"] = [
            {
                "id": "definition_match_rate",
                "criterion": "definition",
                "sql_reference": "queries/audits/current_input_quality_audit.sql",
                "acceptance_rule": "Match rate must be at least 99.5%.",
                "metric": "match_rate_percent",
                "operator": "gte",
                "threshold": 99.5,
                "observed": 98.0,
                "reason": "Independent reconciliation is below the acceptance threshold.",
                "evidence": "Observed aggregate match rate was 98.0%, below 99.5%.",
                "recommended_action": "Resolve mismatches and rerun the saved audit.",
            }
        ]
        result = self.profile(
            variable,
            raw_rows(["2024-01-01", "2024-02-01", "2024-03-01", "2026-03-01"]),
            audit=audit_row(match_rate_percent=98.0, mismatch_count=2),
        )
        self.assertEqual(result["scores"]["definition"], 0)
        rationale = {item["criterion_key"]: item for item in result["score_rationales"]}
        self.assertEqual(rationale["definition"]["scoring_method"], "Validación ejecutable")

    def test_legacy_spec_is_rejected(self) -> None:
        with self.assertRaisesRegex(ValueError, "legacy"):
            profiler.validate_spec(
                {"version": 3, "variables": [base_variable()]},
                self.rules,
            )
        variables = profiler.validate_spec(
            {"version": 4, "variables": [base_variable()]},
            self.rules,
        )
        self.assertEqual(len(variables), 1)

    def test_workbook_has_one_comment_and_row_per_low_score(self) -> None:
        variable = base_variable("Workbook traceability")
        variable["evidence"]["definition"] = "needs_alignment"
        variable["rationale"] = {
            "definition": {
                "why": "An edge-case business rule is not yet approved.",
                "evidence": "The source has data, but no approved edge-case mapping.",
                "recommended_action": "Approve and document the edge-case mapping.",
            }
        }
        variable["blocker"] = "sensitive_restricted"
        assessment = self.profile(
            variable,
            raw_rows(["2023-01-01", "2023-02-01", "2023-03-01", "2026-01-01"]),
        )
        profile = {"schema_version": 4, "assessments": [assessment]}
        with tempfile.TemporaryDirectory(dir=WORKSPACE_TEMP) as directory:
            template = Path(directory) / "template.xlsx"
            output = Path(directory) / "assessment.xlsx"
            renderer.create_template(template)
            renderer.render(profile, template, output)
            workbook = load_workbook(output, read_only=False, data_only=False)
            self.assertEqual(
                workbook.sheetnames,
                ["Variable assessment", "Coverage", "Score rationale"],
            )
            sheet = workbook["Variable assessment"]
            self.assertEqual(
                [sheet.cell(1, column).value for column in range(1, 24)],
                [
                    "Category", "Variable", "Notes", "Considerations", "Source",
                    "Source Detail 1", "Source Detail 2", "Level data",
                    "History available", "Observation frequency", "Type", "values",
                    "Blocker / special label", "Source", "Definition", "Coverage",
                    "Missingness", "Consistency", "History", "Join", "Effort",
                    "Total", "Result",
                ],
            )
            low_scores = 0
            comments = 0
            for column in range(14, 22):
                if sheet.cell(2, column).value in {0, 1}:
                    low_scores += 1
                if sheet.cell(2, column).comment is not None:
                    comments += 1
            self.assertEqual(comments, low_scores)
            self.assertEqual(
                workbook["Score rationale"].max_row - 1,
                low_scores,
            )
            for column in range(14, 22):
                if sheet.cell(2, column).value == 2:
                    self.assertIsNone(sheet.cell(2, column).comment)
            self.assertEqual(sheet.cell(2, 22).value, sum(assessment["scores"].values()))
            self.assertEqual(sheet.cell(2, 23).value, "Sensible / restringida")
            self.assertIsNotNone(sheet.cell(2, 23).comment)
            self.assertIn("Motivo:", sheet.cell(2, 15).comment.text)
            self.assertIn("Resultado por puntos:", sheet.cell(2, 23).comment.text)

    def test_renderer_rejects_missing_rationale(self) -> None:
        assessment = self.profile(
            base_variable("Renderer guard"),
            raw_rows(["2026-01-01"]),
        )
        assessment["score_rationales"] = []
        profile = {"schema_version": 4, "assessments": [assessment]}
        with tempfile.TemporaryDirectory(dir=WORKSPACE_TEMP) as directory:
            template = Path(directory) / "template.xlsx"
            renderer.create_template(template)
            with self.assertRaisesRegex(ValueError, "rationale coverage mismatch"):
                renderer.render(profile, template, Path(directory) / "bad.xlsx")

    def test_renderer_rejects_missing_score_two_verification(self) -> None:
        assessment = self.profile(
            base_variable("Renderer score two guard"),
            raw_rows(["2026-01-01"]),
        )
        del assessment["score_verification"]["definition"]
        profile = {"schema_version": 4, "assessments": [assessment]}
        with tempfile.TemporaryDirectory(dir=WORKSPACE_TEMP) as directory:
            template = Path(directory) / "template.xlsx"
            renderer.create_template(template)
            with self.assertRaisesRegex(
                ValueError, "score-2 verification coverage mismatch"
            ):
                renderer.render(profile, template, Path(directory) / "bad.xlsx")

    def test_renderer_rejects_legacy_profile(self) -> None:
        profile = {"schema_version": 3, "assessments": []}
        with tempfile.TemporaryDirectory(dir=WORKSPACE_TEMP) as directory:
            template = Path(directory) / "template.xlsx"
            renderer.create_template(template)
            with self.assertRaisesRegex(ValueError, "legacy"):
                renderer.render(profile, template, Path(directory) / "bad.xlsx")



    def test_report_detail_explains_low_definition_score(self) -> None:
        variable = base_variable("Detailed definition")
        variable["evidence"]["definition"] = "needs_alignment"
        variable["rationale"] = {
            "definition": {
                "why": "El campo fuente existe, pero los casos límite de negocio no están aprobados.",
                "evidence": "Coverage puede medirse, pero no ha pasado una validación semántica.",
                "recommended_action": "Aprobar la definición y ejecutar una auditoría agregada guardada.",
            }
        }
        assessment = self.profile(
            variable,
            raw_rows(["2024-01-01", "2024-02-01", "2024-03-01", "2026-03-01"]),
        )
        detail = runner.build_report_detail(
            variable,
            assessment,
            self.rules,
            "unit-report",
            1,
            include_report_data=False,
        )
        definition = {
            item["criterion_key"]: item for item in detail["score_application"]
        }["definition"]
        self.assertEqual(definition["score"], 1)
        self.assertIn("needs_alignment", definition["status_rule"])
        self.assertIn("validación semántica", definition["observed_evidence"])
        self.assertIn("Aprobar", definition["recommended_action"])

    def test_report_html_contains_audit_sections(self) -> None:
        variable = base_variable("HTML detail")
        variable["evidence"]["definition"] = "needs_alignment"
        analysis = {
            "language": "es",
            "semantic_type": "direct_field",
            "definition": "Campo categórico disponible en cada snapshot mensual elegible.",
            "numerator_definition": "Empleados elegibles con el campo categórico poblado.",
            "attrition_use": "Segmentar el riesgo de attrition por contexto de puesto.",
            "prediction_cutoff": "Usar sólo información conocida al cierre del snapshot.",
            "source_reasoning": "El concepto solicitado corresponde al campo directo, pero requiere validar dominio e historia.",
            "validation_focus": [
                "Validamos el dominio y los nulos porque valores inesperados cambian la interpretación.",
                "Validamos el grain empleado-fecha porque duplicados alterarían coverage y join.",
            ],
            "prior_evidence_policy": "Los resultados anteriores son sólo contexto y deben revalidarse.",
        }
        variable["analysis"] = analysis
        runner.apply_spanish_publication_narrative(variable, analysis, self.rules)
        assessment = self.profile(variable, raw_rows(["2024-01-01", "2024-02-01", "2024-03-01", "2026-03-01"]))
        detail = runner.build_report_detail(
            variable,
            assessment,
            self.rules,
            "unit-html",
            1,
            include_report_data=False,
        )
        assessment["analysis"] = analysis
        assessment["analysis_plan"] = runner.build_analysis_plan(
            variable, assessment, self.rules, report_detail=detail
        )
        detail["analysis_plan"] = assessment["analysis_plan"]
        detail["historical_story"] = assessment["analysis_plan"]["historical_story"]
        detail["validation_story"] = assessment["analysis_plan"]["validation_story"]
        assessment["report_detail"] = detail
        runner.validate_public_narrative(assessment, detail)
        with tempfile.TemporaryDirectory(dir=WORKSPACE_TEMP) as directory:
            output = Path(directory) / "detail.html"
            runner.write_html(output, assessment)
            content = output.read_text(encoding="utf-8")
        self.assertIn("Resumen histórico y flujo del análisis", content)
        self.assertIn("Contexto del problema y del portafolio", content)
        self.assertIn("Resumen visual", content)
        self.assertIn("Historia de cada validación", content)
        self.assertIn("Por qué lo validamos", content)
        self.assertIn("Aplicación visible de reglas 2/1/0", content)
        self.assertIn("Definition", content)
        self.assertIn("La definición no puede considerarse confirmada", content)
        self.assertIn("Columnas involucradas", content)
        self.assertLess(
            content.index("Resumen visual"),
            content.index("Aplicación visible de reglas 2/1/0"),
        )

    def test_sensitive_report_values_are_suppressed(self) -> None:
        default_hints = runner.workflow.load_yaml(runner.workflow.DEFAULT_HINTS)
        try:
            runner.configure_publication_policy(default_hints, "presentation")
            safe, reason = runner.is_value_display_safe("gender")
            self.assertFalse(safe)
            self.assertIn("suprimidos", reason)
            safe, reason = runner.is_value_display_safe("job_family")
            self.assertTrue(safe)
            self.assertIn("agregados", reason)
        finally:
            runner.configure_publication_policy(default_hints, "presentation")

    def test_internal_analysis_allows_real_protected_category_labels_only(self) -> None:
        default_hints = runner.workflow.load_yaml(runner.workflow.DEFAULT_HINTS)
        try:
            runner.configure_publication_policy(default_hints, "internal_analysis")
            safe, reason = runner.is_value_display_safe("gender")
            self.assertTrue(safe)
            self.assertIn("análisis interno", reason)
            safe, reason = runner.is_value_display_safe("compensation_grade")
            self.assertTrue(safe)
            safe, reason = runner.is_value_display_safe("annual_base_pay")
            self.assertFalse(safe)
            self.assertIn("suprimidos", reason)
            safe, reason = runner.is_value_display_safe("employee_id")
            self.assertFalse(safe)
        finally:
            runner.configure_publication_policy(default_hints, "presentation")

    def test_public_narrative_rejects_missing_validation_story(self) -> None:
        variable = base_variable("Narrative guard")
        analysis = {
            "language": "es",
            "semantic_type": "direct_field",
            "definition": "Campo directo disponible por empleado y snapshot.",
            "numerator_definition": "Empleados elegibles con el campo poblado.",
            "attrition_use": "Aportar contexto laboral al modelo de attrition.",
            "prediction_cutoff": "Usar únicamente información conocida al snapshot.",
            "source_reasoning": "La fuente contiene el campo directo y permite revisar su dominio de forma agregada.",
            "validation_focus": [
                "Validamos el dominio porque categorías inesperadas cambian el significado del campo."
            ],
            "prior_evidence_policy": "Todo resultado anterior requiere revalidación.",
        }
        variable["analysis"] = analysis
        runner.apply_spanish_publication_narrative(variable, analysis, self.rules)
        assessment = self.profile(
            variable,
            raw_rows(["2024-01-01", "2024-02-01", "2024-03-01", "2026-03-01"]),
        )
        detail = runner.build_report_detail(
            variable, assessment, self.rules, "narrative-guard", 1, False
        )
        assessment["analysis_plan"] = runner.build_analysis_plan(
            variable, assessment, self.rules, report_detail=detail
        )
        detail["analysis_plan"] = assessment["analysis_plan"]
        detail["historical_story"] = assessment["analysis_plan"]["historical_story"]
        detail["validation_story"] = assessment["analysis_plan"]["validation_story"]
        runner.validate_public_narrative(assessment, detail)
        detail["validation_story"] = []
        with self.assertRaisesRegex(ValueError, "validation story count"):
            runner.validate_public_narrative(assessment, detail)

    def test_rehire_contextual_story_cannot_raise_score(self) -> None:
        catalog = profiler.load_yaml(
            runner.PROJECT_ROOT / "config" / "attrition_variable_analysis_v4.yaml"
        )
        stories = catalog["variables"]["Rehire"]["validation_story"]
        contextual = next(
            item for item in stories if item["id"] == "rehire_contextual_job_change"
        )
        self.assertEqual(contextual["evidence_role"], "contextual")
        self.assertIn("no eleva el score", contextual["acceptance_rule"])

    def test_spanish_publication_preserves_override_score(self) -> None:
        spec_path = (
            runner.PROJECT_ROOT
            / "config"
            / "attrition_candidate_variables_v4_reassessment_20260814.yaml"
        )
        spec = profiler.load_yaml(spec_path)
        variables, _ = runner.variables_with_analysis(
            spec, spec_path, self.rules
        )
        applied = next(
            item for item in variables if item["variable"] == "Applied to a job"
        )
        override = applied["score_overrides"]["history"]
        self.assertEqual(override["score"], 1)
        self.assertIn("feed no llega", override["reason"])

    def test_analysis_catalog_must_cover_the_spec_exactly(self) -> None:
        first = base_variable("Catalog one")
        second = base_variable("Catalog two")
        analysis = {
            "semantic_type": "direct_field",
            "attrition_use": "Test a direct snapshot field.",
            "source_reasoning": "The requested concept is the source field.",
            "validation_focus": ["domain", "history", "employee-date grain"],
        }
        with tempfile.TemporaryDirectory(dir=WORKSPACE_TEMP) as directory:
            root = Path(directory)
            spec_path = root / "spec.yaml"
            catalog_path = root / "analysis.yaml"
            spec = {
                "version": 4,
                "analysis_catalog": "analysis.yaml",
                "variables": [first, second],
            }
            spec_path.write_text(
                yaml.safe_dump(spec, sort_keys=False),
                encoding="utf-8",
            )
            catalog_path.write_text(
                yaml.safe_dump(
                    {
                        "version": 1,
                        "variables": {
                            "Catalog one": analysis,
                            "Catalog two": analysis,
                        },
                    },
                    sort_keys=False,
                ),
                encoding="utf-8",
            )

            variables, resolved = runner.variables_with_analysis(
                spec, spec_path, self.rules
            )
            self.assertEqual(len(variables), 2)
            self.assertEqual(resolved, catalog_path.resolve())
            self.assertEqual(
                variables[0]["verification"]["semantic_type"],
                "direct_field",
            )

            catalog_path.write_text(
                yaml.safe_dump(
                    {
                        "version": 1,
                        "variables": {"Catalog one": analysis},
                    },
                    sort_keys=False,
                ),
                encoding="utf-8",
            )
            with self.assertRaisesRegex(
                ValueError, "Analysis catalog mismatch"
            ):
                runner.variables_with_analysis(
                    spec, spec_path, self.rules
                )

    def test_semantic_type_selects_specific_analysis_modules(self) -> None:
        assessment = self.profile(
            base_variable("Semantic modules"),
            raw_rows(["2024-01-01", "2024-02-01", "2024-03-01", "2026-03-01"]),
        )
        variable = base_variable("Semantic modules")
        variable["analysis"] = {
            "semantic_type": "grouped_mapping",
            "attrition_use": "Test a grouped field.",
            "prediction_cutoff": "Known at the snapshot.",
            "source_reasoning": "Detailed and grouped fields require a mapping audit.",
            "validation_focus": ["contradictions", "unmapped values", "mapping changes"],
            "prior_evidence_policy": "Prior results are context only.",
        }
        plan = runner.build_analysis_plan(variable, assessment, self.rules)

        self.assertIn("mapping_contradictions", plan["analysis_modules"])
        self.assertIn("mapping_changes_over_time", plan["analysis_modules"])
        self.assertIn("mapping_summary", plan["planned_visualizations"])
        self.assertIn("score_overview", plan["planned_visualizations"])
        self.assertIn("annual_coverage_by_fiscal_year", plan["planned_visualizations"])
        self.assertEqual(plan["visualization_policy"]["priority"], "graph_first")
        self.assertNotIn("root_semantics", plan["analysis_modules"])

    def test_annual_coverage_uses_denominator_weighted_fiscal_summary(self) -> None:
        assessment = {
            "coverage": [
                {
                    "snapshot_date": "2025-10-01",
                    "fiscal_year": "2026",
                    "eligible_employees": 100,
                    "observable_employees": 100,
                    "missing_employees": 0,
                    "event_employees": 10,
                },
                {
                    "snapshot_date": "2025-11-01",
                    "fiscal_year": "2026",
                    "eligible_employees": 300,
                    "observable_employees": 150,
                    "missing_employees": 150,
                    "event_employees": 15,
                },
            ]
        }

        annual = runner.annual_coverage_by_fiscal_year(assessment)

        self.assertEqual(len(annual), 1)
        self.assertEqual(annual[0]["eligible_employee_snapshots"], 400)
        self.assertEqual(annual[0]["observable_employee_snapshots"], 250)
        self.assertEqual(annual[0]["observable_coverage_percent"], 62.5)
        self.assertEqual(annual[0]["missingness_percent"], 37.5)
        self.assertEqual(annual[0]["event_prevalence_percent"], 10.0)

    def test_relationship_summary_splits_distinct_metrics_for_read_only_session(self) -> None:
        queries = runner.build_single_pair_summary_queries(
            base_variable("Relación de columnas"),
            "job_family",
            "job_family_group",
            self.rules,
        )

        self.assertEqual(len(queries), 5)
        self.assertEqual(
            {metric for metric, _ in queries},
            {
                "employee_count",
                "left_distinct_values",
                "right_distinct_values",
                "left_values_with_multiple_right_values",
                "right_values_with_multiple_left_values",
            },
        )
        for _, sql in queries:
            self.assertTrue(sql.startswith("SELECT"))
            self.assertNotIn("UNION ALL", sql)
            self.assertNotIn("CROSS JOIN", sql)
        summary = runner.relationship_summary_row(
            "job_family",
            "job_family_group",
            [
                {"metric": "employee_count", "metric_value": 100},
                {"metric": "left_distinct_values", "metric_value": 10},
                {"metric": "right_distinct_values", "metric_value": 3},
                {
                    "metric": "left_values_with_multiple_right_values",
                    "metric_value": 1,
                },
                {
                    "metric": "right_values_with_multiple_left_values",
                    "metric_value": 2,
                },
            ],
        )
        self.assertEqual(summary["row_type"], "summary")
        self.assertEqual(summary["employee_count"], 100)

    def test_partial_report_failure_is_not_hidden_by_later_success(self) -> None:
        detail = {
            "evidence": {"status": "not_applicable", "query_results": {}},
            "column_relationships": {"notes": []},
        }
        with patch.object(
            runner.profiler,
            "run_aggregate_query",
            side_effect=[RuntimeError("first query failed"), [{"row_count": 1}]],
        ):
            first = runner.run_report_query(
                detail, "SELECT 1", 10, "queries/profiles/first.sql"
            )
            second = runner.run_report_query(
                detail, "SELECT 1", 10, "queries/profiles/second.sql"
            )

        self.assertIsNone(first)
        self.assertEqual(second, [{"row_count": 1}])
        self.assertEqual(detail["evidence"]["status"], "partial_failed")
        self.assertEqual(len(detail["evidence"]["errors"]), 1)
        self.assertIn("first query failed", detail["evidence"]["errors"][0]["error"])

    def test_report_queries_are_hashed_as_contextual_ledger_evidence(self) -> None:
        reference = "queries/audits/current_input_quality_audit.sql"
        detail = {
            "evidence": {
                "query_results": {
                    reference: [{"row_count": 100, "missing_count": 0}]
                }
            }
        }
        ledger = {
            "schema_version": 1,
            "variable": "Ledger report query",
            "queries": [],
        }

        merged = runner.merge_report_queries_into_ledger(
            base_variable("Ledger report query"), detail, ledger
        )

        self.assertEqual(merged["query_count"], 1)
        record = merged["queries"][0]
        self.assertTrue(record["sql_sha256"])
        self.assertTrue(record["result_sha256"])
        self.assertEqual(record["evidence_role"], "contextual")
        self.assertFalse(record["score_eligible"])

    def test_variable_outputs_include_report_detail_json(self) -> None:
        assessment = self.profile(base_variable("Output detail"), raw_rows(["2024-01-01", "2024-02-01", "2024-03-01", "2026-03-01"]))
        variable = base_variable("Output detail")
        variable["analysis"] = {
            "semantic_type": "direct_field",
            "attrition_use": "Test direct-field reporting.",
            "prediction_cutoff": "Known at the snapshot.",
            "source_reasoning": "The fixture represents a direct snapshot field.",
            "validation_focus": ["domain", "history", "join"],
            "prior_evidence_policy": "Fixture evidence only.",
        }
        assessment["report_detail"] = runner.build_report_detail(
            variable,
            assessment,
            self.rules,
            "unit-output",
            1,
            include_report_data=False,
        )
        assessment["variable_spec"] = variable
        assessment["analysis_plan"] = runner.build_analysis_plan(
            variable, assessment, self.rules
        )
        assessment["report_detail"]["analysis_plan"] = assessment["analysis_plan"]
        profile = {"schema_version": 4, "assessments": [assessment]}
        with tempfile.TemporaryDirectory(dir=WORKSPACE_TEMP) as directory:
            artifacts = runner.write_variable_outputs(profile, Path(directory) / "vars")
            report_path = Path(artifacts[0]["report_detail_json"])
            if not report_path.is_absolute():
                report_path = runner.PROJECT_ROOT / report_path
            self.assertTrue(report_path.exists())
            detail = profiler.load_json(report_path)
            analysis_path = runner.PROJECT_ROOT / artifacts[0]["analysis_plan"]
            checkpoint_path = runner.PROJECT_ROOT / artifacts[0]["checkpoint"]
            notebook_path = runner.PROJECT_ROOT / artifacts[0]["notebook"]
            self.assertIn("score_application", detail)
            self.assertTrue(analysis_path.exists())
            self.assertTrue(checkpoint_path.exists())
            notebook = profiler.load_json(notebook_path)
            code_cells = [
                cell for cell in notebook["cells"] if cell["cell_type"] == "code"
            ]
            self.assertTrue(code_cells)
            self.assertTrue(all(cell["execution_count"] for cell in code_cells))
            self.assertTrue(
                any(
                    "image/png" in (output.get("data") or {})
                    for cell in code_cells
                    for output in (cell.get("outputs") or [])
                )
            )

    def test_cumulative_documentation_preserves_manual_content_and_updates_managed_section(self) -> None:
        variable = base_variable("Documentation variable")
        variable["analysis"] = {
            "language": "es",
            "semantic_type": "direct_field",
            "intended_use": "Aportar contexto laboral al caso de uso de RH.",
            "prediction_cutoff": "Cierre del snapshot mensual.",
            "source_reasoning": "El campo directo requiere validar dominio y grain.",
            "validation_focus": [
                "Validamos dominio y grain porque cambian la interpretación del campo."
            ],
            "documentation_path": "shared/README.md",
        }
        assessment = self.profile(
            variable,
            raw_rows(["2024-01-01", "2024-02-01", "2024-03-01", "2026-03-01"]),
        )
        detail = runner.build_report_detail(
            variable, assessment, self.rules, "documentation-unit", 1, False
        )
        portfolio_context = {
            "problem_context": {
                "problem_statement": "Entender qué variables de RH ayudan a priorizar una decisión de retención.",
                "decision_or_use_case": "Seleccionar variables para un análisis longitudinal.",
                "outcome": "Attrition posterior al snapshot.",
                "population": "Empleados regulares activos.",
                "knowledge_boundary": "Cierre del snapshot.",
            },
            "portfolio_overview": {},
            "variables": [
                {
                    "variable": "Documentation variable",
                    "context_alignment": {
                        "role": "candidate_variable",
                        "relevance": "Aporta contexto laboral.",
                        "relationships": [],
                    },
                }
            ],
        }
        assessment["analysis_plan"] = runner.build_analysis_plan(
            variable,
            assessment,
            self.rules,
            report_detail=detail,
            portfolio_context=portfolio_context,
        )
        detail["analysis_plan"] = assessment["analysis_plan"]
        detail["historical_story"] = assessment["analysis_plan"]["historical_story"]
        detail["validation_story"] = assessment["analysis_plan"]["validation_story"]
        assessment["report_detail"] = detail
        assessment["evidence_ledger"] = {"queries": []}
        assessment["variable_spec"] = variable

        with tempfile.TemporaryDirectory(dir=WORKSPACE_TEMP) as directory:
            root = Path(directory) / "docs"
            target = root / "shared" / "README.md"
            target.parent.mkdir(parents=True)
            target.write_text(
                "# Manual\n\nEste contenido fue escrito por una persona.\n",
                encoding="utf-8",
            )
            first = runner.update_variable_documentation(
                variable, assessment, root, "run-one"
            )
            assessment["result"] = "Resultado actualizado"
            second = runner.update_variable_documentation(
                variable, assessment, root, "run-two"
            )
            content = target.read_text(encoding="utf-8")

        self.assertEqual(first["path"], second["path"])
        self.assertIn("Este contenido fue escrito por una persona.", content)
        self.assertEqual(
            content.count("BEGIN evaluate-attrition-variable:documentation-variable"),
            1,
        )
        self.assertIn("`run-two`", content)
        self.assertIn("`run-one`", content)
        self.assertEqual(
            content.count("BEGIN evaluate-attrition-variable-run:documentation-variable:run-one"),
            1,
        )
        self.assertEqual(
            content.count("BEGIN evaluate-attrition-variable-run:documentation-variable:run-two"),
            1,
        )
        self.assertIn("sustituye la revalidación", content)


    def test_visual_html_is_dynamic_and_includes_every_story_and_notebook_chart(self) -> None:
        scores = {
            "source": 1,
            "definition": 1,
            "coverage": 1,
            "missingness": 1,
            "consistency": 1,
            "history": 1,
            "join": 1,
            "effort": 1,
        }
        score_application = [
            {
                "criterion_key": criterion,
                "criterion": criterion,
                "score": score,
                "status_rule": "Regla de prueba",
                "why_below_2": "La evidencia de prueba sigue pendiente.",
                "supporting_evidence": "Resultado agregado de prueba.",
                "recommended_action": "Completar la validación de prueba.",
                "outcome": "warning",
            }
            for criterion, score in scores.items()
        ]
        stories = [
            {
                "id": f"story_{index}",
                "criterion": "definition",
                "question": chr(191) + f"Pregunta de validación {index}?",
                "why": "La comprobación puede cambiar la interpretación.",
                "method": "Auditoría agregada de prueba.",
                "acceptance_rule": "La métrica debe cumplir el límite aprobado.",
                "fields": ["job_family"],
                "population": "Empleados regulares activos.",
                "grain": "Empleado-snapshot",
                "outcome": "warning",
                "observed_evidence": "Resultado agregado de prueba.",
                "observed_metrics": {"anomaly_count": index},
                "interpretation": "Todavía existe una condición pendiente.",
                "score_effect": "Mantiene Definition en 1.",
                "recommended_action": "Completar la validación enfocada.",
                "sql_references": ["queries/audits/current_input_quality_audit.sql"],
            }
            for index in range(1, 6)
        ]
        assessment = {
            "variable": "Visual dynamic test",
            "result": "Media - utilizable con trabajo",
            "total": 8,
            "scores": scores,
            "definition_text": "Campo de prueba para explicar una evaluación dinámica.",
            "source_table": "hrda.test_snapshot",
            "source_columns": ["job_family"],
            "observation_frequency": "Mensual",
            "level_data": "Empleado-snapshot",
            "measurement_type": "snapshot",
            "diagnostics": {},
            "score_verification": {},
        }
        report = {
            "score_application": score_application,
            "validation_story": stories,
            "columns": [],
            "value_distributions": [],
            "numeric_bands": [],
            "column_relationships": {},
            "privacy_note": "Solo se publican agregados seguros.",
            "evidence": {},
        }
        decision = {
            "final_result": assessment["result"],
            "blocker": {"label": "Ninguno"},
            "recommended_actions": [],
        }
        with tempfile.TemporaryDirectory(dir=WORKSPACE_TEMP) as directory:
            root = Path(directory)
            charts_dir = root / "charts"
            charts_dir.mkdir()
            charts = []
            for index in range(1, 4):
                relative = f"charts/chart-{index:02d}.png"
                (root / relative).write_bytes(b"not-a-real-image-needed-for-contract-test")
                charts.append({"file": relative, "title": f"Gráfico {index}"})
            paths = {}
            for name, value in {
                "assessment": assessment,
                "decision": decision,
                "report": report,
                "coverage": [],
                "ledger": {"queries": []},
            }.items():
                path = root / f"{name}.json"
                path.write_text(json.dumps(value, ensure_ascii=False), encoding="utf-8")
                paths[name] = path
            bundle = {
                "assessment": assessment,
                "decision": decision,
                "report": report,
                "coverage": [],
                "ledger": {"queries": []},
                "private_metrics": {},
                "charts": charts,
                "paths": paths,
                "variable_dir": root,
            }
            template = visual_renderer.resolve_template(None)
            document = visual_renderer.build_document(
                bundle, template.read_text(encoding="utf-8")
            )
            errors = visual_renderer.validate_document(document, bundle)
            expected_slides = visual_renderer.expected_slide_count(bundle)

        self.assertEqual(expected_slides, 16)
        self.assertFalse(errors, errors)
        self.assertIn("Validación 5", document)
        self.assertIn('src="charts/chart-03.png"', document)
        self.assertIn('data-slide="15"', document)


class PortfolioWorkflowTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls.rules_path = SKILL_ROOT / "references" / "scoring-rules.yaml"
        cls.rules = profiler.load_yaml(cls.rules_path)
        cls.catalog_spec = (
            runner.PROJECT_ROOT
            / "config"
            / "attrition_candidate_variables_v4_reassessment_20260814.yaml"
        )

    def test_project_visual_html_adapter_is_portable_and_resolved(self) -> None:
        settings = runner.project_configuration(
            runner.PROJECT_ROOT / "config" / "variable_assessment_project.yaml"
        )
        visual = runner.visual_html_v2_configuration(settings)
        self.assertTrue(visual["enabled"])
        self.assertTrue(Path(visual["renderer"]["script"]).is_file())
        self.assertTrue(Path(visual["metrics_builder"]["script"]).is_file())
        self.assertEqual(visual["metrics_builder"]["only_when_blocker"], "sensitive_restricted")
        self.assertTrue(visual["dynamic_slide_count"])
        self.assertGreaterEqual(visual["minimum_slides"], 10)

    def test_semantic_hint_registry_covers_representative_patterns(self) -> None:
        hints = runner.workflow.load_yaml(runner.workflow.DEFAULT_HINTS)
        cases = {
            "Comp ratio": "ratio_rate_percent",
            "Talent index score": "index_score_average",
            "Time since last promotion": "count_window_time_since",
            "Job level group": "group_category_level",
            "Layers from CEO manager": "manager_team_layers",
            "Rehire event flag": "flag_event",
            "External QVoice survey": "survey_external",
        }
        for name, expected in cases.items():
            variable = {
                "variable": name,
                "definition": name,
                "source": {"columns": []},
                "analysis": {"validation_focus": [name]},
            }
            ids = {
                item["id"]
                for item in runner.workflow.detected_cues(variable, hints)
            }
            self.assertIn(expected, ids, name)

    def test_full_catalog_preflight_has_43_unique_variables_and_grouped_questions(self) -> None:
        prepared = runner.prepare_portfolio(
            self.catalog_spec,
            self.rules_path,
            "portfolio-unit",
            execute_metadata=False,
            persist_metadata_sql=False,
        )
        context = prepared["portfolio_context"]
        names = [item["variable"] for item in context["variables"]]
        self.assertEqual(len(names), 43)
        self.assertEqual(len(set(names)), 43)
        self.assertEqual(context["metadata"]["status"], "not_executed")
        self.assertTrue(context["problem_context"]["problem_statement"])
        self.assertEqual(
            sum(context["portfolio_overview"]["semantic_type_counts"].values()),
            43,
        )
        self.assertEqual(
            sum(context["portfolio_overview"]["context_role_counts"].values()),
            43,
        )
        self.assertLess(
            context["evaluation_order"].index("Comp ratio"),
            context["evaluation_order"].index("Within team compensation range"),
        )
        self.assertEqual(
            context["dependencies"]["Within team compensation range"],
            ["Comp ratio"],
        )
        ratio_question = next(
            item
            for item in context["material_questions"]
            if "Comp ratio" in item["variables"]
            and item["id"] == "semantic:ratio_rate_percent"
        )
        self.assertEqual(
            ratio_question["question"],
            "Parece una división; ¿qué constituye el numerador y el denominador?",
        )
        self.assertLess(
            len(context["material_questions"]),
            len(context["variables"]),
        )

    def test_portable_preflight_uses_project_schema_hints_and_intended_use(self) -> None:
        variable = base_variable("Training completion rate")
        variable["category"] = "Learning"
        variable["definition"] = "Tasa de finalización de capacitación por periodo."
        variable["source"]["table"] = "learning_monthly"
        variable["source"]["columns"] = ["completion_rate"]
        variable["measurement"]["observable_columns"] = ["completion_rate"]
        variable["analysis"] = {
            "language": "es",
            "definition": "Tasa mensual de finalización de capacitación elegible.",
            "numerator_definition": "Empleados con cursos elegibles completados en el periodo.",
            "semantic_type": "direct_field",
            "intended_use": "Evaluar preparación de capacidades y contexto de desarrollo.",
            "source_reasoning": "La métrica procede de una plataforma de aprendizaje distinta a HRDA.",
            "validation_focus": [
                "Validamos numerador, denominador y elegibilidad porque cambian la interpretación de la tasa."
            ],
        }

        hints = {
            "version": 1,
            "language": "es",
            "cues": [
                {
                    "id": "learning_completion_rate",
                    "labels": ["capacitación"],
                    "match_any": ["training", "learning", "completion", "capacitación"],
                    "likely_semantic_types": ["direct_field", "precomputed_metric"],
                    "required_topics": ["numerator", "denominator", "eligibility"],
                    "material_question": "¿Qué numerador, denominador y elegibilidad definen la tasa?",
                    "why_it_matters": "Cambia la población y la comparabilidad entre periodos.",
                    "recommended_hypothesis": "Tratarla como métrica precalculada hasta reconciliar la fórmula.",
                }
            ],
            "publication_policies": {"minimum_shared_cohort": 10},
        }
        metadata = {
            "rows": [
                {
                    "table_schema": "learning",
                    "table_name": "learning_monthly",
                    "column_name": "completion_rate",
                    "data_type": "numeric",
                    "ordinal_position": 4,
                }
            ]
        }

        with tempfile.TemporaryDirectory(dir=WORKSPACE_TEMP) as directory:
            root = Path(directory)
            hints_path = root / "portable-hints.yaml"
            spec_path = root / "portable-spec.yaml"
            metadata_path = root / "portable-metadata.json"
            hints_path.write_text(
                yaml.safe_dump(hints, sort_keys=False, allow_unicode=True),
                encoding="utf-8",
            )
            spec_path.write_text(
                yaml.safe_dump(
                    {
                        "version": 4,
                        "semantic_hint_registry": hints_path.name,
                        "runtime": {
                            "metadata": {
                                "default_schema": "learning",
                                "search_schemas": ["learning", "people_analytics"],
                            }
                        },
                        "variables": [variable],
                    },
                    sort_keys=False,
                    allow_unicode=True,
                ),
                encoding="utf-8",
            )
            metadata_path.write_text(
                json.dumps(metadata, ensure_ascii=False, indent=2),
                encoding="utf-8",
            )

            prepared = runner.prepare_portfolio(
                spec_path,
                self.rules_path,
                "portable-project",
                metadata_results_path=metadata_path,
                execute_metadata=False,
                persist_metadata_sql=False,
            )

        context = prepared["portfolio_context"]
        self.assertNotIn("hrda", prepared["metadata_sql"].lower())
        self.assertIn("'learning'", prepared["metadata_sql"])
        self.assertIn("'people_analytics'", prepared["metadata_sql"])
        self.assertEqual(context["metadata"]["default_schema"], "learning")
        self.assertEqual(context["variables"][0]["metadata"]["found"], ["completion_rate"])
        self.assertEqual(context["material_questions"][0]["id"], "semantic:learning_completion_rate")
        self.assertEqual(
            prepared["run_request"]["variables"][0]["intended_use"],
            "Evaluar preparación de capacidades y contexto de desarrollo.",
        )

    def test_problem_context_drives_portfolio_relationships_and_variable_alignment(self) -> None:
        fixtures = SKILL_ROOT / "tests" / "fixtures"
        prepared = runner.prepare_portfolio(
            None,
            self.rules_path,
            "context-unit",
            request_path=fixtures / "offline-three-run-request.yaml",
            catalog_spec_path=self.catalog_spec,
            execute_metadata=False,
            persist_metadata_sql=False,
        )
        context = prepared["portfolio_context"]
        self.assertIn("compensación relativa", context["problem_context"]["problem_statement"])
        self.assertEqual(
            context["portfolio_overview"]["outcome"],
            "Attrition posterior al snapshot mensual.",
        )
        relationship = next(
            item
            for item in context["contextual_relationships"]
            if item["type"] == "problem_context"
        )
        self.assertEqual(
            relationship["variables"],
            ["Comp ratio", "Within team compensation range"],
        )
        by_name = {item["variable"]: item for item in context["variables"]}
        self.assertEqual(
            by_name["Comp ratio"]["context_alignment"]["role"],
            "métrica base de compensación",
        )
        html = runner.workflow.render_portfolio_html(context)
        self.assertIn("Contexto del problema", html)
        self.assertIn("Overview visual del portafolio", html)
        self.assertIn("Mapa contextual de relaciones", html)

    def test_engine_scripts_do_not_name_catalog_variables_or_hrda_schema(self) -> None:
        scripts = "\n".join(
            path.read_text(encoding="utf-8")
            for path in (SKILL_ROOT / "scripts").glob("*.py")
        )
        for forbidden in (
            "Comp ratio",
            "Within team compensation range",
            "Rehire",
            "c.table_schema = 'hrda'",
        ):
            self.assertNotIn(forbidden, scripts)

    def test_dependency_cycles_are_reported_with_stable_original_order(self) -> None:
        names = ["A", "B", "C"]
        dependencies = {"A": ["B"], "B": ["A"], "C": ["A"]}
        order, cycles = runner.workflow.stable_dependency_order(
            names, dependencies
        )
        self.assertEqual(order, ["A", "B", "C"])
        self.assertEqual(cycles, [["A", "B"]])

    def test_pre_execution_plan_requires_explicit_story_seed_and_rule(self) -> None:
        variable = base_variable("No story")
        variable["analysis"] = {
            "semantic_type": "direct_field",
            "attrition_use": "Uso de prueba.",
            "source_reasoning": "Campo directo de prueba.",
            "validation_focus": [],
        }
        with self.assertRaisesRegex(ValueError, "explicit validation_story"):
            runner.workflow.build_pre_execution_plan(
                variable, {"dependencies": []}
            )

    def test_strict_analysis_plan_requires_authored_story_and_revision_reason(self) -> None:
        variable = base_variable("Explicit plan")
        variable["analysis"] = {
            "semantic_type": "direct_field",
            "attrition_use": "Uso de prueba.",
            "source_reasoning": "Campo directo de prueba.",
        }
        prepared = {"variables": [variable]}
        with self.assertRaisesRegex(ValueError, "requires one --analysis-plan"):
            runner.apply_explicit_analysis_plans(prepared, [], strict=True)

        question = chr(191) + "Se confirma el dominio del campo?"
        plan = {
            "variable": "Explicit plan",
            "created_before_execution": True,
            "plan_version": 1,
            "semantic_type": "direct_field",
            "validation_story": [{
                "id": "domain",
                "criterion": "definition",
                "evidence_role": "defining",
                "question": question,
                "why": "El dominio puede cambiar el significado del campo.",
                "method": "Auditoria agregada del dominio.",
                "sql_reference": "queries/audits/current_input_quality_audit.sql",
                "acceptance_rule": "El dominio debe coincidir con la definicion aprobada.",
            }],
        }
        with tempfile.TemporaryDirectory(dir=WORKSPACE_TEMP) as directory:
            path = Path(directory) / "analysis-plan-v1.yaml"
            path.write_text(
                yaml.safe_dump(plan, sort_keys=False, allow_unicode=True),
                encoding="utf-8",
            )
            records = runner.apply_explicit_analysis_plans(
                prepared, [path], strict=True
            )
            self.assertEqual(records[0]["plan_version"], 1)
            self.assertEqual(
                prepared["variables"][0]["analysis"]["plan_authorship"],
                "codex_explicit",
            )
            revision_path = Path(directory) / "analysis-plan-v2.yaml"
            revision = {**plan, "plan_version": 2}
            revision_path.write_text(
                yaml.safe_dump(revision, sort_keys=False, allow_unicode=True),
                encoding="utf-8",
            )
            with self.assertRaisesRegex(ValueError, "plan_revision_reason"):
                runner.apply_explicit_analysis_plans(
                    prepared, [revision_path], strict=True
                )

    def test_legacy_validation_focus_can_be_materialized_only_when_explicitly_allowed(self) -> None:
        variable = base_variable("Historia en español")
        variable["analysis"] = {
            "semantic_type": "derived_feature",
            "source_reasoning": "La señal requiere un cálculo móvil reproducible.",
            "validation_focus": [
                "Validamos exactamente el límite mayor que tres porque tres y cuatro solicitudes deben clasificarse distinto."
            ],
        }

        with self.assertRaisesRegex(ValueError, "explicit validation_story"):
            runner.workflow.materialize_validation_stories(variable)
        story = runner.workflow.materialize_validation_stories(
            variable, allow_legacy_focus=True
        )[0]

        self.assertEqual(
            story["question"],
            "¿Cómo validamos exactamente el límite mayor que tres?",
        )
        self.assertEqual(
            story["why"],
            "Esta comprobación es necesaria porque tres y cuatro solicitudes deben clasificarse distinto.",
        )
        self.assertTrue(runner.text_has_spanish_context(story["why"]))

    def test_compa_ratio_values_are_suppressed_but_bands_remain_supported(self) -> None:
        safe, reason = runner.is_value_display_safe("compa_ratio")
        self.assertFalse(safe)
        self.assertIn("suprimidos", reason)
        spec = profiler.load_yaml(self.catalog_spec)
        catalog = profiler.load_yaml(
            runner.PROJECT_ROOT / "config" / "attrition_variable_analysis_v4.yaml"
        )
        self.assertIn("numeric_bands", catalog["variables"]["Comp ratio"])
        self.assertEqual(spec["version"], 4)

    def test_hr_privacy_baseline_and_project_additions_are_suppressed(self) -> None:
        default_hints = runner.workflow.load_yaml(runner.workflow.DEFAULT_HINTS)
        try:
            runner.configure_publication_policy(
                {
                    "publication_policies": {
                        "exact_value_suppressed_columns": [
                            "learning_assessment_score"
                        ]
                    }
                }
            )
            for column in (
                "employee_id",
                "gender",
                "compa_ratio",
                "home_city",
                "et_1_leader",
                "learning_assessment_score",
            ):
                safe, reason = runner.is_value_display_safe(column)
                self.assertFalse(safe, column)
                self.assertIn("suprimidos", reason)
            safe, _ = runner.is_value_display_safe("learning_status")
            self.assertTrue(safe)
        finally:
            runner.configure_publication_policy(default_hints)

    def test_evidence_reuse_identity_requires_exact_claim_population_grain_window_and_hash(self) -> None:
        variable = base_variable("Reuse identity")
        variable["analysis"] = {"prediction_cutoff": "Snapshot t"}
        item = {"acceptance_rule": "No mismatches."}
        first = runner.workflow.evidence_identity(variable, item, "abc")
        same = runner.workflow.evidence_identity(variable, item, "abc")
        changed_claim = runner.workflow.evidence_identity(
            variable, {"acceptance_rule": "No duplicates."}, "abc"
        )
        variable["population"]["display"] = "Otra población"
        changed_population = runner.workflow.evidence_identity(
            variable, item, "abc"
        )
        self.assertEqual(first, same)
        self.assertNotEqual(first, changed_claim)
        self.assertNotEqual(first, changed_population)

    def test_adaptive_gap_adds_a_new_plan_version(self) -> None:
        variable = base_variable("Adaptive")
        variable["analysis"] = {
            "semantic_type": "direct_field",
            "attrition_use": "Uso de prueba.",
            "source_reasoning": "Campo directo de prueba.",
            "prediction_cutoff": "Snapshot t.",
            "validation_focus": ["Validamos dominio porque puede cambiar."],
            "adaptive_validations": [
                {
                    "trigger": {
                        "query_reference": "queries/audits/current_input_quality_audit.sql",
                        "metric": "anomaly_count",
                        "operator": "gt",
                        "threshold": 0,
                    },
                    "story": {
                        "id": "focused_anomaly",
                        "criterion": "consistency",
                        "evidence_role": "defining",
                        "question": "¿Cuál es la causa de la anomalía observada?",
                        "why": "La anomalía no prevista puede cambiar el score.",
                        "method": "Auditoría agregada enfocada de la anomalía.",
                        "sql_reference": "queries/audits/current_input_quality_audit.sql",
                        "acceptance_rule": "La anomalía debe quedar explicada o corregida.",
                    },
                }
            ],
        }
        ledger = {
            "queries": [
                {
                    "query_reference": "queries/audits/current_input_quality_audit.sql",
                    "metrics": {"anomaly_count": 2},
                }
            ]
        }
        stories = runner.triggered_adaptive_stories(variable, ledger)
        self.assertEqual([item["id"] for item in stories], ["focused_anomaly"])
        variable["analysis"]["validation_story"] = (
            runner.workflow.materialize_validation_stories(
                variable, allow_legacy_focus=True
            ) + stories
        )
        plan = runner.workflow.build_pre_execution_plan(
            variable, {"dependencies": []}, version=2
        )
        self.assertEqual(plan["plan_version"], 2)
        self.assertEqual(len(plan["validation_story"]), 2)

    def test_offline_three_variable_integration_persists_contracts_and_workbook_order(self) -> None:
        fixtures = SKILL_ROOT / "tests" / "fixtures"
        request_path = fixtures / "offline-three-run-request.yaml"
        aggregate_path = fixtures / "offline-three-aggregates.json"

        with tempfile.TemporaryDirectory(dir=WORKSPACE_TEMP) as directory:
            root = Path(directory)
            prepared = runner.prepare_portfolio(
                None,
                self.rules_path,
                "offline-three",
                request_path=request_path,
                catalog_spec_path=self.catalog_spec,
                execute_metadata=False,
                persist_metadata_sql=False,
            )
            plans: list[str] = []
            documentation_root = root / "documentation"
            (documentation_root / "compensation").mkdir(parents=True)
            (documentation_root / "compensation" / "README.md").write_text(
                "# Compensation\n\nContexto manual previo.\n",
                encoding="utf-8",
            )

            def document_completed_variable(
                number: int, total: int, assessment: dict
            ) -> None:
                assessment["documentation"] = runner.update_variable_documentation(
                    assessment["variable_spec"],
                    assessment,
                    documentation_root,
                    "offline-three",
                )

            with patch.object(runner, "REPORT_SQL_ROOT", root / "generated-queries"):
                profile = runner.build_profile(
                    Path(prepared["source_spec_path"]),
                    self.rules_path,
                    aggregate_path,
                    "offline-three",
                    include_report_data=False,
                    on_plan=lambda number, total, plan: plans.append(
                        str(plan["variable"])
                    ),
                    on_assessment=document_completed_variable,
                    prepared_portfolio=prepared,
                    documentation_root=documentation_root,
                )
                variables_dir = root / "vars"
                artifacts = runner.write_variable_outputs(profile, variables_dir)
                runner.validate_variable_artifacts(profile, artifacts)
                workbook = root / "assessment.xlsx"
                renderer.render(profile, renderer.DEFAULT_TEMPLATE, workbook)
                renderer.validate_workbook(workbook, profile)

            self.assertEqual(
                profile["original_order"],
                [
                    "Within team compensation range",
                    "Comp ratio",
                    "Rehire",
                ],
            )
            self.assertEqual(
                profile["evaluation_order"],
                [
                    "Comp ratio",
                    "Within team compensation range",
                    "Rehire",
                ],
            )
            self.assertEqual(plans, profile["evaluation_order"])
            self.assertEqual(
                profile["problem_context"]["outcome"],
                "Attrition posterior al snapshot mensual.",
            )
            by_name = {
                item["variable"]: item for item in profile["assessments"]
            }
            inherited = by_name["Within team compensation range"][
                "inherited_conditions"
            ]
            self.assertEqual(inherited[0]["dependency"], "Comp ratio")
            self.assertEqual(inherited[0]["status"], "limited")
            shared_context = by_name["Within team compensation range"][
                "pre_execution_plan"
            ]["shared_context"]
            self.assertEqual(
                shared_context["problem_context"]["outcome"],
                "Attrition posterior al snapshot mensual.",
            )
            self.assertEqual(
                shared_context["accumulated_knowledge"][0]["variable"],
                "Comp ratio",
            )
            self.assertIn(
                "Conocimiento acumulado: Comp ratio",
                shared_context["documentation_context"]["excerpt"],
            )
            self.assertFalse(shared_context["context_is_evidence"])
            rehire_stories = {
                item["id"]: item
                for item in by_name["Rehire"]["report_detail"]["validation_story"]
            }
            self.assertEqual(
                rehire_stories["rehire_employment_sequence"]["outcome"],
                "pass",
            )
            self.assertEqual(
                set(
                    rehire_stories["rehire_employment_sequence"][
                        "observed_metrics"
                    ]
                ),
                {
                    "rehire_before_original_hire_employees",
                    "rehire_vs_max_hire_mismatch_employees",
                },
            )
            for assessment in profile["assessments"]:
                ledger = assessment["evidence_ledger"]
                coverage_ref = ledger["coverage_query_reference"]
                coverage_record = next(
                    item
                    for item in ledger["queries"]
                    if item["query_reference"] == coverage_ref
                )
                self.assertTrue((runner.PROJECT_ROOT / coverage_ref).is_file())
                self.assertTrue(coverage_record["sql_sha256"])
                self.assertTrue(coverage_record["result_sha256"])
            loaded = load_workbook(workbook, read_only=True, data_only=False)
            self.assertEqual(
                loaded.sheetnames,
                ["Variable assessment", "Coverage", "Score rationale"],
            )
            sheet = loaded["Variable assessment"]
            self.assertEqual(
                [sheet.cell(row, 2).value for row in range(2, 5)],
                profile["original_order"],
            )
            loaded.close()

if __name__ == "__main__":
    unittest.main()


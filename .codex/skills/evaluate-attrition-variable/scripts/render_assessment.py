#!/usr/bin/env python
"""Render traceable attrition assessments to the approved three-sheet workbook."""

from __future__ import annotations

import argparse
import json
import re
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation


SKILL_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_TEMPLATE = SKILL_ROOT / "assets" / "variable-assessment-template.xlsx"
ASSESSMENT_SHEET = "Variable assessment"
COVERAGE_SHEET = "Coverage"
RATIONALE_SHEET = "Score rationale"
ASSESSMENT_HEADERS = [
    "Category",
    "Variable",
    "Notes",
    "Considerations",
    "Source",
    "Source Detail 1",
    "Source Detail 2",
    "Level data",
    "History available",
    "Observation frequency",
    "Type",
    "values",
    "Blocker / special label",
    "Source",
    "Definition",
    "Coverage",
    "Missingness",
    "Consistency",
    "History",
    "Join",
    "Effort",
    "Total",
    "Result",
]
COVERAGE_HEADERS = [
    "Variable",
    "Source table",
    "Source columns",
    "Measurement type",
    "Snapshot date",
    "Fiscal year",
    "Eligible employees",
    "Observable employees",
    "Missing employees",
    "Observable coverage %",
    "Coverage score",
    "Event employees",
    "Event prevalence %",
    "Stable >=95%",
    "First observed date",
    "Usable since date",
    "Population filter",
    "Numerator definition",
    "Notes",
]
RATIONALE_HEADERS = [
    "Category",
    "Variable",
    "Criterion",
    "Score",
    "Status / rule",
    "Why score is below 2",
    "Supporting evidence",
    "Recommended action",
    "Scoring method",
    "Source table",
    "Source column(s)",
]
CRITERIA = [
    "source",
    "definition",
    "coverage",
    "missingness",
    "consistency",
    "history",
    "join",
    "effort",
]
QUALITATIVE_CRITERIA = {"source", "definition", "consistency", "join", "effort"}

NAVY = "183642"
TEAL = "0F6B78"
LIGHT_TEAL = "D9EDEF"
GREEN = "D9EAD3"
AMBER = "FFF2CC"
RED = "F4CCCC"
GRAY = "E7EAEC"
WHITE = "FFFFFF"
GRID = Side(style="thin", color="C7CDD1")


def slugify(value: str) -> str:
    slug = re.sub(r"[^a-z0-9]+", "-", value.lower()).strip("-")
    return slug or "variable"


def header_style(cell, fill_color: str) -> None:
    cell.fill = PatternFill("solid", fgColor=fill_color)
    cell.font = Font(color=WHITE, bold=True)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = Border(bottom=Side(style="medium", color=TEAL))


def style_template_sheet(sheet, headers: list[str], widths: list[float]) -> None:
    sheet.append(headers)
    sheet.sheet_view.showGridLines = False
    sheet.row_dimensions[1].height = 34
    for index, (header, width) in enumerate(zip(headers, widths), start=1):
        cell = sheet.cell(1, index)
        header_style(cell, NAVY)
        sheet.column_dimensions[cell.column_letter].width = width
    sheet.auto_filter.ref = f"A1:{sheet.cell(1, len(headers)).column_letter}1"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_title_rows = "1:1"
    sheet.oddFooter.center.text = "Evaluación de disponibilidad de variables de attrition"
    sheet.oddFooter.right.text = "Página &P de &N"


def create_template(path: Path) -> None:
    workbook = Workbook()
    default = workbook.active
    workbook.remove(default)
    assessment = workbook.create_sheet(ASSESSMENT_SHEET)
    coverage = workbook.create_sheet(COVERAGE_SHEET)
    rationale = workbook.create_sheet(RATIONALE_SHEET)

    style_template_sheet(
        assessment,
        ASSESSMENT_HEADERS,
        [28, 38, 34, 52, 24, 48, 62, 25, 72, 30, 24, 64, 27, 11, 12, 11, 13, 13, 11, 9, 9, 9, 28],
    )
    assessment.freeze_panes = "N2"
    assessment.page_setup.orientation = "landscape"
    score_validation = DataValidation(
        type="whole", operator="between", formula1="0", formula2="2", allow_blank=False
    )
    score_validation.promptTitle = "Score"
    score_validation.prompt = "Ingresa 0, 1 o 2."
    score_validation.error = "Sólo se permite 0, 1 o 2."
    score_validation.errorTitle = "Score inválido"
    score_validation.showErrorMessage = True
    assessment.add_data_validation(score_validation)
    score_validation.add("N2:U5000")
    assessment.conditional_formatting.add(
        "N2:U5000", CellIsRule(operator="equal", formula=["2"], fill=PatternFill("solid", fgColor=GREEN))
    )
    assessment.conditional_formatting.add(
        "N2:U5000", CellIsRule(operator="equal", formula=["1"], fill=PatternFill("solid", fgColor=AMBER))
    )
    assessment.conditional_formatting.add(
        "N2:U5000", CellIsRule(operator="equal", formula=["0"], fill=PatternFill("solid", fgColor=RED))
    )

    style_template_sheet(
        coverage,
        COVERAGE_HEADERS,
        [34, 38, 55, 22, 15, 13, 18, 20, 18, 22, 15, 17, 20, 17, 18, 18, 46, 62, 62],
    )
    coverage.freeze_panes = "A2"
    coverage.page_setup.orientation = "landscape"
    coverage.conditional_formatting.add(
        "K2:K5000", CellIsRule(operator="equal", formula=["2"], fill=PatternFill("solid", fgColor=GREEN))
    )
    coverage.conditional_formatting.add(
        "K2:K5000", CellIsRule(operator="equal", formula=["1"], fill=PatternFill("solid", fgColor=AMBER))
    )
    coverage.conditional_formatting.add(
        "K2:K5000", CellIsRule(operator="equal", formula=["0"], fill=PatternFill("solid", fgColor=RED))
    )

    style_template_sheet(
        rationale,
        RATIONALE_HEADERS,
        [28, 36, 18, 10, 36, 58, 72, 58, 28, 42, 58],
    )
    rationale.freeze_panes = "A2"
    rationale.page_setup.orientation = "landscape"
    rationale.conditional_formatting.add(
        "D2:D5000", CellIsRule(operator="equal", formula=["1"], fill=PatternFill("solid", fgColor=AMBER))
    )
    rationale.conditional_formatting.add(
        "D2:D5000", CellIsRule(operator="equal", formula=["0"], fill=PatternFill("solid", fgColor=RED))
    )

    workbook.properties.title = "Evaluación de disponibilidad de variables de attrition"
    workbook.properties.subject = (
        "Evidencia de scoring, coverage longitudinal y racionales por criterio"
    )
    workbook.properties.creator = "evaluate-attrition-variable skill"
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def parse_date(value: Any) -> date | None:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return date.fromisoformat(str(value)[:10])


def load_profile(path: Path) -> dict[str, Any]:
    data = json.loads(path.read_text(encoding="utf-8"))
    assessments = data.get("assessments")
    if not isinstance(assessments, list) or not assessments:
        raise ValueError("Profile JSON must contain a non-empty assessments list.")
    return data


def clear_data_rows(sheet) -> None:
    if sheet.max_row > 1:
        sheet.delete_rows(2, sheet.max_row - 1)


def set_body_style(sheet, row_number: int, maximum_column: int) -> None:
    fill = PatternFill("solid", fgColor="F7F9FA" if row_number % 2 == 0 else WHITE)
    for column in range(1, maximum_column + 1):
        cell = sheet.cell(row_number, column)
        cell.fill = fill
        cell.border = Border(bottom=GRID)
        cell.alignment = Alignment(vertical="top", wrap_text=True)
    sheet.row_dimensions[row_number].height = 45


def label_fill(value: str) -> PatternFill:
    normalized = value.lower()
    if normalized.startswith("high"):
        return PatternFill("solid", fgColor=GREEN)
    if normalized.startswith("medium") or normalized == "tbd":
        return PatternFill("solid", fgColor=AMBER)
    if normalized == "none":
        return PatternFill("solid", fgColor=GRAY)
    return PatternFill("solid", fgColor=RED)


def join_notes(*values: Any) -> str:
    parts = [str(value).strip() for value in values if value not in (None, "")]
    return " ".join(parts)


def assessment_considerations(assessment: dict[str, Any]) -> str:
    """Return a compact, factual summary of constraints that affect use."""
    items: list[str] = []
    for condition in assessment.get("inherited_conditions") or []:
        if isinstance(condition, dict):
            text = condition.get("reason") or condition.get("condition")
        else:
            text = condition
        if text:
            items.append(str(text))
    plan = assessment.get("analysis_plan") or {}
    for gap in ((plan.get("gap_review") or {}).get("unresolved") or []):
        if isinstance(gap, dict):
            text = gap.get("reason") or gap.get("description") or gap.get("criterion")
        else:
            text = gap
        if text:
            items.append(str(text))
    for rationale in assessment.get("score_rationales") or []:
        text = rationale.get("why_below_2") if isinstance(rationale, dict) else None
        if text:
            items.append(str(text))
    unique = list(dict.fromkeys(" ".join(item.split()) for item in items if item))
    return " | ".join(unique[:4]) or "Sin consideraciones adicionales registradas."


def assessment_semantic_type(assessment: dict[str, Any]) -> str:
    plan = assessment.get("analysis_plan") or {}
    return str(
        plan.get("semantic_type")
        or (assessment.get("variable_spec") or {}).get("analysis", {}).get("semantic_type")
        or "Por confirmar"
    )


def assessment_values_summary(assessment: dict[str, Any]) -> str:
    """Return only the public aggregate value-domain summary for the matrix."""
    report = assessment.get("report_detail") or {}
    summaries: list[str] = []
    for item in report.get("value_distributions") or []:
        column = str(item.get("column") or "campo")
        if item.get("suppressed"):
            summaries.append(f"{column}: suprimido por privacidad")
            continue
        values = item.get("values") or []
        safe_values = []
        for value in values[:8]:
            label = value.get("value_label")
            count = value.get("employee_count")
            if label is not None and count is not None:
                safe_values.append(f"{label} ({count})")
        if safe_values:
            summaries.append(f"{column}: " + ", ".join(safe_values))
    for band in (report.get("numeric_bands") or [])[:8]:
        label = band.get("band_label")
        count = band.get("employee_count")
        if label is not None and count is not None:
            summaries.append(f"Banda {label}: {count}")
    return " | ".join(summaries) or "Sin valores agregados publicables."


def validate_assessment_rationales(
    assessment: dict[str, Any],
) -> dict[str, dict[str, Any]]:
    variable = str(assessment.get("variable") or "Variable")
    scores = assessment.get("scores") or {}
    rationales = assessment.get("score_rationales")
    if not isinstance(rationales, list):
        raise ValueError(f"{variable}: score_rationales must be a list.")

    by_criterion: dict[str, dict[str, Any]] = {}
    required_fields = (
        "criterion",
        "score",
        "status_rule",
        "why_below_2",
        "supporting_evidence",
        "recommended_action",
        "scoring_method",
        "source_table",
        "source_columns",
    )
    for rationale in rationales:
        if not isinstance(rationale, dict):
            raise ValueError(f"{variable}: each score rationale must be an object.")
        criterion = str(
            rationale.get("criterion_key")
            or rationale.get("criterion")
            or ""
        ).strip().lower()
        if criterion not in CRITERIA:
            raise ValueError(f"{variable}: invalid rationale criterion {criterion!r}.")
        if criterion in by_criterion:
            raise ValueError(f"{variable}: duplicate rationale for {criterion}.")
        missing = [
            field
            for field in required_fields
            if rationale.get(field) is None or str(rationale.get(field)).strip() == ""
        ]
        if missing:
            raise ValueError(
                f"{variable}: rationale for {criterion} is missing "
                + ", ".join(missing)
                + "."
            )
        rationale_score = int(rationale["score"])
        if rationale_score not in {0, 1}:
            raise ValueError(
                f"{variable}: rationale for {criterion} has score "
                f"{rationale_score}; only 0/1 scores may have rationale rows."
            )
        if int(scores.get(criterion, -1)) != rationale_score:
            raise ValueError(
                f"{variable}: rationale score for {criterion} does not match assessment."
            )
        by_criterion[criterion] = rationale

    expected = {
        criterion for criterion in CRITERIA if int(scores.get(criterion, -1)) in {0, 1}
    }
    actual = set(by_criterion)
    if actual != expected:
        raise ValueError(
            f"{variable}: rationale coverage mismatch; missing={sorted(expected - actual)}, "
            f"extra={sorted(actual - expected)}."
        )

    verification = assessment.get("score_verification")
    if not isinstance(verification, dict):
        raise ValueError(f"{variable}: score_verification must be an object.")
    expected_verified = {
        criterion
        for criterion in QUALITATIVE_CRITERIA
        if int(scores.get(criterion, -1)) == 2
    }
    actual_verified = set(verification)
    if actual_verified != expected_verified:
        raise ValueError(
            f"{variable}: score-2 verification coverage mismatch; "
            f"missing={sorted(expected_verified - actual_verified)}, "
            f"extra={sorted(actual_verified - expected_verified)}."
        )
    required_verification_fields = (
        "method",
        "evidence",
        "result",
        "as_of_date",
        "query_reference",
        "acceptance_rule",
        "outcome",
        "semantic_type",
    )
    for criterion, item in verification.items():
        if not isinstance(item, dict):
            raise ValueError(
                f"{variable}: score-2 verification for {criterion} must be an object."
            )
        missing = [
            field
            for field in required_verification_fields
            if item.get(field) is None or str(item.get(field)).strip() == ""
        ]
        if missing:
            raise ValueError(
                f"{variable}: score-2 verification for {criterion} is missing "
                + ", ".join(missing)
                + "."
            )
        if str(item["outcome"]).strip().lower() != "pass":
            raise ValueError(
                f"{variable}: score-2 verification for {criterion} must have "
                "outcome 'pass'."
            )
        if criterion in {"definition", "consistency", "join"}:
            metrics = item.get("metrics")
            if not isinstance(metrics, dict) or not metrics:
                raise ValueError(
                    f"{variable}: score-2 verification for {criterion} requires "
                    "aggregate metrics."
                )
    return by_criterion


def brief_text(value: Any, limit: int = 180) -> str:
    text = " ".join(str(value or "").split())
    if len(text) <= limit:
        return text
    return text[: limit - 1].rstrip() + "…"


def score_comment(rationale: dict[str, Any]) -> Comment:
    message = (
        f"Score {int(rationale['score'])} - "
        f"{brief_text(rationale['recommended_action'], 140)}\n"
        f"Motivo: {brief_text(rationale['why_below_2'])}\n"
        f"Evidencia: {brief_text(rationale['supporting_evidence'])}\n"
        'Consulta "Score rationale" para ver el detalle.'
    )
    return Comment(message, "evaluate-attrition-variable")


def append_assessment(sheet, assessment: dict[str, Any]) -> None:
    scores = assessment.get("scores") or {}
    missing_scores = [criterion for criterion in CRITERIA if criterion not in scores]
    if missing_scores:
        raise ValueError(
            f"{assessment.get('variable')}: missing score(s): {', '.join(missing_scores)}"
        )
    rationale_by_criterion = validate_assessment_rationales(assessment)
    values = [
        assessment.get("category"),
        assessment.get("variable"),
        assessment.get("notes"),
        assessment_considerations(assessment),
        assessment.get("source_type"),
        assessment.get("source_table"),
        assessment.get("source_columns"),
        assessment.get("level_data"),
        assessment.get("history_available"),
        assessment.get("observation_frequency"),
        assessment_semantic_type(assessment),
        assessment_values_summary(assessment),
        assessment.get("blocker_label"),
        *[int(scores[criterion]) for criterion in CRITERIA],
        int(assessment.get("total", sum(int(scores[c]) for c in CRITERIA))),
        assessment.get("result"),
    ]
    sheet.append(values)
    row = sheet.max_row
    set_body_style(sheet, row, len(ASSESSMENT_HEADERS))
    for column in range(14, 22):
        sheet.cell(row, column).alignment = Alignment(horizontal="center", vertical="center")
    sheet.cell(row, 22).alignment = Alignment(horizontal="center", vertical="center")
    sheet.cell(row, 22).font = Font(bold=True)
    sheet.cell(row, 13).fill = label_fill(str(assessment.get("blocker_label") or ""))
    sheet.cell(row, 23).fill = label_fill(str(assessment.get("result") or ""))
    sheet.cell(row, 23).font = Font(bold=True)

    for criterion, rationale in rationale_by_criterion.items():
        column = 14 + CRITERIA.index(criterion)
        sheet.cell(row, column).comment = score_comment(rationale)
    point_result = assessment.get("point_result")
    result = assessment.get("result")
    if point_result and result != point_result:
        sheet.cell(row, 23).comment = Comment(
            f"Resultado por puntos: {point_result}. El blocker o label especial prevalece sobre los puntos.",
            "evaluate-attrition-variable",
        )


def append_coverage_row(
    sheet, assessment: dict[str, Any], coverage: dict[str, Any] | None
) -> None:
    warnings = assessment.get("warnings") or []
    if coverage is None:
        values = [
            assessment.get("variable"),
            assessment.get("source_table"),
            assessment.get("source_columns"),
            assessment.get("measurement_type"),
            None,
            None,
            None,
            None,
            None,
            None,
            int((assessment.get("scores") or {}).get("coverage", 0)),
            None,
            None,
            "No",
            parse_date(assessment.get("first_observed_date")),
            parse_date(assessment.get("usable_since_date")),
            assessment.get("population_filter"),
            assessment.get("numerator_definition"),
            join_notes(
                assessment.get("definition_text"),
                "No existen observaciones validadas de coverage.",
                " ".join(warnings),
            ),
        ]
    else:
        coverage_percent = coverage.get("observable_coverage_percent")
        event_percent = coverage.get("event_prevalence_percent")
        values = [
            assessment.get("variable"),
            assessment.get("source_table"),
            assessment.get("source_columns"),
            assessment.get("measurement_type"),
            parse_date(coverage.get("snapshot_date")),
            coverage.get("fiscal_year"),
            coverage.get("eligible_employees"),
            coverage.get("observable_employees"),
            coverage.get("missing_employees"),
            None if coverage_percent is None else float(coverage_percent) / 100.0,
            coverage.get("coverage_score"),
            coverage.get("event_employees"),
            None if event_percent is None else float(event_percent) / 100.0,
            "Sí" if coverage.get("stable_at_threshold") else "No",
            parse_date(coverage.get("first_observed_date")),
            parse_date(coverage.get("usable_since_date")),
            assessment.get("population_filter"),
            assessment.get("numerator_definition"),
            join_notes(
                coverage.get("notes"),
                " ".join(warnings) if coverage is assessment.get("coverage", [None])[0] else None,
            ),
        ]
    sheet.append(values)
    row = sheet.max_row
    set_body_style(sheet, row, len(COVERAGE_HEADERS))
    for column in (5, 15, 16):
        sheet.cell(row, column).number_format = "yyyy-mm-dd"
    for column in (10, 13):
        sheet.cell(row, column).number_format = "0.0%"
    for column in range(7, 15):
        sheet.cell(row, column).alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )
    sheet.row_dimensions[row].height = 36


def append_rationale_row(sheet, rationale: dict[str, Any]) -> None:
    values = [
        rationale.get("category"),
        rationale.get("variable"),
        rationale.get("criterion"),
        int(rationale.get("score")),
        rationale.get("status_rule"),
        rationale.get("why_below_2"),
        rationale.get("supporting_evidence"),
        rationale.get("recommended_action"),
        rationale.get("scoring_method"),
        rationale.get("source_table"),
        rationale.get("source_columns"),
    ]
    sheet.append(values)
    row = sheet.max_row
    set_body_style(sheet, row, len(RATIONALE_HEADERS))
    sheet.cell(row, 4).alignment = Alignment(horizontal="center", vertical="center")
    sheet.cell(row, 4).font = Font(bold=True)
    sheet.row_dimensions[row].height = 72


def render(profile: dict[str, Any], template: Path, output: Path) -> None:
    schema_version = int(profile.get("schema_version") or 0)
    if schema_version < 4:
        raise ValueError(
            f"Profile schema version 4 or newer is required; version "
            f"{schema_version or 'missing'} is legacy and lacks v4 executable "
            "derivation and scoring evidence."
        )
    workbook = load_workbook(template)
    expected_sheets = [ASSESSMENT_SHEET, COVERAGE_SHEET, RATIONALE_SHEET]
    if workbook.sheetnames != expected_sheets:
        raise ValueError(
            "Template must contain exactly: Variable assessment, Coverage, "
            "Score rationale."
        )
    assessment_sheet = workbook[ASSESSMENT_SHEET]
    coverage_sheet = workbook[COVERAGE_SHEET]
    rationale_sheet = workbook[RATIONALE_SHEET]
    clear_data_rows(assessment_sheet)
    clear_data_rows(coverage_sheet)
    clear_data_rows(rationale_sheet)

    for assessment in profile["assessments"]:
        append_assessment(assessment_sheet, assessment)
        coverage_rows = assessment.get("coverage") or []
        if coverage_rows:
            for coverage in coverage_rows:
                append_coverage_row(coverage_sheet, assessment, coverage)
        else:
            append_coverage_row(coverage_sheet, assessment, None)
        for rationale in assessment.get("score_rationales") or []:
            append_rationale_row(rationale_sheet, rationale)

    assessment_sheet.auto_filter.ref = f"A1:W{assessment_sheet.max_row}"
    coverage_sheet.auto_filter.ref = f"A1:S{coverage_sheet.max_row}"
    rationale_sheet.auto_filter.ref = f"A1:K{rationale_sheet.max_row}"
    assessment_sheet.sheet_properties.pageSetUpPr.fitToPage = True
    coverage_sheet.sheet_properties.pageSetUpPr.fitToPage = True
    rationale_sheet.sheet_properties.pageSetUpPr.fitToPage = True
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    validate_workbook(output, profile)


def validate_workbook(path: Path, profile: dict[str, Any] | None = None) -> None:
    workbook = load_workbook(path, read_only=False, data_only=False)
    if workbook.sheetnames != [
        ASSESSMENT_SHEET,
        COVERAGE_SHEET,
        RATIONALE_SHEET,
    ]:
        raise ValueError(f"Unexpected sheet order in {path}")
    assessment_headers = [
        workbook[ASSESSMENT_SHEET].cell(1, column).value
        for column in range(1, len(ASSESSMENT_HEADERS) + 1)
    ]
    coverage_headers = [
        workbook[COVERAGE_SHEET].cell(1, column).value
        for column in range(1, len(COVERAGE_HEADERS) + 1)
    ]
    rationale_headers = [
        workbook[RATIONALE_SHEET].cell(1, column).value
        for column in range(1, len(RATIONALE_HEADERS) + 1)
    ]
    if assessment_headers != ASSESSMENT_HEADERS:
        raise ValueError("Variable assessment headers do not match the required layout.")
    if coverage_headers != COVERAGE_HEADERS:
        raise ValueError("Coverage headers do not match the required layout.")
    if rationale_headers != RATIONALE_HEADERS:
        raise ValueError("Score rationale headers do not match the required layout.")
    if profile is None:
        return

    assessments = profile["assessments"]
    assessment_sheet = workbook[ASSESSMENT_SHEET]
    coverage_sheet = workbook[COVERAGE_SHEET]
    rationale_sheet = workbook[RATIONALE_SHEET]
    expected_coverage_rows = 1 + sum(
        len(assessment.get("coverage") or []) or 1
        for assessment in assessments
    )
    expected_rationales = [
        rationale
        for assessment in assessments
        for rationale in (assessment.get("score_rationales") or [])
    ]
    if assessment_sheet.max_row != len(assessments) + 1:
        raise ValueError("Variable assessment row count does not match the profile.")
    if coverage_sheet.max_row != expected_coverage_rows:
        raise ValueError("Coverage row count does not match the profile.")
    if rationale_sheet.max_row != len(expected_rationales) + 1:
        raise ValueError("Score rationale row count does not match the profile.")

    for row, assessment in enumerate(assessments, start=2):
        scores = assessment["scores"]
        if assessment_sheet.cell(row, 2).value != assessment.get("variable"):
            raise ValueError(f"Assessment row {row} has an unexpected variable.")
        expected_total = sum(int(scores[criterion]) for criterion in CRITERIA)
        if assessment_sheet.cell(row, 22).value != expected_total:
            raise ValueError(f"Assessment row {row} has an invalid total.")
        if assessment_sheet.cell(row, 23).value != assessment.get("result"):
            raise ValueError(f"Assessment row {row} has an invalid result label.")
        validate_assessment_rationales(assessment)
        for offset, criterion in enumerate(CRITERIA, start=14):
            score = int(scores[criterion])
            cell = assessment_sheet.cell(row, offset)
            if cell.value != score:
                raise ValueError(
                    f"Assessment row {row} has an invalid {criterion} score."
                )
            if score in {0, 1} and cell.comment is None:
                raise ValueError(
                    f"Assessment row {row} {criterion} score lacks a comment."
                )
            if score == 2 and cell.comment is not None:
                raise ValueError(
                    f"Assessment row {row} {criterion} score 2 has a comment."
                )

    for row, rationale in enumerate(expected_rationales, start=2):
        expected_values = [
            rationale.get("category"),
            rationale.get("variable"),
            rationale.get("criterion"),
            int(rationale.get("score")),
            rationale.get("status_rule"),
            rationale.get("why_below_2"),
            rationale.get("supporting_evidence"),
            rationale.get("recommended_action"),
            rationale.get("scoring_method"),
            rationale.get("source_table"),
            rationale.get("source_columns"),
        ]
        actual_values = [
            rationale_sheet.cell(row, column).value
            for column in range(1, len(RATIONALE_HEADERS) + 1)
        ]
        if actual_values != expected_values:
            raise ValueError(f"Score rationale row {row} does not match the profile.")


def default_output_path(profile: dict[str, Any], output_dir: Path) -> Path:
    assessments = profile["assessments"]
    slug = (
        slugify(str(assessments[0].get("variable") or "variable"))
        if len(assessments) == 1
        else "multiple-variables"
    )
    timestamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    return output_dir / f"{slug}_assessment_{timestamp}.xlsx"


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Render profiled attrition variables to the approved Excel layout."
    )
    parser.add_argument("--profile", help="Profile JSON from profile_variable.py.")
    parser.add_argument("--template", default=str(DEFAULT_TEMPLATE))
    parser.add_argument("--output", help="Exact output workbook path.")
    parser.add_argument("--output-dir", default="output/assessments")
    parser.add_argument(
        "--create-template",
        action="store_true",
        help="Create or replace the three-sheet workbook template, then exit.",
    )
    args = parser.parse_args()

    template = Path(args.template).resolve()
    if args.create_template:
        create_template(template)
        validate_workbook(template)
        print(f"Created template: {template}")
        return
    if not args.profile:
        raise SystemExit("--profile is required unless --create-template is used.")
    if not template.exists():
        raise SystemExit(f"Template does not exist: {template}")

    profile = load_profile(Path(args.profile))
    output = (
        Path(args.output).resolve()
        if args.output
        else default_output_path(profile, Path(args.output_dir).resolve())
    )
    render(profile, template, output)
    print(f"Created workbook: {output}")


if __name__ == "__main__":
    try:
        main()
    except Exception as exc:
        raise SystemExit(f"Workbook rendering failed: {exc}") from exc
